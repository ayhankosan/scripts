"""FFA project locations: Fabric version 4 to version 8 change analysis.

Same cells as FFA_Fabric_V4_V8_Change_Analysis_v3.ipynb. Run the whole file with
    python FFA_Fabric_V4_V8_Change_Analysis_v3.py
or open it in Visual Studio Code and use "Run Cell" on each "# %%" block. Edit only the Configuration cell.
"""
from __future__ import annotations

try:                                    # tables show as formatted output in VS Code, as plain text in a console
    from IPython.display import display
except ImportError:
    display = print


# %% [markdown]
# # FFA project locations: Fabric V4 → V8 change analysis
#
# **Goal.** For the six Federal Funding Account (FFA) projects below, and only these, show how moving from Fabric Version 4 to Version 8 changed (a) the applicant-submitted location lists and (b) the FFA-eligible Fabric locations, and explain how the two changes relate.
#
# | Project | PR | V4 applicant submission | V8 applicant submission |
# |---|---|---|---|
# | San Luis Obispo Affordable Broadband Network | PR-5578 | All_448_PIDs… CSV | Q3 2026 report |
# | Oakland Connect – Last Mile Connectivity for Oaklanders | PR-4651 | All_448_PIDs… CSV | Q2 2026 report |
# | Coloma Broadband Project (El Dorado) | PR-5631 | All_448_PIDs… CSV | Q3 2026 report |
# | Indian Creek Broadband Project (El Dorado) | PR-5676 | All_448_PIDs… CSV | Q3 2026 report |
# | Garden Valley Broadband Project (El Dorado) | PR-5674 | All_448_PIDs… CSV | Q3 2026 report |
# | Georgetown Broadband Project (El Dorado) | PR-5644 | All_448_PIDs… CSV | Q3 2026 report |
#
# **Eligibility rule.** Served = not eligible, Unserved = eligible. V4: `FFAv4Dec23`, V8: `status_code`; both coded **1 = served, 0 = unserved (eligible)**.
#
# **Analysis structure**
#
# | # | Comparison | Question answered |
# |---|---|---|
# | 1 | Initial grantee submission ↔ Fabric version 4 eligibility data | Were the IDs in the initial grantee submission valid Fabric version 4 locations, inside the project area, and eligible? |
# | 2 | 2026 Q2 and Q3 data ↔ Fabric version 8 eligibility data | Same for the IDs in the 2026 Q2 and Q3 data against the Fabric version 8 eligibility data |
# | 3 | Initial grantee submission ↔ 2026 Q2 and Q3 data | Which locations are in both, not carried over or new? |
# | 4 | Fabric version 4 eligibility data ↔ Fabric version 8 eligibility data (project area) | How did total and eligible Fabric locations in each project area change, and why (one cause category per location)? |
# | 5 | **Synthesis** | For every submitted location, the Fabric change that explains its submission change, plus eligible-location coverage per version |
#
# Edit **only the Configuration cell**, then *Run All*. Section 3a prints the distinct values of each eligibility field so the eligible-value sets can be checked before trusting the numbers.

# %%

import logging
import re
import shutil
import warnings
from pathlib import Path

import duckdb
import geopandas as gpd
import numpy as np
import pandas as pd
import pyogrio
import requests
import shapely

logging.basicConfig(level=logging.INFO, format="%(asctime)s  %(levelname)s  %(message)s", datefmt="%H:%M:%S")
log = logging.getLogger("ffa_v4_v8")
warnings.filterwarnings("ignore", category=FutureWarning)   # pandas deprecation notices only
pd.set_option("display.max_columns", 60)
pd.set_option("display.width", 200)

# %% [markdown]
# ## Configuration
# Paths, field names and eligibility values. Every other cell reads from here.

# %%
# ---------------------------------------------------------------- folders
BASE_DIR = Path(r"C:\Change_Analysis")
INPUT_DIR = BASE_DIR / "inputs"
OUTPUT_DIR = BASE_DIR / "outputs"
OUTPUT_GDB = OUTPUT_DIR / "FFA_V4_V8_Change_Analysis.gdb"
SUMMARY_CSV = OUTPUT_DIR / "FFA_V4_V8_project_summary.csv"

# ---------------------------------------------------------------- which outputs to write
# The results workbooks (all projects and one per project), the PDF briefs and the geodatabases are always written.
WRITE_CSV = True         # every results tab also as a CSV file (outputs/csv and each project's csv folder)
WRITE_MARKDOWN = True    # Markdown reports: all projects, and one per project
WRITE_HTML = False       # True also writes the HTML report

# ---------------------------------------------------------------- the six projects (only these are analysed)
PROJECTS = {
    "SLO":       {"label": "San Luis Obispo Affordable Broadband Network",
                  "name_like": "%San Luis Obispo Affordable Broadband Network%"},
    "OAK":       {"label": "Oakland Connect - Last Mile Connectivity for Oaklanders",
                  "name_like": "%Oakland Connect%Last Mile Connectivity for Oaklanders%"},
    "ED_COLOMA": {"label": "Coloma Broadband Project",
                  "name_like": "%Coloma Broadband Project%"},
    "ED_INDCRK": {"label": "Indian Creek Broadband Project",
                  "name_like": "%Indian Creek Broadband Project%"},
    "ED_GRDNVL": {"label": "Garden Valley Broadband Project",
                  "name_like": "%Garden Valley Broadband Project%"},
    "ED_GRGTWN": {"label": "Georgetown Broadband Project",
                  "name_like": "%Georgetown Broadband Project%"},
}
PR_NUMBERS = {"SLO": "PR-5578", "OAK": "PR-4651", "ED_COLOMA": "PR-5631",
              "ED_INDCRK": "PR-5676", "ED_GRDNVL": "PR-5674", "ED_GRGTWN": "PR-5644"}

# ---------------------------------------------------------------- Fabric eligibility inputs
# CSV or file geodatabase (.gdb → set the *_LAYER name)
V4_PATH = INPUT_DIR / "v4_eligibility_locations.csv"   # or .../Combined_EligibilityV4_Updated_02192025.gdb
V4_LAYER = None
V8_PATH = INPUT_DIR / "FFA_v8_D25_eligibility_full.csv"
V8_LAYER = None

# Field roles in each Fabric file ("address" and "units" may be None; a field missing from the file is skipped).
# V4 eligibility = FFA vintage December 2023. If the exact name differs, the notebook
# looks for a single field containing both "FFA" and "Dec23" and uses it.
V4_FIELDS = {"id": "location_id", "lat": "latitude", "lon": "longitude",
             "elig": "FFAv4Dec23", "bsl": "bsl_flag", "address": "address_primary", "units": "unit_count"}
V8_FIELDS = {"id": "location_id", "lat": "latitude", "lon": "longitude",
             "elig": "status_code", "bsl": "bsl_flag", "address": None, "units": "unit_count"}

# Values in the eligibility field that mean "eligible" (unserved).
# V4 FFAv4Dec23 and V8 status_code are both coded 1 = served, 0 = unserved.
# Section 3a prints the actual distinct values: confirm these sets against it.
V4_ELIGIBLE_VALUES = {"0"}
V8_ELIGIBLE_VALUES = {"0"}

# False = every Fabric location counts, whatever its bsl_flag.
# True  = only bsl_flag = True locations count in totals.
BSL_ONLY = False

# ---------------------------------------------------------------- applicant submissions
# V4 submission: GMS application locations (Fabric V4 IDs)
SUB_V4_PATH = INPUT_DIR / "All_448_PIDs_Combined_Filtered_Only_113_Reward_Projects.csv"
SUB_V4_FIELDS = {"id": "location_id", "pr": "GMS_PID", "title": "Project_Title",
                 "lat": "Latitude", "lon": "Longitude", "address": "Address_primary"}

# V8 submission: 2026 quarterly reports (Fabric V8 IDs), one file per project, found by
# PR number (or project name) in the file name. Prefixes are tried in order.
SUB_V8_DIR = INPUT_DIR
SUB_V8_PREFERENCE = ["Q3 2026", "Q3", "Q2 2026", "Q2"]
SUB_V8_FIELDS = {"id": "Fabric_ID_Number__c", "lat": "Latitude__c", "lon": "Longitude__c"}
HEADER_SEARCH_ROWS = 50        # rows searched for the header in each report sheet

# ---------------------------------------------------------------- project boundaries
# Local file geodatabase / GeoPackage / shapefile, or an ArcGIS feature service layer URL
# ending in /FeatureServer/<n> or /MapServer/<n>.
BOUNDARY_SOURCE = INPUT_DIR / "FFA_Projects_Awarded.gdb"
BOUNDARY_LAYER = None                              # None = read every layer that matches a project (one-layer-per-project gdb)
BOUNDARY_NAME_FIELDS = ["Project_Name", "Name"]    # name_like is tested against each
ARCGIS_TOKEN = None                                # only for secured services

# ---------------------------------------------------------------- tuning
DUCKDB_MEMORY_LIMIT = None      # e.g. "8GB"; None = DuckDB default (80% of RAM)
BBOX_PAD_DEG = 0.05            # read-window padding around the project areas (degrees)
METRIC_CRS = "EPSG:3310"       # California Albers, for distances
WGS84 = "EPSG:4326"

# %% [markdown]
# ## Helper functions

# %%
def like_to_regex(pattern: str) -> re.Pattern:
    """Converts a SQL LIKE pattern (% and _) into a case-insensitive regex."""
    out = []
    for ch in pattern:
        if ch == "%":
            out.append(".*")
        elif ch == "_":
            out.append(".")
        else:
            out.append(re.escape(ch))
    return re.compile("^" + "".join(out) + "$", re.IGNORECASE | re.DOTALL)


def to_location_id(series: pd.Series) -> pd.Series:
    """Converts IDs such as '1012345678', 1012345678.0 or ' 1012345678 ' to Int64."""
    cleaned = series.astype("string").str.strip().str.replace(r"\.0+$", "", regex=True)
    cleaned = cleaned.str.replace(",", "", regex=False)
    num = pd.to_numeric(cleaned, errors="coerce")
    whole = num.where(np.isfinite(num) & (num == np.floor(num)))
    return whole.astype("Int64")


BSL_TRUE = {"1", "true", "t", "y", "yes"}


CLEANING: list[dict] = []            # one record per cleaned list: rows read, IDs dropped, duplicates removed
DUPLICATES: list[pd.DataFrame] = []   # every duplicated ID with how often it appeared


def standardize_ids(df: pd.DataFrame, source: str, prefer_col: str | None = None,
                    dataset: str | None = None, project_key: str | None = None, note: str = "") -> pd.DataFrame:
    """Converts location_id to whole numbers and removes duplicates.

    When an ID appears more than once, the record where ``prefer_col`` is True
    (the eligible record) is kept. With ``dataset``, the counts and the
    duplicated IDs are recorded in CLEANING / DUPLICATES for the outputs.
    """
    n0 = len(df)
    df = df.copy()
    raw = df["location_id"].astype("string").str.strip()
    blank = int((raw.isna() | raw.eq("")).sum())                 # no ID at all
    df["location_id"] = to_location_id(df["location_id"])
    bad = int(df["location_id"].isna().sum())                    # blank + not a whole number
    df = df.dropna(subset=["location_id"])
    counts = df["location_id"].value_counts()
    dup_ids = counts[counts > 1]
    if prefer_col is not None:
        df = df.sort_values(prefer_col, ascending=False, kind="stable")
    df = df.drop_duplicates("location_id", keep="first")
    n_dup = n0 - bad - len(df)
    log.info("%s: %s rows → %s unique IDs (%s non-numeric IDs dropped, %s duplicates removed)",
             source, f"{n0:,}", f"{len(df):,}", bad, n_dup)
    if dataset is not None:
        CLEANING.append({"dataset": dataset, "project_key": project_key, "rows_read": n0,
                         "non_numeric_rows_skipped": bad, "blank_id_rows": blank, "invalid_id_rows": bad - blank,
                         "ids_listed_more_than_once": len(dup_ids),
                         "duplicate_rows_removed": n_dup, "unique_ids_used": len(df), "note": note})
        if len(dup_ids):
            DUPLICATES.append(pd.DataFrame({"dataset": dataset, "project_key": project_key,
                                            "location_id": dup_ids.index.astype("int64"),
                                            "times_listed": dup_ids.to_numpy()}))
    return df.reset_index(drop=True)


def points_from_lonlat(df: pd.DataFrame, lon: str = "lon", lat: str = "lat") -> gpd.GeoDataFrame:
    """Builds a WGS84 point GeoDataFrame; rows without coordinates get empty geometry."""
    x = pd.to_numeric(df[lon], errors="coerce")
    y = pd.to_numeric(df[lat], errors="coerce")
    geom = gpd.points_from_xy(x, y)
    geom[(x.isna() | y.isna()).to_numpy()] = None
    return gpd.GeoDataFrame(df, geometry=geom, crs=WGS84)


def check_gdb(path: Path) -> Path:
    """Validates a file geodatabase folder and returns the path GDAL can open.

    Handles the nested folder created when a .gdb.zip is extracted on Windows
    (X.gdb\\X.gdb) and explains the usual causes of "Permission denied".
    Non-.gdb paths are returned unchanged.
    """
    path = Path(path)
    if path.suffix.lower() != ".gdb":
        return path
    if path.is_file():
        raise ValueError(f"{path} is a file, not a geodatabase folder. If it is a zip archive, extract it first.")
    if not path.exists():
        zipped = path.with_name(path.name + ".zip")
        hint = f" A zipped copy exists ({zipped.name}): extract it." if zipped.exists() else ""
        raise FileNotFoundError(f"{path} does not exist.{hint}")
    try:
        if any(path.glob("*.gdbtable")):
            return path
        nested = [p for p in path.rglob("*.gdb") if p.is_dir() and any(p.glob("*.gdbtable"))]
        contents = sorted(p.name for p in path.iterdir())[:15]
    except PermissionError as e:
        raise PermissionError(f"Windows denies access to {path}. Check that the folder is fully downloaded "
                              "(OneDrive/SharePoint: right-click → Always keep on this device) and readable "
                              "by this user.") from e
    if len(nested) == 1:
        log.warning("Using nested geodatabase %s (typical after extracting a .gdb.zip)", nested[0])
        return nested[0]
    raise PermissionError(
        f"{path} contains no .gdbtable files, so GDAL cannot open it as a geodatabase "
        f"(Windows reports this as 'Permission denied'). Folder contents: {contents}. "
        "Typical causes: the zip was extracted into a nested or wrong folder, the folder is an "
        "online-only OneDrive/SharePoint copy, or the path points to the parent folder instead of the .gdb.")


def pct_change(old: float, new: float) -> float:
    """Percent change from old to new; NaN when old is zero."""
    return round((new - old) / old * 100, 2) if old else np.nan


def table_columns(path: Path, layer: str | None) -> list[str]:
    """Field names of a CSV or GDAL-readable table without reading its rows."""
    path = check_gdb(path)
    if path.suffix.lower() in {".gdb", ".gpkg", ".shp"}:
        return list(pyogrio.read_info(path, layer=layer)["fields"])
    return list(pd.read_csv(path, nrows=0, encoding="utf-8-sig").columns)


def resolve_elig_field(path: Path, layer: str | None, wanted: str, hint: tuple[str, ...] = ()) -> str:
    """Returns the eligibility field name, falling back to a unique case-insensitive hint match."""
    cols = table_columns(path, layer)
    if wanted in cols:
        return wanted
    lower = {c.lower(): c for c in cols}
    if wanted.lower() in lower:
        return lower[wanted.lower()]
    if hint:
        cands = [c for c in cols if all(h.lower() in c.lower().replace("_", "") for h in hint)]
        if len(cands) == 1:
            log.warning("Field %r not in %s; using %r instead", wanted, path.name, cands[0])
            return cands[0]
    ffa = [c for c in cols if "ffa" in c.lower()]
    raise KeyError(f"Eligibility field {wanted!r} not found in {path.name}. FFA-related fields: {ffa}")

# %% [markdown]
# ## 1. Project boundaries
# Each project polygon is pulled by name; multi-part projects are dissolved into one area. Matched names are printed for checking.

# %%
def read_feature_service(url: str, token: str | None = None) -> gpd.GeoDataFrame:
    """Reads every feature of an ArcGIS REST layer as WGS84 GeoJSON, page by page."""
    url = url.rstrip("/")
    feats, offset = [], 0
    while True:
        params = {"where": "1=1", "outFields": "*", "outSR": 4326, "f": "geojson",
                  "resultOffset": offset, "resultRecordCount": 1000}
        if token:
            params["token"] = token
        r = requests.get(f"{url}/query", params=params, timeout=120)
        r.raise_for_status()
        payload = r.json()
        if "error" in payload:
            raise RuntimeError(f"Feature service error: {payload['error']}")
        batch = payload.get("features", [])
        feats.extend(batch)
        exceeded = payload.get("exceededTransferLimit") or payload.get("properties", {}).get("exceededTransferLimit")
        if not batch or not exceeded:
            break
        offset += len(batch)
    return gpd.GeoDataFrame.from_features(feats, crs=WGS84)


def count_project_matches(df: pd.DataFrame) -> int:
    """Number of the six projects whose name_like matches a value in the name fields or "_layer"."""
    fields = [f for f in BOUNDARY_NAME_FIELDS + ["_layer"] if f in df.columns]
    values = {str(v) for f in fields for v in df[f].dropna()}
    return sum(any(like_to_regex(p["name_like"]).match(v) for v in values) for p in PROJECTS.values())


def read_boundary_layers(path: Path) -> gpd.GeoDataFrame:
    """Reads every polygon layer whose name fields or layer name match one of the six projects.

    Works both for one combined project layer and for geodatabases that store one
    layer per project (e.g. P3616_San_Luis_Obispo_Affordable_Broadband_Network).
    The layer name, with underscores as spaces, is kept in field "_layer" and matched too.
    """
    found, parts = [], []
    for name, geom_type in pyogrio.list_layers(path):
        if geom_type is None or "polygon" not in str(geom_type).lower():
            continue
        fields = [f for f in BOUNDARY_NAME_FIELDS if f in pyogrio.read_info(path, layer=name)["fields"]]
        lname = name.replace("_", " ")
        vals = pyogrio.read_dataframe(path, layer=name, columns=fields, read_geometry=False) if fields else pd.DataFrame()
        n = count_project_matches(pd.concat([vals, pd.DataFrame({"_layer": [lname]})], ignore_index=True))
        if n:
            gdf = gpd.read_file(path, layer=name)
            if gdf.crs is None:
                log.warning("Layer %r has no CRS; skipped", name)
                continue
            gdf["_layer"] = lname
            parts.append(gdf.to_crs(WGS84))
            found.append((name, n, len(gdf)))
    if not parts:
        raise ValueError(f"No polygon layer in {path.name} matches any project name in {BOUNDARY_NAME_FIELDS} "
                         "or in the layer names. Set BOUNDARY_LAYER / BOUNDARY_NAME_FIELDS in the Configuration cell.")
    print(pd.DataFrame(found, columns=["layer used", "projects_matched", "features"]).to_string(index=False))
    return gpd.GeoDataFrame(pd.concat(parts, ignore_index=True), crs=WGS84)


def polygonal_2d(geom) -> shapely.MultiPolygon:
    """Keeps only the polygon parts of a geometry, drops Z/M, and returns a MultiPolygon.

    make_valid / union_all can return a GeometryCollection holding stray lines or
    points; the file geodatabase driver rejects that geometry type.
    """
    polys, stack = [], [shapely.force_2d(geom)]
    while stack:
        g = stack.pop()
        if g is None or g.is_empty:
            continue
        if g.geom_type == "Polygon":
            polys.append(g)
        elif g.geom_type in ("MultiPolygon", "GeometryCollection"):
            stack.extend(shapely.get_parts(g))
    return shapely.MultiPolygon(polys)


def load_boundaries() -> gpd.GeoDataFrame:
    """Returns one dissolved WGS84 polygon per project key, with matched names."""
    src = str(BOUNDARY_SOURCE)
    if src.lower().startswith("http"):
        raw = read_feature_service(src, ARCGIS_TOKEN)
    else:
        src_path = check_gdb(Path(BOUNDARY_SOURCE))
        raw = gpd.read_file(src_path, layer=BOUNDARY_LAYER) if BOUNDARY_LAYER else read_boundary_layers(src_path)
    if raw.crs is None:
        raise ValueError("Boundary layer has no CRS; define it before running.")
    raw = raw.to_crs(WGS84)
    raw["geometry"] = shapely.make_valid(raw.geometry.values)
    name_fields = [f for f in BOUNDARY_NAME_FIELDS + ["_layer"] if f in raw.columns]
    if not name_fields:
        raise KeyError(f"None of {BOUNDARY_NAME_FIELDS} found in boundary fields: {list(raw.columns)}")

    rows = []
    for key, p in PROJECTS.items():
        rx = like_to_regex(p["name_like"])
        hit = pd.Series(False, index=raw.index)
        for f in name_fields:
            hit |= raw[f].map(lambda v: bool(rx.match(str(v))) if pd.notna(v) else False).astype(bool)
        sub = raw[hit]
        if sub.empty:
            log.error("%s: no boundary matched %r in %s", key, p["name_like"], name_fields)
            continue
        names = sorted({str(v) for f in name_fields for v in sub[f].dropna()})
        print(f"{key:<10} {len(sub)} part(s) matched → {names}")
        rows.append({"project_key": key, "label": p["label"], "pr_number": PR_NUMBERS[key],
                     "n_parts": len(sub), "matched_names": " | ".join(names)[:500],
                     "geometry": polygonal_2d(shapely.union_all(sub.geometry.values))})
    missing = set(PROJECTS) - {r["project_key"] for r in rows}
    if missing:
        raise ValueError(f"No boundary found for: {sorted(missing)}. Check name_like / BOUNDARY_NAME_FIELDS.")
    return gpd.GeoDataFrame(rows, crs=WGS84)


areas = load_boundaries()
areas["area_sq_mi"] = (areas.to_crs(METRIC_CRS).area / 2_589_988.11).round(2)
areas[["project_key", "label", "n_parts", "area_sq_mi", "matched_names"]]

# %% [markdown]
# ## 2. Grantee submissions
# * **Initial grantee submission** – rows of `All_448_PIDs_Combined_Filtered_Only_113_Reward_Projects.csv` for the six projects (matched on `GMS_PID`; falls back to `Project_Title` if a PR number finds nothing).
# * **2026 Q2 and Q3 data** – the 2026 quarterly report per project (Q3 2026, or Q2 when Q3 is missing), found by PR number or project name in the file name.

# %%
def pr_digits(value) -> str:
    """Extracts the numeric part of a PR number ('PR-5578' → '5578')."""
    m = re.search(r"(\d+)", str(value))
    return m.group(1) if m else ""


def read_table(path: Path, header_field: str | None = None) -> pd.DataFrame:
    """Reads a CSV or Excel table with every column as text.

    With ``header_field``, the header row is the first row (in any sheet) that
    contains that field name. This handles report templates with title, label,
    "Required" and help-text rows above or below the header; those extra rows
    are dropped later because their IDs are not numeric.
    """
    is_excel = path.suffix.lower() in {".xlsx", ".xlsm", ".xls"}
    if header_field is None:
        return (pd.read_excel(path, dtype=str) if is_excel
                else pd.read_csv(path, dtype=str, encoding="utf-8-sig", low_memory=False))
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message=".*extension is not supported.*")
        sheets = (pd.read_excel(path, sheet_name=None, header=None, dtype=str) if is_excel
                  else {"csv": pd.read_csv(path, header=None, dtype=str, encoding="utf-8-sig", low_memory=False)})
    for sheet, raw in sheets.items():
        top = raw.head(HEADER_SEARCH_ROWS)
        hits = top.index[top.apply(lambda r: r.astype("string").str.strip().eq(header_field).any(), axis=1)]
        if len(hits):
            h = hits[0]
            df = raw.iloc[h + 1:].copy()
            df.columns = [str(c).strip() for c in raw.iloc[h]]
            df = df.dropna(how="all")
            log.info("%s: header found on sheet %r, row %s", path.name, sheet, h + 1)
            return df.reset_index(drop=True)
    raise KeyError(f"Field {header_field!r} not found in the first {HEADER_SEARCH_ROWS} rows of any sheet in {path.name}")


def load_sub_v4() -> pd.DataFrame:
    """V4 applicant submission for the six projects: project_key, location_id, lat, lon, address."""
    raw = read_table(SUB_V4_PATH)
    f = SUB_V4_FIELDS
    raw["_pr"] = raw[f["pr"]].map(pr_digits) if f.get("pr") in raw else ""
    frames = []
    for key, pr in PR_NUMBERS.items():
        sub = raw[raw["_pr"] == pr_digits(pr)]
        if sub.empty and f.get("title") in raw:
            rx = like_to_regex(PROJECTS[key]["name_like"])
            sub = raw[raw[f["title"]].fillna("").map(lambda t: bool(rx.match(t)))]
            log.info("%s: no rows for %s in %s; matched %s rows by project title", key, pr, f["pr"], len(sub))
        out = pd.DataFrame({"project_key": key, "location_id": sub[f["id"]],
                            "lat": pd.to_numeric(sub[f["lat"]], errors="coerce"),
                            "lon": pd.to_numeric(sub[f["lon"]], errors="coerce"),
                            "address": sub[f["address"]] if f.get("address") in sub else pd.NA})
        frames.append(standardize_ids(out, f"V4 submission {key} ({pr})", dataset="Initial grantee submission",
                                      project_key=key, note=SUB_V4_PATH.name))
    return pd.concat(frames, ignore_index=True)


def find_sub_v8_file(key: str) -> Path | None:
    """Picks the 2026 quarterly report for a project, honouring SUB_V8_PREFERENCE order."""
    pat = re.compile(rf"PR[-_ ]?0*{pr_digits(PR_NUMBERS[key])}(?!\d)", re.IGNORECASE)
    label_rx = like_to_regex(PROJECTS[key]["name_like"])
    files = [p for p in SUB_V8_DIR.iterdir()
             if p.suffix.lower() in {".csv", ".xlsx", ".xlsm", ".xls"} and p.resolve() != Path(SUB_V4_PATH).resolve()]
    cands = [p for p in files if pat.search(p.stem)] or [p for p in files if label_rx.match(p.stem)]
    for prefix in SUB_V8_PREFERENCE:
        hits = sorted(p for p in cands if p.stem.lower().startswith(prefix.lower()))
        if hits:
            return hits[-1]
    return sorted(cands)[-1] if cands else None


def load_sub_v8() -> tuple[pd.DataFrame, dict[str, str]]:
    """V8 applicant submission (2026 quarterly reports) plus the file used per project."""
    f = SUB_V8_FIELDS
    frames, used = [], {}
    for key in PROJECTS:
        path = find_sub_v8_file(key)
        if path is None:
            log.warning("%s: no 2026 quarterly report found in %s", key, SUB_V8_DIR)
            used[key] = ""
            continue
        used[key] = path.name
        raw = read_table(path, header_field=f["id"])
        out = pd.DataFrame({"project_key": key, "location_id": raw[f["id"]],
                            "lat": pd.to_numeric(raw.get(f["lat"]), errors="coerce"),
                            "lon": pd.to_numeric(raw.get(f["lon"]), errors="coerce"), "address": pd.NA})
        frames.append(standardize_ids(out, f"V8 submission {key} ← {path.name}", dataset="2026 Q2 and Q3 data",
                                      project_key=key,
                                      note=f"{path.name}; skipped rows include the template label, Required and help-text rows"))
    return pd.concat(frames, ignore_index=True), used


sub_v4 = load_sub_v4()
sub_v8, sub_v8_files = load_sub_v8()
submitted_ids = set(sub_v4["location_id"].astype("int64")) | set(sub_v8["location_id"].astype("int64"))
pd.DataFrame({"Initial grantee submission": sub_v4.groupby("project_key").size(),
              "2026 Q2 and Q3 data": sub_v8.groupby("project_key").size(),
              "V8 file": pd.Series(sub_v8_files)}).reindex(list(PROJECTS)).fillna(0)

# %% [markdown]
# ## 3. Fabric version 4 and version 8 eligibility data – DuckDB
# Each statewide file is scanned once by DuckDB, reading only the needed columns and keeping only rows inside the padded windows around the six project areas, plus any submitted ID wherever it is. In the same query, IDs are converted to whole numbers, eligibility and BSL flags are derived, and duplicate IDs are removed; when the copies of an ID disagree on eligibility, the ID is marked as *conflicting source records* and left out of every count (never resolved in favour of eligibility). Only the filtered rows come back into pandas.
#
# CSV files are read with `read_csv`; a `.gdb` / `.gpkg` is read with the DuckDB spatial extension (`ST_Read`), falling back to pyogrio if the extension cannot be loaded.

# %%
read_windows = [(a - BBOX_PAD_DEG, b - BBOX_PAD_DEG, c + BBOX_PAD_DEG, d + BBOX_PAD_DEG)
                for a, b, c, d in areas.geometry.bounds.to_numpy()]

V4_FIELDS["elig"] = resolve_elig_field(V4_PATH, V4_LAYER, V4_FIELDS["elig"], hint=("ffa", "dec23"))
V8_FIELDS["elig"] = resolve_elig_field(V8_PATH, V8_LAYER, V8_FIELDS["elig"])
print("V4 eligibility field:", V4_FIELDS["elig"], "| V8 eligibility field:", V8_FIELDS["elig"])

con = duckdb.connect()
if DUCKDB_MEMORY_LIMIT:
    con.execute(f"SET memory_limit = '{DUCKDB_MEMORY_LIMIT}'")
con.register("submitted", pd.DataFrame({"location_id": sorted(submitted_ids)}, dtype="int64"))


def q(name: str) -> str:
    """Quotes a column name as a SQL identifier."""
    return '"' + name.replace('"', '""') + '"'


def lit(text: str) -> str:
    """Quotes a value as a SQL string literal."""
    return "'" + str(text).replace("'", "''") + "'"


def norm_sql(expr: str) -> str:
    """SQL expression for the lower-cased, trimmed text of a value, with a trailing '.0' removed."""
    return f"regexp_replace(lower(trim(CAST({expr} AS VARCHAR))), '\\.0+$', '')"


def fabric_source(path: Path, layer: str | None, cols: list[str]) -> str:
    """Returns a SQL FROM-source for a Fabric file (CSV, or gdb/gpkg via spatial / pyogrio)."""
    path = check_gdb(path)
    if path.suffix.lower() not in {".gdb", ".gpkg"}:
        return f"read_csv({lit(path.as_posix())}, header = true, all_varchar = true)"
    try:
        con.execute("INSTALL spatial; LOAD spatial;")
        layer_arg = f", layer = {lit(layer)}" if layer else ""
        return f"ST_Read({lit(path.as_posix())}{layer_arg})"
    except duckdb.Error as e:
        log.warning("DuckDB spatial extension unavailable (%s); reading %s with pyogrio", e, path.name)
        con.register("gdb_src", pyogrio.read_dataframe(path, layer=layer, columns=cols, read_geometry=False))
        return "gdb_src"


def load_fabric(path: Path, layer: str | None, fields: dict, eligible_values: set[str], version: str,
                ids_table: str | None = None) -> pd.DataFrame:
    """Reads one Fabric eligibility file → location_id, lon, lat, eligible, bsl, counted, address, units, elig_raw.

    Default: rows inside the read windows plus every submitted ID. With ids_table: only the IDs in that
    registered table (used to look up IDs that sit outside the read windows), without cleaning records.
    """
    present = set(table_columns(path, layer))
    for opt in ("bsl", "address", "units"):
        if fields.get(opt) and fields[opt] not in present:
            log.warning("%s: field '%s' (%s) not in %s; skipped", version, fields[opt], opt, path.name)
            fields[opt] = None
    cols = sorted({c for c in (fields["id"], fields["lat"], fields["lon"], fields["elig"],
                               fields.get("bsl"), fields.get("address"), fields.get("units")) if c})
    src = fabric_source(path, layer, cols)
    lon, lat = f"TRY_CAST({q(fields['lon'])} AS DOUBLE)", f"TRY_CAST({q(fields['lat'])} AS DOUBLE)"
    window_sql = " OR ".join(f"(lon BETWEEN {a} AND {c} AND lat BETWEEN {b} AND {d})" for a, b, c, d in read_windows)
    elig_list = ", ".join(lit(v.lower()) for v in eligible_values)
    bsl_list = ", ".join(lit(v) for v in BSL_TRUE)
    bsl_expr = f"{norm_sql(q(fields['bsl']))} IN ({bsl_list})" if fields.get("bsl") else "TRUE"
    addr_expr = f"CAST({q(fields['address'])} AS VARCHAR)" if fields.get("address") else "CAST(NULL AS VARCHAR)"
    units_expr = f"TRY_CAST({q(fields['units'])} AS DOUBLE)" if fields.get("units") else "CAST(NULL AS DOUBLE)"
    where_sql = (f"id_num = floor(id_num) AND CAST(id_num AS BIGINT) IN (SELECT location_id FROM {ids_table})"
                 if ids_table else
                 f"({window_sql}) OR (id_num = floor(id_num) AND CAST(id_num AS BIGINT) IN (SELECT location_id FROM submitted))")
    table = f"fab_{version}" + ("_extra" if ids_table else "")

    log.info("DuckDB scan of %s (%s) …", version, path.name)
    con.execute(f"""
        CREATE OR REPLACE TEMP TABLE {table} AS
        WITH src AS (
            SELECT TRY_CAST(replace(trim(CAST({q(fields['id'])} AS VARCHAR)), ',', '') AS DOUBLE) AS id_num,
                   {lon} AS lon, {lat} AS lat,
                   CAST({q(fields['elig'])} AS VARCHAR) AS elig_raw,
                   {norm_sql(q(fields['elig']))} IN ({elig_list}) AS eligible,
                   {bsl_expr} AS bsl,
                   {addr_expr} AS address,
                   {units_expr} AS units
            FROM {src}
        )
        SELECT * FROM src
        WHERE {where_sql}
    """)
    if ids_table:
        df = con.execute(f"""
            SELECT CAST(id_num AS BIGINT) AS location_id, lon, lat, elig_raw, eligible, bsl, address, units,
                   bool_or(eligible) OVER w AND NOT bool_and(eligible) OVER w AS conflict
            FROM {table}
            WINDOW w AS (PARTITION BY id_num)
            QUALIFY row_number() OVER (PARTITION BY id_num ORDER BY eligible) = 1""").df()
        df["location_id"] = df["location_id"].astype("Int64")
        df["eligible"] = df["eligible"].fillna(False).astype(bool)
        df["bsl"] = df["bsl"].fillna(False).astype(bool)
        df["conflict"] = df["conflict"].fillna(False).astype(bool)
        df["counted"] = (df["bsl"] | (not BSL_ONLY)) & ~df["conflict"]
        return df
    n_rows, n_bad, n_unique = con.execute(f"""
        SELECT count(*),
               count(*) FILTER (WHERE id_num IS NULL OR id_num <> floor(id_num)),
               count(DISTINCT id_num) FILTER (WHERE id_num = floor(id_num))
        FROM fab_{version}""").fetchone()
    df = con.execute(f"""
        SELECT CAST(id_num AS BIGINT) AS location_id, lon, lat, elig_raw, eligible, bsl, address, units,
               bool_or(eligible) OVER w AND NOT bool_and(eligible) OVER w AS conflict
        FROM fab_{version}
        WHERE id_num = floor(id_num)
        WINDOW w AS (PARTITION BY id_num)
        QUALIFY row_number() OVER (PARTITION BY id_num ORDER BY eligible) = 1
        ORDER BY location_id""").df()
    log.info("Fabric %s: %s rows in scope → %s unique IDs (%s non-numeric IDs dropped, %s duplicates removed)",
             version, f"{n_rows:,}", f"{n_unique:,}", n_bad, n_rows - n_bad - n_unique)
    fabric_dups[version] = con.execute(f"""
        SELECT CAST(id_num AS BIGINT) AS location_id, count(*) AS times_listed,
               bool_or(eligible) AND NOT bool_and(eligible) AS eligibility_conflict
        FROM fab_{version}
        WHERE id_num = floor(id_num)
        GROUP BY id_num HAVING count(*) > 1""").df()
    CLEANING.append({"dataset": f"Fabric version {version[1]} eligibility data (read window, all projects)", "project_key": None,
                     "rows_read": int(n_rows), "non_numeric_rows_skipped": int(n_bad),
                     "ids_listed_more_than_once": len(fabric_dups[version]),
                     "duplicate_rows_removed": int(n_rows - n_bad - n_unique), "unique_ids_used": int(n_unique),
                     "note": f"{path.name}; duplicated IDs whose copies disagree on eligibility are marked as "
                             "conflicting source records and left out of every count"})
    df["location_id"] = df["location_id"].astype("Int64")
    df["eligible"] = df["eligible"].fillna(False).astype(bool)
    df["bsl"] = df["bsl"].fillna(False).astype(bool)
    df["conflict"] = df["conflict"].fillna(False).astype(bool)
    df["counted"] = (df["bsl"] | (not BSL_ONLY)) & ~df["conflict"]       # conflicting records are never counted
    log.info("Fabric %s: %s IDs with conflicting source records (left out of the counts)", version,
             f"{int(df['conflict'].sum()):,}")
    return df


fabric_dups: dict[str, pd.DataFrame] = {}
fabric = {
    "v4": load_fabric(V4_PATH, V4_LAYER, V4_FIELDS, V4_ELIGIBLE_VALUES, "v4"),
    "v8": load_fabric(V8_PATH, V8_LAYER, V8_FIELDS, V8_ELIGIBLE_VALUES, "v8"),
}
ATTR_COLS = ["lon", "lat", "eligible", "bsl", "counted", "conflict", "address", "units"]
attrs = {v: df.set_index("location_id")[ATTR_COLS] for v, df in fabric.items()}

# %% [markdown]
# ### 3a. Check eligibility values ⚠️
# If a value that should count as eligible shows `counted_as_eligible = 0`, edit `V4_ELIGIBLE_VALUES` / `V8_ELIGIBLE_VALUES` and re-run.

# %%
for v, df in fabric.items():
    fields, values = (V4_FIELDS, V4_ELIGIBLE_VALUES) if v == "v4" else (V8_FIELDS, V8_ELIGIBLE_VALUES)
    vc = (df.groupby(df["elig_raw"].astype("string").fillna("<null>"))["eligible"]
            .agg(records="size", counted_as_eligible="sum").sort_values("records", ascending=False))
    print(f"\n{v.upper()} field '{fields['elig']}' → eligible values {sorted(values)}")
    display(vc.head(25))
    if df["eligible"].sum() == 0:
        warnings.warn(f"{v.upper()}: no record matched the eligible-value set; check the configuration.")

# %% [markdown]
# ## 4. Spatial match (point-in-polygon, WGS84)
# Every Fabric location inside each project area, per version.

# %%
fabric_pts = {v: points_from_lonlat(df) for v, df in fabric.items()}
in_area = {}
for v, pts in fabric_pts.items():
    j = gpd.sjoin(pts[["location_id", "geometry"]].dropna(subset=["geometry"]),
                  areas[["project_key", "geometry"]], predicate="intersects", how="inner")
    in_area[v] = j[["project_key", "location_id"]].drop_duplicates().reset_index(drop=True)
area_sets = {v: {k: set(g["location_id"]) for k, g in in_area[v].groupby("project_key")} for v in fabric}

# An ID inside a project area in one version but missing from the other version's read window is looked
# up in the whole file of the other version, so "not in the Fabric" always means absent statewide.
SOURCES = {"v4": (V4_PATH, V4_LAYER, V4_FIELDS, V4_ELIGIBLE_VALUES), "v8": (V8_PATH, V8_LAYER, V8_FIELDS, V8_ELIGIBLE_VALUES)}
for v, other in (("v4", "v8"), ("v8", "v4")):
    need = sorted(set().union(*area_sets[other].values()) - set(attrs[v].index.astype("int64")))
    if not need:
        continue
    con.register(f"need_{v}", pd.DataFrame({"location_id": need}, dtype="int64"))
    extra = load_fabric(*SOURCES[v], v, ids_table=f"need_{v}")
    log.info("Fabric %s: %s of %s IDs from %s project areas found outside the read window",
             v, len(extra), len(need), other)
    if len(extra):
        fabric[v] = pd.concat([fabric[v], extra], ignore_index=True)
        attrs[v] = fabric[v].set_index("location_id")[ATTR_COLS]

pd.DataFrame({f"{v.upper()} locations in area": in_area[v].groupby("project_key").size() for v in fabric}).reindex(list(PROJECTS))

# %%
ST_NOT_IN = "Not in Fabric"
ST_NOT_BSL = "Not a BSL"
ST_ELIG = "Eligible"
ST_NOT_ELIG = "Not eligible"
ST_ELIG_OUT = "Eligible, outside project area"
ST_NOT_ELIG_OUT = "Not eligible, outside project area"
ST_CONFLICT = "Conflicting source records"


def fabric_status(key: str, lid: int, v: str) -> str:
    """Status of one location ID in one Fabric version, relative to one project area."""
    a = attrs[v]
    if lid not in a.index:
        return ST_NOT_IN
    if a.at[lid, "conflict"]:
        return ST_CONFLICT
    if not a.at[lid, "counted"]:
        return ST_NOT_BSL
    elig = bool(a.at[lid, "eligible"])
    if lid in area_sets[v].get(key, set()):
        return ST_ELIG if elig else ST_NOT_ELIG
    return ST_ELIG_OUT if elig else ST_NOT_ELIG_OUT

# %% [markdown]
# ## Comparisons 1 & 2 – each submission against its own Fabric version
# * **Comparison 1:** initial grantee submission ↔ Fabric version 4 eligibility data
# * **Comparison 2:** 2026 Q2 and Q3 data ↔ Fabric version 8 eligibility data
#
# Per project: submitted, found in the Fabric, found inside the project area, eligible, not eligible, not found.

# %%
for df, v in ((sub_v4, "v4"), (sub_v8, "v8")):
    df["fabric_status"] = [fabric_status(k, int(i), v) for k, i in zip(df["project_key"], df["location_id"])]

rows = []
for df, v, label in ((sub_v4, "v4", "1: Initial grantee submission vs Fabric version 4 eligibility data"), (sub_v8, "v8", "2: 2026 Q2 and Q3 data vs Fabric version 8 eligibility data")):
    for key in PROJECTS:
        st = df.loc[df["project_key"] == key, "fabric_status"]
        n = len(st)
        found = int((st != ST_NOT_IN).sum())
        elig = int(st.isin([ST_ELIG, ST_ELIG_OUT]).sum())
        rows.append({"comparison": label, "project_key": key, "submitted": n, "found": found,
                     "not_found": n - found, "in_project_area": int(st.isin([ST_ELIG, ST_NOT_ELIG]).sum()),
                     "outside_area": int(st.isin([ST_ELIG_OUT, ST_NOT_ELIG_OUT]).sum()),
                     "not_bsl": int((st == ST_NOT_BSL).sum()), "conflicting": int((st == ST_CONFLICT).sum()),
                     "eligible": elig,
                     "not_eligible": found - elig - int(st.isin([ST_NOT_BSL, ST_CONFLICT]).sum()),
                     "pct_found": round(found / n * 100, 1) if n else np.nan,
                     "pct_eligible_of_found": round(elig / found * 100, 1) if found else np.nan})
cmp12 = pd.DataFrame(rows)
cmp12

# %% [markdown]
# ## Comparison 3 – initial grantee submission ↔ 2026 Q2 and Q3 data
# Every location submitted in either version gets one category:
#
# | Category | Meaning |
# |---|---|
# | Kept (same ID) | ID in both submissions |
# | Dropped from submission | ID only in the initial grantee submission |
# | Added to submission | ID only in the 2026 Q2 and Q3 data |
#
# IDs are compared as they are. A dropped ID and a nearby added ID are not treated as the same building: closeness alone cannot show that.
#
# Coordinates come from the submission files, falling back to the matching Fabric version.

# %%
SC_KEPT = "Kept (same ID)"
SC_DROPPED = "Dropped from submission"
SC_ADDED = "Added to submission"
SUB_CATEGORY_ORDER = [SC_KEPT, SC_DROPPED, SC_ADDED]


def with_fallback_coords(df: pd.DataFrame, v: str) -> pd.DataFrame:
    """Fills missing submission coordinates from the Fabric version the IDs belong to."""
    df = df.copy()
    miss = (df["lat"].isna() | df["lon"].isna()) & df["location_id"].isin(attrs[v].index)
    idx = df.loc[miss, "location_id"]
    df.loc[miss, "lat"] = attrs[v].loc[idx, "lat"].to_numpy()
    df.loc[miss, "lon"] = attrs[v].loc[idx, "lon"].to_numpy()
    return df


s4 = with_fallback_coords(sub_v4, "v4")
s8 = with_fallback_coords(sub_v8, "v8")

frames = []
for key in PROJECTS:
    a = s4[s4["project_key"] == key].set_index("location_id")
    b = s8[s8["project_key"] == key].set_index("location_id")
    ids = sorted(set(a.index) | set(b.index))
    d = pd.DataFrame({"project_key": key, "location_id": pd.array(ids, dtype="Int64")})
    d["in_sub_v4"] = d["location_id"].isin(a.index).to_numpy()
    d["in_sub_v8"] = d["location_id"].isin(b.index).to_numpy()
    for side, src in (("v4", a), ("v8", b)):
        r = src.reindex(d["location_id"])
        d[f"sub_lon_{side}"], d[f"sub_lat_{side}"] = r["lon"].to_numpy(), r["lat"].to_numpy()
    d["sub_address_v4"] = a["address"].reindex(d["location_id"]).to_numpy()

    d["sub_change"] = np.select([d["in_sub_v4"] & d["in_sub_v8"], d["in_sub_v4"]], [SC_KEPT, SC_DROPPED],
                                default=SC_ADDED)
    frames.append(d)

subs = pd.concat(frames, ignore_index=True)

# Position shift between submissions for kept IDs
kept = subs["sub_change"].eq(SC_KEPT)
p4 = gpd.GeoSeries(gpd.points_from_xy(subs["sub_lon_v4"], subs["sub_lat_v4"]), crs=WGS84).to_crs(METRIC_CRS)
p8 = gpd.GeoSeries(gpd.points_from_xy(subs["sub_lon_v8"], subs["sub_lat_v8"]), crs=WGS84).to_crs(METRIC_CRS)
subs["sub_shift_m"] = np.where(kept, p4.distance(p8).round(1), np.nan)

cmp3 = (subs.groupby(["project_key", "sub_change"]).size().unstack(fill_value=0)
            .reindex(index=list(PROJECTS), columns=SUB_CATEGORY_ORDER, fill_value=0))
cmp3["V4 submitted"] = subs.groupby("project_key")["in_sub_v4"].sum()
cmp3["V8 submitted"] = subs.groupby("project_key")["in_sub_v8"].sum()
cmp3["net change"] = cmp3["V8 submitted"] - cmp3["V4 submitted"]
cmp3

# %% [markdown]
# ## Comparison 4 – Fabric version 4 ↔ version 8 eligibility data (project areas)
# Every Fabric location counted in a project area in either version gets exactly one cause category:
#
# | Category | Interpretation |
# |---|---|
# | Eligible → not eligible | Deployment or availability data changed between Fabric versions |
# | Not eligible → eligible | Deployment or availability data changed between Fabric versions |
# | ID only in older version (retired) | Location ID is not in the newer Fabric |
# | ID only in newer version (new) | Location ID is not in the older Fabric |
# | Moved into or out of the project area | Location coordinates shifted between Fabric versions |
# | BSL status changed | Only possible when `BSL_ONLY = True` |
# | Unchanged | No difference |
#
# `delta_total` and `delta_eligible` per row sum to exactly each project's net change, so the category table accounts for the whole change.

# %%
CAT_ELIG_LOST = "Eligible -> not eligible"
CAT_ELIG_GAINED = "Not eligible -> eligible"
CAT_RETIRED = "ID only in older version (retired)"
CAT_NEW = "ID only in newer version (new)"
CAT_MOVED = "Moved into or out of the project area"
CAT_BSL = "BSL status changed"
CAT_SAME = "Unchanged"
CAT_CONFLICT = "Conflicting source records"
CATEGORY_ORDER = [CAT_CONFLICT, CAT_ELIG_LOST, CAT_ELIG_GAINED, CAT_RETIRED, CAT_NEW, CAT_MOVED, CAT_BSL, CAT_SAME]


def compare_fabric(key: str) -> pd.DataFrame:
    """One row per Fabric location counted in the project area in V4 and/or V8, with its category."""
    counted, conflicted = {}, {}
    for v in ("v4", "v8"):
        m = pd.Index(sorted(area_sets[v].get(key, set())), dtype="Int64")
        counted[v] = set(m[attrs[v].loc[m, "counted"].to_numpy()])
        conflicted[v] = set(m[attrs[v].loc[m, "conflict"].to_numpy()])
    ids = counted["v4"] | counted["v8"] | conflicted["v4"] | conflicted["v8"]
    df = pd.DataFrame({"project_key": key, "location_id": pd.array(sorted(ids), dtype="Int64")})
    for v in ("v4", "v8"):
        a = attrs[v].reindex(df["location_id"])
        df[f"in_{v}"] = df["location_id"].isin(attrs[v].index).to_numpy()
        df[f"in_area_{v}"] = df["location_id"].isin(area_sets[v].get(key, set())).to_numpy()
        df[f"counted_{v}"] = df["location_id"].isin(counted[v]).to_numpy()
        df[f"conflict_{v}"] = df["location_id"].isin(conflicted[v]).to_numpy()
        df[f"bsl_{v}"] = a["bsl"].fillna(False).astype(bool).to_numpy()
        df[f"elig_{v}"] = a["eligible"].fillna(False).astype(bool).to_numpy()
        df[f"lon_{v}"], df[f"lat_{v}"] = a["lon"].to_numpy(), a["lat"].to_numpy()
        df[f"address_{v}"] = a["address"].to_numpy()
    conds = [df["conflict_v4"] | df["conflict_v8"], df["in_v4"] & ~df["in_v8"], ~df["in_v4"] & df["in_v8"],
             df["in_area_v4"] != df["in_area_v8"], df["counted_v4"] != df["counted_v8"],
             df["elig_v4"] & ~df["elig_v8"], ~df["elig_v4"] & df["elig_v8"]]
    df["category"] = np.select(conds, [CAT_CONFLICT, CAT_RETIRED, CAT_NEW, CAT_MOVED, CAT_BSL, CAT_ELIG_LOST,
                                       CAT_ELIG_GAINED],
                               default=CAT_SAME)
    df["delta_total"] = df["counted_v8"].astype(int) - df["counted_v4"].astype(int)
    df["delta_eligible"] = ((df["counted_v8"] & df["elig_v8"]).astype(int)
                            - (df["counted_v4"] & df["elig_v4"]).astype(int))
    df["in_sub_v4"] = df["location_id"].isin(set(sub_v4.loc[sub_v4["project_key"] == key, "location_id"])).to_numpy()
    df["in_sub_v8"] = df["location_id"].isin(set(sub_v8.loc[sub_v8["project_key"] == key, "location_id"])).to_numpy()
    return df


changes = pd.concat([compare_fabric(k) for k in PROJECTS], ignore_index=True)

# Coordinate shift (m) for IDs present in both Fabric versions
both = changes["in_v4"] & changes["in_v8"]
q4 = gpd.GeoSeries(gpd.points_from_xy(changes["lon_v4"], changes["lat_v4"]), crs=WGS84).to_crs(METRIC_CRS)
q8 = gpd.GeoSeries(gpd.points_from_xy(changes["lon_v8"], changes["lat_v8"]), crs=WGS84).to_crs(METRIC_CRS)
changes["shift_m"] = np.where(both, q4.distance(q8).round(1), np.nan)

# Address change (only when both Fabric files carry an address field)
if V4_FIELDS.get("address") and V8_FIELDS.get("address"):
    norm = lambda s: s.astype("string").str.upper().str.replace(r"[^A-Z0-9]", "", regex=True)
    changes["address_changed"] = (both & (norm(changes["address_v4"]) != norm(changes["address_v8"])).fillna(False)).astype(bool)
else:
    changes["address_changed"] = pd.NA

cmp4 = (changes.groupby(["project_key", "category"])
        .agg(locations=("location_id", "size"), delta_total=("delta_total", "sum"),
             delta_eligible=("delta_eligible", "sum"),
             in_v4_submission=("in_sub_v4", "sum"), in_v8_submission=("in_sub_v8", "sum"))
        .reset_index())
cmp4["category"] = pd.Categorical(cmp4["category"], CATEGORY_ORDER, ordered=True)
cmp4 = cmp4.sort_values(["project_key", "category"]).reset_index(drop=True)
cmp4["category"] = cmp4["category"].astype(str)
cmp4

# %% [markdown]
# ## 5. Synthesis – how the submission changes and the Fabric changes fit together
#
# **a) Explanation per submitted location.** Each location from Comparison 3 is linked to its Fabric status in V4 and V8 (Comparisons 1, 2 and 4):
#
# | Submission change | Explanation (first that applies) |
# |---|---|
# | Kept | eligible in both · lost eligibility in V8 · gained eligibility in V8 · not eligible in either |
# | Dropped | ID retired in V8 Fabric · outside the project area in V8 · not eligible in V8 · still eligible in V8 (applicant decision) |
# | Added | new location in V8 Fabric · moved into the area in V8 · became eligible in V8 · already eligible in V4 (applicant decision) |
#
# **b) Coverage.** Of the eligible Fabric locations in each project area, how many were submitted, per version, and how many submitted locations were not eligible.

# %%
def explain(r: pd.Series) -> str:
    """Links one submitted location's change to the Fabric change behind it."""
    s4_, s8_ = r["fabric_status_v4"], r["fabric_status_v8"]
    e4, e8 = s4_ in (ST_ELIG, ST_ELIG_OUT), s8_ in (ST_ELIG, ST_ELIG_OUT)
    c = r["sub_change"]
    if ST_CONFLICT in (s4_, s8_):
        return "Conflicting source records in the Fabric"
    if c == SC_KEPT:
        if e4 and e8:
            return "Kept: eligible in V4 and V8"
        if e4:
            return "Kept: lost eligibility in V8 (now served)" if s8_ != ST_NOT_IN else "Kept: ID not in V8 Fabric"
        if e8:
            return "Kept: gained eligibility in V8"
        return "Kept: not eligible in V4 or V8"
    if c == SC_DROPPED:
        if s8_ == ST_NOT_IN:
            return "Dropped: ID retired in V8 Fabric"
        if s8_ in (ST_ELIG_OUT, ST_NOT_ELIG_OUT):
            return "Dropped: outside project area in V8"
        if not e8:
            return "Dropped: not eligible in V8 (served)"
        return "Dropped: still eligible in V8 (applicant decision)"
    # Added
    if s4_ == ST_NOT_IN:
        return "Added: new location in V8 Fabric"
    if s4_ in (ST_ELIG_OUT, ST_NOT_ELIG_OUT):
        return "Added: moved into project area in V8"
    if not e4:
        return "Added: became eligible in V8"
    return "Added: already eligible in V4 (applicant decision)"


subs["fabric_status_v4"] = [fabric_status(k, int(i), "v4") for k, i in zip(subs["project_key"], subs["location_id"])]
subs["fabric_status_v8"] = [fabric_status(k, int(i), "v8") for k, i in zip(subs["project_key"], subs["location_id"])]
subs["explanation"] = subs.apply(explain, axis=1)
subs["v8_submitted_not_eligible"] = subs["in_sub_v8"] & ~subs["fabric_status_v8"].isin([ST_ELIG, ST_ELIG_OUT])

synthesis = (subs.groupby(["project_key", "sub_change", "explanation"]).size().rename("locations").reset_index())
synthesis["sub_change"] = pd.Categorical(synthesis["sub_change"], SUB_CATEGORY_ORDER, ordered=True)
synthesis = synthesis.sort_values(["project_key", "sub_change", "locations"], ascending=[True, True, False]).reset_index(drop=True)
synthesis["sub_change"] = synthesis["sub_change"].astype(str)
display(synthesis)

# Coverage of eligible Fabric locations by the submission of the same version
cov_rows = []
for key in PROJECTS:
    ch = changes[changes["project_key"] == key]
    r = {"project_key": key}
    for v in ("v4", "v8"):
        elig_ids = ch.loc[ch[f"counted_{v}"] & ch[f"elig_{v}"], "location_id"]
        sub_ids = set((sub_v4 if v == "v4" else sub_v8).query("project_key == @key")["location_id"])
        r[f"{v}_elig_in_area"] = len(elig_ids)
        r[f"{v}_elig_submitted"] = int(elig_ids.isin(sub_ids).sum())
        r[f"{v}_elig_not_submitted"] = r[f"{v}_elig_in_area"] - r[f"{v}_elig_submitted"]
        r[f"{v}_pct_elig_submitted"] = round(r[f"{v}_elig_submitted"] / r[f"{v}_elig_in_area"] * 100, 1) if r[f"{v}_elig_in_area"] else np.nan
    cov_rows.append(r)
coverage = pd.DataFrame(cov_rows)
coverage

# %% [markdown]
# ## 5b. Plain-language categories (shared by every map, page, spreadsheet and report)
#
# The four data sources are named the same way in every output:
#
# | Source | Name used |
# |---|---|
# | Fabric V4 eligibility file (`FFAv4Dec23`) | **Fabric version 4 eligibility data** |
# | Fabric V8 eligibility file (`status_code`) | **Fabric version 8 eligibility data** |
# | All_448_PIDs application file | **initial grantee submission** |
# | 2026 quarterly reports | **2026 Q2 and Q3 data** (per project: **2026 Q3 data**, or **2026 Q2 data** where Q3 is missing) |
#
# Every location gets one **category** in one of three **tiers**, and each category has one colour everywhere:
#
# | Tier | Meaning |
# |---|---|
# | **Needs follow-up** | Something in the 2026 Q2 and Q3 data (or missing from it) that is not explained by the observed Fabric change, or a source-data conflict |
# | **Consistent with Fabric change** | Matches an observed change between the Fabric version 4 and version 8 eligibility data (eligibility classification changed, ID present in only one version, point moved). The data shows association, not cause |
# | **No change** | Nothing changed |
#
# * **Eligibility changes (Fabric version 4 → version 8 eligibility data):** every Fabric location in the project area.
# * **Grantee submission changes (initial grantee submission → 2026 Q2 and Q3 data):** every listed location. Follow-up checks come first: a location in the 2026 Q2 and Q3 data that is outside the area, missing from the Fabric version 8 eligibility data or not eligible is flagged even if it was also in the initial grantee submission.
# * **Follow-up:** the follow-up categories only; each issue (missing ID, outside the area, not eligible, not a serviceable location, removed but eligible, added without a Fabric change, conflicting records) has its own category.

# %%
# Names used for the four data sources in every output
T_V4 = "Fabric version 4 eligibility data"
T_V8 = "Fabric version 8 eligibility data"
T_SUB4 = "initial grantee submission"
T_SUB8 = "2026 Q2 and Q3 data"          # all projects; one project: qr_term(key) → "2026 Q3 data" / "2026 Q2 data"
T_SUB8_SHORT = "2026 Q2/Q3 data"        # compact form for category labels


def qr_term(key: str) -> str:
    """Name of one project's 2026 data, from the quarter of the report file used."""
    name = sub_v8_files.get(key, "").upper()
    for q in ("Q3", "Q2"):
        if name.startswith(q):
            return f"2026 {q} data"
    return T_SUB8


def cap(text: str) -> str:
    """First letter upper-case, the rest unchanged."""
    return text[:1].upper() + text[1:]


# One colour per meaning, shared by every map, page, spreadsheet and report
RGB = {"darkred": (128, 0, 38), "darkorange": (217, 95, 2), "pink": (197, 27, 125),
       "gold": (230, 171, 2), "purple": (118, 42, 131), "lilac": (175, 141, 195), "brown": (140, 81, 10),
       "blue": (33, 102, 172), "teal": (1, 133, 113), "slate": (99, 99, 99),
       "skyblue": (86, 180, 233), "grey": (150, 150, 150), "lightgrey": (205, 205, 205),
       "charcoal": (45, 45, 45), "olive": (128, 128, 0)}
TIER_FU, TIER_EX, TIER_NC = "Needs follow-up", "Consistent with Fabric change", "No change"
TIER_ORDER = [TIER_FU, TIER_EX, TIER_NC]

# code: (plain label, what it means, RGB colour, tier). Labels state what the data shows, not why it happened.
FABRIC_CLASSES = {
    "A10": ("Conflicting source records",
            "The copies of this location ID in one Fabric file disagree on eligibility; it is left out of every count "
            "until the source is corrected.", RGB["olive"], TIER_FU),
    "A1": ("Eligible to not eligible",
           f"Eligibility classification changed from eligible in the {T_V4} to not eligible in the {T_V8}.",
           RGB["gold"], TIER_EX),
    "A2": ("Not eligible to eligible",
           f"Eligibility classification changed from not eligible in the {T_V4} to eligible in the {T_V8}.",
           RGB["blue"], TIER_EX),
    "A4": ("ID present only in Fabric version 4",
           f"This location ID is in the {T_V4} but not in the {T_V8}. That does not show the building is gone.",
           RGB["brown"], TIER_EX),
    "A5": ("ID present only in Fabric version 8",
           f"This location ID is in the {T_V8} but not in the {T_V4}. That does not show the building is new.",
           RGB["teal"], TIER_EX),
    "A6": ("Map point moved across boundary",
           f"Same ID in both; its point moved into or out of the project area between the {T_V4} and the {T_V8} "
           "(a map edit, not a building that moved).", RGB["lilac"], TIER_EX),
    "A7": ("Serviceable-location status changed",
           f"Counted as a broadband serviceable location in only one of the {T_V4} and the {T_V8}.", RGB["slate"], TIER_EX),
    "A8": ("No change – eligible", f"Same ID, inside the area, eligible in both the {T_V4} and the {T_V8}.",
           RGB["skyblue"], TIER_NC),
    "A9": ("No change – not eligible",
           f"Same ID, inside the area, not eligible in either the {T_V4} or the {T_V8}.", RGB["lightgrey"], TIER_NC),
}
SUB_CLASSES = {
    "B01": ("In 2026 Q2/Q3 data, missing from Fabric version 8",
            f"In the {T_SUB8}, but the location ID is not in the {T_V8}.", RGB["darkred"], TIER_FU),
    "B13": ("In 2026 Q2/Q3 data, outside the award area",
            f"In the {T_SUB8}, but its point in the {T_V8} is outside the project area.", RGB["charcoal"], TIER_FU),
    "B02": ("In 2026 Q2/Q3 data, not eligible in Fabric version 8",
            f"In the {T_SUB8}, but the {T_V8} classifies it as not eligible.", RGB["darkorange"], TIER_FU),
    "B14": ("In 2026 Q2/Q3 data, not a broadband serviceable location",
            f"In the {T_SUB8}, but the {T_V8} does not count it as a broadband serviceable location.",
            RGB["slate"], TIER_FU),
    "B03": ("Not carried into 2026 Q2/Q3 data, still eligible",
            f"In the {T_SUB4} but not the {T_SUB8}, although still eligible inside the area in the {T_V8}.",
            RGB["purple"], TIER_FU),
    "B04": ("New in 2026 Q2/Q3 data, already eligible in Fabric version 4",
            f"In the {T_SUB8} but not the {T_SUB4}, although already eligible in the {T_V4}.", RGB["pink"], TIER_FU),
    "B12": ("Conflicting source records in the Fabric",
            "The copies of this location ID in a Fabric file disagree on eligibility, so its status cannot be read.",
            RGB["olive"], TIER_FU),
    "B05": ("Not carried into 2026 Q2/Q3 data, now not eligible",
            f"In the {T_SUB4} only; the {T_V8} classifies it as not eligible.", RGB["gold"], TIER_EX),
    "B06": ("Not carried into 2026 Q2/Q3 data, ID present only in Fabric version 4",
            f"In the {T_SUB4} only; the ID is not in the {T_V8}.", RGB["brown"], TIER_EX),
    "B07": ("Not carried over or new, point moved across boundary",
            f"The point moved out of or into the project area between the {T_V4} and the {T_V8}.", RGB["lilac"], TIER_EX),
    "B09": ("New in 2026 Q2/Q3 data, not eligible to eligible",
            f"In the {T_SUB8} only; not eligible in the {T_V4}, eligible in the {T_V8}.", RGB["blue"], TIER_EX),
    "B10": ("New in 2026 Q2/Q3 data, ID present only in Fabric version 8",
            f"In the {T_SUB8} only; the location ID is in the {T_V8} but not the {T_V4}.", RGB["teal"], TIER_EX),
    "B11": ("In both submissions, eligible in Fabric version 8",
            f"Same ID in the {T_SUB4} and the {T_SUB8}; inside the area and eligible in the {T_V8}.", RGB["grey"], TIER_NC),
}
FOLLOW_UP_CLASSES = ["B01", "B13", "B02", "B14", "B03", "B04", "B12"]
FOLLOW_CLASSES = {**{k: SUB_CLASSES[k] for k in FOLLOW_UP_CLASSES},
                  "C0": ("Other listed location", "Listed location that needs no follow-up.", RGB["lightgrey"], TIER_NC)}
NO_CHANGE_CLASSES = {c for c, v in {**FABRIC_CLASSES, **SUB_CLASSES}.items() if v[3] == TIER_NC} | {"C0"}

# Plain wording for the synthesis explanations (Section 5)
WHY_PLAIN = {
    "Dropped: ID retired in V8 Fabric": f"ID present only in the {T_V4}",
    "Dropped: outside project area in V8": f"Outside the project area in the {T_V8}",
    "Dropped: not eligible in V8 (served)": f"Not eligible in the {T_V8}",
    "Dropped: still eligible in V8 (applicant decision)": f"Still eligible in the {T_V8} – grantee decision",
    "Added: new location in V8 Fabric": f"ID present only in the {T_V8}",
    "Added: moved into project area in V8": f"Moved into the project area in the {T_V8}",
    "Added: became eligible in V8": f"Not eligible to eligible in the {T_V8}",
    "Added: already eligible in V4 (applicant decision)": f"Already eligible in the {T_V4} – grantee decision",
    "Kept: eligible in V4 and V8": "No change",
    "Kept: lost eligibility in V8 (now served)": f"Eligible to not eligible in the {T_V8}",
    "Kept: ID not in V8 Fabric": f"ID missing from the {T_V8}",
    "Kept: gained eligibility in V8": f"Not eligible to eligible in the {T_V8}",
    "Conflicting source records in the Fabric": "Conflicting source records in the Fabric",
    "Kept: not eligible in V4 or V8": f"Not eligible in the {T_V4} or the {T_V8}",
}
STATUS_PLAIN = {ST_NOT_IN: "Not in this Fabric version", ST_NOT_BSL: "Not a serviceable location",
                ST_ELIG: "Eligible", ST_NOT_ELIG: "Not eligible", ST_CONFLICT: "Conflicting source records",
                ST_ELIG_OUT: "Eligible, but outside project area", ST_NOT_ELIG_OUT: "Not eligible, outside project area"}
LABEL = {p: v["label"] for p, v in PROJECTS.items()}


def fmt_id(v) -> str:
    """Location ID as text, or empty."""
    return "" if pd.isna(v) else str(int(v))


# ---- Eligibility changes: Fabric locations in the project areas (one point per building)
cat = changes["category"]
cls = cat.map({CAT_CONFLICT: "A10", CAT_ELIG_LOST: "A1", CAT_ELIG_GAINED: "A2", CAT_MOVED: "A6", CAT_BSL: "A7"})
cls = cls.mask(cat.eq(CAT_RETIRED), "A4").mask(cat.eq(CAT_NEW), "A5")
same = cat.eq(CAT_SAME)
cls = cls.mask(same & changes["elig_v8"], "A8").mask(same & ~changes["elig_v8"], "A9")
changes["map_class"] = cls
changes["map_show"] = True
changes["map_label"] = cls.map(lambda c: FABRIC_CLASSES[c][0])
changes["what_it_means"] = cls.map(lambda c: FABRIC_CLASSES[c][1])
changes["tier"] = cls.map(lambda c: FABRIC_CLASSES[c][3])
detail = pd.Series("", index=changes.index)
moved_words = pd.Series(np.where(changes["in_area_v8"], "Moved into the area by ", "Moved out of the area by "),
                        index=changes.index)
detail = detail.mask(cls.eq("A6"), moved_words + changes["shift_m"].round(0).astype("Int64").astype(str) + " m")
changes["detail"] = detail
changes["project"] = changes["project_key"].map(LABEL)

# ---- Grantee list changes: follow-up checks first, then the map explanation
s4, s8 = subs["fabric_status_v4"], subs["fabric_status_v8"]
chg = subs["sub_change"]
in_list8 = subs["in_sub_v8"]
changed_list = chg.isin([SC_ADDED, SC_DROPPED])
non_bsl8 = (pd.Series(attrs["v8"]["bsl"].reindex(subs["location_id"].astype("int64")).eq(False).to_numpy(),
                      index=subs.index) & ~s8.isin([ST_NOT_IN, ST_CONFLICT])) | s8.eq(ST_NOT_BSL)
subs["map_class"] = np.select(
    [in_list8 & s8.eq(ST_NOT_IN),
     in_list8 & s8.eq(ST_CONFLICT),
     in_list8 & s8.isin([ST_ELIG_OUT, ST_NOT_ELIG_OUT]),
     in_list8 & non_bsl8,
     in_list8 & s8.eq(ST_NOT_ELIG),
     changed_list & (s4.eq(ST_CONFLICT) | s8.eq(ST_CONFLICT)),
     chg.eq(SC_ADDED) & s4.eq(ST_ELIG),
     chg.eq(SC_ADDED) & s4.isin([ST_ELIG_OUT, ST_NOT_ELIG_OUT]),
     chg.eq(SC_ADDED) & s4.eq(ST_NOT_IN),
     chg.eq(SC_ADDED),
     chg.eq(SC_DROPPED) & s8.eq(ST_NOT_IN),
     chg.eq(SC_DROPPED) & s8.isin([ST_ELIG_OUT, ST_NOT_ELIG_OUT]),
     chg.eq(SC_DROPPED) & s8.eq(ST_ELIG),
     chg.eq(SC_DROPPED)],
    ["B01", "B12", "B13", "B14", "B02", "B12", "B04", "B07", "B10", "B09", "B06", "B07", "B03", "B05"], default="B11")
subs["map_show"] = True
subs["map_label"] = subs["map_class"].map(lambda c: SUB_CLASSES[c][0])
subs["what_it_means"] = subs["map_class"].map(lambda c: SUB_CLASSES[c][1])
subs["tier"] = subs["map_class"].map(lambda c: SUB_CLASSES[c][3])
subs["needs_follow_up"] = subs["map_class"].isin(FOLLOW_UP_CLASSES)
subs["why"] = subs["explanation"].map(WHY_PLAIN).fillna(subs["explanation"])
subs["status_v4_plain"] = s4.map(STATUS_PLAIN)
subs["status_v8_plain"] = s8.map(STATUS_PLAIN)
subs["detail"] = ""
# Listed locations are drawn at the Fabric point of their ID (version 8, else version 4); the grantee's own
# coordinates are used only when the ID is in neither Fabric version
_lid = subs["location_id"].astype("int64")
_a8, _a4 = attrs["v8"].reindex(_lid), attrs["v4"].reindex(_lid)
subs["fab_lon"] = (_a8["lon"].fillna(_a4["lon"]).to_numpy())
subs["fab_lat"] = (_a8["lat"].fillna(_a4["lat"]).to_numpy())
subs["point_source"] = np.where(_a8["lon"].notna().to_numpy(), "Fabric version 8",
                       np.where(_a4["lon"].notna().to_numpy(), "Fabric version 4", "Grantee coordinates (ID not in the Fabric)"))
subs["fab_lon"] = subs["fab_lon"].fillna(subs["sub_lon_v8"]).fillna(subs["sub_lon_v4"])
subs["fab_lat"] = subs["fab_lat"].fillna(subs["sub_lat_v8"]).fillna(subs["sub_lat_v4"])
subs["project"] = subs["project_key"].map(LABEL)

legend = pd.DataFrame(
    [{"page": pg, "map_class": k, "map_label": v[0], "what_it_means": v[1], "tier": v[3],
      "red": v[2][0], "green": v[2][1], "blue": v[2][2]}
     for pg, classes in (("A", FABRIC_CLASSES), ("B", SUB_CLASSES), ("C", FOLLOW_CLASSES))
     for k, v in classes.items()])

plain_a = changes[changes["map_show"]].pivot_table(index="project_key", columns="map_class", values="location_id",
                                                   aggfunc="size", fill_value=0)
plain_a = plain_a.reindex(index=list(PROJECTS), columns=list(FABRIC_CLASSES), fill_value=0)
plain_b = subs[subs["map_show"]].pivot_table(index="project_key", columns="map_class", values="location_id",
                                             aggfunc="size", fill_value=0)
plain_b = plain_b.reindex(index=list(PROJECTS), columns=list(SUB_CLASSES), fill_value=0)
display(plain_a.rename(columns=lambda c: FABRIC_CLASSES[c][0]))
display(plain_b.rename(columns=lambda c: SUB_CLASSES[c][0]))

# ---- Data cleaning: what was removed before counting, per list and project
fab_rows = []
for v in ("v4", "v8"):
    d = fabric_dups.get(v, pd.DataFrame(columns=["location_id", "times_listed", "eligibility_conflict"]))
    m = in_area[v].astype({"location_id": "int64"}).merge(d, on="location_id")
    if len(m):
        DUPLICATES.append(m.assign(dataset=f"Fabric version {v[1]} eligibility data (project area)")[
            ["dataset", "project_key", "location_id", "times_listed", "eligibility_conflict"]])
    for key in PROJECTS:
        mk = m[m["project_key"] == key]
        n_ids = len(area_sets[v].get(key, set()))
        fab_rows.append({"dataset": f"Fabric version {v[1]} eligibility data (project area)", "project_key": key,
                         "rows_read": n_ids + int((mk["times_listed"] - 1).sum()), "non_numeric_rows_skipped": 0,
                         "ids_listed_more_than_once": len(mk),
                         "duplicate_rows_removed": int((mk["times_listed"] - 1).sum()), "unique_ids_used": n_ids,
                         "note": f"{int(mk['eligibility_conflict'].sum())} duplicated IDs had conflicting eligibility; "
                                 "they are marked as conflicting source records and left out of the counts"})
cleaning = (pd.concat([pd.DataFrame(CLEANING), pd.DataFrame(fab_rows)], ignore_index=True)
            .drop_duplicates(["dataset", "project_key"], keep="last").reset_index(drop=True))
cleaning["project"] = cleaning["project_key"].map(LABEL).fillna("All projects (statewide read window)")
dup_detail = (pd.concat(DUPLICATES, ignore_index=True) if DUPLICATES
              else pd.DataFrame(columns=["dataset", "project_key", "location_id", "times_listed", "eligibility_conflict"]))
dup_detail = (dup_detail.reindex(columns=["dataset", "project_key", "location_id", "times_listed", "eligibility_conflict"])
              .drop_duplicates(["dataset", "project_key", "location_id"]).reset_index(drop=True))
dup_detail["project"] = dup_detail["project_key"].map(LABEL)
display(cleaning[["project", "dataset", "rows_read", "non_numeric_rows_skipped", "ids_listed_more_than_once",
                  "duplicate_rows_removed", "unique_ids_used", "note"]])

# %% [markdown]
# ## 6. Project summary (Map Series index fields) and plain-English page stories

# %%
CAT_FIELD = {CAT_ELIG_LOST: "f_elig_lost", CAT_ELIG_GAINED: "f_elig_gain", CAT_RETIRED: "f_retired",
             CAT_NEW: "f_new", CAT_MOVED: "f_moved", CAT_BSL: "f_bsl_chg"}
SUB_FIELD = {SC_KEPT: "s_kept", SC_DROPPED: "s_dropped", SC_ADDED: "s_added"}


def project_row(key: str) -> dict:
    """Summary metrics for one project (all four comparisons + coverage)."""
    ch = changes[changes["project_key"] == key]
    sb = subs[subs["project_key"] == key]
    r = {"project_key": key, "sub_v8_file": sub_v8_files.get(key, "")}
    # Comparison 4: Fabric totals in project area
    for v in ("v4", "v8"):
        r[f"{v}_total"] = int(ch[f"counted_{v}"].sum())
        r[f"{v}_elig"] = int((ch[f"counted_{v}"] & ch[f"elig_{v}"]).sum())
    r["total_chg"], r["elig_chg"] = r["v8_total"] - r["v4_total"], r["v8_elig"] - r["v4_elig"]
    r["total_pct"], r["elig_pct"] = pct_change(r["v4_total"], r["v8_total"]), pct_change(r["v4_elig"], r["v8_elig"])
    for cat, fld in CAT_FIELD.items():
        r[fld] = int((ch["category"] == cat).sum())
    # Comparisons 1 & 2
    for v, label in (("v4", "1: Initial grantee submission vs Fabric version 4 eligibility data"), ("v8", "2: 2026 Q2 and Q3 data vs Fabric version 8 eligibility data")):
        c = cmp12[(cmp12["comparison"] == label) & (cmp12["project_key"] == key)].iloc[0]
        r[f"sub{v[1]}_n"], r[f"sub{v[1]}_found"] = int(c["submitted"]), int(c["found"])
        r[f"sub{v[1]}_in_area"], r[f"sub{v[1]}_elig"] = int(c["in_project_area"]), int(c["eligible"])
    r["sub_chg"] = r["sub8_n"] - r["sub4_n"]
    # Comparison 3
    for cat, fld in SUB_FIELD.items():
        r[fld] = int((sb["sub_change"] == cat).sum())
    # Synthesis
    for fld, pattern in (("x_retired", "retired in V8 Fabric"), ("x_served", "not eligible in V8|lost eligibility"),
                         ("x_applicant", "applicant decision"), ("x_new_fabric", "new location in V8|new ID only in V8")):
        r[fld] = int(sb["explanation"].str.contains(pattern, regex=True).sum())
    r["sub8_not_elig"] = int(sb["v8_submitted_not_eligible"].sum())
    return r



summary = (areas.drop(columns="geometry")
           .merge(pd.DataFrame([project_row(k) for k in PROJECTS]), on="project_key")
           .merge(coverage, on="project_key"))

LIST_SHORT_RATIO, LIST_LONG_RATIO = 0.5, 2.0   # 2026 data vs initial submission sizes that trigger a caveat


def plural(n: int, word: str) -> str:
    """'1 location' / '3 locations'."""
    return f"{n:,} {word}" + ("" if n == 1 else "s")


def cleaning_value(dataset: str, key: str, col: str) -> int:
    """One cleaning count for a dataset and project (0 when not recorded)."""
    hit = cleaning[(cleaning["dataset"] == dataset) & (cleaning["project_key"] == key)]
    return int(hit[col].iloc[0]) if len(hit) else 0


for pre, ds in (("sub4", "Initial grantee submission"), ("sub8", "2026 Q2 and Q3 data"),
                ("f4", "Fabric version 4 eligibility data (project area)"), ("f8", "Fabric version 8 eligibility data (project area)")):
    summary[f"{pre}_dup_ids"] = summary["project_key"].map(lambda k: cleaning_value(ds, k, "ids_listed_more_than_once"))
    summary[f"{pre}_dup_rows"] = summary["project_key"].map(lambda k: cleaning_value(ds, k, "duplicate_rows_removed"))
    summary[f"{pre}_bad_rows"] = summary["project_key"].map(lambda k: cleaning_value(ds, k, "non_numeric_rows_skipped"))
for c in FOLLOW_UP_CLASSES:
    summary[f"fu_{c.lower()}"] = summary["project_key"].map(plain_b[c])
summary["fu_total"] = summary[[f"fu_{c.lower()}" for c in FOLLOW_UP_CLASSES]].sum(axis=1)


def change_words(n: int, noun: str) -> str:
    """'12 fewer eligible' / '3 more eligible' / 'no change in eligible'."""
    if n < 0:
        return f"{-n:,} fewer {noun}"
    if n > 0:
        return f"{n:,} more {noun}"
    return f"no change in {noun}"


def lower_first(text: str) -> str:
    """First letter lower-case unless the text starts with a name such as Fabric or FCC."""
    return text if text.split(" ")[0] in ("Fabric", "FCC") else text[:1].lower() + text[1:]


def bottom_line(r: pd.Series) -> str:
    """Three-sentence summary for the project's first page and the map series."""
    qr = qr_term(r["project_key"])
    d = int(r["elig_chg"])
    vs = f"{abs(d):,} {'fewer' if d < 0 else 'more'} than in the {T_V4}" if d else f"the same as in the {T_V4}"
    fu = (f"{plural(r['fu_total'], 'listed location')} need{'s' if r['fu_total'] == 1 else ''} follow-up."
          if r["fu_total"] else "No listed location needs follow-up.")
    return (f"In the {T_V8}, {r['v8_elig']:,} of the {r['v8_total']:,} locations in the project area are eligible ({vs}). "
            f"The {qr} lists {plural(r['sub8_n'], 'location')}, "
            + (f"{r['sub8_v8_in_elig']:,} of them eligible inside the project area in the {T_V8}. "
               if "sub8_v8_in_elig" in r.index and pd.notna(r["sub8_v8_in_elig"])
               else f"{r['sub8_elig']:,} of them eligible in the {T_V8}. ") + fu)


def story_a(r: pd.Series) -> str:
    """Plain-English summary for the eligibility-changes map."""
    n = plain_a.loc[r["project_key"]]
    s = (f"In the {T_V4} this project area had {r['v4_total']:,} locations, {r['v4_elig']:,} of them eligible. "
         f"In the {T_V8} it has {r['v8_total']:,}, {r['v8_elig']:,} eligible "
         f"({change_words(int(r['elig_chg']), 'eligible')}).")
    parts = []
    if n["A1"]:
        parts.append(f"{n['A1']:,} changed from eligible to not eligible")
    if n["A2"]:
        parts.append(f"{n['A2']:,} changed from not eligible to eligible")
    if n["A4"]:
        parts.append(f"{plural(n['A4'], 'ID')} present only in the {T_V4}")
    if n["A5"]:
        parts.append(f"{plural(n['A5'], 'ID')} present only in the {T_V8}")
    if n["A6"]:
        parts.append(f"{n['A6']:,} map points moved across the boundary")
    if n["A7"]:
        parts.append(f"{n['A7']:,} changed serviceable-location status")
    if parts:
        s += " Observed Fabric changes: " + "; ".join(parts) + "."
    if n["A10"]:
        s += f" {plural(n['A10'], 'ID')} with conflicting source records {'is' if n['A10'] == 1 else 'are'} left out of the counts."
    if r["f4_dup_ids"] or r["f8_dup_ids"]:
        s += (f" Data note: {r['f4_dup_ids']:,} IDs in the {T_V4} and {r['f8_dup_ids']:,} in the {T_V8} appeared more "
              "than once here; each is counted once, and IDs whose copies disagree on eligibility are left out "
              "as conflicting source records.")
    return s


def story_b(r: pd.Series) -> str:
    """Plain-English summary for the grantee-submission-changes map."""
    n = plain_b.loc[r["project_key"]]
    qr = qr_term(r["project_key"])
    s = (f"The {T_SUB4} listed {r['sub4_n']:,} locations; the {qr} lists {r['sub8_n']:,}: "
         f"{r['s_kept']:,} in both, {r['s_dropped']:,} not carried over, "
         f"{r['s_added']:,} new.")
    fu = [f"{n[c]:,} {lower_first(SUB_CLASSES[c][0]).replace(T_SUB8_SHORT, qr)}" for c in FOLLOW_UP_CLASSES if n[c]]
    s += (" Needs follow-up: " + "; ".join(fu) + "." if fu else " Nothing in the submissions needs follow-up.")
    if r["sub4_n"] and r["sub8_n"] < LIST_SHORT_RATIO * r["sub4_n"]:
        s += (f" Caution: the {qr} plans far fewer locations than the {T_SUB4}; confirm the reduced scope with "
              "the grantee.")
    elif r["sub4_n"] and r["sub8_n"] > LIST_LONG_RATIO * r["sub4_n"]:
        s += (f" Caution: the {qr} plans far more locations than the {T_SUB4}; check that the added locations "
              "are inside the awarded area and eligible.")
    notes = []
    if r["sub4_dup_ids"]:
        notes.append(f"the {T_SUB4} repeated {plural(r['sub4_dup_ids'], 'location')} "
                     f"({plural(r['sub4_dup_rows'], 'extra row')} removed)")
    if r["sub8_dup_ids"]:
        notes.append(f"the {qr} repeated {plural(r['sub8_dup_ids'], 'location')} "
                     f"({plural(r['sub8_dup_rows'], 'extra row')} removed)")
    if r["sub4_bad_rows"]:
        notes.append(f"{plural(r['sub4_bad_rows'], 'row')} of the {T_SUB4} without a valid ID skipped")
    if notes:
        s += " Data note: " + "; ".join(notes) + "."
    return s


def story_c(r: pd.Series) -> str:
    """Plain-English summary for the follow-up map."""
    n = plain_b.loc[r["project_key"]]
    qr = qr_term(r["project_key"])
    if not r["fu_total"]:
        return "No listed location needs follow-up: every list change is consistent with the observed Fabric change."
    head = f"{plural(r['fu_total'], 'listed location')} need{'s' if r['fu_total'] == 1 else ''} follow-up:"
    items = [f"{n[c]:,} {lower_first(SUB_CLASSES[c][0]).replace(T_SUB8_SHORT, qr)}" for c in FOLLOW_UP_CLASSES if n[c]]
    return head + " " + "; ".join(items) + ". Location IDs are in the project workbook, tab 'Follow-up list'."


summary["page_title"] = summary["label"] + " (" + summary["pr_number"] + ")"
summary["bottom_line"] = summary.apply(bottom_line, axis=1)
summary["story_a"] = summary.apply(story_a, axis=1)
summary["story_b"] = summary.apply(story_b, axis=1)
summary["story_c"] = summary.apply(story_c, axis=1)
summary["narrative"] = summary["bottom_line"] + " " + summary["story_a"] + " " + summary["story_b"]
for _, r in summary.iterrows():
    print(f"■ {r['page_title']}\n  Bottom line: {r['bottom_line']}\n  A: {r['story_a']}\n  B: {r['story_b']}\n  C: {r['story_c']}\n")
summary.drop(columns=["narrative", "bottom_line", "story_a", "story_b", "story_c", "matched_names"])

# %% [markdown]
# ## 6a. Key numbers, from-to tables and change patterns
#
# Every location is placed with the **Fabric's own coordinates** for the version it is checked against (the grantee's
# latitude/longitude are only used for the coordinate check). Locations are matched by **location ID only**.
#
# | Block | Population | Checked against |
# |---|---|---|
# | A – Fabric in the project area | Fabric locations inside the polygon | Fabric version 4 and version 8 eligibility data |
# | B – Grantee lists | Initial grantee submission and 2026 Q2 and Q3 data | Each list against **both** Fabric versions (own version = main comparison; other version = cross-check) |
# | C – From-to tables | Fabric locations in the area; each grantee list | Status in version 4 → status in version 8 |
# | D – Coverage | Eligible locations inside the polygon | Share listed by the grantee (version 4 with the initial grantee submission, version 8 with the 2026 data) |
#
# Status of a location in one Fabric version, built from separate facts (present in the Fabric, inside the award area,
# eligible, conflicting source records): **Eligible, inside** · **Not eligible, inside** · **Eligible, outside** ·
# **Not eligible, outside** · **Not in the Fabric** · **Conflicting source records**.
# Every block reconciles: listed = matched + not found; matched = inside + outside; each of those = eligible + not eligible.
# The 2026 Q2 and Q3 data is the grantee's **planned** location list.
#
# **Change patterns** come from analyst-defined review thresholds (all configurable below and listed, with the count and
# denominator behind each, in the *Review rules* tab). A project can show several; when none is triggered it reads
# *No configured material-change threshold triggered* — a statement about the thresholds, not a finding that nothing
# changed.

# %%
# ---- thresholds for the change patterns
MIN_LOCATIONS = 10              # a pattern needs at least this many locations behind it
SERVICE_DROP_SHARE = 0.05
LIST_CHANGE_SHARE = 0.10
DROPPED_ELIGIBLE_SHARE = 0.05
OUTSIDE_SHARE = 0.05
ID_CHURN_SHARE = 0.05
COORD_GAP_M, COORD_GAP_SHARE = 100, 0.05
EDGE_M = 50                     # "near the boundary": Fabric point within this distance of the polygon edge
OUT_BANDS_M = (100, 1000)       # outside distance bands: < 100 m, 100 m – 1 km, > 1 km

S_EI, S_NI, S_EO, S_NO, S_NOT, S_CONF, S_NB = (
    "Eligible, inside", "Not eligible, inside", "Eligible, outside", "Not eligible, outside", "Not in the Fabric",
    "Conflicting source records", "Not counted (not a broadband serviceable location)")
STATES = [S_EI, S_NI, S_EO, S_NO, S_NOT, S_CONF]
OUT = [S_EO, S_NO]
STATE_OF = {ST_ELIG: S_EI, ST_NOT_ELIG: S_NI, ST_ELIG_OUT: S_EO, ST_NOT_ELIG_OUT: S_NO, ST_NOT_IN: S_NOT,
            ST_CONFLICT: S_CONF, ST_NOT_BSL: S_NB}


def fsum(m: pd.DataFrame, a, b) -> int:
    """Sum of from-to cells; a and b are one status or a list of statuses (missing ones count 0)."""
    rows = [x for x in ([a] if isinstance(a, str) else a) if x in m.index]
    cols = [x for x in ([b] if isinstance(b, str) else b) if x in m.columns]
    return int(m.loc[rows, cols].to_numpy().sum()) if rows and cols else 0
HAS_UNITS = {v: attrs[v]["units"].notna().any() for v in ("v4", "v8")}
POLY_M = areas.set_index("project_key").to_crs(METRIC_CRS).geometry


def to_metric(a: pd.DataFrame) -> gpd.GeoSeries:
    """Fabric points (lon/lat columns) in metres; rows without coordinates stay empty."""
    ok = a["lon"].notna() & a["lat"].notna()
    g = gpd.GeoSeries([None] * len(a), index=a.index, crs=WGS84)
    g[ok] = gpd.points_from_xy(a.loc[ok, "lon"], a.loc[ok, "lat"])
    return g.to_crs(METRIC_CRS)


def fabric_block(key: str, v: str) -> dict:
    """Block A: Fabric locations inside the polygon in one version."""
    ids = sorted(i for i in area_sets[v].get(key, set()) if attrs[v].at[i, "counted"])
    a = attrs[v].loc[ids]
    e = a["eligible"].astype(bool)
    r = {"loc": len(a), "elig": int(e.sum()), "nelig": int((~e).sum()),
         "bsl": int(a["bsl"].sum()), "nonbsl": int((~a["bsl"]).sum()),
         "conflict": sum(1 for i in area_sets[v].get(key, set()) if attrs[v].at[i, "conflict"])}
    if HAS_UNITS[v]:
        r["units"] = int(a["units"].fillna(0).sum())
        r["units_elig"] = int(a.loc[e, "units"].fillna(0).sum())
    return r


def list_block(key: str, sub: pd.DataFrame, v: str, own: bool) -> tuple[dict, pd.Series]:
    """Block B: one grantee list checked against one Fabric version (Fabric coordinates decide inside/outside)."""
    sub = sub.drop_duplicates("location_id")
    ids = pd.Index(sub["location_id"].astype("int64"))
    st = pd.Series([fabric_status(key, int(i), v) for i in ids], index=ids)
    a = attrs[v].reindex(ids)
    matched = st.ne(ST_NOT_IN)
    inside = st.isin([ST_ELIG, ST_NOT_ELIG])
    outside = st.isin([ST_ELIG_OUT, ST_NOT_ELIG_OUT])
    elig = a["eligible"].fillna(False).astype(bool)
    pts = to_metric(a[matched])
    poly = POLY_M[key]
    dist = pts.distance(poly)
    edge = pts.distance(poly.boundary)
    o = outside[matched]
    r = {"listed": len(ids), "matched": int(matched.sum()), "notfound": int((~matched).sum()),
         "conflict": int(st.eq(ST_CONFLICT).sum()),
         "in_elig": int((inside & elig).sum()), "in_nelig": int((inside & ~elig).sum()),
         "out_elig": int((outside & elig).sum()), "out_nelig": int((outside & ~elig).sum()),
         "bsl": int((matched & a["bsl"].fillna(False).astype(bool)).sum()),
         "nonbsl": int((matched & ~a["bsl"].fillna(False).astype(bool)).sum()),
         "out_near": int((o & (dist < OUT_BANDS_M[0])).sum()),
         "out_mid": int((o & (dist >= OUT_BANDS_M[0]) & (dist <= OUT_BANDS_M[1])).sum()),
         "out_far": int((o & (dist > OUT_BANDS_M[1])).sum()),
         "near_edge": int((edge <= EDGE_M).sum())}
    if HAS_UNITS[v]:
        r["units"] = int(a.loc[matched, "units"].fillna(0).sum())
    if own:                                   # grantee coordinates against the Fabric point of the same version
        g = sub.set_index(sub["location_id"].astype("int64")).reindex(ids)[["lon", "lat"]][matched]
        gap = to_metric(g).distance(pts)
        r["coord_checked"] = int(gap.notna().sum())
        r["coord_gap"] = int((gap > COORD_GAP_M).sum())
    return r, st.map(STATE_OF)


def crosstab(frm: pd.Series, to: pd.Series) -> pd.DataFrame:
    order = STATES + ([S_NB] if S_NB in set(frm) | set(to) else [])
    return pd.crosstab(frm, to).reindex(index=order, columns=order, fill_value=0)


def pct(new: float, old: float) -> float:
    return round((new - old) / old * 100, 1) if old else np.nan


KEY = {}            # project_key → flat dict of every key number
TRANS = []          # long from-to table
for key in PROJECTS:
    k = {}
    for v in ("v4", "v8"):
        for m, val in fabric_block(key, v).items():
            k[f"a_{v}_{m}"] = val
    s4 = sub_v4[sub_v4["project_key"] == key]
    s8 = sub_v8[sub_v8["project_key"] == key]
    states = {}
    for lst, sub in (("sub4", s4), ("sub8", s8)):
        for v in ("v4", "v8"):
            r, states[(lst, v)] = list_block(key, sub, v, own=(lst == "sub4") == (v == "v4"))
            for m, val in r.items():
                k[f"{lst}_{v}_{m}"] = val
    # Coverage: eligible inside the area that the matching list includes
    k["d_v4_listed"], k["d_v8_listed"] = k["sub4_v4_in_elig"], k["sub8_v8_in_elig"]
    k["d_v4_unlisted"], k["d_v8_unlisted"] = k["a_v4_elig"] - k["d_v4_listed"], k["a_v8_elig"] - k["d_v8_listed"]
    k["d_v4_nelig_listed"], k["d_v8_nelig_listed"] = k["sub4_v4_in_nelig"], k["sub8_v8_in_nelig"]
    k["d_v4_nelig_unlisted"] = k["a_v4_nelig"] - k["d_v4_nelig_listed"]
    k["d_v8_nelig_unlisted"] = k["a_v8_nelig"] - k["d_v8_nelig_listed"]
    # From-to tables
    ch = changes[changes["project_key"] == key]
    f4 = pd.Series([STATE_OF[fabric_status(key, int(i), "v4")] for i in ch["location_id"]])
    f8 = pd.Series([STATE_OF[fabric_status(key, int(i), "v8")] for i in ch["location_id"]])
    pops = {"Fabric locations in the project area": (f4, f8),
            cap(T_SUB4): (states[("sub4", "v4")], states[("sub4", "v8")]),
            T_SUB8: (states[("sub8", "v4")], states[("sub8", "v8")])}
    added = ~states[("sub8", "v4")].index.isin(states[("sub4", "v4")].index)
    pops["Added in the 2026 data"] = (states[("sub8", "v4")][added], states[("sub8", "v8")][added])
    for pop, (frm, to) in pops.items():
        m = crosstab(frm, to)
        for f in m.index:
            for t in m.columns:
                TRANS.append({"project_key": key, "population": pop, "from_fabric_v4": f, "to_fabric_v8": t,
                              "locations": int(m.at[f, t])})
    KEY[key] = k

transitions = pd.DataFrame(TRANS)
key_numbers = pd.DataFrame.from_dict(KEY, orient="index")
key_numbers.index.name = "project_key"


def tmat(key: str, pop: str) -> pd.DataFrame:
    t = transitions[(transitions["project_key"] == key) & (transitions["population"] == pop)]
    m = t.pivot(index="from_fabric_v4", columns="to_fabric_v8", values="locations")
    order = [x for x in STATES + [S_NB] if x in m.index]
    return m.reindex(index=order, columns=order)


# ---- plain-English explanation of how the change came together
def join_parts(parts: list[str]) -> str:
    parts = [p for p in parts if p]
    return ", ".join(parts[:-1]) + (" and " if len(parts) > 1 else "") + parts[-1] if parts else ""


def story_why(key: str) -> str:
    k, qr = KEY[key], qr_term(key)
    f = tmat(key, "Fabric locations in the project area")
    e4, e8 = k["a_v4_elig"], k["a_v8_elig"]
    kept = fsum(f, S_EI, S_EI)
    lost = join_parts([f"{fsum(f, S_EI, S_NI):,} changed to not eligible" if fsum(f, S_EI, S_NI) else "",
                       f"{fsum(f, S_EI, S_NOT):,} are IDs present only in version 4" if fsum(f, S_EI, S_NOT) else "",
                       f"{fsum(f, S_EI, OUT):,} have points now outside the area" if fsum(f, S_EI, OUT) else "",
                       f"{fsum(f, S_EI, S_CONF):,} have conflicting source records" if fsum(f, S_EI, S_CONF) else ""])
    won = join_parts([f"{fsum(f, S_NI, S_EI):,} changed from not eligible" if fsum(f, S_NI, S_EI) else "",
                      f"{fsum(f, S_NOT, S_EI):,} are IDs present only in version 8" if fsum(f, S_NOT, S_EI) else "",
                      f"{fsum(f, OUT, S_EI):,} have points that moved into the area" if fsum(f, OUT, S_EI) else "",
                      f"{fsum(f, S_CONF, S_EI):,} had conflicting source records in version 4" if fsum(f, S_CONF, S_EI) else ""])
    s = (f"Inside the project area, eligible locations went from {e4:,} in the {T_V4} to {e8:,} in the {T_V8} "
         f"({change_words(e8 - e4, 'eligible')}). Of the {e4:,}, {kept:,} are still eligible"
         + (f"; {lost}" if lost else "") + "." + (f" Eligible in version 8 but not before: {won}." if won else ""))
    i8 = tmat(key, cap(T_SUB4)).sum(axis=0)
    g = lambda st: int(sum(i8.get(x, 0) for x in ([st] if isinstance(st, str) else st)))
    n4 = k["sub4_v4_listed"]
    s += (f" Of the {n4:,} locations in the {T_SUB4}, " + join_parts([
        f"{g(S_EI):,} are eligible inside the area in the {T_V8}",
        f"{g(S_NI):,} are not eligible inside the area" if g(S_NI) else "",
        f"{g(OUT):,} lie outside the area" if g(OUT) else "",
        f"{g(S_NOT):,} are not in the {T_V8}" if g(S_NOT) else "",
        f"{g(S_CONF):,} have conflicting source records" if g(S_CONF) else ""]) + ".")
    a4 = tmat(key, "Added in the 2026 data").sum(axis=1)
    h = lambda st: int(sum(a4.get(x, 0) for x in ([st] if isinstance(st, str) else st)))
    was = lambda n: "was" if n == 1 else "were"
    r = summary.set_index("project_key").loc[key]
    s += (f" The {qr} keeps {r['s_kept']:,} of them, does not carry over {r['s_dropped']:,} and adds {r['s_added']:,}"
          + (f" (in the {T_V4}: " + join_parts([
              f"{h(S_EI):,} {was(h(S_EI))} already eligible inside the area" if h(S_EI) else "",
              f"{h(S_NI):,} {was(h(S_NI))} not eligible" if h(S_NI) else "",
              f"{h(OUT):,} {was(h(OUT))} outside the area" if h(OUT) else "",
              f"{h(S_NOT):,} {was(h(S_NOT))} not in the Fabric" if h(S_NOT) else "",
              f"{h(S_CONF):,} had conflicting source records" if h(S_CONF) else ""]) + ")" if r["s_added"] else "")
          + ".")
    return s


# ---- change patterns: analyst-defined review thresholds, each recorded with its count and denominator
P_NONE, P_SERVICE, P_FEWER, P_MORE, P_IDS, P_DQ = (
    "No configured material-change threshold triggered", "Listed locations changed to not eligible",
    "Fewer planned locations", "More or shifted planned locations", "ID changes", "Data validation issues")
RULE_ROWS: list[dict] = []


def big(n: int, base: int, share: float) -> bool:
    return n >= MIN_LOCATIONS and n >= share * base


def rule(key: str, pattern: str, name: str, count: int, denom: int, share: float | None, test: bool) -> bool:
    """Records one review-threshold check (always, triggered or not) and returns whether it was triggered."""
    RULE_ROWS.append({"project_key": key, "Pattern": pattern, "Review rule": name, "Count": int(count),
                      "Denominator": int(denom) if denom is not None else None,
                      "Share (percent)": round(count / denom * 100, 1) if denom else None,
                      "Threshold": (f"{share:.0%} of the denominator and at least {MIN_LOCATIONS}" if share is not None
                                    else "any"), "Triggered": "Yes" if test else "No",
                      "Kind": "Analyst-defined review threshold"})
    return test


def patterns(key: str) -> tuple[list[str], str]:
    k, qr = KEY[key], qr_term(key)
    r = summary.set_index("project_key").loc[key]
    n4, n8 = k["sub4_v4_listed"], k["sub8_v8_listed"]
    found, notes = [], []
    served = fsum(tmat(key, cap(T_SUB4)), S_EI, S_NI)
    if rule(key, P_SERVICE, "Initial-list locations eligible inside the area in version 4, not eligible in version 8",
            served, n4, SERVICE_DROP_SHARE, big(served, n4, SERVICE_DROP_SHARE)):
        found.append(P_SERVICE)
        notes.append(f"{served:,} of {n4:,} locations in the {T_SUB4} were eligible inside the area in the {T_V4} and "
                     f"are classified as not eligible in the {T_V8}.")
    dropped_elig = int(plain_b.at[key, "B03"])
    shorter = rule(key, P_FEWER, "2026 list shorter than the initial list", max(n4 - n8, 0), n4, LIST_CHANGE_SHARE,
                   n4 - n8 >= MIN_LOCATIONS and n8 <= (1 - LIST_CHANGE_SHARE) * n4)
    dropped = rule(key, P_FEWER, "Still-eligible locations not carried over", dropped_elig, n4, DROPPED_ELIGIBLE_SHARE,
                   big(dropped_elig, n4, DROPPED_ELIGIBLE_SHARE))
    if shorter or dropped:
        found.append(P_FEWER)
        notes.append(f"The {qr} plans {n8:,} locations against {n4:,} in the {T_SUB4}; {dropped_elig:,} still-eligible "
                     "locations were not carried over.")
    outside8 = k["sub8_v8_out_elig"] + k["sub8_v8_out_nelig"]
    longer = rule(key, P_MORE, "2026 list longer than the initial list", max(n8 - n4, 0), n4, LIST_CHANGE_SHARE,
                  n8 - n4 >= MIN_LOCATIONS and n8 >= (1 + LIST_CHANGE_SHARE) * n4)
    outside = rule(key, P_MORE, "2026 locations outside the award area", outside8, n8, OUTSIDE_SHARE,
                   big(outside8, n8, OUTSIDE_SHARE))
    if longer or outside:
        found.append(P_MORE)
        notes.append(f"The {qr} plans {n8:,} locations against {n4:,}; {outside8:,} of {n8:,} lie outside the project "
                     f"area in the {T_V8}.")
    gone = k["sub4_v8_notfound"]
    if rule(key, P_IDS, "Initial-list IDs not present in version 8", gone, n4, ID_CHURN_SHARE, big(gone, n4, ID_CHURN_SHARE)):
        found.append(P_IDS)
        notes.append(f"{gone:,} of {n4:,} IDs in the {T_SUB4} are not present in the {T_V8}.")
    dq = []
    dups = int(r["sub4_dup_rows"] + r["sub8_dup_rows"])
    bad = int(r["sub4_bad_rows"] + r["sub8_bad_rows"])
    conflicts = k["sub8_v8_conflict"] + k["a_v8_conflict"]
    gaps = k.get("sub4_v4_coord_gap", 0) + k.get("sub8_v8_coord_gap", 0)
    checked = k.get("sub4_v4_coord_checked", 0) + k.get("sub8_v8_coord_checked", 0)
    if rule(key, P_DQ, "Duplicate rows in the grantee lists", dups, None, None, dups > 0):
        dq.append(plural(dups, "duplicate row"))
    if rule(key, P_DQ, "Rows without a valid location ID", bad, None, None, bad > 0):
        dq.append(plural(bad, "row") + " without a valid location ID")
    if rule(key, P_DQ, "2026 IDs not present in version 8", k["sub8_v8_notfound"], n8, None, k["sub8_v8_notfound"] > 0):
        dq.append(plural(k['sub8_v8_notfound'], "ID") + f" in the {qr} not in the {T_V8}")
    if rule(key, P_DQ, "Conflicting source records (2026 list and project area, version 8)", conflicts, None, None,
            conflicts > 0):
        dq.append(plural(conflicts, "ID") + " with conflicting source records")
    if rule(key, P_DQ, f"Grantee coordinates more than {COORD_GAP_M} m from the Fabric point", gaps, checked,
            COORD_GAP_SHARE, big(gaps, checked, COORD_GAP_SHARE)):
        dq.append(f"{gaps:,} of {checked:,} grantee coordinates more than {COORD_GAP_M} m from the Fabric point")
    if dq:
        found.append(P_DQ)
        notes.append("Data validation: " + join_parts(dq) + ".")
    if not [p for p in found if p != P_DQ]:
        found.insert(0, P_NONE)
        notes.insert(0, "None of the configured scope or eligibility thresholds was reached; see the Review rules tab "
                        "for the counts behind each.")
    return found, " ".join(notes)


summary["story_why"] = summary["project_key"].map(story_why)
pat = {key: patterns(key) for key in PROJECTS}
summary["patterns"] = summary["project_key"].map(lambda k_: " · ".join(pat[k_][0]))
summary["pattern_notes"] = summary["project_key"].map(lambda k_: pat[k_][1])
summary = summary.merge(key_numbers.reset_index(), on="project_key", how="left")
summary["bottom_line"] = summary.apply(bottom_line, axis=1)
summary["narrative"] = summary["bottom_line"] + " " + summary["story_a"] + " " + summary["story_b"]

# ---- reconciliation of every block
for key, k in KEY.items():
    for v in ("v4", "v8"):
        assert k[f"a_{v}_loc"] == k[f"a_{v}_elig"] + k[f"a_{v}_nelig"]
        for lst in ("sub4", "sub8"):
            p = f"{lst}_{v}_"
            assert k[p + "listed"] == k[p + "matched"] + k[p + "notfound"], (key, p)
            inside_outside = k[p + "in_elig"] + k[p + "in_nelig"] + k[p + "out_elig"] + k[p + "out_nelig"]
            assert inside_outside <= k[p + "matched"], (key, p)
            assert k[p + "out_near"] + k[p + "out_mid"] + k[p + "out_far"] == k[p + "out_elig"] + k[p + "out_nelig"]
    assert fsum(tmat(key, "Fabric locations in the project area"), S_EI, STATES + [S_NB]) == k["a_v4_elig"]
    assert fsum(tmat(key, "Fabric locations in the project area"), STATES + [S_NB], S_EI) == k["a_v8_elig"]
    assert k["sub4_v4_listed"] == summary.set_index("project_key").at[key, "sub4_n"]
    assert k["sub8_v8_listed"] == summary.set_index("project_key").at[key, "sub8_n"]
review_rules = pd.DataFrame(RULE_ROWS)
display(summary[["project_key", "patterns", "pattern_notes"]])

def fmt_n(v) -> str:
    return "–" if v is None or (isinstance(v, float) and np.isnan(v)) else f"{int(v):,}"


def fmt_d(v) -> str:
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return "–"
    return f"{int(v):+,}" if int(v) else "0"


def fmt_p(v) -> str:
    return "–" if v is None or (isinstance(v, float) and np.isnan(v)) else f"{v:+.1f}%"


def key_tables(key: str) -> dict[str, pd.DataFrame]:
    """Plain tables for one project: Fabric in the area, grantee lists, coverage, from-to tables."""
    k, qr = KEY[key], qr_term(key)
    g = lambda name: k.get(name, np.nan)
    rows_a = [("Locations", "loc"), ("Eligible", "elig"), ("Not eligible", "nelig")]
    if HAS_UNITS["v4"] and HAS_UNITS["v8"]:
        rows_a += [("Units at the locations (unit count)", "units"), ("Units at eligible locations", "units_elig")]
    rows_a += [("Broadband serviceable locations", "bsl"), ("Not broadband serviceable locations", "nonbsl"),
               ("Conflicting source records (not counted above)", "conflict")]
    A = pd.DataFrame([{"Fabric locations inside the project area": lab, "Version 4": g(f"a_v4_{m}"),
                       "Version 8": g(f"a_v8_{m}"), "Change": g(f"a_v8_{m}") - g(f"a_v4_{m}")} for lab, m in rows_a])
    rows_b = [("Listed", "listed"), ("Matched to the Fabric", "matched"), ("Not found in the Fabric", "notfound"),
              ("Conflicting source records", "conflict"),
              ("Inside the area – eligible", "in_elig"), ("Inside the area – not eligible", "in_nelig"),
              ("Outside the area – eligible", "out_elig"), ("Outside the area – not eligible", "out_nelig"),
              ("Outside: under 100 m", "out_near"), ("Outside: 100 m to 1 km", "out_mid"), ("Outside: over 1 km", "out_far"),
              (f"Within {EDGE_M} m of the boundary", "near_edge")]
    if HAS_UNITS["v4"] and HAS_UNITS["v8"]:
        rows_b.insert(3, ("Units at matched locations", "units"))
    rows_b += [("Matched and broadband serviceable", "bsl"), ("Matched and not broadband serviceable", "nonbsl"),
               (f"Grantee point over {COORD_GAP_M} m from Fabric point", "coord_gap")]
    B = pd.DataFrame([{"Grantee lists checked against the Fabric": lab,
                       f"{cap(T_SUB4)} · version 4": g(f"sub4_v4_{m}"), f"{qr} · version 8": g(f"sub8_v8_{m}"),
                       "Change": g(f"sub8_v8_{m}") - g(f"sub4_v4_{m}"),
                       f"{cap(T_SUB4)} · version 8": g(f"sub4_v8_{m}"), f"{qr} · version 4": g(f"sub8_v4_{m}")}
                      for lab, m in rows_b])
    cov = [("Eligible inside the area", "a_{v}_elig"), ("… listed by the grantee", "d_{v}_listed"),
           ("… not listed", "d_{v}_unlisted"), ("Not eligible inside the area", "a_{v}_nelig"),
           ("… listed by the grantee ", "d_{v}_nelig_listed"), ("… not listed ", "d_{v}_nelig_unlisted")]
    D = pd.DataFrame([{"Coverage of locations inside the area": lab, "Version 4": g(f.format(v="v4")),
                       "Version 8": g(f.format(v="v8"))} for lab, f in cov])
    D["Change"] = D["Version 8"] - D["Version 4"]
    out = {"Fabric in the area": A, "Grantee lists": B, "Coverage": D}
    for pop, name in (("Fabric locations in the project area", "From-to Fabric"), (cap(T_SUB4), "From-to initial list"),
                      (T_SUB8, "From-to 2026 list")):
        m = tmat(key, pop)
        m = m.loc[[s for s in m.index if m.loc[s].sum() or s in STATES], [s for s in m.columns if m[s].sum() or s in STATES]]
        m.insert(len(m.columns), "Total", m.sum(axis=1))
        m.loc["Total"] = m.sum(axis=0)
        out[name] = m.rename_axis(f"{pop}: version 4 → version 8").reset_index()
    return out

SUB4, SUB8 = "Initial grantee submission", "2026 Q2 and Q3 data"
FIELD_GUIDE = {
    "project_key": ("project_key", "Project key"), "label": ("project_name", "Project name"),
    "pr_number": ("pr_number", "PR number"), "page_title": ("page_title", "Page title"),
    "n_parts": ("project_area_polygons", "Polygons in the project area"),
    "matched_names": ("matched_project_area_names", "Project area layers or names matched"),
    "area_sq_mi": ("project_area_square_miles", "Project area (square miles)"),
    "sub_v8_file": ("quarterly_report_file", "2026 quarterly report file used"),
    "v4_total": ("version_4_locations_in_area", "Fabric version 4 – locations in the project area"),
    "v4_elig": ("version_4_eligible_in_area", "Fabric version 4 – eligible locations in the project area"),
    "v8_total": ("version_8_locations_in_area", "Fabric version 8 – locations in the project area"),
    "v8_elig": ("version_8_eligible_in_area", "Fabric version 8 – eligible locations in the project area"),
    "total_chg": ("change_in_locations_in_area", "Change in locations in the project area (version 8 minus version 4)"),
    "elig_chg": ("change_in_eligible_in_area", "Change in eligible locations in the project area (version 8 minus version 4)"),
    "total_pct": ("percent_change_locations_in_area", "Percent change in locations in the project area"),
    "elig_pct": ("percent_change_eligible_in_area", "Percent change in eligible locations in the project area"),
    "f_elig_lost": ("area_lost_eligibility", "Locations in the area whose classification changed from eligible to not eligible"),
    "f_elig_gain": ("area_became_eligible", "Locations in the area that became eligible"),
    "f_retired": ("area_ids_removed_in_version_8", "IDs in the area that are not in Fabric version 8"),
    "f_new": ("area_new_ids_in_version_8", "IDs in the area that are new in Fabric version 8"),
    "f_moved": ("area_points_moved_across_boundary", "Locations whose point moved across the project boundary"),
    "f_bsl_chg": ("area_serviceable_status_changed", "Locations whose broadband serviceable status changed"),
    "sub4_n": ("initial_submission_locations", f"{SUB4} – locations listed"),
    "sub4_found": ("initial_submission_found_in_version_4", f"{SUB4} – found in Fabric version 4"),
    "sub4_in_area": ("initial_submission_inside_area_version_4", f"{SUB4} – inside the area in Fabric version 4"),
    "sub4_elig": ("initial_submission_eligible_version_4", f"{SUB4} – eligible in Fabric version 4, inside or outside the area"),
    "sub8_n": ("data_2026_locations", f"{SUB8} – locations listed"),
    "sub8_found": ("data_2026_found_in_version_8", f"{SUB8} – found in Fabric version 8"),
    "sub8_in_area": ("data_2026_inside_area_version_8", f"{SUB8} – inside the area in Fabric version 8"),
    "sub8_elig": ("data_2026_eligible_version_8", f"{SUB8} – eligible in Fabric version 8, inside or outside the area"),
    "sub_chg": ("change_in_listed_locations", f"Change in listed locations ({SUB8} minus {SUB4.lower()})"),
    "s_kept": ("in_both_lists", "Locations in both lists (same ID)"),
    "s_dropped": ("not_carried_into_2026_data", f"Locations not carried into the {SUB8}"),
    "s_added": ("new_in_2026_data", f"Locations new in the {SUB8}"),
    "x_retired": ("not_carried_id_not_in_version_8", "Not carried over because the ID is not in Fabric version 8"),
    "x_served": ("listed_locations_now_not_eligible", "Listed or not-carried-over locations whose classification changed from eligible to not eligible"),
    "x_applicant": ("changes_not_explained_by_fabric", "List changes not explained by observed Fabric change"),
    "x_new_fabric": ("added_new_in_version_8", "Added locations that are new in Fabric version 8"),
    "sub8_not_elig": ("data_2026_not_eligible_version_8", f"{SUB8} – not eligible in Fabric version 8"),
    "v4_elig_in_area": ("synthesis_version_4_eligible_in_area", "Coverage: eligible locations in the area (version 4)"),
    "v4_elig_submitted": ("synthesis_version_4_eligible_listed", f"Coverage: of those, listed in the {SUB4.lower()}"),
    "v4_elig_not_submitted": ("synthesis_version_4_eligible_not_listed", f"Coverage: of those, not listed in the {SUB4.lower()}"),
    "v4_pct_elig_submitted": ("synthesis_version_4_percent_listed", "Coverage: percent of eligible locations listed (version 4)"),
    "v8_elig_in_area": ("synthesis_version_8_eligible_in_area", "Coverage: eligible locations in the area (version 8)"),
    "v8_elig_submitted": ("synthesis_version_8_eligible_listed", f"Coverage: of those, listed in the {SUB8}"),
    "v8_elig_not_submitted": ("synthesis_version_8_eligible_not_listed", f"Coverage: of those, not listed in the {SUB8}"),
    "v8_pct_elig_submitted": ("synthesis_version_8_percent_listed", "Coverage: percent of eligible locations listed (version 8)"),
    "fu_total": ("follow_up_total", "Needs follow-up – total"),
    "far_points": ("listed_locations_over_5_km_away", "Listed locations more than 5 km from the project area (not on the map)"),
    "bottom_line": ("bottom_line", "Bottom line"), "story_a": ("eligibility_summary", "Eligibility summary"),
    "story_b": ("submission_summary", "Grantee submission summary"), "story_c": ("follow_up_summary", "Follow-up summary"),
    "narrative": ("full_summary", "Full summary"), "patterns": ("change_pattern", "Change pattern"),
    "pattern_notes": ("change_pattern_notes", "Change pattern notes"),
    "story_why": ("how_the_change_came_together", "How the change came together"),
}
FU_NAMES = {"B01": "follow_up_2026_not_in_version_8", "B13": "follow_up_2026_outside_area",
            "B02": "follow_up_2026_not_eligible_version_8", "B14": "follow_up_2026_not_serviceable_location",
            "B03": "follow_up_not_carried_over_still_eligible", "B04": "follow_up_new_already_eligible_version_4",
            "B12": "follow_up_conflicting_source_records"}
for c_ in FOLLOW_UP_CLASSES:
    FIELD_GUIDE[f"fu_{c_.lower()}"] = (FU_NAMES[c_], "Needs follow-up – " + SUB_CLASSES[c_][0].replace(T_SUB8_SHORT, SUB8))
for pre, (fname, lab) in {"sub4": ("initial_submission", SUB4), "sub8": ("data_2026", SUB8),
                          "f4": ("fabric_version_4_area", "Fabric version 4 (project area)"),
                          "f8": ("fabric_version_8_area", "Fabric version 8 (project area)")}.items():
    FIELD_GUIDE[f"{pre}_dup_ids"] = (f"{fname}_ids_listed_more_than_once", f"{lab} – IDs listed more than once")
    FIELD_GUIDE[f"{pre}_dup_rows"] = (f"{fname}_duplicate_rows_removed", f"{lab} – duplicate rows removed")
    FIELD_GUIDE[f"{pre}_bad_rows"] = (f"{fname}_rows_without_valid_id", f"{lab} – rows without a valid ID")
# Key numbers (section 6a)
A_NAMES = {"loc": ("locations", "locations"), "elig": ("eligible", "eligible"), "nelig": ("not_eligible", "not eligible"),
           "units": ("units", "units at the locations (unit count)"), "units_elig": ("units_at_eligible", "units at eligible locations"),
           "bsl": ("serviceable", "broadband serviceable locations"), "nonbsl": ("not_serviceable", "not broadband serviceable locations"),
           "conflict": ("conflicting_records", "conflicting source records (left out of the counts)")}
B_NAMES = {"listed": ("listed", "listed"), "matched": ("matched", "matched to the Fabric"),
           "notfound": ("not_found", "not found in the Fabric"), "units": ("units", "units at matched locations"),
           "conflict": ("conflicting_records", "conflicting source records"),
           "in_elig": ("inside_eligible", "inside the area, eligible"),
           "in_nelig": ("inside_not_eligible", "inside the area, not eligible"),
           "out_elig": ("outside_eligible", "outside the area, eligible"),
           "out_nelig": ("outside_not_eligible", "outside the area, not eligible"),
           "out_near": ("outside_under_100_m", "outside, under 100 m from the area"),
           "out_mid": ("outside_100_m_to_1_km", "outside, 100 m to 1 km from the area"),
           "out_far": ("outside_over_1_km", "outside, over 1 km from the area"),
           "near_edge": ("within_50_m_of_boundary", f"within {EDGE_M} m of the boundary"),
           "bsl": ("serviceable", "matched and broadband serviceable"), "nonbsl": ("not_serviceable", "matched and not broadband serviceable"),
           "coord_checked": ("grantee_points_checked", "grantee coordinates checked"),
           "coord_gap": ("grantee_point_over_100_m_off", f"grantee point over {COORD_GAP_M} m from the Fabric point")}
for v in ("v4", "v8"):
    n_ = v[1]
    for m, (fn, lab) in A_NAMES.items():
        FIELD_GUIDE[f"a_{v}_{m}"] = (f"area_version_{n_}_{fn}", f"Fabric version {n_} inside the area – {lab}")
    for lst, (fn, lab) in {"sub4": ("initial_submission", SUB4), "sub8": ("data_2026", SUB8)}.items():
        for m, (mf, ml) in B_NAMES.items():
            FIELD_GUIDE[f"{lst}_{v}_{m}"] = (f"{fn}_vs_version_{n_}_{mf}", f"{lab} checked against Fabric version {n_} – {ml}")
    for m, (fn, lab) in {"listed": ("eligible_listed", "eligible inside the area, listed"),
                         "unlisted": ("eligible_not_listed", "eligible inside the area, not listed"),
                         "nelig_listed": ("not_eligible_listed", "not eligible inside the area, listed"),
                         "nelig_unlisted": ("not_eligible_not_listed", "not eligible inside the area, not listed")}.items():
        FIELD_GUIDE[f"d_{v}_{m}"] = (f"coverage_version_{n_}_{fn}", f"Coverage, version {n_} – {lab}")
assert all(len(f) <= 64 for f, _ in FIELD_GUIDE.values()), "geodatabase field name too long"
assert len({f for f, _ in FIELD_GUIDE.values()}) == len(FIELD_GUIDE), "duplicate geodatabase field name"


def readable(df: pd.DataFrame) -> pd.DataFrame:
    """Columns renamed to plain labels (unknown columns keep their name)."""
    return df.rename(columns={c: FIELD_GUIDE[c][1] for c in df.columns if c in FIELD_GUIDE})


def field_names(df: pd.DataFrame) -> pd.DataFrame:
    """Columns renamed to geodatabase field names without abbreviations."""
    return df.rename(columns={c: FIELD_GUIDE[c][0] for c in df.columns if c in FIELD_GUIDE})


field_guide = pd.DataFrame([{"source_column": c, "field_name": f, "label": lab} for c, (f, lab) in FIELD_GUIDE.items()])
missing = [c for c in summary.columns if c not in FIELD_GUIDE and c != "geometry"]
if missing:
    log.warning("Summary columns without a plain label: %s", missing)

# %% [markdown]
# ### 6b. Page extents (whole project visible on every map)

# %%
# Page extents: project outline plus every point that belongs to the project (points farther than
# FAR_POINT_KM from the outline are left out of the frame and reported instead), padded and stretched to
# the map frame's shape so the whole project is always visible.
MAP_FRAME_W_IN, MAP_FRAME_H_IN = 9.0, 6.75      # map frame on the 14 x 8.5 in page
FAR_POINT_KM = 5
PAD_FRACTION = 0.06
MERC = "EPSG:3857"


def project_points(key: str) -> dict[str, gpd.GeoDataFrame]:
    """The three point sets drawn for one project, in Web Mercator."""
    a = changes[(changes["project_key"] == key) & changes["map_show"]].copy()
    a["lon"], a["lat"] = a["lon_v8"].fillna(a["lon_v4"]), a["lat_v8"].fillna(a["lat_v4"])
    b = subs[(subs["project_key"] == key) & subs["map_show"]].copy()
    b["lon"], b["lat"] = b["fab_lon"], b["fab_lat"]
    c = b.copy()
    c["map_class"] = c["map_class"].where(c["needs_follow_up"], "C0")
    out = {}
    for name, df in (("A", a), ("B", b), ("C", c)):
        g = points_from_lonlat(df).dropna(subset=["geometry"])
        out[name] = g[~g.geometry.is_empty].to_crs(MERC)
    return out


def page_extent(key: str) -> tuple[tuple[float, float, float, float], int]:
    """Frame-shaped extent (Web Mercator) covering the project and its points; also the count of far points."""
    area = areas.loc[areas["project_key"] == key].to_crs(MERC).geometry.iloc[0]
    pts = pd.concat([g.geometry for g in project_points(key).values()])
    km = np.cos(np.radians(gpd.GeoSeries([area.centroid], crs=MERC).to_crs(WGS84).iloc[0].y)) / 1000
    dist_km = pts.distance(area) * km
    near = pts[dist_km <= FAR_POINT_KM]
    x0, y0, x1, y1 = gpd.GeoSeries(list(near) + [area], crs=MERC).total_bounds
    w, h = max(x1 - x0, 500), max(y1 - y0, 500)
    cx, cy = (x0 + x1) / 2, (y0 + y1) / 2
    w, h = w * (1 + 2 * PAD_FRACTION), h * (1 + 2 * PAD_FRACTION)
    aspect = MAP_FRAME_W_IN / MAP_FRAME_H_IN
    if w / h > aspect:
        h = w / aspect
    else:
        w = h * aspect
    return (cx - w / 2, cy - h / 2, cx + w / 2, cy + h / 2), int((dist_km > FAR_POINT_KM).sum())


extents, far_points = {}, {}
for key in PROJECTS:
    extents[key], far_points[key] = page_extent(key)
summary["far_points"] = summary["project_key"].map(far_points)
extent_rows = [{"project_key": k, "geometry": shapely.box(*extents[k])} for k in PROJECTS]
page_extents = (gpd.GeoDataFrame(extent_rows, crs=MERC).to_crs(WGS84)
                .merge(summary[["project_key", "page_title", "pr_number", "bottom_line", "story_a", "story_b",
                                "story_c", "fu_total", "far_points"]], on="project_key"))
print({k: f"{far_points[k]} far points" for k in PROJECTS})

# Shaded density instead of dots for the large "no change" group on the eligibility map: thousands of
# grey dots hide the changes, so these locations are counted in hexagons and shaded light to darker grey.
SHADED_CLASSES = {"A9"}                 # No change – not eligible
HEX_ACROSS = 160                        # hexagons across the map frame width
SHADE_RGB = {1: (204, 204, 204), 2: (170, 170, 170), 3: (130, 130, 130)}
SHADE_LABEL = {1: "Fewer locations", 2: "More locations", 3: "Most locations"}


def hex_cells(x: np.ndarray, y: np.ndarray, size: float) -> pd.DataFrame:
    """Counts points per pointy-top hexagon of circumradius `size`; returns centre x/y and count per cell."""
    q = (np.sqrt(3) / 3 * x - y / 3) / size
    r = (2 / 3 * y) / size
    cx, cz = q, r
    cy = -cx - cz
    rx, ry, rz = np.round(cx), np.round(cy), np.round(cz)
    dx, dy, dz = np.abs(rx - cx), np.abs(ry - cy), np.abs(rz - cz)
    fix_x = (dx > dy) & (dx > dz)
    fix_y = ~fix_x & (dy > dz)
    rx = np.where(fix_x, -ry - rz, rx)
    rz = np.where(~fix_x & ~fix_y, -rx - ry, rz)
    cells = pd.DataFrame({"q": rx.astype(int), "r": rz.astype(int)}).value_counts().rename("locations").reset_index()
    cells["x"] = size * np.sqrt(3) * (cells["q"] + cells["r"] / 2)
    cells["y"] = size * 1.5 * cells["r"]
    return cells


def hexagon(x: float, y: float, size: float) -> shapely.Polygon:
    ang = np.radians(np.arange(6) * 60 + 30)
    return shapely.Polygon(np.column_stack([x + size * np.cos(ang), y + size * np.sin(ang)]))


density_rows = []
for key in PROJECTS:
    a = changes[(changes["project_key"] == key) & changes["map_class"].isin(SHADED_CLASSES)]
    if a.empty:
        continue
    g = points_from_lonlat(a.assign(lon=a["lon_v8"].fillna(a["lon_v4"]), lat=a["lat_v8"].fillna(a["lat_v4"])))
    g = g[~g.geometry.is_empty].to_crs(MERC)
    size = (extents[key][2] - extents[key][0]) / HEX_ACROSS / np.sqrt(3)
    for code in sorted(SHADED_CLASSES):
        gc = g[g["map_class"] == code]
        if gc.empty:
            continue
        cells = hex_cells(gc.geometry.x.to_numpy(), gc.geometry.y.to_numpy(), size)
        q1, q2 = np.quantile(cells["locations"], [1 / 3, 2 / 3])
        cells["shade"] = 1 + (cells["locations"] > q1).astype(int) + (cells["locations"] > q2).astype(int)
        for c in cells.itertuples():
            density_rows.append({"project_key": key, "map_class": code, "locations": int(c.locations),
                                 "shade": int(c.shade), "shade_label": SHADE_LABEL[int(c.shade)],
                                 "geometry": hexagon(c.x, c.y, size)})
density = gpd.GeoDataFrame(density_rows, columns=["project_key", "map_class", "locations", "shade", "shade_label",
                                                  "geometry"], crs=MERC)
density["page_title"] = density["project_key"].map(summary.set_index("project_key")["page_title"])
# Every shaded location sits in exactly one hexagon
assert density.groupby("project_key")["locations"].sum().reindex(list(PROJECTS), fill_value=0).equals(
    changes[changes["map_class"].isin(SHADED_CLASSES)].groupby("project_key").size()
    .reindex(list(PROJECTS), fill_value=0)), "density cells do not account for every shaded location"

# %% [markdown]
# ## 6b. What changed, how it fits the observed Fabric change, and what to review
#
# Three kinds of output are kept apart: **source facts** (what the lists and the Fabric files contain), **analytical
# interpretations** (what the rules below read from those facts) and **analyst decisions** (blank columns for the
# analyst). An interpretation says whether a change is *consistent with* the observed Fabric change; it never states
# the cause.
#
# **Source facts per listed location** are separate fields, never merged: present in the Fabric, inside the award
# area, eligible, broadband serviceable location, and conflicting source records, each for version 4 and version 8,
# plus whether the location is in each grantee list. "Not eligible" means the eligibility classification is not
# eligible; it is not read as "now has service".
#
# **Conflicting source records.** When copies of one location ID in a Fabric file disagree on eligibility, the ID is
# marked as conflicting, left out of every eligibility count, reported as an unresolved count, and never resolved in
# favour of eligibility.
#
# **Three review dimensions per location:** list continuity (Retained / Removed / Added), current status (Acceptable /
# Review required / Not on the current list) and the observed Fabric transition. Retained does not mean no action.
#
# **Scope change.** Retained, removed and added, net change, gross turnover (removed + added) and turnover rate
# (turnover ÷ initial locations; it can pass 100 %).
#
# **Eligibility change.** Eligible in both versions, not eligible in both, eligible → not eligible, not eligible →
# eligible, eligible IDs present only in version 4 or only in version 8, eligible points moved out of or into the
# area, and moves to or from conflicting records or non-serviceable status. Gross loss and gross gain reconcile the
# version 4 eligible total to the version 8 total.
#
# **Observed explanation of each added or removed location** (with its percent of gross turnover):
#
# | Class | When |
# |---|---|
# | Consistent with Fabric change | Removed: ID present only in version 4, point now outside, or eligible → not eligible. Added: ID present only in version 8, point moved in, or not eligible → eligible |
# | Not explained by observed Fabric change | Removed although still eligible inside the area; added although already eligible in version 4, not eligible, or outside the area |
# | Data validation issue | Conflicting source records, or the ID is missing from the Fabric version it belongs to |
# | Indeterminate | The ID is in neither Fabric version |
#
# **Flags, in priority order** (every flag is recorded: *All flags*, *All recommended actions*; the first gives
# *Priority*, *Priority reason* and *Primary action*):
#
# | Priority | Flag | Action |
# |---|---|---|
# | 1 | Current location ID not found in Fabric version 8 | Correct data |
# | 2 | Current Fabric point outside the award area | Verify boundary |
# | 3 | Current location not eligible in Fabric version 8 | Verify eligibility |
# | 4 | Current record not a broadband serviceable location | Verify eligibility |
# | 5 | Removed but still eligible inside the area | Confirm scope change |
# | 6 | Added without an observed Fabric transition | Confirm scope change |
# | 7 | Conflicting source records | Correct data |
# | 7 | Grantee coordinates more than `COORD_GAP_M` m from the Fabric point (data validation) | Correct data |
#
# Priorities 1–7 except the coordinate flag are the locations that **need follow-up** and match the follow-up map.
#
# **Evidence strength** describes the evidence, not certainty about a conclusion. *Strong*: exact ID in both Fabric
# versions. *Moderate*: exact ID in one version only. *Limited*: ID in neither version, within `EDGE_M` m of the
# boundary, a coordinate discrepancy, or conflicting source records. Each basis is its own Yes/No field.
#
# **Current exceptions** (2026 data against Fabric version 8): IDs not found, outside, not eligible, not serviceable
# locations, conflicting records, coordinate discrepancies, locations with more than one issue, locations usable after
# data validation (eligible inside, only a data issue) and acceptable locations. **Coverage rate** is reported with its
# numerator and denominator, by locations and by units.
#
# **Spatial findings.** Near-boundary counts, points moved into or out of the area, point displacement (median, 95th
# percentile, maximum), grantee coordinate gaps (count, rate, median, 95th percentile, maximum) and concentration
# (share in the busiest 10 % of `CONCENTRATION_CELL_M` m hexagons). Moved coordinates are map coordinate edits.
#
# **Five determinations per project:** scope change, eligibility change, observed explanation, current exceptions and
# analyst disposition (suggested: Accept and document, Request clarification, Request corrected data, or Further
# review), each with a blank *Analyst decision* column.
#
# **Review rules.** The change-pattern thresholds are analyst-defined review thresholds set in the notebook settings.
# Every check is recorded with its count, denominator, share, threshold and whether it was triggered. When none is
# triggered the pattern reads "No configured material-change threshold triggered".

# %%
CONCENTRATION_CELL_M = 400       # hexagon size (circumradius, metres) for geographic concentration
CONCENTRATION_TOP = 0.10         # busiest share of occupied hexagons
CONCENTRATED_SHARE = 0.50        # share of locations in the busiest hexagons that counts as concentrated

C_CONSISTENT, C_NOT_EXPLAINED, C_VALIDATION, C_INDETERMINATE, C_RETAINED = (
    "Consistent with Fabric change", "Not explained by observed Fabric change", "Data validation issue",
    "Indeterminate", "Retained (in both lists)")
CLASSES = [C_CONSISTENT, C_NOT_EXPLAINED, C_VALIDATION, C_INDETERMINATE]
A_NONE, A_DOC, A_ELIG, A_BOUNDARY, A_SCOPE, A_DATA = (
    "No action", "Document change", "Verify eligibility", "Verify boundary", "Confirm scope change", "Correct data")
K_FU, K_DV = "Needs follow-up", "Data validation"
FLAGS = {   # code: (issue, action, question, priority (1 = first), kind)
    "missing": ("Current location ID not found in Fabric version 8", A_DATA,
                "What is the Fabric version 8 location ID for this location?", 1, K_FU),
    "outside": ("Current Fabric point outside the award area", A_BOUNDARY,
                "Why is a location outside the awarded area in the 2026 data?", 2, K_FU),
    "not_eligible": ("Current location not eligible in Fabric version 8", A_ELIG,
                     "How does a location classified as not eligible in Fabric version 8 qualify?", 3, K_FU),
    "not_bsl": ("Current record not a broadband serviceable location", A_ELIG,
                "Is this location a broadband serviceable location?", 4, K_FU),
    "dropped_eligible": ("Removed from the submission but still eligible inside the area", A_SCOPE,
                         "Why was this eligible location taken off the planned list?", 5, K_FU),
    "added_unexplained": ("Added without an observed Fabric transition (already eligible in version 4)", A_SCOPE,
                          "Why was this already-eligible location not in the initial grantee submission?", 6, K_FU),
    "conflict": ("Conflicting source records in the Fabric", A_DATA,
                 "Which record is right for this location? The Fabric copies disagree on eligibility.", 7, K_FU),
    "coordinates": (f"Current coordinate discrepancy (over {COORD_GAP_M} m)", A_DATA,
                    "Please confirm the location; the coordinates you reported are far from the FCC point.", 7, K_DV),
}
FLAG_ORDER = sorted(FLAGS, key=lambda f_: FLAGS[f_][3])
FOLLOW_UP_FLAGS = [f_ for f_ in FLAG_ORDER if FLAGS[f_][4] == K_FU]
ST_OUT = {ST_ELIG_OUT, ST_NOT_ELIG_OUT}


def grantee_gap_m(sub: pd.DataFrame, v: str) -> pd.Series:
    """Distance (m) from the grantee's own coordinates to the Fabric point of the same ID; NaN when either is missing."""
    s = sub.drop_duplicates("location_id").set_index(sub.drop_duplicates("location_id")["location_id"].astype("int64"))
    a = attrs[v].reindex(s.index)
    ok = s["lon"].notna() & s["lat"].notna() & a["lon"].notna()
    out = pd.Series(np.nan, index=s.index)
    if ok.any():
        g = gpd.GeoSeries(gpd.points_from_xy(s.loc[ok, "lon"], s.loc[ok, "lat"]), index=s.index[ok], crs=WGS84).to_crs(METRIC_CRS)
        f = gpd.GeoSeries(gpd.points_from_xy(a.loc[ok, "lon"], a.loc[ok, "lat"]), index=s.index[ok], crs=WGS84).to_crs(METRIC_CRS)
        out[ok] = g.distance(f).round(1)
    return out


# ---- per listed location: separate source facts first, then interpretations
s4, s8 = subs["fabric_status_v4"], subs["fabric_status_v8"]
chg = subs["sub_change"]
lid = subs["location_id"].astype("int64")
for v in ("v4", "v8"):                       # source facts, one field per dimension
    a = attrs[v].reindex(lid)
    subs[f"present_in_fabric_{v}"] = a["lon"].notna().to_numpy() | lid.isin(attrs[v].index).to_numpy()
    subs[f"inside_award_area_{v}"] = [i in area_sets[v].get(k, set()) for k, i in zip(subs["project_key"], lid)]
    subs[f"conflicting_records_{v}"] = a["conflict"].fillna(False).astype(bool).to_numpy()
    subs[f"eligible_{v}"] = (a["eligible"].fillna(False).astype(bool) & ~a["conflict"].fillna(False).astype(bool)).to_numpy()
    subs[f"broadband_serviceable_{v}"] = a["bsl"].fillna(False).astype(bool).to_numpy()
gap4 = pd.concat([grantee_gap_m(sub_v4[sub_v4["project_key"] == k], "v4").rename(k) for k in PROJECTS], axis=1)
gap8 = pd.concat([grantee_gap_m(sub_v8[sub_v8["project_key"] == k], "v8").rename(k) for k in PROJECTS], axis=1)
subs["grantee_gap_v4_m"] = [gap4.at[i, k] if i in gap4.index and v else np.nan
                            for k, i, v in zip(subs["project_key"], lid, subs["in_sub_v4"])]
subs["grantee_gap_v8_m"] = [gap8.at[i, k] if i in gap8.index and v else np.nan
                            for k, i, v in zip(subs["project_key"], lid, subs["in_sub_v8"])]
subs["grantee_gap_m"] = subs[["grantee_gap_v4_m", "grantee_gap_v8_m"]].max(axis=1)

# Cause of each list change (association with the observed Fabric change, not proof of cause)
neither = s4.eq(ST_NOT_IN) & s8.eq(ST_NOT_IN)
conflict_any = s4.eq(ST_CONFLICT) | s8.eq(ST_CONFLICT)
removed, added = chg.eq(SC_DROPPED), chg.eq(SC_ADDED)
cause_rules = [
    (~(removed | added), C_RETAINED, "In both lists"),
    (neither, C_INDETERMINATE, "Location ID in neither Fabric version"),
    (conflict_any, C_VALIDATION, "Conflicting source records in the Fabric"),
    (removed & s4.eq(ST_NOT_IN), C_VALIDATION, "Initial ID not present in Fabric version 4"),
    (removed & s8.eq(ST_NOT_IN), C_CONSISTENT, "ID present only in Fabric version 4"),
    (removed & s8.isin(ST_OUT), C_CONSISTENT, "Point now outside the area"),
    (removed & s8.isin([ST_NOT_ELIG, ST_NOT_BSL]), C_CONSISTENT, "Eligible to not eligible"),
    (removed & s8.eq(ST_ELIG), C_NOT_EXPLAINED, "Still eligible inside the area"),
    (added & s8.eq(ST_NOT_IN), C_VALIDATION, "ID not present in Fabric version 8"),
    (added & s8.isin(ST_OUT), C_NOT_EXPLAINED, "Added outside the area"),
    (added & s8.isin([ST_NOT_ELIG, ST_NOT_BSL]), C_NOT_EXPLAINED, "Added although not eligible"),
    (added & s4.eq(ST_NOT_IN), C_CONSISTENT, "ID present only in Fabric version 8"),
    (added & s4.isin(ST_OUT), C_CONSISTENT, "Point moved into the area"),
    (added & s4.isin([ST_NOT_ELIG, ST_NOT_BSL]), C_CONSISTENT, "Not eligible to eligible"),
]
subs["cause_class"] = np.select([c for c, _, _ in cause_rules], [k for _, k, _ in cause_rules], default=C_NOT_EXPLAINED)
subs["cause_reason"] = np.select([c for c, _, _ in cause_rules], [w for _, _, w in cause_rules],
                                 default="Already eligible in Fabric version 4")

# Flags: every issue is recorded; nothing is hidden behind the first one
in8 = subs["in_sub_v8"]
not_bsl = in8 & non_bsl8
flag_cols = {"missing": in8 & s8.eq(ST_NOT_IN), "outside": in8 & s8.isin(ST_OUT),
             "not_eligible": in8 & s8.isin([ST_NOT_ELIG, ST_NOT_ELIG_OUT]), "not_bsl": not_bsl,
             "dropped_eligible": removed & s8.eq(ST_ELIG), "added_unexplained": added & s4.eq(ST_ELIG),
             "conflict": (in8 & s8.eq(ST_CONFLICT)) | ((removed | added) & conflict_any),
             "coordinates": subs["grantee_gap_m"].gt(COORD_GAP_M)}
for f_, m in flag_cols.items():
    subs[f"flag_{f_}"] = m.fillna(False).to_numpy()
subs["needs_follow_up_flag"] = subs[[f"flag_{f_}" for f_ in FOLLOW_UP_FLAGS]].any(axis=1)
subs["data_validation_flag"] = subs[[f"flag_{f_}" for f_ in FLAG_ORDER if FLAGS[f_][4] == K_DV]].any(axis=1)
subs["flag_count"] = subs[[f"flag_{f_}" for f_ in FLAGS]].sum(axis=1)
on = lambda r: [f_ for f_ in FLAG_ORDER if r[f"flag_{f_}"]]
subs["all_flags"] = subs.apply(lambda r: "; ".join(FLAGS[f_][0] for f_ in on(r)), axis=1)
subs["all_recommended_actions"] = subs.apply(lambda r: "; ".join(dict.fromkeys(FLAGS[f_][1] for f_ in on(r))), axis=1)
subs["priority"] = subs.apply(lambda r: FLAGS[on(r)[0]][3] if on(r) else np.nan, axis=1)
subs["priority_reason"] = subs.apply(lambda r: FLAGS[on(r)[0]][0] if on(r) else "", axis=1)
subs["primary_action"] = subs.apply(
    lambda r: FLAGS[on(r)[0]][1] if on(r) else (A_DOC if r["cause_class"] != C_RETAINED else A_NONE), axis=1)
subs["questions"] = subs.apply(lambda r: " ".join(dict.fromkeys(FLAGS[f_][2] for f_ in on(r))), axis=1)

# Three separate review dimensions: list continuity, current compliance, observed Fabric transition
subs["list_continuity"] = chg.map({SC_KEPT: "Retained", SC_DROPPED: "Removed", SC_ADDED: "Added"})
subs["current_status"] = np.where(~in8, "Not on the current list",
                                  np.where(s8.eq(ST_ELIG) & subs["flag_count"].eq(0), "Acceptable", "Review required"))
subs["usable_after_validation"] = in8 & s8.eq(ST_ELIG) & ~subs["needs_follow_up_flag"] & subs["data_validation_flag"]
subs["fabric_transition"] = np.select(
    [conflict_any, neither, s8.eq(ST_NOT_IN), s4.eq(ST_NOT_IN), s4.ne(s8)],
    ["Conflicting source records", "In neither version", "ID present only in Fabric version 4",
     "ID present only in Fabric version 8", "Status or geography changed"], default="No observed change")

# Evidence strength (quality of the evidence, not certainty that a conclusion is right), with its basis
cur_v = np.where(in8, "v8", "v4")
edge = pd.Series(np.inf, index=subs.index)
for key in PROJECTS:
    bnd = POLY_M[key].boundary
    for v in ("v4", "v8"):
        m = subs["project_key"].eq(key) & (cur_v == v)
        a = attrs[v].reindex(lid[m])
        ok = a["lon"].notna().to_numpy()
        if ok.any():
            pts = gpd.GeoSeries(gpd.points_from_xy(a["lon"][ok], a["lat"][ok]), crs=WGS84).to_crs(METRIC_CRS)
            edge.loc[subs.index[m][ok]] = pts.distance(bnd).to_numpy()
subs["distance_to_boundary_m"] = edge.replace(np.inf, np.nan).round(1)
subs["basis_exact_id_both_versions"] = ~s4.eq(ST_NOT_IN) & ~s8.eq(ST_NOT_IN)
subs["basis_exact_id_one_version"] = s4.eq(ST_NOT_IN) ^ s8.eq(ST_NOT_IN)
subs["basis_boundary_sensitive"] = subs["distance_to_boundary_m"].le(EDGE_M)
subs["basis_coordinate_discrepancy"] = subs["flag_coordinates"]
subs["basis_conflicting_records"] = conflict_any
limited = (neither | subs["basis_boundary_sensitive"] | subs["basis_coordinate_discrepancy"]
           | subs["basis_conflicting_records"])
subs["evidence_strength"] = np.where(limited, "Limited",
                                     np.where(subs["basis_exact_id_both_versions"], "Strong", "Moderate"))


def evidence(r: pd.Series) -> str:
    parts = [f"Fabric version 4: {STATUS_PLAIN[r['fabric_status_v4']].lower()}",
             f"Fabric version 8: {STATUS_PLAIN[r['fabric_status_v8']].lower()}",
             "in both lists" if r["in_sub_v4"] and r["in_sub_v8"] else
             ("only in the initial grantee submission" if r["in_sub_v4"] else "only in the 2026 data")]
    if r["in_sub_v8"] and not r["broadband_serviceable_v8"] and r["present_in_fabric_v8"]:
        parts.append("not a broadband serviceable location in version 8")
    if pd.notna(r["distance_to_boundary_m"]) and r["distance_to_boundary_m"] <= EDGE_M:
        parts.append(f"{r['distance_to_boundary_m']:.0f} m from the boundary")
    if pd.notna(r["grantee_gap_m"]) and r["grantee_gap_m"] > COORD_GAP_M:
        parts.append(f"grantee point {r['grantee_gap_m']:,.0f} m from the Fabric point")
    return "; ".join(parts)


subs["evidence"] = subs.apply(evidence, axis=1)

# Every location the follow-up map shows is flagged for follow-up, and nothing else is
assert subs["needs_follow_up_flag"].equals(subs["needs_follow_up"].astype(bool)), "follow-up flags differ from the map"


# ---- per project: turnover, eligibility gains and losses, causes, flags, confidence, spatial findings
def concentration(x: np.ndarray, y: np.ndarray) -> tuple[float, str]:
    """Share of points in the busiest CONCENTRATION_TOP of occupied hexagons, and that cell's centre (lat, lon)."""
    if len(x) < MIN_LOCATIONS:
        return np.nan, ""
    cells = hex_cells(x, y, CONCENTRATION_CELL_M).sort_values("locations", ascending=False)
    n_top = max(1, int(np.ceil(len(cells) * CONCENTRATION_TOP)))
    share = round(cells["locations"].head(n_top).sum() / len(x) * 100, 1)
    top = gpd.GeoSeries([shapely.Point(cells["x"].iat[0], cells["y"].iat[0])], crs=METRIC_CRS).to_crs(WGS84).iloc[0]
    return share, f"{top.y:.4f}, {top.x:.4f}"


def metric_xy(lon: pd.Series, lat: pd.Series) -> tuple[np.ndarray, np.ndarray]:
    ok = lon.notna() & lat.notna()
    g = gpd.GeoSeries(gpd.points_from_xy(lon[ok], lat[ok]), crs=WGS84).to_crs(METRIC_CRS)
    return g.x.to_numpy(), g.y.to_numpy()


def pct_of(n: float, d: float) -> float:
    return round(n / d * 100, 1) if d else np.nan


CAUSE_NAMES = dict(zip(("consistent", "not_explained", "validation", "indeterminate"), CLASSES))
DISPOSITIONS = ("Accept and document", "Request clarification", "Request corrected data", "Further review")


def p95(s: pd.Series) -> float:
    s = s.dropna()
    return round(float(s.quantile(0.95)), 1) if len(s) else np.nan


DETERMINATIONS: dict[str, list[dict]] = {}
for key in PROJECTS:
    k = KEY[key]
    sb = subs[subs["project_key"] == key]
    cur = sb[sb["in_sub_v8"]]
    ch = changes[changes["project_key"] == key]
    f = tmat(key, "Fabric locations in the project area")
    # Scope change
    k["turn_retained"], k["turn_removed"], k["turn_added"] = (int(sb["sub_change"].eq(c).sum())
                                                            for c in (SC_KEPT, SC_DROPPED, SC_ADDED))
    k["turn_net"] = k["turn_added"] - k["turn_removed"]
    k["turn_gross"] = k["turn_added"] + k["turn_removed"]
    k["turn_rate"] = pct_of(k["turn_gross"], k["sub4_v4_listed"])
    # Eligibility change, reconciled (conflicting source records are a separate state, never counted as eligible)
    k["e_stable_eligible"], k["e_stable_not_eligible"] = fsum(f, S_EI, S_EI), fsum(f, S_NI, S_NI)
    k["e_to_not_eligible"], k["e_to_eligible"] = fsum(f, S_EI, S_NI), fsum(f, S_NI, S_EI)
    k["e_removed"], k["e_new"] = fsum(f, S_EI, S_NOT), fsum(f, S_NOT, S_EI)
    k["e_moved_out"], k["e_moved_in"] = fsum(f, S_EI, OUT), fsum(f, OUT, S_EI)
    k["e_to_conflict"], k["e_from_conflict"] = fsum(f, S_EI, S_CONF), fsum(f, S_CONF, S_EI)
    k["e_to_non_bsl"], k["e_from_non_bsl"] = fsum(f, S_EI, S_NB), fsum(f, S_NB, S_EI)
    k["e_gross_loss"] = (k["e_to_not_eligible"] + k["e_removed"] + k["e_moved_out"] + k["e_to_conflict"]
                         + k["e_to_non_bsl"])
    k["e_gross_gain"] = (k["e_to_eligible"] + k["e_new"] + k["e_moved_in"] + k["e_from_conflict"]
                         + k["e_from_non_bsl"])
    k["e_net"] = k["e_gross_gain"] - k["e_gross_loss"]
    assert k["a_v4_elig"] + k["e_net"] == k["a_v8_elig"], f"{key}: eligible totals do not reconcile"
    # Observed explanation of list changes, with the percent of gross turnover in each class
    changed = sb[sb["cause_class"] != C_RETAINED]
    for name, c in CAUSE_NAMES.items():
        k[f"cause_{name}"] = int(changed["cause_class"].eq(c).sum())
        k[f"cause_{name}_rate"] = pct_of(k[f"cause_{name}"], k["turn_gross"])
    # Flags, follow-up, evidence strength
    for fl in FLAGS:
        k[f"flag_{fl}"] = int(sb[f"flag_{fl}"].sum())
    k["follow_up_locations"] = int(sb["needs_follow_up_flag"].sum())
    k["data_validation_locations"] = int(sb["data_validation_flag"].sum())
    for c in ("Strong", "Moderate", "Limited"):
        k[f"evidence_{c.lower()}"] = int(sb["evidence_strength"].eq(c).sum())
    # Current exceptions (the 2026 data against Fabric version 8)
    k["cur_listed"] = len(cur)
    k["cur_not_found"] = int(cur["flag_missing"].sum())
    k["cur_outside"] = int(cur["flag_outside"].sum())
    k["cur_not_eligible"] = int(cur["flag_not_eligible"].sum())
    k["cur_non_bsl"] = int(cur["flag_not_bsl"].sum())
    k["cur_conflict"] = int(cur["fabric_status_v8"].eq(ST_CONFLICT).sum())
    k["cur_coordinates"] = int(cur["flag_coordinates"].sum())
    k["cur_multiple"] = int(cur["flag_count"].ge(2).sum())
    k["cur_acceptable"] = int(cur["current_status"].eq("Acceptable").sum())
    k["cur_review"] = int(cur["current_status"].eq("Review required").sum())
    k["cur_usable_after_validation"] = int(cur["usable_after_validation"].sum())
    # Coverage: eligible locations inside the area that the 2026 data includes, as numerator and denominator
    k["cov_numerator"], k["cov_denominator"] = k["sub8_v8_in_elig"], k["a_v8_elig"]
    k["cov_rate"] = pct_of(k["cov_numerator"], k["cov_denominator"])
    if HAS_UNITS["v8"]:
        ids = cur.loc[cur["fabric_status_v8"].eq(ST_ELIG), "location_id"].astype("int64")
        k["cov_units_numerator"] = int(attrs["v8"]["units"].reindex(ids).fillna(0).sum())
        k["cov_units_denominator"] = int(k.get("a_v8_units_elig", 0))
        k["cov_units_rate"] = pct_of(k["cov_units_numerator"], k["cov_units_denominator"])
    else:
        k["cov_units_numerator"] = k["cov_units_denominator"] = k["cov_units_rate"] = np.nan
    # Spatial findings
    both = ch[ch["in_v4"] & ch["in_v8"]]
    moved = both["shift_m"][both["shift_m"] > 1]
    k["s_points_moved"] = len(moved)
    k["s_displacement_median_m"] = round(float(moved.median()), 1) if len(moved) else 0.0
    k["s_displacement_p95_m"] = p95(moved) if len(moved) else 0.0
    k["s_displacement_max_m"] = round(float(moved.max()), 1) if len(moved) else 0.0
    k["s_moved_into_area"] = fsum(f, OUT, [S_EI, S_NI])
    k["s_moved_out_of_area"] = fsum(f, [S_EI, S_NI], OUT)
    gaps = sb["grantee_gap_m"].dropna()
    k["s_grantee_gap_checked"] = len(gaps)
    k["s_grantee_gap_over"] = int(gaps.gt(COORD_GAP_M).sum())
    k["s_grantee_gap_rate"] = pct_of(k["s_grantee_gap_over"], len(gaps))
    k["s_grantee_gap_median_m"] = round(float(gaps.median()), 1) if len(gaps) else np.nan
    k["s_grantee_gap_p95_m"] = p95(gaps)
    k["s_grantee_gap_max_m"] = round(float(gaps.max()), 1) if len(gaps) else np.nan
    k["s_near_boundary"] = int(sb["distance_to_boundary_m"].le(EDGE_M).sum())
    groups = {"additions": sb[sb["sub_change"].eq(SC_ADDED)], "removals": sb[sb["sub_change"].eq(SC_DROPPED)],
              "follow_up": sb[sb["needs_follow_up_flag"]]}
    for name, g in groups.items():
        k[f"conc_{name}"], k[f"conc_{name}_where"] = concentration(*metric_xy(g["fab_lon"], g["fab_lat"]))
    ec = ch[ch["category"].isin([CAT_ELIG_LOST, CAT_ELIG_GAINED])]
    k["conc_eligibility"], k["conc_eligibility_where"] = concentration(
        *metric_xy(ec["lon_v8"].fillna(ec["lon_v4"]), ec["lat_v8"].fillna(ec["lat_v4"])))
    # Recommended actions, highest priority first
    acts = [FLAGS[fl][1] for fl in FLAG_ORDER if k[f"flag_{fl}"]]
    k["actions"] = (", ".join(dict.fromkeys(acts)) if acts else
                    (A_DOC if k["turn_gross"] or k["e_gross_loss"] or k["e_gross_gain"] else A_NONE))
    # Analyst disposition (suggested; the analyst records the decision)
    validation = (k["cur_not_found"] + k["cur_conflict"] + k["flag_conflict"] + k["flag_coordinates"]
                  + k["cause_validation"])
    if validation:
        k["disposition"] = DISPOSITIONS[2]
        why_disp = "IDs not found, conflicting source records or coordinate discrepancies need corrected data first."
    elif k["follow_up_locations"]:
        k["disposition"] = DISPOSITIONS[1]
        why_disp = "Locations need follow-up with the grantee."
    elif k["cause_indeterminate"]:
        k["disposition"] = DISPOSITIONS[3]
        why_disp = "Some list changes cannot be compared with the Fabric (ID in neither version)."
    else:
        k["disposition"] = DISPOSITIONS[0]
        why_disp = "No follow-up, data validation issue or indeterminate change."
    # Five determinations: source facts, analytical interpretation, suggested value, analyst decision
    qr = qr_term(key)
    exc = join_parts([plural(k["cur_not_found"], "ID") + f" not found in the {T_V8}" if k["cur_not_found"] else "",
                      f"{k['cur_outside']:,} outside the award area" if k["cur_outside"] else "",
                      f"{k['cur_not_eligible']:,} not eligible" if k["cur_not_eligible"] else "",
                      f"{k['cur_non_bsl']:,} not broadband serviceable locations" if k["cur_non_bsl"] else "",
                      f"{k['cur_conflict']:,} with conflicting source records" if k["cur_conflict"] else "",
                      f"{k['cur_coordinates']:,} with a coordinate discrepancy" if k["cur_coordinates"] else ""])
    DETERMINATIONS[key] = [
        {"Determination": "1. Scope change",
         "Source facts": f"{k['sub4_v4_listed']:,} locations in the {T_SUB4}, {k['sub8_v8_listed']:,} in the {qr}: "
                         f"{k['turn_retained']:,} retained, {k['turn_removed']:,} removed, {k['turn_added']:,} added.",
         "Analytical interpretation": (f"Gross turnover {k['turn_gross']:,} ({k['turn_rate']} % of the initial list), "
                                       f"net change {k['turn_net']:+,}." if k["turn_gross"] else "No locations added or removed."),
         "Suggested value": "Scope changed" if k["turn_gross"] else "No scope change"},
        {"Determination": "2. Eligibility change",
         "Source facts": f"Eligible locations inside the area: {k['a_v4_elig']:,} in the {T_V4}, {k['a_v8_elig']:,} in "
                         f"the {T_V8}; gross loss {k['e_gross_loss']:,}, gross gain {k['e_gross_gain']:,}.",
         "Analytical interpretation": (f"Net change {k['e_net']:+,}; the eligibility classification changed from eligible "
                                       f"to not eligible for {k['e_to_not_eligible']:,} locations and from not eligible to "
                                       f"eligible for {k['e_to_eligible']:,}."),
         "Suggested value": "Eligible count changed" if k["e_gross_loss"] or k["e_gross_gain"] else "No eligibility change"},
        {"Determination": "3. Observed explanation",
         "Source facts": f"{k['turn_gross']:,} list changes compared with the observed Fabric transition of each ID.",
         "Analytical interpretation": (join_parts([f"{k[f'cause_{n}']:,} {lower_first(c)} ({k[f'cause_{n}_rate']} %)"
                                                   for n, c in CAUSE_NAMES.items() if k[f"cause_{n}"]]) + "."
                                       if k["turn_gross"] else "No list changes to explain."),
         "Suggested value": ("No list changes" if not k["turn_gross"] else
                             "All consistent with Fabric change" if k["cause_consistent"] == k["turn_gross"] else
                             "Partly not explained by observed Fabric change")},
        {"Determination": "4. Current exceptions",
         "Source facts": f"{k['cur_listed']:,} locations in the {qr} checked against the {T_V8}"
                         + (f": {exc}." if exc else "."),
         "Analytical interpretation": (f"{k['cur_review']:,} require review ({k['cur_multiple']:,} with more than one "
                                       f"issue; {k['cur_usable_after_validation']:,} usable after data validation); "
                                       f"{k['cur_acceptable']:,} acceptable."),
         "Suggested value": "Current exceptions present" if k["cur_review"] else "No current exceptions"},
        {"Determination": "5. Analyst disposition",
         "Source facts": f"{k['follow_up_locations']:,} locations need follow-up; {k['flag_missing']:,} current IDs not "
                         f"found, {k['flag_conflict']:,} with conflicting source records, {k['flag_coordinates']:,} "
                         f"with a coordinate discrepancy, {k['cause_indeterminate']:,} indeterminate list changes.",
         "Analytical interpretation": why_disp + " Suggested from the determinations above; it is not a decision.",
         "Suggested value": k["disposition"]},
    ]
    for d in DETERMINATIONS[key]:
        d.update({"Analyst decision": "", "Analyst notes": ""})


def spatial_sentence(key: str) -> str:
    k, qr = KEY[key], qr_term(key)
    s = (f"Spatially, {k['cur_outside']:,} of the {k['sub8_v8_listed']:,} locations in the {qr} lie outside the project "
         f"area and {plural(k['s_near_boundary'], 'listed location')} sit within {EDGE_M} m of its boundary; "
         f"{plural(k['s_moved_into_area'], 'Fabric point')} moved into the area and {k['s_moved_out_of_area']:,} out of it "
         "(map coordinate edits).")
    conc = [n for n in ("additions", "removals", "eligibility", "follow_up")
            if pd.notna(k[f"conc_{n}"]) and k[f"conc_{n}"] >= CONCENTRATED_SHARE * 100]
    if conc:
        s += " Concentrated: " + ", ".join(n.replace("_", "-") for n in conc) + "."
    return s


summary = summary.drop(columns=[c for c in summary.columns if c in KEY[next(iter(KEY))] and c != "project_key"])
summary = summary.merge(pd.DataFrame.from_dict(KEY, orient="index").rename_axis("project_key").reset_index(),
                        on="project_key", how="left")
summary["spatial_line"] = summary["project_key"].map(spatial_sentence)
summary["bottom_line_full"] = (summary["bottom_line"] + " " + summary["spatial_line"]
                               + " Recommended action: " + summary["actions"].str.lower() + "."
                               + " Suggested disposition: " + summary["disposition"].str.lower() + ".")
determinations = pd.DataFrame([{"project_key": k_, **d} for k_, ds in DETERMINATIONS.items() for d in ds])

# ---- plain labels for the new measures (geodatabase name, label)
FIELD_GUIDE.update({
    "turn_retained": ("list_retained", "Locations retained (in both lists; retained does not mean no action)"),
    "turn_removed": ("list_removed", "Locations removed (not carried into the 2026 data)"),
    "turn_added": ("list_added", "Locations added (new in the 2026 data)"),
    "turn_net": ("list_net_change", "Net change in listed locations (added minus removed)"),
    "turn_gross": ("list_gross_turnover", "Gross turnover (removed plus added)"),
    "turn_rate": ("list_turnover_rate_percent", "Turnover rate, percent (gross turnover divided by initial locations)"),
    "e_stable_eligible": ("eligible_both_versions", "Eligible inside the area in both versions"),
    "e_stable_not_eligible": ("not_eligible_both_versions", "Not eligible inside the area in both versions"),
    "e_to_not_eligible": ("eligible_to_not_eligible", "Eligibility classification changed from eligible to not eligible"),
    "e_to_eligible": ("not_eligible_to_eligible", "Eligibility classification changed from not eligible to eligible"),
    "e_removed": ("eligible_ids_only_in_v4", "Eligible location IDs present only in Fabric version 4"),
    "e_new": ("eligible_ids_only_in_v8", "Eligible location IDs present only in Fabric version 8"),
    "e_moved_out": ("eligible_points_moved_out", "Eligible points moved out of the area"),
    "e_moved_in": ("eligible_points_moved_in", "Eligible points moved into the area"),
    "e_to_conflict": ("eligible_to_conflicting", "Eligible in version 4, conflicting source records in version 8"),
    "e_from_conflict": ("conflicting_to_eligible", "Conflicting source records in version 4, eligible in version 8"),
    "e_to_non_bsl": ("eligible_to_not_serviceable", "Eligible in version 4, not a serviceable location in version 8"),
    "e_from_non_bsl": ("not_serviceable_to_eligible", "Not a serviceable location in version 4, eligible in version 8"),
    "e_gross_loss": ("eligible_gross_loss", "Gross loss of eligible locations"),
    "e_gross_gain": ("eligible_gross_gain", "Gross gain of eligible locations"),
    "e_net": ("eligible_net_change", "Net change in eligible locations (gross gain minus gross loss)"),
    **{f"cause_{n}": (f"changes_{n}", f"List changes – {lower_first(c)}") for n, c in CAUSE_NAMES.items()},
    **{f"cause_{n}_rate": (f"changes_{n}_percent", f"List changes – {lower_first(c)} (percent of gross turnover)")
       for n, c in CAUSE_NAMES.items()},
    **{f"flag_{fl}": (f"flag_{fl}", "Flag – " + FLAGS[fl][0]) for fl in FLAGS},
    "follow_up_locations": ("follow_up_locations", "Locations that need follow-up"),
    "data_validation_locations": ("data_validation_locations", "Locations with a data validation issue"),
    "evidence_strong": ("evidence_strong", "Listed locations with strong evidence"),
    "evidence_moderate": ("evidence_moderate", "Listed locations with moderate evidence"),
    "evidence_limited": ("evidence_limited", "Listed locations with limited evidence"),
    "cur_listed": ("current_listed", "Locations in the 2026 data"),
    "cur_not_found": ("current_not_found_in_v8", "Current location IDs not found in Fabric version 8"),
    "cur_outside": ("current_outside_area", "Current locations outside the award area"),
    "cur_not_eligible": ("current_not_eligible", "Current locations not eligible in Fabric version 8"),
    "cur_non_bsl": ("current_not_serviceable", "Current records that are not broadband serviceable locations"),
    "cur_conflict": ("current_conflicting_records", "Current locations with conflicting source records (unresolved)"),
    "cur_coordinates": ("current_coordinate_discrepancy", f"Current locations with a coordinate discrepancy over {COORD_GAP_M} m"),
    "cur_multiple": ("current_multiple_issues", "Current locations with more than one issue"),
    "cur_acceptable": ("current_acceptable", "Current locations acceptable (eligible inside the area, no flag)"),
    "cur_review": ("current_review_required", "Current locations that require review"),
    "cur_usable_after_validation": ("current_usable_after_validation",
                                    "Current locations usable after data validation (eligible inside, only a data issue)"),
    "cov_numerator": ("coverage_numerator", "Coverage numerator: eligible locations inside the area in the 2026 data"),
    "cov_denominator": ("coverage_denominator", "Coverage denominator: eligible locations inside the area in Fabric version 8"),
    "cov_rate": ("coverage_rate_percent", "Coverage rate, percent (numerator divided by denominator)"),
    "cov_units_numerator": ("coverage_units_numerator", "Coverage by units: units at eligible listed locations inside the area"),
    "cov_units_denominator": ("coverage_units_denominator", "Coverage by units: units at eligible locations inside the area"),
    "cov_units_rate": ("coverage_units_rate_percent", "Coverage rate by units, percent"),
    "s_points_moved": ("points_moved_between_versions", "Fabric points that moved between versions (more than 1 m)"),
    "s_displacement_median_m": ("point_move_median_m", "Median distance a moved Fabric point moved (m)"),
    "s_displacement_p95_m": ("point_move_95th_percentile_m", "95th percentile distance a moved Fabric point moved (m)"),
    "s_displacement_max_m": ("point_move_maximum_m", "Maximum distance a moved Fabric point moved (m)"),
    "s_moved_into_area": ("points_moved_into_area", "Fabric points moved into the area"),
    "s_moved_out_of_area": ("points_moved_out_of_area", "Fabric points moved out of the area"),
    "s_grantee_gap_checked": ("grantee_points_checked", "Grantee coordinates checked against the Fabric point"),
    "s_grantee_gap_over": ("grantee_point_gaps_over_threshold",
                           f"Grantee coordinates more than {COORD_GAP_M} m from the Fabric point"),
    "s_grantee_gap_rate": ("grantee_point_gap_rate_percent", "Coordinate discrepancy rate, percent of coordinates checked"),
    "s_grantee_gap_median_m": ("grantee_point_gap_median_m", "Median distance from grantee coordinates to the Fabric point (m)"),
    "s_grantee_gap_p95_m": ("grantee_point_gap_95th_percentile_m",
                            "95th percentile distance from grantee coordinates to the Fabric point (m)"),
    "s_grantee_gap_max_m": ("grantee_point_gap_maximum_m", "Maximum distance from grantee coordinates to the Fabric point (m)"),
    "s_near_boundary": ("listed_near_boundary", f"Listed locations within {EDGE_M} m of the boundary"),
    **{f"conc_{n}": (f"concentration_{n}_percent", f"Concentration of {lab}: percent in the busiest 10 % of "
                     f"{CONCENTRATION_CELL_M} m cells") for n, lab in
       (("additions", "additions"), ("removals", "removals"), ("eligibility", "eligibility changes"),
        ("follow_up", "follow-up locations"))},
    **{f"conc_{n}_where": (f"concentration_{n}_centre", f"Busiest cell for {lab} (latitude, longitude)") for n, lab in
       (("additions", "additions"), ("removals", "removals"), ("eligibility", "eligibility changes"),
        ("follow_up", "follow-up locations"))},
    "actions": ("recommended_action", "Recommended action"),
    "disposition": ("suggested_disposition", "Suggested analyst disposition (the analyst records the decision)"),
    "spatial_line": ("spatial_findings", "Spatial findings"),
    "bottom_line_full": ("bottom_line_full", "Bottom line (scope, eligibility, spatial change, action and disposition)"),
})
assert all(len(f_) <= 64 for f_, _ in FIELD_GUIDE.values())
assert len({f_ for f_, _ in FIELD_GUIDE.values()}) == len(FIELD_GUIDE)
field_guide = pd.DataFrame([{"source_column": c, "field_name": f_, "label": lab} for c, (f_, lab) in FIELD_GUIDE.items()])
display(summary[["project_key", "turn_gross", "turn_rate", "e_gross_loss", "e_gross_gain", "follow_up_locations",
                 "data_validation_locations", "cur_review", "disposition"]])

# %% [markdown]
# ## 6d. All projects at a glance
# One row per project: eligible locations and the main reasons they changed, the grantee lists with turnover, coverage, follow-up, change pattern, recommended action and a suggested next step.

# %%
FLOWS = [(S_EI, S_NI, -1, "eligible to not eligible"), (S_EI, S_NOT, -1, "IDs present only in version 4"),
         (S_EI, OUT, -1, "points moved out of the area"), (S_EI, S_CONF, -1, "conflicting source records"),
         (S_NI, S_EI, +1, "not eligible to eligible"), (S_NOT, S_EI, +1, "IDs present only in version 8"),
         (OUT, S_EI, +1, "points moved into the area"), (S_CONF, S_EI, +1, "conflict in version 4 only")]
FU_SHORT = {"B01": "not in Fabric version 8", "B13": "outside the award area", "B02": "not eligible in Fabric version 8",
            "B14": "not a broadband serviceable location", "B03": "eligible, not carried over",
            "B04": "new, already eligible in version 4", "B12": "conflicting source records"}
assert set(FOLLOW_UP_CLASSES) <= set(FU_SHORT), "every follow-up category needs a short label"


def main_reasons(key: str, top: int = 3) -> str:
    """Largest flows behind the change in eligible locations, e.g. '−14 eligible to not eligible'."""
    f = tmat(key, "Fabric locations in the project area")
    flows = [(sign * fsum(f, a, b), words) for a, b, sign, words in FLOWS if fsum(f, a, b)]
    flows.sort(key=lambda t: -abs(t[0]))
    return ", ".join(f"{n:+,} {w}".replace("-", "−") for n, w in flows[:top]) or "No change"


def next_step(key: str) -> str:
    """One or two plain questions for the grantee, from the change patterns and the key numbers."""
    k, r = KEY[key], summary.set_index("project_key").loc[key]
    n8 = k["sub8_v8_listed"]
    bad8 = k["sub8_v8_in_nelig"] + k["sub8_v8_out_elig"] + k["sub8_v8_out_nelig"]
    steps = []
    if bad8 >= MIN_LOCATIONS and bad8 >= 0.5 * n8:
        steps.append(f"Ask the grantee to clarify why {bad8:,} of the {n8:,} planned locations are not eligible "
                     "inside the area.")
    if P_FEWER in r["patterns"]:
        steps.append(f"Confirm the reduced scope ({k['sub4_v4_listed'] - n8:,} fewer planned locations; "
                     f"{int(plain_b.at[key, 'B03']):,} still-eligible locations dropped).")
    if P_MORE in r["patterns"]:
        out8 = k["sub8_v8_out_elig"] + k["sub8_v8_out_nelig"]
        steps.append(f"Confirm the added locations ({out8:,} lie outside the area).")
    if P_SERVICE in r["patterns"]:
        served = fsum(tmat(key, cap(T_SUB4)), S_EI, S_NI)
        steps.append(f"Ask how the {served:,} listed locations now classified as not eligible affect the plan.")
    if P_IDS in r["patterns"]:
        steps.append(f"Ask the grantee to update {k['sub4_v8_notfound']:,} IDs to Fabric version 8.")
    if P_DQ in r["patterns"] and len(steps) < 2:
        steps.append("Request corrected data (duplicates, invalid or missing IDs, conflicting records, coordinates).")
    return " ".join(steps[:2]) or "No action needed; note the review in the project file."


def md_pct(v) -> str:
    return "–" if pd.isna(v) else f"{v:.0f} %"


def arrow(a: int, b: int) -> str:
    return f"{a:,} → {b:,} ({b - a:+,})".replace("(-", "(−")


GLANCE_COLS = ["Project", "Eligible locations in the project area: Fabric version 4 → version 8",
               "Main reasons for the change in eligible locations",
               "Locations listed: initial grantee submission → 2026 data (gross turnover, turnover rate)",
               "Listed and eligible inside the area: initial grantee submission (version 4) → 2026 data (version 8)",
               "2026 locations that are not eligible or lie outside the area (version 8)",
               "Location IDs not found in the Fabric: initial grantee submission / 2026 data",
               "Eligible locations inside the area that the 2026 data includes (version 8)",
               "Needs follow-up (largest reason)", "Change pattern", "Recommended action", "Suggested next step"]
glance_rows = []
for key in PROJECTS:
    k, r = KEY[key], summary.set_index("project_key").loc[key]
    fu = {c: int(plain_b.at[key, c]) for c in FOLLOW_UP_CLASSES}
    big_fu = max(fu, key=fu.get)
    bad8 = k["sub8_v8_in_nelig"] + k["sub8_v8_out_elig"] + k["sub8_v8_out_nelig"]
    glance_rows.append({
        "Project": f"{r['label']} ({r['pr_number']})",
        GLANCE_COLS[1]: arrow(k["a_v4_elig"], k["a_v8_elig"]),
        GLANCE_COLS[2]: main_reasons(key),
        GLANCE_COLS[3]: arrow(k["sub4_v4_listed"], k["sub8_v8_listed"])
                        + f"; turnover {k['turn_gross']:,} ({md_pct(k['turn_rate'])})",
        GLANCE_COLS[4]: f"{k['sub4_v4_in_elig']:,} → {k['sub8_v8_in_elig']:,}",
        GLANCE_COLS[5]: f"{bad8:,} of {k['sub8_v8_listed']:,}",
        GLANCE_COLS[6]: f"{k['sub4_v4_notfound']:,} / {k['sub8_v8_notfound']:,}",
        GLANCE_COLS[7]: f"{k['d_v8_listed']:,} of {k['a_v8_elig']:,}",
        GLANCE_COLS[8]: f"{int(r['fu_total']):,}" + (f" ({fu[big_fu]:,} {FU_SHORT.get(big_fu, lower_first(SUB_CLASSES[big_fu][0]))})" if r["fu_total"] else ""),
        GLANCE_COLS[9]: r["patterns"],
        GLANCE_COLS[10]: k["actions"],
        GLANCE_COLS[11]: next_step(key)})
glance = pd.DataFrame(glance_rows, columns=GLANCE_COLS)
GLANCE_CSV = OUTPUT_DIR / "FFA_V4_V8_all_projects_at_a_glance.csv"
if WRITE_CSV:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    glance.to_csv(GLANCE_CSV, index=False, encoding="utf-8-sig")
display(glance)

# %% [markdown]
# ## 7. Write outputs
# * **File geodatabase** (`OUTPUT_GDB`)
#   * `project_areas` – six project polygons with every summary field → **Map Series index layer** (page name field: `page_title`; summary text fields `story_a`, `story_b`)
#   * `eligibility_v4_points`, `eligibility_v8_points` – Fabric locations inside the project areas, per version
#   * `submitted_locations` – every V4/V8 submitted location with Comparison 3 category, Fabric status in V4 and V8 and explanation
#   * `fabric_change_points` – every project-area Fabric location whose category is not *Unchanged* (V8 position, else V4)
#   * `map_A_eligibility_changes`, `map_B_list_changes`, `map_C_follow_up`, `page_extents`, `map_legend` – plain-labelled layers and the map series index for ArcGIS Pro (`build_map_series.py`)
#   * tables: `summary_projects`, `cmp1_2_submission_vs_fabric`, `cmp3_submission_changes`, `cmp4_fabric_changes`, `synthesis_explanations`, `synthesis_coverage`
#   * `data_cleaning`, `duplicate_ids` – rows skipped and duplicates removed per list and project
# * **Summary CSV** (`SUMMARY_CSV`)
#
# Requires GDAL ≥ 3.6 (OpenFileGDB write). 64-bit integer fields need ArcGIS Pro 3.2+; on older Pro set `GDB_OPTS = {}`.

# %%
GDB_OPTS = {"TARGET_ARCGIS_VERSION": "ARCGIS_PRO_3_2_OR_LATER"}

OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
if OUTPUT_GDB.exists():
    if OUTPUT_GDB.suffix.lower() != ".gdb":
        raise ValueError(f"Refusing to delete {OUTPUT_GDB}: not a .gdb folder")
    shutil.rmtree(OUTPUT_GDB)


def gdb_ready(df: pd.DataFrame) -> pd.DataFrame:
    """Casts columns to types the FileGDB driver accepts and trims long text."""
    out = df.copy()
    for c in out.columns:
        if c == "geometry":
            continue
        if out[c].dtype == bool or str(out[c].dtype) == "boolean":
            out[c] = out[c].astype("Int16")
        elif str(out[c].dtype) == "category":
            out[c] = out[c].astype(str)
        elif out[c].dtype == object or str(out[c].dtype).startswith(("string", "str")):
            out[c] = out[c].astype("string").str.slice(0, 2000)
    return out


def write(df: pd.DataFrame, layer: str, gdb: Path | None = None, quiet: bool = False) -> None:
    """Writes one layer or table to a file geodatabase (the unified results gdb by default)."""
    pyogrio.write_dataframe(gdb_ready(df), gdb or OUTPUT_GDB, layer=layer, driver="OpenFileGDB", **GDB_OPTS)
    if not quiet:
        log.info("wrote %-30s %8s rows", layer, f"{len(df):,}")


# Polygon parts only (the geodatabase rejects GeometryCollections)
areas["geometry"] = [polygonal_2d(g) for g in areas.geometry]

# Every output layer and table is built once here, keyed by its geodatabase name. It is written in full to the
# unified results gdb, and filtered by project_key into each project's own gdb (Section 9).
GDB_LAYERS: dict[str, pd.DataFrame] = {}
GDB_LAYERS["project_areas"] = areas[["project_key", "geometry"]].merge(summary, on="project_key")

for v in ("v4", "v8"):
    pts = fabric_pts[v].merge(in_area[v], on="location_id")
    pts = pts[["project_key", "location_id", "eligible", "bsl", "counted", "elig_raw", "geometry"]]
    GDB_LAYERS[f"eligibility_{v}_points"] = gpd.GeoDataFrame(pts, geometry="geometry", crs=WGS84)

sl = subs.copy()
sl["lon"] = sl["sub_lon_v8"].fillna(sl["sub_lon_v4"])
sl["lat"] = sl["sub_lat_v8"].fillna(sl["sub_lat_v4"])
GDB_LAYERS["submitted_locations"] = points_from_lonlat(sl)

chg = changes[changes["category"] != CAT_SAME].copy()
chg["lon"] = chg["lon_v8"].fillna(chg["lon_v4"])
chg["lat"] = chg["lat_v8"].fillna(chg["lat_v4"])
GDB_LAYERS["fabric_change_points"] = points_from_lonlat(chg.drop(columns=["address_v4", "address_v8"]))


# Map layers for the ArcGIS Pro map series (plain fields only); page_title drives each layer's Page Query
TITLE = summary.set_index("project_key")["page_title"]
ma = changes[changes["map_show"] & ~changes["map_class"].isin(SHADED_CLASSES)].copy()
ma["page_title"] = ma["project_key"].map(TITLE)
ma["lon"] = ma["lon_v8"].fillna(ma["lon_v4"])
ma["lat"] = ma["lat_v8"].fillna(ma["lat_v4"])
ma = ma.rename(columns={"elig_v4": "eligible_fabric_v4", "elig_v8": "eligible_fabric_v8"})
GDB_LAYERS["map_A_eligibility_changes"] = points_from_lonlat(ma[[
    "project_key", "page_title", "project", "location_id", "map_class", "map_label", "tier", "what_it_means",
    "detail", "eligible_fabric_v4", "eligible_fabric_v8", "shift_m", "lon", "lat"]])

mb = subs[subs["map_show"]].copy()
mb["page_title"] = mb["project_key"].map(TITLE)
mb["lon"], mb["lat"] = mb["fab_lon"], mb["fab_lat"]
mb = mb.rename(columns={"in_sub_v4": "in_initial_submission", "in_sub_v8": "in_2026_q2_q3_data",
                        "status_v4_plain": "status_fabric_v4", "status_v8_plain": "status_fabric_v8",
    "broadband_serviceable_v4": "serviceable_location_fabric_v4", "broadband_serviceable_v8": "serviceable_location_fabric_v8",
    "conflicting_records_v4": "conflicting_records_fabric_v4", "conflicting_records_v8": "conflicting_records_fabric_v8",
    "cause_class": "observed_explanation", "cause_reason": "observed_explanation_reason",
    "questions": "questions_for_grantee"})
cols_b = ["project_key", "page_title", "project", "location_id", "map_class", "map_label", "tier", "what_it_means", "why",
          "detail", "needs_follow_up", "in_initial_submission", "in_2026_q2_q3_data", "status_fabric_v4", "status_fabric_v8",
          "point_source", "lon", "lat"]
GDB_LAYERS["map_B_list_changes"] = points_from_lonlat(mb[cols_b])
mc = mb.copy()
mc["map_class"] = mc["map_class"].where(mc["needs_follow_up"], "C0")
mc["map_label"] = mc["map_class"].map(lambda c: FOLLOW_CLASSES[c][0])
GDB_LAYERS["map_C_follow_up"] = points_from_lonlat(mc[cols_b])
GDB_LAYERS["map_A_no_change_density"] = density.to_crs(WGS84)
GDB_LAYERS["key_numbers_from_to"] = transitions
GDB_LAYERS["field_guide"] = field_guide
GDB_LAYERS["page_extents"] = page_extents
GDB_LAYERS["map_legend"] = legend                      # no project_key: copied whole into every project gdb
GDB_LAYERS["data_cleaning"] = cleaning
GDB_LAYERS["duplicate_ids"] = dup_detail

GDB_LAYERS["summary_projects"] = summary
GDB_LAYERS["cmp1_2_submission_vs_fabric"] = cmp12
GDB_LAYERS["cmp3_submission_changes"] = cmp3.reset_index().rename(columns=lambda c: re.sub(r"\W+", "_", c).strip("_").lower())
GDB_LAYERS["cmp4_fabric_changes"] = cmp4
GDB_LAYERS["synthesis_explanations"] = synthesis
GDB_LAYERS["synthesis_coverage"] = coverage

for name, df in GDB_LAYERS.items():
    write(df, name)
if WRITE_CSV:
    readable(summary.drop(columns="geometry", errors="ignore")).to_csv(SUMMARY_CSV, index=False, encoding="utf-8-sig")
    log.info("summary CSV → %s", SUMMARY_CSV)

# %% [markdown]
# ## 7b. Plain-language Excel workbook and HTML report (all six projects)
# * **`FFA_V4_V8_Results_Plain.xlsx`** – plain tab and column names, one row per project, category counts with colour swatches, the follow-up list and every location.
# * **`FFA_V4_V8_Results_Explained.html`** – a readable report: what was compared, results at a glance, a story per project with the same colours as the maps.

# %%
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
import html as html_lib

XLSX_PATH = OUTPUT_DIR / "FFA_V4_V8_Results_Plain.xlsx"
HTML_PATH = OUTPUT_DIR / "FFA_V4_V8_Results_Explained.html"
RUN_DATE = pd.Timestamp.today().strftime("%B %d, %Y")
V4_FIELD_NAME, V8_FIELD_NAME = V4_FIELDS["elig"], V8_FIELDS["elig"]
ALL_CLASSES = {**FABRIC_CLASSES, **SUB_CLASSES, **FOLLOW_CLASSES}
COLOUR = {k: v[2] for k, v in ALL_CLASSES.items()}
hexc = lambda rgb: "{:02X}{:02X}{:02X}".format(*rgb)
yes_no = lambda s: s.map({True: "Yes", False: "No"})

GLOSSARY = [
    ("Fabric", "The FCC's national list of every place where broadband could be installed (homes, businesses), each with an ID and a map point."),
    (cap(T_V4), "Fabric version 4 with FFA eligibility (field FFAv4Dec23): 0 = unserved (eligible), 1 = served. The version used when projects applied."),
    (cap(T_V8), "Fabric version 8 with FFA eligibility (field status_code): 0 = unserved (eligible), 1 = served. The current version."),
    ("Location ID", "The FCC's number for one location. Between Fabric versions an ID can be removed and new IDs added."),
    ("Eligible", "Coded 0 (unserved) in the eligibility field: no cable or fiber at 25/3 Mbps or better reported."),
    ("Not eligible", "Coded 1 (served) in the eligibility field. The code alone does not say why it changed."),
    ("Conflicting source records", "The copies of a location ID in one Fabric file disagree on eligibility; the ID is "
                                   "left out of every count, never resolved in favour of eligibility."),
    (cap(T_SUB4), "The locations the grantee listed when applying (All_448_PIDs file), using Fabric version 4 IDs."),
    (T_SUB8, "The locations in the grantee's 2026 quarterly report, Q3 or Q2 where Q3 is missing, using Fabric version 8 IDs."),
    (TIER_FU, f"A listed location not explained by the observed Fabric change: in the {T_SUB8} but missing from the {T_V8}, outside the area, not eligible or not a serviceable location; not carried over although still eligible; new although already eligible in the {T_V4}; or conflicting source records."),
    (TIER_EX, f"A change that matches an observed difference between the {T_V4} and the {T_V8}: eligibility classification changed, ID present in only one version, or point moved. Association, not cause."),
]

# ---- tables with plain column names (shared with the per-project workbooks)
PLAIN_COLS = {"project": "Project", "location_id": "Location ID", "map_label": "Category", "tier": "Tier",
              "what_it_means": "What it means", "why": "Why", "detail": "Detail",
              "elig_v4": f"Eligible – {T_V4}", "elig_v8": f"Eligible – {T_V8}",
              "in_sub_v4": f"In {T_SUB4}", "in_sub_v8": f"In {T_SUB8}",
              "status_v4_plain": f"Status – {T_V4}", "status_v8_plain": f"Status – {T_V8}"}


def plain_summary_table(sm: pd.DataFrame) -> pd.DataFrame:
    """One row per project with plain column names."""
    return pd.DataFrame({
        "Project": sm["label"], "PR number": sm["pr_number"], "Bottom line": sm["bottom_line"],
        "Change pattern": sm["patterns"], "How the change came together": sm["story_why"],
        f"Locations in area – {T_V4}": sm["v4_total"], f"Locations in area – {T_V8}": sm["v8_total"],
        f"Eligible in area – {T_V4}": sm["v4_elig"], f"Eligible in area – {T_V8}": sm["v8_elig"],
        "Change in eligible": sm["elig_chg"], "Change in eligible (%)": sm["elig_pct"],
        f"Locations in {T_SUB4}": sm["sub4_n"], f"Locations in {T_SUB8}": sm["sub8_n"],
        f"Of those in {T_SUB8}, eligible in {T_V8}": sm["sub8_elig"],
        "In both submissions (same ID)": sm["s_kept"],
        f"Not carried into {T_SUB8}": sm["s_dropped"], f"New in {T_SUB8}": sm["s_added"],
        "Needs follow-up (total)": sm["fu_total"],
        **{f"Follow-up: {SUB_CLASSES[c][0]}": sm[f"fu_{c.lower()}"] for c in FOLLOW_UP_CLASSES},
        "2026 data file used": sm["sub_v8_file"],
        f"Rows removed as duplicates – {T_SUB4}": sm["sub4_dup_rows"],
        f"Rows removed as duplicates – {T_SUB8}": sm["sub8_dup_rows"],
        "Eligibility summary": sm["story_a"], "Grantee submission summary": sm["story_b"],
    })


def location_tables(ch: pd.DataFrame, sb: pd.DataFrame) -> dict[str, pd.DataFrame]:
    """Follow-up list, eligibility changes and listed locations with plain columns (map_class kept for swatches)."""
    fu = sb[sb["needs_follow_up"] & sb["map_show"]].sort_values(["project", "map_class", "location_id"])
    fu = fu[["project", "location_id", "map_class", "map_label", "why", "in_sub_v4", "in_sub_v8",
             "status_v4_plain", "status_v8_plain", "detail"]].assign(
        in_sub_v4=lambda d: yes_no(d["in_sub_v4"]), in_sub_v8=lambda d: yes_no(d["in_sub_v8"]))
    el = ch[~ch["map_class"].isin(NO_CHANGE_CLASSES)].sort_values(["project", "map_class", "location_id"])
    el = el[["project", "location_id", "map_class", "map_label", "what_it_means", "detail", "elig_v4", "elig_v8"]].assign(
        elig_v4=lambda d: yes_no(d["elig_v4"]), elig_v8=lambda d: yes_no(d["elig_v8"]))
    ls = sb.sort_values(["project", "map_class", "location_id"])
    ls = ls[["project", "location_id", "map_class", "map_label", "tier", "why", "in_sub_v4", "in_sub_v8",
             "status_v4_plain", "status_v8_plain", "detail"]].assign(
        in_sub_v4=lambda d: yes_no(d["in_sub_v4"]), in_sub_v8=lambda d: yes_no(d["in_sub_v8"]))
    counts = []
    for page, df, classes in (("Eligibility changes", ch[ch["map_show"]], FABRIC_CLASSES),
                              ("Grantee submission changes", sb[sb["map_show"]], SUB_CLASSES)):
        g = df.groupby(["project", "map_class"]).size().rename("Locations").reset_index()
        g["map_label"] = g["map_class"].map(lambda c: classes[c][0])
        g["tier"] = g["map_class"].map(lambda c: classes[c][3])
        g.insert(1, "Map page", page)
        counts.append(g)
    ct = pd.concat(counts, ignore_index=True)[["project", "Map page", "map_class", "map_label", "tier", "Locations"]]
    return {"Follow-up list": fu.rename(columns=PLAIN_COLS), "Category counts": ct.rename(columns=PLAIN_COLS),
            "Eligibility changes": el.rename(columns=PLAIN_COLS), "Submitted locations": ls.rename(columns=PLAIN_COLS)}


def cleaning_table(cl: pd.DataFrame) -> pd.DataFrame:
    """Data cleaning counts with plain column names."""
    cl = cl.reindex(columns=list(cl.columns) + [c for c in ("blank_id_rows", "invalid_id_rows") if c not in cl.columns])
    cl = cl.assign(blank_id_rows=cl["blank_id_rows"].fillna(0).astype(int),
                   invalid_id_rows=cl["invalid_id_rows"].fillna(cl["non_numeric_rows_skipped"]).astype(int))
    return cl[["project", "dataset", "rows_read", "blank_id_rows", "invalid_id_rows", "non_numeric_rows_skipped",
               "ids_listed_more_than_once", "duplicate_rows_removed", "unique_ids_used", "note"]].rename(columns={
        "project": "Project", "dataset": "List", "rows_read": "Rows read",
        "blank_id_rows": "Rows with no location ID (skipped)",
        "invalid_id_rows": "Rows with an invalid location ID (skipped)",
        "non_numeric_rows_skipped": "Rows skipped in total",
        "ids_listed_more_than_once": "Location IDs listed more than once", "duplicate_rows_removed": "Duplicate rows removed",
        "unique_ids_used": "Locations counted", "note": "Note"})


WRAP_SHEETS = {"At a glance", "Key numbers", "Column guide"}
WRAP_COLS = {"What to ask the grantee", "What the category means", "Why the list changed", "Meaning", "Measure",
             "Bottom line", "How the change came together", "Next step", "Main reasons", "Change pattern", "Pattern notes", "Eligibility summary", "Grantee submission summary", "What it means", "Why", "Detail", "Note"}


def write_workbook(path: Path, readme: list[tuple[str, str]], sheets: dict[str, pd.DataFrame]) -> None:
    """Writes a formatted workbook: read-me tab, navy headers, colour swatch on Category cells, filters."""
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        pd.DataFrame(readme, columns=["A", "B"]).to_excel(xw, sheet_name="Read me", index=False, header=False)
        ws = xw.sheets["Read me"]
        ws["A1"].font = Font(bold=True, size=14)
        for row in ws.iter_rows():
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
            if row[0].value and not row[1].value or row[0].value in ("Tab", "Key term"):
                row[0].font = Font(bold=True)
        ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 34, 110
        head_fill = PatternFill("solid", fgColor="1F3864")
        for name, df in sheets.items():
            out = df.drop(columns=["map_class"], errors="ignore")
            wrap_all = name in WRAP_SHEETS
            out.to_excel(xw, sheet_name=name[:31], index=False)
            ws = xw.sheets[name[:31]]
            cols = list(out.columns)
            for j, col in enumerate(cols, start=1):
                cell = ws.cell(row=1, column=j)
                cell.font, cell.fill = Font(bold=True, color="FFFFFF"), head_fill
                cell.alignment = Alignment(wrap_text=True, vertical="center")
                body = out[col].astype(str).str.len()
                wrap = col in WRAP_COLS or (wrap_all and j > 1)
                width = (70 if col in WRAP_COLS and col.endswith(("summary", "line")) else
                         60 if col == "Measure" else 40 if wrap else
                         min(42, max(12, len(str(col)) + 2, int(body.quantile(0.9)) + 2 if len(body) else 12)))
                ws.column_dimensions[get_column_letter(j)].width = width
                if wrap:
                    for i in range(2, len(out) + 2):
                        ws.cell(row=i, column=j).alignment = Alignment(wrap_text=True, vertical="top")
            if "map_class" in df.columns and "Category" in cols:
                j = cols.index("Category") + 1
                for i, code in enumerate(df["map_class"], start=2):
                    rgb = COLOUR.get(code)
                    if rgb:
                        ws.cell(row=i, column=j).fill = PatternFill("solid", fgColor=hexc(rgb))
                        if sum(rgb) < 380:
                            ws.cell(row=i, column=j).font = Font(color="FFFFFF")
            ws.freeze_panes = "B2"
            ws.auto_filter.ref = ws.dimensions


tables = location_tables(changes, subs)
readme_all = ([(f"FFA project locations: {T_V4} compared with {T_V8}", ""), (f"Run on {RUN_DATE}", ""), ("", ""),
               ("What this compares", f"For six FFA projects: how eligibility changed inside each project area between the "
                f"{T_V4} and the {T_V8}, and how the grantee's locations changed between the {T_SUB4} and the {T_SUB8}."),
               ("Eligibility fields", f"{cap(T_V4)}: {V4_FIELD_NAME}. {cap(T_V8)}: {V8_FIELD_NAME}. Both coded 1 = served, 0 = unserved (eligible)."), ("", ""),
               ("Tab", "What it holds"),
               ("At a glance", "One row per project: eligible locations and why they changed, the grantee lists, coverage, follow-up, change pattern and a suggested next step."),
               ("Project summary", "One row per project: bottom line, totals, follow-up counts, plain-English summaries."),
               ("Key numbers", "Every key-number table of every project in one long table."),
               ("Follow-up list", "Every listed location that needs follow-up, with the reason."),
               ("Category counts", "How many locations fall in each category, per project and map page."),
               ("Eligibility changes", f"Every Fabric location in a project area whose eligibility, ID or position changed between the {T_V4} and the {T_V8}."),
               ("Submitted locations", f"Every location in the {T_SUB4} or the {T_SUB8}, with its category and reason."),
               ("Data cleaning", "Per source and project: rows read, rows skipped, duplicates removed, locations counted."),
               ("Duplicate IDs", "Every location ID that appeared more than once, and how often."),
               ("", ""), ("Key term", "Meaning")] + GLOSSARY)
def all_key_numbers() -> pd.DataFrame:
    """Every key-number table of every project in one long table: project, table, measure, column, value."""
    rows = []
    for key in PROJECTS:
        for name, df in key_tables(key).items():
            first = df.columns[0]
            for _, row in df.iterrows():
                for col in df.columns[1:]:
                    rows.append({"Project": LABEL[key], "Table": name, "Measure": row[first], "Column": col,
                                 "Value": row[col]})
    return pd.DataFrame(rows)



RESULTS_XLSX = OUTPUT_DIR / "FFA_Fabric_version_4_to_8_results_all_projects.xlsx"

COLUMN_MEANINGS: dict[str, str] = {}
KEY_SECTIONS = [
    ("Fabric locations inside the project area",
     ["a_v4_loc", "a_v8_loc", "total_chg", "a_v4_elig", "a_v8_elig", "elig_chg", "a_v4_nelig", "a_v8_nelig",
      "a_v4_bsl", "a_v8_bsl", "a_v4_nonbsl", "a_v8_nonbsl", "a_v4_units", "a_v8_units", "a_v4_units_elig",
      "a_v8_units_elig", "f_elig_lost", "f_elig_gain", "f_retired", "f_new", "f_moved"]),
    *[(title, [f"{lst}_{v}_{m}" for m in ("listed", "matched", "notfound", "in_elig", "in_nelig", "out_elig",
                                           "out_nelig", "out_near", "out_mid", "out_far", "near_edge", "bsl",
                                           "nonbsl", "units", "coord_gap")])
      for title, lst, v in (("Initial grantee submission checked against Fabric version 4", "sub4", "v4"),
                            ("2026 Q2 and Q3 data checked against Fabric version 8", "sub8", "v8"),
                            ("Check: initial grantee submission looked up in Fabric version 8", "sub4", "v8"),
                            ("Check: 2026 Q2 and Q3 data looked up in Fabric version 4", "sub8", "v4"))],
    ("How the grantee's list changed", ["sub_chg", "s_kept", "s_dropped", "s_added"]),
    ("Coverage of locations inside the area",
     [f"d_{v}_{m}" for v in ("v4", "v8") for m in ("listed", "unlisted", "nelig_listed", "nelig_unlisted")]),
    ("Needs follow-up", [f"fu_{c.lower()}" for c in FOLLOW_UP_CLASSES] + ["fu_total"]),
    ("Data cleaning", ["sub4_dup_rows", "sub4_bad_rows", "sub8_dup_rows", "sub8_bad_rows", "f4_dup_rows", "f8_dup_rows"]),
    ("Summary", ["patterns", "pattern_notes", "bottom_line", "story_why", "story_a", "story_b", "story_c",
                 "sub_v8_file"]),
]
ASK = {"B01": "Ask for the Fabric version 8 location ID of this location in the 2026 Q2 and Q3 data.",
       "B13": "Ask why a location outside the awarded area is in the 2026 Q2 and Q3 data.",
       "B02": "Ask how a location that the Fabric version 8 eligibility data classifies as not eligible qualifies.",
       "B14": "Ask whether this location is a broadband serviceable location.",
       "B03": "Ask why an eligible location in the initial grantee submission is not in the 2026 data.",
       "B04": "Ask why this already-eligible location was not in the initial grantee submission.",
       "B12": "Request corrected data: the Fabric copies of this location ID disagree on eligibility."}
assert set(FOLLOW_UP_CLASSES) <= set(ASK), "every follow-up category needs a question"
YES_NO_COLS = {"in_sub_v4", "in_sub_v8", "elig_v4", "elig_v8"}
COLUMN_HEADS = {"project": "Project", "location_id": "Location ID", "map_label": "Category", "tier": "Tier",
                "what_it_means": "What the category means", "why": "Why the list changed", "detail": "Detail",
                "elig_v4": "Eligible in Fabric version 4", "elig_v8": "Eligible in Fabric version 8",
                "in_sub_v4": "In the initial grantee submission", "in_sub_v8": "In the 2026 Q2 and Q3 data",
                "status_v4_plain": "Status in Fabric version 4", "status_v8_plain": "Status in Fabric version 8",
                "ask": "What to ask the grantee"}


def key_numbers_by_project(keys: list[str]) -> pd.DataFrame:
    """Measures down the side, one column per project."""
    sm = summary.set_index("project_key")
    rows = []
    for section, cols in KEY_SECTIONS:
        for c in cols:
            if c in sm.columns:
                rows.append({"Section": section, "Measure": FIELD_GUIDE.get(c, (c, c))[1],
                             **{LABEL[k]: sm.at[k, c] for k in keys}})
    return pd.DataFrame(rows)


def stacked(keys: list[str], name: str) -> pd.DataFrame:
    """One key-number table for several projects, stacked with a Project column."""
    parts = []
    for k in keys:
        df = key_tables(k)[name].copy()
        df.columns = [c.replace(qr_term(k), T_SUB8) for c in df.columns]      # same headers for every project
        df.insert(0, "Project", LABEL[k])
        parts.append(df)
    out = pd.concat(parts, ignore_index=True)
    return out.rename(columns=lambda c: c.replace(" · ", ", checked against Fabric ")
                      .replace("Version 4", "Fabric version 4").replace("Version 8", "Fabric version 8"))


def from_to(keys: list[str]) -> pd.DataFrame:
    """Status in Fabric version 4 (rows) against status in Fabric version 8 (columns), with totals."""
    parts = []
    for k in keys:
        for pop in ("Fabric locations in the project area", cap(T_SUB4), T_SUB8, "Added in the 2026 data"):
            m = tmat(k, pop)
            m["Total"] = m.sum(axis=1)
            m.loc["Total"] = m.sum(axis=0)
            m = m.reset_index().rename(columns={"from_fabric_v4": "Status in Fabric version 4"})
            m.columns = ["Status in Fabric version 4"] + [f"Now (Fabric version 8): {c}" if c != "Total" else "Total"
                                                         for c in m.columns[1:]]
            m.insert(0, "Locations", pop)
            m.insert(0, "Project", LABEL[k])
            parts.append(m)
    return pd.concat(parts, ignore_index=True).fillna(0)


def location_sheets(keys: list[str]) -> dict[str, pd.DataFrame]:
    """Follow-up list, eligibility changes, submitted locations and category counts for the given projects."""
    ch = changes[changes["project_key"].isin(keys)]
    sb = subs[subs["project_key"].isin(keys)].assign(ask=lambda d: d["map_class"].map(ASK).fillna(""))
    yn = lambda d: d.assign(**{c: yes_no(d[c]) for c in YES_NO_COLS if c in d.columns})
    fu = sb[sb["needs_follow_up"]].sort_values(["project", "map_class", "location_id"])
    fu = yn(fu[["project", "location_id", "map_class", "map_label", "ask", "in_sub_v4", "in_sub_v8",
                "status_v4_plain", "status_v8_plain", "detail"]])
    el = ch[~ch["map_class"].isin(NO_CHANGE_CLASSES)].sort_values(["project", "map_class", "location_id"])
    el = yn(el[["project", "location_id", "map_class", "map_label", "what_it_means", "elig_v4", "elig_v8", "detail"]])
    ls = sb.sort_values(["project", "map_class", "location_id"])
    ls = yn(ls[["project", "location_id", "map_class", "map_label", "tier", "why", "in_sub_v4", "in_sub_v8",
                "status_v4_plain", "status_v8_plain", "detail"]])
    counts = []
    for page, df, classes in (("Eligibility changes in the project area", ch, FABRIC_CLASSES),
                              ("Grantee submission changes", sb, SUB_CLASSES)):
        g = df.groupby(["project", "map_class"]).size().rename("Locations").reset_index()
        g["map_label"] = g["map_class"].map(lambda c: classes[c][0])
        g["tier"] = g["map_class"].map(lambda c: classes[c][3])
        g.insert(1, "Map", page)
        counts.append(g)
    ct = pd.concat(counts, ignore_index=True)[["project", "Map", "map_class", "map_label", "tier", "Locations"]]
    return {"Follow-up list": fu.rename(columns=COLUMN_HEADS),
            "Eligibility changes": el.rename(columns=COLUMN_HEADS),
            "Submitted locations": ls.rename(columns=COLUMN_HEADS),
            "Category counts": ct.rename(columns=COLUMN_HEADS)}


def results_readme(keys: list[str]) -> list[tuple[str, str]]:
    one = len(keys) == 1
    title = LABEL[keys[0]] + f" ({summary.set_index('project_key').at[keys[0], 'pr_number']})" if one else "All six projects"
    return [(f"{title}: Fabric version 4 to version 8 review", ""), (f"Run on {RUN_DATE}", ""), ("", ""),
            ("What this compares", "How eligibility changed inside each Federal Funding Account (FFA) project area "
             f"between the {T_V4} and the {T_V8}, and how the grantee's locations changed between the {T_SUB4} and "
             f"the {T_SUB8} (the grantee's planned list)."),
            ("How locations are placed", "Every listed location is placed at the Fabric point of its location ID; "
             "the grantee's own coordinates are only used to check them. Lists are matched to the Fabric by location "
             "ID only."),
            ("Eligibility fields", f"{cap(T_V4)}: {V4_FIELD_NAME}. {cap(T_V8)}: {V8_FIELD_NAME}. Both coded "
             "1 = served, 0 = unserved (eligible)."), ("", ""),
            ("Tab", "What it holds"),
            ("At a glance", "One row per project: eligible locations and why they changed, the grantee lists, "
             "coverage, follow-up, change pattern and a suggested next step."),
            ("Key numbers", "Every measure for " + ("this project" if one else "each project, one column per project")
             + ", grouped in sections."),
            ("Fabric in the project area", "Fabric locations inside the project area: version 4, version 8, change."),
            ("Grantee lists and the Fabric", "Each grantee list checked against its own Fabric version and, as a "
             "check, against the other version."),
            ("Coverage", "Eligible and not eligible locations inside the area: listed by the grantee or not."),
            ("From version 4 to version 8", "Where locations went: status in Fabric version 4 against status in "
             "Fabric version 8, with totals."),
            ("Determinations", "Five determinations per project (scope change, eligibility change, observed "
             "explanation, current exceptions, analyst disposition): source facts, interpretation, suggested value and "
             "a blank analyst decision."),
            ("Follow-up list", "One row per flag, highest priority first: location ID, issue, evidence, observed "
             "explanation, recommended action, evidence strength and the question to ask the grantee."),
            ("Change drivers", "Counts and percentages by reason: why locations were removed or added, and why eligible "
             "locations were lost or gained."),
            ("Review rules", "Every analyst-defined review threshold with its count, denominator and whether it was "
             "triggered."),
            ("Spatial findings", "Boundary, point movement, coordinate gaps and where changes concentrate."),
            ("Eligibility changes", "Every Fabric location in the area whose eligibility, ID or position changed."),
            ("Submitted locations", f"Every location in the {T_SUB4} or the {T_SUB8}, with its category, observed explanation, "
             "source facts per Fabric version, list continuity, current status, flags, priority, actions, evidence strength "
             "and blank analyst columns."),
            ("Category counts", "How many locations fall in each map category."),
            ("Data cleaning", "Rows read, rows skipped, duplicates removed and locations counted, per list."),
            ("Duplicate location IDs", "Every location ID that appeared more than once, and how often."),
            ("Column guide", "Every column header in this workbook and what it means."),
            ("", ""), ("Key term", "Meaning")] + GLOSSARY


def column_guide(sheets: dict[str, pd.DataFrame]) -> pd.DataFrame:
    """Every column header with its meaning (key-number measures are explained by their own wording)."""
    meaning = {"Project": "Project name.", "Location ID": "The FCC's number for one location.",
               "Category": "What happened to the location; the cell colour matches the maps.",
               "Tier": "Needs follow-up, Consistent with Fabric change, or No change.",
               "Section": "Group of related measures.", "Measure": "What is counted, written out in full.",
               "Locations": "Number of locations, or which group of locations the rows describe.",
               "Total": "Row or column total.", "Map": "The map page the category belongs to.",
               "Change": "Fabric version 8 minus Fabric version 4 (or 2026 data minus initial grantee submission).",
               "Status in Fabric version 4": "Where the location stood in Fabric version 4: eligible inside the area, "
                                             "not eligible inside, outside the area, or not in the Fabric.",
               }
    for c in GLOSSARY:
        meaning.setdefault(c[0], c[1])
    rows = []
    for name, df in sheets.items():
        for col in df.columns:
            if col == "map_class":
                continue
            text = meaning.get(col, COLUMN_MEANINGS.get(col, ""))
            if not text and col.startswith("Now (Fabric version 8): "):
                text = f"Locations whose status in Fabric version 8 is: {col.split(': ', 1)[1].lower()}."
            if not text and col in LABEL.values():
                text = "Value of the measure for this project."
            rows.append({"Tab": name, "Column": col, "Meaning": text or col + "."})
    return pd.DataFrame(rows)


def result_sheets(keys: list[str]) -> dict[str, pd.DataFrame]:
    """Every results tab for the given projects (all six, or one)."""
    sheets = {"At a glance": glance[glance["Project"].str.startswith(tuple(LABEL[k] for k in keys))],
              "Key numbers": key_numbers_by_project(keys),
              "Fabric in the project area": stacked(keys, "Fabric in the area"),
              "Grantee lists and the Fabric": stacked(keys, "Grantee lists"),
              "Coverage": stacked(keys, "Coverage"),
              "From version 4 to version 8": from_to(keys),
              **location_sheets(keys),
              "Data cleaning": cleaning_table(cleaning[cleaning["project_key"].isin(keys)
                                                       | (cleaning["project_key"].isna() & (len(keys) > 1))]),
              "Duplicate location IDs": dup_detail[dup_detail["project_key"].isin(keys)][
                  ["project", "dataset", "location_id", "times_listed", "eligibility_conflict"]].assign(
                  eligibility_conflict=lambda d: d["eligibility_conflict"].map({True: "Yes", False: "No"}).fillna(""))
                  .rename(columns={"project": "Project", "dataset": "List", "location_id": "Location ID",
                                   "times_listed": "Times listed",
                                   "eligibility_conflict": "Eligible in one copy, not eligible in another"})}
    sheets["Column guide"] = column_guide(sheets)
    return sheets


def results_workbook(path: Path, keys: list[str]) -> None:
    """The results workbook for the given projects (all six, or one)."""
    write_workbook(path, results_readme(keys), result_sheets(keys))

KEY_SECTIONS[1:1] = [
    ("Eligibility change in detail", ["e_stable_eligible", "e_stable_not_eligible", "e_to_not_eligible", "e_to_eligible",
                                      "e_removed", "e_new", "e_moved_out", "e_moved_in", "e_to_conflict",
                                      "e_from_conflict", "e_to_non_bsl", "e_from_non_bsl", "e_gross_loss",
                                      "e_gross_gain", "e_net"])]
KEY_SECTIONS.extend([
    ("Scope change", ["turn_retained", "turn_removed", "turn_added", "turn_net", "turn_gross", "turn_rate"]),
    ("Observed explanation of the list changes", [f"cause_{n}{r}" for n in CAUSE_NAMES for r in ("", "_rate")]),
    ("Current exceptions (2026 data against Fabric version 8)",
     ["cur_listed", "cur_not_found", "cur_outside", "cur_not_eligible", "cur_non_bsl", "cur_conflict",
      "cur_coordinates", "cur_multiple", "cur_review", "cur_usable_after_validation", "cur_acceptable"]),
    ("Coverage rate", ["cov_numerator", "cov_denominator", "cov_rate", "cov_units_numerator",
                       "cov_units_denominator", "cov_units_rate"]),
    ("Flags and action", [f"flag_{fl}" for fl in FLAG_ORDER] + ["follow_up_locations", "data_validation_locations",
                                                                "actions"]),
    ("Evidence strength", ["evidence_strong", "evidence_moderate", "evidence_limited"]),
    ("Spatial findings", ["s_points_moved", "s_displacement_median_m", "s_displacement_p95_m", "s_displacement_max_m",
                          "s_moved_into_area", "s_moved_out_of_area", "s_near_boundary", "s_grantee_gap_checked",
                          "s_grantee_gap_over", "s_grantee_gap_rate", "s_grantee_gap_median_m",
                          "s_grantee_gap_p95_m", "s_grantee_gap_max_m"]
     + [f"conc_{n}{w}" for n in ("additions", "removals", "eligibility", "follow_up") for w in ("", "_where")]),
    ("Analyst disposition", ["disposition"]),
    ("Bottom line", ["bottom_line_full"]),
])


def change_drivers(keys: list[str]) -> pd.DataFrame:
    """Counts and percentages by reason: list changes by observed explanation, eligibility changes by flow."""
    rows = []
    for key in keys:
        k = KEY[key]
        sb = subs[(subs["project_key"] == key) & (subs["cause_class"] != C_RETAINED)]
        for (cls, why, chg_), n in sb.groupby(["cause_class", "cause_reason", "sub_change"]).size().items():
            rows.append({"Project": LABEL[key], "What changed": "Grantee list: " + (
                "removed" if chg_ == SC_DROPPED else "added"), "Observed explanation": cls, "Reason": why,
                "Locations": int(n), "Percent of gross turnover or of the gross eligible change": pct_of(n, k["turn_gross"])})
        for name, kind in (("e_to_not_eligible", "loss"), ("e_removed", "loss"), ("e_moved_out", "loss"),
                           ("e_to_conflict", "loss"), ("e_to_non_bsl", "loss"), ("e_to_eligible", "gain"),
                           ("e_new", "gain"), ("e_moved_in", "gain"), ("e_from_conflict", "gain"),
                           ("e_from_non_bsl", "gain")):
            if k[name]:
                rows.append({"Project": LABEL[key], "What changed": f"Eligible locations in the area: gross {kind}",
                             "Observed explanation": "Observed Fabric transition", "Reason": FIELD_GUIDE[name][1],
                             "Locations": k[name],
                             "Percent of gross turnover or of the gross eligible change": pct_of(k[name], k[f"e_gross_{kind}"])})
    return pd.DataFrame(rows)


def spatial_findings(keys: list[str]) -> pd.DataFrame:
    cols = KEY_SECTIONS[[t for t, _ in KEY_SECTIONS].index("Spatial findings")][1]
    return pd.DataFrame([{"Project": LABEL[k], "Finding": FIELD_GUIDE[c][1], "Value": KEY[k][c]}
                         for k in keys for c in cols] + [{"Project": LABEL[k], "Finding": "In words",
                                                         "Value": summary.set_index("project_key").at[k, "spatial_line"]}
                                                        for k in keys])


def follow_up_table(keys: list[str]) -> pd.DataFrame:
    """One row per flag: location ID, priority, issue, evidence, action, evidence strength and the question to ask."""
    sb = subs[subs["project_key"].isin(keys)]
    rows = []
    for _, r in sb.iterrows():
        for fl in FLAG_ORDER:
            if r[f"flag_{fl}"]:
                rows.append({"Project": r["project"], "Location ID": int(r["location_id"]),
                             "Priority": FLAGS[fl][3], "Issue": FLAGS[fl][0], "Kind": FLAGS[fl][4],
                             "Evidence": r["evidence"], "List continuity": r["list_continuity"],
                             "Observed explanation": r["cause_class"], "Reason": r["cause_reason"],
                             "Recommended action": FLAGS[fl][1], "Evidence strength": r["evidence_strength"],
                             "What to ask the grantee": FLAGS[fl][2]})
    cols = ["Project", "Location ID", "Priority", "Issue", "Kind", "Evidence", "List continuity",
            "Observed explanation", "Reason", "Recommended action", "Evidence strength", "What to ask the grantee"]
    return pd.DataFrame(rows, columns=cols).sort_values(["Project", "Priority", "Issue", "Location ID"], ignore_index=True)


LOCATION_FACTS = {   # source facts first, then interpretations, then the analyst's columns
    "project": "Project", "location_id": "Location ID",
    "in_sub_v4": "In the initial grantee submission", "in_sub_v8": "In the 2026 Q2 and Q3 data",
    "present_in_fabric_v4": "Present in Fabric version 4", "inside_award_area_v4": "Inside the award area in Fabric version 4",
    "eligible_v4": "Eligible in Fabric version 4", "broadband_serviceable_v4": "Broadband serviceable location in Fabric version 4",
    "conflicting_records_v4": "Conflicting source records in Fabric version 4",
    "present_in_fabric_v8": "Present in Fabric version 8", "inside_award_area_v8": "Inside the award area in Fabric version 8",
    "eligible_v8": "Eligible in Fabric version 8", "broadband_serviceable_v8": "Broadband serviceable location in Fabric version 8",
    "conflicting_records_v8": "Conflicting source records in Fabric version 8",
    "distance_to_boundary_m": "Distance to the boundary (m)", "grantee_gap_m": "Grantee point to Fabric point (m)",
    "list_continuity": "List continuity", "fabric_transition": "Observed Fabric transition",
    "current_status": "Current status", "map_label": "Category", "cause_class": "Observed explanation",
    "cause_reason": "Reason", "all_flags": "All flags", "flag_count": "Number of flags", "priority": "Priority",
    "priority_reason": "Priority reason", "primary_action": "Primary action",
    "all_recommended_actions": "All recommended actions", "evidence_strength": "Evidence strength",
    "basis_exact_id_both_versions": "Basis: exact ID in both versions",
    "basis_exact_id_one_version": "Basis: exact ID in one version only",
    "basis_boundary_sensitive": "Basis: boundary-sensitive", "basis_coordinate_discrepancy": "Basis: coordinate discrepancy",
    "basis_conflicting_records": "Basis: conflicting source records", "evidence": "Evidence",
}
FACT_YES_NO = [c for c in LOCATION_FACTS if c.startswith(("in_sub_", "present_", "inside_", "eligible_", "broadband_",
                                                           "conflicting_", "basis_"))]


def listed_locations(keys: list[str]) -> pd.DataFrame:
    sb = subs[subs["project_key"].isin(keys)].sort_values(["project", "priority", "map_class", "location_id"])
    out = sb[["map_class", *LOCATION_FACTS]].copy()
    for c in FACT_YES_NO:
        out[c] = out[c].fillna(False).astype(bool).map({True: "Yes", False: "No"})
    out["priority"] = out["priority"].astype("Int64")
    out["Analyst decision"], out["Analyst notes"] = "", ""
    return out.rename(columns=LOCATION_FACTS)


def determinations_table(keys: list[str]) -> pd.DataFrame:
    d = determinations[determinations["project_key"].isin(keys)].copy()
    d.insert(0, "Project", d.pop("project_key").map(LABEL))
    return d


def review_rules_table(keys: list[str]) -> pd.DataFrame:
    r = review_rules[review_rules["project_key"].isin(keys)].copy()
    r.insert(0, "Project", r.pop("project_key").map(LABEL))
    return r


_base_location_sheets = location_sheets


def location_sheets(keys: list[str]) -> dict[str, pd.DataFrame]:
    """Determinations, follow-up list (one row per flag), drivers, review rules and every listed location."""
    base = _base_location_sheets(keys)
    return {"Determinations": determinations_table(keys), "Follow-up list": follow_up_table(keys),
            "Change drivers": change_drivers(keys), "Review rules": review_rules_table(keys),
            "Spatial findings": spatial_findings(keys), "Eligibility changes": base["Eligibility changes"],
            "Submitted locations": listed_locations(keys), "Category counts": base["Category counts"]}


WRAP_COLS.update({"Issue", "Evidence", "Reason", "Recommended action", "Observed explanation", "Finding", "Value",
                  "All flags", "All recommended actions", "Source facts", "Analytical interpretation", "Priority reason",
                  "Review rule", "Analyst notes"})
WRAP_SHEETS.add("Determinations")
COLUMN_MEANINGS.update({
    "Issue": "What the flag is about.",
    "Kind": "Needs follow-up (also on the follow-up map) or Data validation.",
    "Priority": "1 = look at first: 1 ID not found in Fabric version 8, 2 outside the award area, 3 not eligible, "
                "4 not a broadband serviceable location, 5 removed but still eligible, 6 added without an observed "
                "Fabric transition, 7 conflicting source records or coordinate discrepancy.",
    "Priority reason": "The highest-priority flag on the location.",
    "Evidence": "The source facts behind the flags: status in both Fabric versions, which lists hold the location, "
                "boundary distance and coordinate gap where relevant.",
    "Observed explanation": "Analytical interpretation of an added or removed location: Consistent with Fabric change, "
                            "Not explained by observed Fabric change, Data validation issue, or Indeterminate. "
                            "Consistent means the observed Fabric transition fits the list change; it does not prove "
                            "the cause.",
    "Reason": "The observed transition behind the explanation.",
    "Recommended action": "Verify eligibility, Verify boundary, Confirm scope change, Correct data, Document change, "
                          "or No action.",
    "Primary action": "The action for the highest-priority flag (Document change or No action when there is no flag).",
    "All recommended actions": "Every action that applies to the location.",
    "All flags": "Every flag on the location, highest priority first.", "Number of flags": "How many flags apply.",
    "Evidence strength": "Strength of the evidence (not certainty of a conclusion). Strong: exact ID in both Fabric "
                         "versions, no limiting factor. Moderate: exact ID in one version only. Limited: ID in neither "
                         "version, boundary-sensitive, coordinate discrepancy, or conflicting source records.",
    "List continuity": "Retained (in both lists), Removed, or Added. Retained does not mean no action.",
    "Observed Fabric transition": "What changed for the ID between Fabric version 4 and version 8.",
    "Current status": "Acceptable (in the 2026 data, eligible inside the area, no flag), Review required, or "
                      "Not on the current list.",
    "What changed": "Which count changed: the grantee list or the eligible locations in the area.",
    "Percent of gross turnover or of the gross eligible change": "Share of the gross turnover (removed plus added), "
                                                                  "or of the gross loss or gross gain of eligible locations.",
    "Finding": "Spatial measure, written out.", "Value": "Value of the finding.",
    "Distance to the boundary (m)": "Distance from the Fabric point to the project boundary.",
    "Grantee point to Fabric point (m)": "Distance from the grantee's coordinates to the Fabric point.",
    "Determination": "One of five determinations: scope change, eligibility change, observed explanation, current "
                     "exceptions, analyst disposition.",
    "Source facts": "Counts taken directly from the grantee lists and the Fabric.",
    "Analytical interpretation": "What the analysis reads from the facts.",
    "Suggested value": "The analysis's suggestion. For the disposition: Accept and document, Request clarification, "
                       "Request corrected data, or Further review.",
    "Analyst decision": "Left blank for the analyst's decision.", "Analyst notes": "Left blank for the analyst.",
    "Pattern": "The change pattern the review rule belongs to.", "Review rule": "What the rule counts.",
    "Count": "Locations or rows counted.", "Denominator": "What the count is divided by.",
    "Share (percent)": "Count divided by the denominator.",
    "Threshold": "Analyst-defined review threshold (configurable in the notebook settings).",
    "Triggered": "Yes if the count reached the threshold.",
    **{v: f"Yes or No: {v[0].lower() + v[1:]}." for c, v in LOCATION_FACTS.items() if c in FACT_YES_NO},
    "Rows with no location ID (skipped)": "Rows where the location ID cell was empty.",
    "Rows with an invalid location ID (skipped)": "Rows where the location ID was not a whole number.",
    "Rows skipped in total": "Rows with no or an invalid location ID.",
    GLANCE_COLS[1]: "Eligible (unserved) Fabric locations inside the project area in each version, and the change.",
    GLANCE_COLS[2]: "The largest flows behind the change in eligible locations, from the From version 4 to version 8 tab.",
    GLANCE_COLS[3]: "Locations in each grantee list, the net change, gross turnover (removed plus added) and turnover "
                    "rate (turnover divided by the initial list; can pass 100 %).",
    GLANCE_COLS[4]: "Listed locations that are eligible and inside the area, each list checked against its own Fabric version.",
    GLANCE_COLS[5]: "Locations in the 2026 data that are not eligible or lie outside the area in Fabric version 8.",
    GLANCE_COLS[6]: "Listed location IDs that are not in the matching Fabric version.",
    GLANCE_COLS[7]: "How many of the eligible locations inside the area the 2026 data includes.",
    GLANCE_COLS[8]: "Locations that need follow-up, with the most common reason.",
    "Change pattern": "One-line classification of the project's change (rules in the Read me and the notebook).",
    GLANCE_COLS[11]: "One or two questions for the grantee, generated from the numbers; review before sending.",
    "Coverage of locations inside the area": "Eligible and not eligible locations inside the area, and whether the "
                                             "grantee lists them.",
    "Fabric version 4": "Count in the Fabric version 4 eligibility data (with the initial grantee submission where "
                        "lists are involved).",
    "Fabric version 8": "Count in the Fabric version 8 eligibility data (with the 2026 data where lists are involved).",
    "Fabric locations inside the project area": "What is counted inside the project area.",
    "Grantee lists checked against the Fabric": "What is counted for each grantee list.",
    "Initial grantee submission, checked against Fabric version 4": "The initial grantee submission against its own Fabric version.",
    "2026 Q2 and Q3 data, checked against Fabric version 8": "The 2026 data against its own Fabric version.",
    "Initial grantee submission, checked against Fabric version 8": "Check: the initial grantee submission looked up in Fabric version 8.",
    "2026 Q2 and Q3 data, checked against Fabric version 4": "Check: the 2026 data looked up in Fabric version 4.",
    "List": "Which list or Fabric file the row describes.", "Rows read": "Rows in the source before cleaning.",
    "Location IDs listed more than once": "IDs that appeared on more than one row.",
    "Duplicate rows removed": "Extra rows removed so each location counts once.",
    "Locations counted": "Unique locations used in the analysis.", "Note": "How the source was cleaned.",
    "Times listed": "How many rows carried this ID.",
    "Eligible in one copy, not eligible in another": "Yes when the copies disagree on eligibility; the ID is then marked "
                                                     "as conflicting source records and left out of the eligibility counts.",
    "What the category means": "Plain explanation of the category.", "Detail": "Extra detail, such as how far a point moved.",
    "Eligible in Fabric version 4": "Yes if eligible (unserved) in Fabric version 4.",
    "Eligible in Fabric version 8": "Yes if eligible (unserved) in Fabric version 8.",
    "What to ask the grantee": "The question to put to the grantee for this flag.",
    "In the initial grantee submission": "Yes if the location is in the initial grantee submission.",
    "In the 2026 Q2 and Q3 data": "Yes if the location is in the 2026 data.",
    "Status in Fabric version 8": "Where the location stands in Fabric version 8: eligible inside the area, not "
                                  "eligible inside, outside the area, or not in the Fabric.",
})


# ---- Markdown and CSV result files
def md_escape(v) -> str:
    if isinstance(v, float):
        return "–" if np.isnan(v) else (f"{v:,.1f}" if v != int(v) else f"{int(v):,}")
    if isinstance(v, (int, np.integer)):
        return f"{int(v):,}"
    return str(v).replace("|", "/").replace("\n", " ")


def md_table(df: pd.DataFrame) -> str:
    if df.empty:
        return "_None._\n"
    head = "| " + " | ".join(map(str, df.columns)) + " |\n|" + "---|" * len(df.columns) + "\n"
    return head + "".join("| " + " | ".join(md_escape(v) for v in row) + " |\n" for row in df.itertuples(index=False))


def project_markdown(key: str, full_follow_up: bool) -> str:
    r, k = summary.set_index("project_key").loc[key], KEY[key]
    g = lambda c: md_escape(k.get(c, r.get(c, np.nan)))
    metrics = pd.DataFrame([
        ("Locations listed: initial grantee submission → 2026 data", f"{g('sub4_v4_listed')} → {g('sub8_v8_listed')}"),
        ("Retained / removed / added", f"{g('turn_retained')} / {g('turn_removed')} / {g('turn_added')}"),
        ("Net change / gross turnover / turnover rate", f"{g('turn_net')} / {g('turn_gross')} / {g('turn_rate')} %"),
        ("Eligible locations in the area: version 4 → version 8", f"{g('a_v4_elig')} → {g('a_v8_elig')}"),
        ("Eligibility gross loss / gross gain / net", f"{g('e_gross_loss')} / {g('e_gross_gain')} / {g('e_net')}"),
        *[(f"List changes: {lower_first(c)}", f"{g(f'cause_{n}')} ({g(f'cause_{n}_rate')} % of gross turnover)")
          for n, c in CAUSE_NAMES.items()],
        ("Current: not found in version 8 / outside / not eligible / not serviceable / conflicting",
         f"{g('cur_not_found')} / {g('cur_outside')} / {g('cur_not_eligible')} / {g('cur_non_bsl')} / {g('cur_conflict')}"),
        ("Current: review required / multiple issues / usable after validation / acceptable",
         f"{g('cur_review')} / {g('cur_multiple')} / {g('cur_usable_after_validation')} / {g('cur_acceptable')}"),
        ("Coverage rate (numerator / denominator)", f"{g('cov_rate')} % ({g('cov_numerator')} / {g('cov_denominator')})"),
        ("Coverage rate by units (numerator / denominator)",
         f"{g('cov_units_rate')} % ({g('cov_units_numerator')} / {g('cov_units_denominator')})"),
        ("Needs follow-up / data validation", f"{g('follow_up_locations')} / {g('data_validation_locations')}"),
        ("Evidence strength strong / moderate / limited",
         f"{g('evidence_strong')} / {g('evidence_moderate')} / {g('evidence_limited')}"),
    ], columns=["Measure", "Value"])
    fu = follow_up_table([key])
    rules = review_rules_table([key]).drop(columns=["Project", "Kind"])
    out = [f"## {r['label']} ({r['pr_number']})\n",
           f"**Bottom line.** {r['bottom_line_full']}\n",
           f"**Change pattern:** {r['patterns']}  \n**Recommended action:** {r['actions']}  \n"
           f"**Suggested disposition:** {r['disposition']} (analyst decision recorded separately)\n",
           "### Determinations\n", md_table(determinations_table([key]).drop(columns=["Project", "Analyst decision",
                                                                                        "Analyst notes"])),
           "### Key metrics\n", md_table(metrics),
           "### Change drivers\n", md_table(change_drivers([key]).drop(columns="Project")),
           "### Review rules (analyst-defined review thresholds)\n", md_table(rules),
           "### Spatial findings\n", r["spatial_line"] + "\n",
           md_table(spatial_findings([key]).drop(columns="Project").iloc[:-1]),
           "### How the change came together\n", r["story_why"] + "\n",
           f"### Follow-up ({len(fu):,} flags)\n"]
    if full_follow_up:
        out.append(md_table(fu.drop(columns="Project")[["Location ID", "Priority", "Issue", "Evidence",
                                                         "Recommended action", "Evidence strength",
                                                         "What to ask the grantee"]]))
    else:
        out.append(md_table(fu.groupby(["Priority", "Issue", "Recommended action"]).size().rename("Locations")
                            .reset_index()))
        out.append(f"Location-level list: `projects/{key}_{r['pr_number']}/{key}_{r['pr_number']}_results.md`.\n")
    return "\n".join(out)


def write_markdown(path: Path, keys: list[str]) -> None:
    one = len(keys) == 1
    head = [f"# FFA project locations: Fabric version 4 to version 8 review{'' if one else ' – all projects'}\n",
            f"Run on {RUN_DATE}. {cap(T_V4)} (field {V4_FIELD_NAME}) and {T_V8} (field {V8_FIELD_NAME}), "
            f"0 = unserved (eligible), 1 = served. The {T_SUB8} is the grantee's planned list. Locations are matched "
            "by location ID and placed at the Fabric point of that ID; closeness is never used to treat two IDs as one "
            "building. Moved coordinates are map coordinate edits. Source facts, analytical interpretations and the "
            "analyst's decisions are kept apart: the explanations say whether a change is consistent with the observed "
            "Fabric change, not what caused it.\n"]
    if not one:
        head += ["## At a glance\n", md_table(glance)]
    body = [project_markdown(k, full_follow_up=one) for k in keys]
    path.write_text("\n".join(head + body), encoding="utf-8")


def write_csvs(folder: Path, keys: list[str]) -> None:
    folder.mkdir(parents=True, exist_ok=True)
    for name, df in result_sheets(keys).items():
        df.drop(columns=["map_class"], errors="ignore").to_csv(
            folder / (re.sub(r"[^A-Za-z0-9]+", "_", name).strip("_").lower() + ".csv"), index=False, encoding="utf-8-sig")


RESULTS_MD = OUTPUT_DIR / "FFA_Fabric_version_4_to_8_results_all_projects.md"

results_workbook(RESULTS_XLSX, list(PROJECTS))
log.info("Results workbook → %s", RESULTS_XLSX)
if WRITE_MARKDOWN:
    write_markdown(RESULTS_MD, list(PROJECTS))
    log.info("Markdown report → %s", RESULTS_MD)
if WRITE_CSV:
    write_csvs(OUTPUT_DIR / "csv", list(PROJECTS))
    log.info("CSV files → %s", OUTPUT_DIR / "csv")


# ---- HTML report
esc = html_lib.escape
CSS = """
body{font-family:Segoe UI,Calibri,Arial,sans-serif;color:#1d1d1f;background:#fff;max-width:1060px;margin:32px auto;padding:0 20px;line-height:1.5}
h1{font-size:26px;margin:0 0 4px}h2{font-size:20px;margin:36px 0 10px;border-bottom:2px solid #1F3864;padding-bottom:4px}
h3{font-size:17px;margin:0 0 8px}.muted{color:#5f6368;font-size:14px}
table{border-collapse:collapse;width:100%;font-size:13.5px;margin:8px 0 16px}th,td{border:1px solid #d9d9d9;padding:6px 8px;text-align:left;vertical-align:top}
th{background:#1F3864;color:#fff;font-weight:600}td.num{text-align:right;font-variant-numeric:tabular-nums}
.card{border:1px solid #d9d9d9;border-radius:8px;padding:16px 18px;margin:18px 0}
.grid{display:grid;grid-template-columns:1fr 1fr;gap:18px}@media(max-width:800px){.grid{grid-template-columns:1fr}}
.story{background:#f5f7fb;border-left:4px solid #1F3864;padding:10px 12px;font-size:14.5px;margin-bottom:10px}
.bottom{background:#fff8e6;border-left:4px solid #E6AB02;padding:10px 12px;font-size:15px;margin:6px 0 14px}
.sw{display:inline-block;width:12px;height:12px;border-radius:50%;margin-right:6px;vertical-align:-1px;border:1px solid #888}
.caution{color:#a33}.tier{font-weight:600;background:#f0f0f0}
"""


def sw(code: str) -> str:
    return f'<span class="sw" style="background:rgb{COLOUR[code]}"></span>'


def class_table(counts: pd.Series, classes: dict) -> str:
    rows = ""
    for tier in TIER_ORDER:
        items = [c for c in classes if classes[c][3] == tier and counts.get(c, 0)]
        if items:
            rows += f"<tr><td class='tier' colspan='2'>{esc(tier)}</td></tr>"
            rows += "".join(f"<tr><td>{sw(c)}{esc(classes[c][0])}</td><td class='num'>{int(counts[c]):,}</td></tr>"
                            for c in items)
    return f"<table><tr><th>Category</th><th>Locations</th></tr>{rows}</table>"


def signed(n) -> str:
    return f"{int(n):+,}"


clean_rows = "".join(
    f"<tr><td>{esc(r['project'])}</td><td>{esc(r['dataset'])}</td><td class='num'>{r['rows_read']:,}</td>"
    f"<td class='num'>{r['non_numeric_rows_skipped']:,}</td><td class='num'>{r['ids_listed_more_than_once']:,}</td>"
    f"<td class='num'>{r['duplicate_rows_removed']:,}</td><td class='num'>{r['unique_ids_used']:,}</td></tr>"
    for _, r in cleaning.iterrows())
glance_html = "".join("<tr>" + "".join(f"<td>{esc(str(v))}</td>" for v in row) + "</tr>"
                      for row in glance.itertuples(index=False))
cards = "".join(
    f"<div class='card'><h3>{esc(r['page_title'])}</h3><div class='bottom'>{esc(r['bottom_line'])}</div><div class='grid'>"
    f"<div><b>Eligibility changes ({esc(T_V4)} → {esc(T_V8)})</b><div class='story'>{esc(r['story_a'])}</div>"
    f"{class_table(plain_a.loc[r['project_key']], FABRIC_CLASSES)}</div>"
    f"<div><b>Grantee submission changes ({esc(T_SUB4)} → {esc(qr_term(r['project_key']))})</b><div class='story'>"
    f"{esc(r['story_b']).replace('Caution:', '<span class=caution><b>Caution:</b></span>')}</div>"
    f"{class_table(plain_b.loc[r['project_key']], SUB_CLASSES)}</div></div></div>" for _, r in summary.iterrows())
legend_rows = lambda classes: "".join(f"<tr><td>{sw(c)}{esc(v[0])}</td><td>{esc(v[3])}</td><td>{esc(v[1])}</td></tr>"
                                      for c, v in classes.items())
gloss = "".join(f"<tr><td><b>{esc(t)}</b></td><td>{esc(m)}</td></tr>" for t, m in GLOSSARY)

report = f"""<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>FFA Fabric Version 4 to 8 Results</title><style>{CSS}</style></head><body>
<h1>FFA project locations: what changed from the {T_V4} to the {T_V8}</h1>
<p class="muted">Six Federal Funding Account projects · run on {RUN_DATE} · {cap(T_V4)}: field {esc(V4_FIELD_NAME)}; {T_V8}: field {esc(V8_FIELD_NAME)}; both 0 = unserved (eligible), 1 = served</p>
<h2>Results at a glance</h2>
<table><tr>{"".join(f"<th>{esc(c)}</th>" for c in GLANCE_COLS)}</tr>{glance_html}</table>
<h2>Project by project</h2>{cards}
<h2>Categories and colours</h2>
<p><b>Eligibility changes</b> – every Fabric location in the project area, {T_V4} compared with {T_V8}.</p>
<table><tr><th>Category</th><th>Tier</th><th>What it means</th></tr>{legend_rows(FABRIC_CLASSES)}</table>
<p><b>Grantee submission changes</b> – every location in the {T_SUB4} or the {T_SUB8}. Follow-up checks are applied first.</p>
<table><tr><th>Category</th><th>Tier</th><th>What it means</th></tr>{legend_rows(SUB_CLASSES)}</table>
<h2>Data cleaning</h2>
<p>Before counting, each list was cleaned: rows without a valid location ID were skipped and every location counted once. When an eligibility file listed the same ID twice with different eligibility, the ID was marked as conflicting source records and left out of the counts.</p>
<table><tr><th>Project</th><th>List</th><th>Rows read</th><th>Rows skipped (no valid ID)</th><th>IDs listed more than once</th><th>Duplicate rows removed</th><th>Locations counted</th></tr>{clean_rows}</table>
<h2>Key terms</h2><table>{gloss}</table>
<h2>Cautions</h2><ul>
<li>An ID removed from the Fabric is not proof that a building disappeared; check the area before quoting it as a loss.</li>
<li>The {T_SUB8} is the grantee's planned location list, so a shorter list than the {T_SUB4} means fewer planned locations.</li>
<li>"Needs follow-up" is a question to ask the grantee, not a finding of fault.</li></ul>
</body></html>"""
if WRITE_HTML:
    HTML_PATH.write_text(report, encoding="utf-8")
    log.info("HTML report → %s", HTML_PATH)

# %% [markdown]
# ## 7c. Analyst geodatabase: everything flagged on one layer
#
# `FFA_V4_V8_Analyst.gdb` holds what an analyst needs, in two layers and one key table:
#
# | Layer | One row per | What it carries |
# |---|---|---|
# | `projects` | project (polygon) | bottom line, key numbers, follow-up counts and the plain-English summaries |
# | `locations` | project × location ID (point) | every flag for that location: whether it is in the Fabric version 4 and version 8 eligibility data, eligible and inside the area in each, whether it is in the initial grantee submission and the 2026 Q2/Q3 data, what changed in the Fabric data, what changed in the grantee submission and why, whether it needs follow-up and the reason, and duplicate notes |
# | `category_legend` | category | label, tier, meaning and colour |
#
# Flags are written as Yes / No so attribute tables read plainly. One row per project and location ID: a location both in the Fabric eligibility data for the area and in a grantee submission is one row with both sets of flags. Each project folder gets the same geodatabase filtered to that project (`<key>_<PR>_analyst.gdb`). The full technical geodatabase (`FFA_V4_V8_Change_Analysis.gdb`) keeps every intermediate layer and table for auditing and the ArcGIS Pro map series.

# %%
ANALYST_GDB = OUTPUT_DIR / "FFA_V4_V8_Analyst.gdb"
YN = {True: "Yes", False: "No"}


def yes_no_col(s: pd.Series) -> pd.Series:
    """Booleans as 'Yes' / 'No' (missing → 'No')."""
    return s.fillna(False).astype(bool).map(YN)


# Fabric side: every Fabric location in a project area in the Fabric version 4 and/or version 8 eligibility data
fab = changes[["project_key", "location_id", "map_label", "tier", "detail", "shift_m",
               "lon_v8", "lat_v8", "lon_v4", "lat_v4"]].rename(columns={
    "map_label": "fcc_map_change", "tier": "fcc_map_change_tier", "detail": "fcc_map_detail"})
# Submission side: every location in the initial grantee submission and/or the 2026 Q2 and Q3 data
lst = subs[["project_key", "location_id", "in_sub_v4", "in_sub_v8", "map_label", "tier", "why", "needs_follow_up",
            "status_v4_plain", "status_v8_plain", "detail", "fab_lon", "fab_lat"]].rename(
    columns={"map_label": "list_change", "tier": "list_change_tier", "why": "list_change_why", "detail": "list_detail"})
loc = fab.merge(lst, on=["project_key", "location_id"], how="outer")
loc["location_id"] = loc["location_id"].astype("int64")

# Map-version flags for every row, from the cleaned Fabric records
for v in ("v4", "v8"):
    a = attrs[v]
    ids = loc["location_id"]
    loc[f"in_fcc_map_{v}"] = ids.isin(a.index)
    loc[f"conflicting_records_{v}"] = ids.map(a["conflict"].astype(bool)).fillna(False)
    loc[f"eligible_{v}"] = ids.map(a["eligible"].astype(bool)).fillna(False) & ~loc[f"conflicting_records_{v}"]
    loc[f"broadband_serviceable_{v}"] = ids.map(a["bsl"].astype(bool)).fillna(False)
    loc[f"inside_area_{v}"] = [lid in area_sets[v].get(k, set()) for k, lid in zip(loc["project_key"], ids)]

loc["in_application_list"] = loc["in_sub_v4"].fillna(False)
loc["in_2026_report"] = loc["in_sub_v8"].fillna(False)
loc["on_grantee_list"] = loc["in_application_list"] | loc["in_2026_report"]
loc["needs_follow_up"] = loc["needs_follow_up"].fillna(False).astype(bool)
loc["follow_up_reason"] = np.where(loc["needs_follow_up"], loc["list_change"], "")
loc["fcc_map_change"] = loc["fcc_map_change"].fillna("Outside the project area in both Fabric versions")
loc["list_change"] = loc["list_change"].fillna(f"Not in the {T_SUB4} or the {T_SUB8}")
for c in ("fcc_map_change_tier", "list_change_tier"):
    loc[c] = loc[c].fillna(TIER_NC)
loc["status_v4_plain"] = loc["status_v4_plain"].fillna(
    loc.apply(lambda r: STATUS_PLAIN[fabric_status(r["project_key"], r["location_id"], "v4")], axis=1))
loc["status_v8_plain"] = loc["status_v8_plain"].fillna(
    loc.apply(lambda r: STATUS_PLAIN[fabric_status(r["project_key"], r["location_id"], "v8")], axis=1))

# Duplicate notes: how often the ID appeared in each source before cleaning
dup_note = pd.Series("", index=loc.index)
if len(dup_detail):
    for ds, words in (("Initial grantee submission", T_SUB4), ("2026 Q2 and Q3 data", T_SUB8_SHORT),
                      ("Fabric version 4 eligibility data (project area)", T_V4),
                      ("Fabric version 8 eligibility data (project area)", T_V8)):
        d = dup_detail[dup_detail["dataset"] == ds].set_index(["project_key", "location_id"])["times_listed"]
        n = pd.Series([d.get((k, i), 0) for k, i in zip(loc["project_key"], loc["location_id"])], index=loc.index)
        dup_note = dup_note.where(n == 0, dup_note.where(dup_note == "", dup_note + "; ")
                                  + "listed " + n.astype(str) + "x in the " + words)
loc["listed_more_than_once"] = dup_note.ne("")
loc["duplicate_note"] = dup_note

# One headline per location: follow-up first, then the list change, then the map change
loc["headline"] = np.where(loc["needs_follow_up"], "Needs follow-up: " + loc["follow_up_reason"],
                  np.where(loc["on_grantee_list"], loc["list_change"], loc["fcc_map_change"]))
loc["project"] = loc["project_key"].map(LABEL)
loc["page_title"] = loc["project_key"].map(summary.set_index("project_key")["page_title"])
loc["detail"] = (loc["fcc_map_detail"].fillna("") + "; " + loc["list_detail"].fillna("")).str.strip("; ")
REVIEW_COLS = ["list_continuity", "current_status", "fabric_transition", "cause_class", "cause_reason", "all_flags",
               "priority", "priority_reason", "primary_action", "all_recommended_actions", "evidence_strength",
               "questions"]
loc = loc.merge(subs[["project_key", "location_id", *REVIEW_COLS]], on=["project_key", "location_id"], how="left")
loc["list_continuity"] = loc["list_continuity"].fillna("Not in either list")
loc["current_status"] = loc["current_status"].fillna("Not on the current list")
for c in REVIEW_COLS:
    if c != "priority":
        loc[c] = loc[c].fillna("")
loc["lon"] = loc["lon_v8"].fillna(loc["lon_v4"]).fillna(loc["fab_lon"])
loc["lat"] = loc["lat_v8"].fillna(loc["lat_v4"]).fillna(loc["fab_lat"])

FLAG_COLS = ["needs_follow_up", "on_grantee_list", "in_application_list", "in_2026_report", "in_fcc_map_v4",
             "in_fcc_map_v8", "inside_area_v4", "inside_area_v8", "eligible_v4", "eligible_v8", "listed_more_than_once",
             "conflicting_records_v4", "conflicting_records_v8", "broadband_serviceable_v4", "broadband_serviceable_v8"]
analyst_locations = loc[["project_key", "project", "page_title", "location_id", "headline", "needs_follow_up",
                         "follow_up_reason", "fcc_map_change", "fcc_map_change_tier", "list_change", "list_change_tier",
                         "list_change_why", "on_grantee_list", "in_application_list", "in_2026_report",
                         "in_fcc_map_v4", "in_fcc_map_v8", "inside_area_v4", "inside_area_v8", "eligible_v4",
                         "eligible_v8", "broadband_serviceable_v4", "broadband_serviceable_v8",
                         "conflicting_records_v4", "conflicting_records_v8", "status_v4_plain", "status_v8_plain",
                         *REVIEW_COLS, "shift_m", "detail", "listed_more_than_once", "duplicate_note", "lon", "lat"]].copy()
for c in FLAG_COLS:
    analyst_locations[c] = yes_no_col(analyst_locations[c])
ANALYST_FIELD_NAMES = {
    "fcc_map_change": "fabric_change", "fcc_map_change_tier": "fabric_change_tier", "list_change": "submission_change",
    "list_change_tier": "submission_change_tier", "list_change_why": "submission_change_why",
    "on_grantee_list": "in_any_submission", "in_application_list": "in_initial_submission",
    "in_2026_report": "in_2026_q2_q3_data", "in_fcc_map_v4": "in_fabric_v4_data", "in_fcc_map_v8": "in_fabric_v8_data",
    "inside_area_v4": "inside_area_fabric_v4", "inside_area_v8": "inside_area_fabric_v8",
    "eligible_v4": "eligible_fabric_v4", "eligible_v8": "eligible_fabric_v8",
    "status_v4_plain": "status_fabric_v4", "status_v8_plain": "status_fabric_v8",
    "broadband_serviceable_v4": "serviceable_location_fabric_v4", "broadband_serviceable_v8": "serviceable_location_fabric_v8",
    "conflicting_records_v4": "conflicting_records_fabric_v4", "conflicting_records_v8": "conflicting_records_fabric_v8",
    "cause_class": "observed_explanation", "cause_reason": "observed_explanation_reason",
    "questions": "questions_for_grantee"}
analyst_locations = (points_from_lonlat(analyst_locations.rename(columns=ANALYST_FIELD_NAMES))
                     .sort_values(["project_key", "needs_follow_up", "location_id"], ascending=[True, False, True])
                     .reset_index(drop=True))

PROJECT_COLS = ["project_key", "label", "pr_number", "page_title", "bottom_line", "v4_total", "v8_total", "v4_elig",
                "v8_elig", "elig_chg", "sub4_n", "sub8_n", "sub8_elig", "s_kept",
                "s_dropped", "s_added", "fu_total", *[f"fu_{c.lower()}" for c in FOLLOW_UP_CLASSES], "far_points",
                "sub4_dup_rows", "sub8_dup_rows", "patterns", "pattern_notes", "story_why", "story_a", "story_b",
                "story_c", "sub_v8_file", *key_numbers.columns]
PROJECT_COLS += [c for c in dict.fromkeys([*KEY[next(iter(KEY))], "spatial_line", "bottom_line_full"])
                 if c in summary.columns and c not in PROJECT_COLS]
analyst_projects = field_names(areas[["project_key", "geometry"]].merge(summary[PROJECT_COLS], on="project_key"))
ANALYST_LAYERS = {"projects": analyst_projects, "locations": analyst_locations, "category_legend": legend,
                  "field_guide": field_guide}


def write_analyst_gdb(path: Path, key: str | None = None) -> None:
    """Writes the analyst geodatabase, for all projects or one project."""
    if path.exists():
        if path.suffix.lower() != ".gdb":
            raise ValueError(f"Refusing to delete {path}: not a .gdb folder")
        shutil.rmtree(path)
    for name, df in ANALYST_LAYERS.items():
        part = df if key is None or "project_key" not in df.columns else df[df["project_key"] == key]
        write(part, name, gdb=path, quiet=key is not None)


write_analyst_gdb(ANALYST_GDB)

# Every listed location and every Fabric location in an area appears exactly once per project
assert not analyst_locations.duplicated(["project_key", "location_id"]).any()
assert len(analyst_locations) == len(pd.concat([changes[["project_key", "location_id"]],
                                                subs[["project_key", "location_id"]]]).drop_duplicates())
analyst_locations.groupby("project")[["needs_follow_up"]].agg(lambda s: (s == "Yes").sum()).rename(
    columns={"needs_follow_up": "locations needing follow-up"})

# %% [markdown]
# ## 9. Per-project packages for the project analyst
#
# One folder per project in `outputs/projects/`, e.g. `SLO_PR-5578/`:
#
# | File | Contents |
# |---|---|
# | `<key>_<PR>_brief.pdf` | Page 1 summary (bottom line, key numbers, how eligible locations changed step by step, follow-up counts); pages 2–4 maps on Light Gray Canvas: eligibility changes (Fabric version 4 → version 8 eligibility data), grantee submission changes (initial grantee submission → 2026 Q2/Q3 data), follow-up |
# | `<key>_<PR>_maps_imagery.pdf` | The same three maps on aerial imagery |
# | `<key>_<PR>_results.xlsx` | The project's results workbook: At a glance, Key numbers, Fabric in the project area, Grantee lists and the Fabric, Coverage, From version 4 to version 8, Follow-up list, Eligibility changes, Submitted locations, Category counts, Data cleaning, Duplicate location IDs, Column guide |
# | `<key>_<PR>_analyst.gdb` | The analyst geodatabase for this project: `projects` (outline with all summary fields), `locations` (every location with all its flags), `category_legend` |
# | `<key>_<PR>_summary.csv` | The project's summary: one plain-named measure per row |
#
# The unified results (`FFA_V4_V8_Analyst.gdb`, the technical `FFA_V4_V8_Change_Analysis.gdb`, the all-project Excel workbook, HTML report and summary CSV) stay in `outputs/`. Set `WRITE_FULL_PROJECT_GDB = True` for a per-project copy of every technical layer, or `WRITE_GPKG = True` for a GeoPackage.
#
# Pages are 14 × 8.5 in (legal, landscape). Each map frames the whole project outline plus every point within `FAR_POINT_KM` of it; text sits in fixed panels and shrinks to fit, so nothing overlaps. Basemaps need internet access (Esri tiles via `contextily`); without it the maps draw on a plain background.

# %%
import textwrap
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
from matplotlib.lines import Line2D
from matplotlib.patches import FancyBboxPatch, Rectangle

try:
    import contextily as ctx
except ImportError:                                   # basemaps are optional
    ctx = None
    log.warning("contextily not installed: maps draw without a basemap (pip install contextily)")

PROJECTS_DIR = OUTPUT_DIR / "projects"
PAGE_W, PAGE_H = 14.0, 8.5                            # inches, legal landscape
NAVY, INK, QUIET = "#1F3864", "#1d1d1f", "#5f6368"
BASEMAPS = {"gray": ("Esri.WorldGrayCanvas", "Basemap: Esri, HERE, Garmin, © OpenStreetMap contributors"),
            "imagery": ("Esri.WorldImagery", "Basemap: Esri, Maxar, Earthstar Geographics")}
POINT_SIZE = {TIER_FU: 30, TIER_EX: 20, TIER_NC: 12}  # marker area in pt² (about 5.5, 4.5 and 3.5 pt across)
DESC_SIZE, LEGEND_SIZE = 10.0, 9.0                   # same text sizes on every page
LEGEND_MARKER = 7.5                                  # legend dot size (pt across)
MAP_LEFT_IN, MAP_BOTTOM_IN = 0.35, 0.75
PANEL_X_IN = MAP_LEFT_IN + MAP_FRAME_W_IN + 0.35     # right-hand panel
PANEL_W_IN = PAGE_W - PANEL_X_IN - 0.35
FOOTER = (f"Sources: {cap(T_V4)} (field {{v4}}) and {T_V8} (field status_code), 0 = unserved (eligible), 1 = served; "
          f"CPUC FFA awarded project areas; {T_SUB4} and {T_SUB8}. Eligible = no cable or fiber at 25/3 Mbps or better.")
basemap_ok = {"gray": ctx is not None, "imagery": ctx is not None}


def fx(x_in: float) -> float:
    return x_in / PAGE_W


def fy(y_in: float) -> float:
    return y_in / PAGE_H


def rgb01(rgb) -> tuple:
    return tuple(v / 255 for v in rgb)


def fit_text(fig, text: str, x_in: float, y_top_in: float, w_in: float, h_in: float,
             size: float = 10.5, min_size: float = 7.0, **kw):
    """Draws wrapped text inside a box (inches), shrinking the font until it fits both ways."""
    renderer = fig.canvas.get_renderer()
    color = kw.pop("color", INK)
    s = size
    while s >= min_size:
        chars = int(w_in * 72 / (s * 0.50))
        while chars > 15:
            wrapped = "\n".join(textwrap.fill(p, chars) if p else "" for p in text.split("\n"))
            t = fig.text(fx(x_in), fy(y_top_in), wrapped, fontsize=s, va="top", ha="left", linespacing=1.3,
                         color=color, **kw)
            bb = t.get_window_extent(renderer=renderer)
            if bb.width / fig.dpi <= w_in:
                if bb.height / fig.dpi <= h_in:
                    return t
                t.remove()
                break
            t.remove()
            chars -= 3
        s -= 0.5
    t = fig.text(fx(x_in), fy(y_top_in), textwrap.shorten(text, int(w_in * h_in * 20), placeholder=" …"),
                 fontsize=min_size, va="top", ha="left", wrap=True, color=color)
    return t


def page_frame(fig, title: str, subtitle: str, page_no: int, n_pages: int, credit: str = "") -> None:
    """Title band and footer shared by every page."""
    renderer = fig.canvas.get_renderer()
    for text, y_in, size, min_size, weight, colour in ((title, PAGE_H - 0.42, 17, 11, "bold", NAVY),
                                                       (subtitle, PAGE_H - 0.72, 10.5, 8, "normal", QUIET)):
        t = fig.text(fx(0.35), fy(y_in), text, fontsize=size, fontweight=weight, color=colour, va="center")
        while t.get_window_extent(renderer=renderer).width / fig.dpi > PAGE_W - 0.7 and size > min_size:
            size -= 0.5
            t.set_fontsize(size)
    fig.add_artist(Line2D([fx(0.35), fx(PAGE_W - 0.35)], [fy(PAGE_H - 0.92)] * 2, color=NAVY, lw=1.2))
    fig.add_artist(Line2D([fx(0.35), fx(PAGE_W - 0.35)], [fy(0.52)] * 2, color="#d0d0d0", lw=0.6))
    fit_text(fig, FOOTER.format(v4=V4_FIELD_NAME) + (" " + credit + "." if credit else "") + f" Created {RUN_DATE}.",
             0.35, 0.45, PAGE_W - 1.6, 0.36, size=7, min_size=6, color=QUIET)
    fig.text(fx(PAGE_W - 0.35), fy(0.3), f"Page {page_no} of {n_pages}", fontsize=8, color=QUIET, ha="right", va="center")


def panel_heading(fig, text: str, x_in: float, y_in: float) -> None:
    fig.text(fx(x_in), fy(y_in), text, fontsize=11.5, fontweight="bold", color=NAVY, va="top")


def add_basemap(ax, kind: str, zoom_adjust: int = 0) -> bool:
    """Esri tiles under the map; False (plain background) when tiles cannot be fetched."""
    ax.set_facecolor("#f4f4f2")
    if not basemap_ok.get(kind):
        return False
    try:
        src = ctx.providers
        for part in BASEMAPS[kind][0].split("."):
            src = src[part]
        ctx.add_basemap(ax, crs=MERC, source=src, attribution=False, zoom_adjust=zoom_adjust, reset_extent=True)
        return True
    except Exception as e:                              # offline or tile server refused
        log.warning("Basemap '%s' unavailable (%s); maps draw on a plain background", kind, e)
        basemap_ok[kind] = False
        return False


def nice_length(m: float) -> float:
    exp = 10 ** np.floor(np.log10(m))
    return max(v * exp for v in (1, 2, 5) if v * exp <= m)


def scale_bar(ax, extent, lat: float) -> None:
    """Scale bar (miles) in the lower left of the map, on a white backing."""
    x0, y0, x1, y1 = extent
    ground_per_unit = np.cos(np.radians(lat))           # Web Mercator metres → ground metres
    target_mi = (x1 - x0) * 0.2 * ground_per_unit / 1609.344
    length_mi = nice_length(target_mi) if target_mi >= 0.1 else 0.1
    length_units = length_mi * 1609.344 / ground_per_unit
    bx, by = x0 + (x1 - x0) * 0.03, y0 + (y1 - y0) * 0.04
    ax.add_patch(Rectangle((bx - (x1 - x0) * 0.012, by - (y1 - y0) * 0.022), length_units + (x1 - x0) * 0.024,
                           (y1 - y0) * 0.075, facecolor="white", alpha=0.85, edgecolor="none", zorder=20))
    ax.plot([bx, bx + length_units], [by, by], color=INK, lw=2.2, solid_capstyle="butt", zorder=21)
    ax.text(bx + length_units / 2, by + (y1 - y0) * 0.012, f"{length_mi:g} mi", ha="center", va="bottom",
            fontsize=7.5, color=INK, zorder=21)


def north_arrow(ax) -> None:
    ax.annotate("N", xy=(0.965, 0.955), xytext=(0.965, 0.885), xycoords="axes fraction", ha="center", va="center",
                fontsize=9, fontweight="bold", color=INK,
                arrowprops=dict(arrowstyle="-|>", color=INK, lw=1.2),
                bbox=dict(boxstyle="round,pad=0.25", fc="white", ec="none", alpha=0.85), zorder=22)


def draw_points(ax, pts: gpd.GeoDataFrame, classes: dict, imagery: bool) -> pd.Series:
    """Points by category, quiet classes first so changes sit on top; returns counts per class."""
    counts = pts["map_class"].value_counts()
    edge = "white"
    for tier in reversed(TIER_ORDER):
        for code in [c for c in classes if classes[c][3] == tier]:
            g = pts[pts["map_class"] == code]
            if len(g) and code not in SHADED_CLASSES:
                ax.scatter(g.geometry.x, g.geometry.y, s=POINT_SIZE[tier], c=[rgb01(classes[code][2])],
                           edgecolors=edge, linewidths=0.35 if imagery else 0.25, zorder=10 + TIER_ORDER[::-1].index(tier))
    return counts


def balanced_wrap(text: str, chars: int) -> list[str]:
    """Wraps to the fewest lines, then evens the line lengths so no line holds a lone word or number."""
    lines = textwrap.wrap(text, chars)
    if len(lines) < 2:
        return lines
    for w in range(len(text) // len(lines), chars + 1):
        trial = textwrap.wrap(text, w)
        if len(trial) == len(lines):
            return trial
    return lines


def draw_density(ax, key: str, imagery: bool) -> pd.Series:
    """Hexagons for the shaded classes, under the points; returns the located count per class."""
    d = density[density["project_key"] == key]
    for shade, grp in d.groupby("shade"):
        grp.plot(ax=ax, color=rgb01(SHADE_RGB[shade]), edgecolor="none", alpha=0.6 if imagery else 0.75, zorder=8)
    return d.groupby("map_class")["locations"].sum()


def draw_legend(fig, counts: pd.Series, classes: dict, x_in: float, y_top_in: float, w_in: float, h_in: float,
                rename=lambda t: t) -> float:
    """Legend grouped by tier, only categories present on the page, with counts; long labels wrap to two lines,
    and the font shrinks until everything fits the box."""
    rows = []
    for tier in TIER_ORDER:
        items = [c for c in classes if classes[c][3] == tier and counts.get(c, 0)]
        if items:
            rows.append(("tier", tier, None))
            rows += [("item", c, int(counts[c])) for c in items]
    if not rows:
        return y_top_in
    label_w = w_in - 0.24 - 0.55                         # marker column + count column
    probe = fig.text(0, 0, "", alpha=0)
    renderer = fig.canvas.get_renderer()

    def width_in(text: str, size: float) -> float:
        probe.set_text(text)
        probe.set_fontsize(size)
        return probe.get_window_extent(renderer=renderer).width / fig.dpi

    def wrap_to(text: str, size: float) -> list[str]:
        """Fewest lines that fit label_w (measured), then evened out."""
        words, lines, cur = text.split(), [], ""
        for w in words:
            trial = (cur + " " + w).strip()
            if cur and width_in(trial, size) > label_w:
                lines.append(cur)
                cur = w
            else:
                cur = trial
        lines.append(cur)
        if len(lines) > 1:
            for chars in range(len(text) // len(lines), len(text) + 1):
                even = textwrap.wrap(text, chars)
                if len(even) == len(lines) and all(width_in(l, size) <= label_w for l in even):
                    return even
        return lines

    for size in np.arange(LEGEND_SIZE, 6.9, -0.25):          # shrinks only when the panel cannot hold it
        wrapped = {c: wrap_to(rename(classes[c][0]) + (" (shaded, darker = more)" if c in SHADED_CLASSES else ""), size)
                   for k, c, _ in rows if k == "item"}
        line = size * 1.45 / 72
        need = sum(line * (len(wrapped[c]) + 0.15) if k == "item" else line * 1.2 for k, c, _ in rows)
        if need <= h_in:
            break
    probe.remove()
    y = y_top_in
    for kind, val, n in rows:
        if kind == "tier":
            fig.text(fx(x_in), fy(y), val, fontsize=size, fontweight="bold", color=QUIET, va="top")
            y -= line * 1.2
            continue
        shaded = val in SHADED_CLASSES
        fig.add_artist(Line2D([fx(x_in + 0.09)], [fy(y - line * 0.42)], marker="h" if shaded else "o",
                              markersize=LEGEND_MARKER * (1.3 if shaded else 1.0),
                              markerfacecolor=rgb01(SHADE_RGB[2] if shaded else classes[val][2]),
                              markeredgecolor="white", linestyle="none"))
        fig.text(fx(x_in + 0.24), fy(y), "\n".join(wrapped[val]), fontsize=size, color=INK, va="top", linespacing=1.2)
        fig.text(fx(x_in + w_in), fy(y), f"{n:,}", fontsize=size, color=INK, va="top", ha="right")
        y -= line * len(wrapped[val]) + line * 0.15
    return y


# Places labelled on the small overview map: (name, latitude, longitude, rank); rank 1 is drawn first
PLACES = [
    ("San Francisco", 37.7749, -122.4194, 1), ("Oakland", 37.8044, -122.2712, 1), ("San Jose", 37.3382, -121.8863, 1),
    ("Sacramento", 38.5816, -121.4944, 1), ("Stockton", 37.9577, -121.2908, 1), ("Santa Rosa", 38.4405, -122.7144, 1),
    ("Fresno", 36.7378, -119.7871, 1), ("Bakersfield", 35.3733, -119.0187, 1), ("Santa Barbara", 34.4208, -119.6982, 1),
    ("Reno", 39.5296, -119.8138, 1),
    ("Berkeley", 37.8716, -122.2727, 2), ("Hayward", 37.6688, -122.0808, 2), ("Fremont", 37.5485, -121.9886, 2),
    ("Walnut Creek", 37.9101, -122.0652, 2), ("Richmond", 37.9358, -122.3477, 2), ("Vallejo", 38.1041, -122.2566, 2),
    ("Concord", 37.9780, -122.0311, 2), ("San Mateo", 37.5630, -122.3255, 2), ("Daly City", 37.6879, -122.4702, 2),
    ("San Leandro", 37.7249, -122.1561, 2), ("Alameda", 37.7652, -122.2416, 2), ("Folsom", 38.6780, -121.1761, 2),
    ("Roseville", 38.7521, -121.2880, 2), ("Auburn", 38.8966, -121.0769, 2), ("Placerville", 38.7296, -120.7985, 2),
    ("South Lake Tahoe", 38.9399, -119.9772, 2), ("Truckee", 39.3280, -120.1833, 2), ("Jackson", 38.3488, -120.7741, 2),
    ("Grass Valley", 39.2191, -121.0611, 2), ("Davis", 38.5449, -121.7405, 2), ("Elk Grove", 38.4088, -121.3716, 2),
    ("San Luis Obispo", 35.2828, -120.6596, 2), ("Santa Maria", 34.9530, -120.4357, 2),
    ("Paso Robles", 35.6266, -120.6910, 2), ("Atascadero", 35.4894, -120.6707, 2), ("Lompoc", 34.6391, -120.4579, 2),
    ("Emeryville", 37.8313, -122.2852, 3), ("Piedmont", 37.8244, -122.2317, 3), ("Castro Valley", 37.6941, -122.0864, 3),
    ("Orinda", 37.8771, -122.1797, 3), ("Lafayette", 37.8858, -122.1180, 3), ("San Lorenzo", 37.6810, -122.1244, 3),
    ("Morro Bay", 35.3658, -120.8499, 3), ("Los Osos", 35.3111, -120.8324, 3), ("Pismo Beach", 35.1428, -120.6413, 3),
    ("Arroyo Grande", 35.1186, -120.5907, 3), ("Grover Beach", 35.1217, -120.6210, 3), ("Avila Beach", 35.1797, -120.7319, 3),
    ("Cayucos", 35.4428, -120.8921, 3), ("Templeton", 35.5497, -120.7060, 3), ("Santa Margarita", 35.3897, -120.6088, 3),
    ("Nipomo", 35.0428, -120.4760, 3), ("Coloma", 38.7999, -120.8891, 3), ("Georgetown", 38.9068, -120.8388, 3),
    ("Garden Valley", 38.8524, -120.8580, 3), ("Cool", 38.8874, -121.0147, 3), ("Pollock Pines", 38.7616, -120.5861, 3),
    ("Diamond Springs", 38.6946, -120.8149, 3), ("Shingle Springs", 38.6657, -120.9241, 3),
    ("Cameron Park", 38.6688, -120.9872, 3), ("El Dorado Hills", 38.6857, -121.0822, 3), ("Foresthill", 39.0202, -120.8180, 3),
    ("Camino", 38.7385, -120.6750, 3), ("Somerset", 38.6477, -120.6844, 3),
]
_pl = gpd.GeoDataFrame({"name": [p[0] for p in PLACES], "rank": [p[3] for p in PLACES]},
                       geometry=gpd.points_from_xy([p[2] for p in PLACES], [p[1] for p in PLACES]), crs=WGS84).to_crs(MERC)
LOCATOR_H_IN, LOCATOR_MIN_H_IN = 1.3, 0.65


def locator(fig, key: str, extent, kind: str, h_in: float = LOCATOR_H_IN) -> None:
    """Small overview map with labelled towns and the main map's frame, when a basemap is available."""
    if not basemap_ok.get(kind):
        return
    import matplotlib.patheffects as pe
    w_in = PANEL_W_IN
    ax = fig.add_axes([fx(PANEL_X_IN), fy(MAP_BOTTOM_IN), fx(w_in), fy(h_in)])
    x0, y0, x1, y1 = extent
    cx, cy = (x0 + x1) / 2, (y0 + y1) / 2
    span_h = max((y1 - y0) * 6, (x1 - x0) * 6 * h_in / w_in)
    lx0, lx1 = cx - span_h * w_in / h_in / 2, cx + span_h * w_in / h_in / 2
    ly0, ly1 = cy - span_h / 2, cy + span_h / 2
    ax.set_xlim(lx0, lx1), ax.set_ylim(ly0, ly1)
    if not add_basemap(ax, kind, zoom_adjust=-1):
        ax.remove()
        return
    ax.set_xlim(lx0, lx1), ax.set_ylim(ly0, ly1)
    ax.add_patch(Rectangle((x0, y0), x1 - x0, y1 - y0, fill=False, edgecolor=NAVY, lw=1.6, zorder=6))
    ax.set_xticks([]), ax.set_yticks([])
    for s_ in ax.spines.values():
        s_.set_edgecolor("#999999"), s_.set_linewidth(0.6)
    # Labels: biggest places first; a label is skipped when it would overlap another label, the frame
    # box or the edge of the overview map
    renderer = fig.canvas.get_renderer()
    ax_box = ax.get_window_extent(renderer=renderer)
    halo = [pe.withStroke(linewidth=2.2, foreground="white")]
    near = _pl.cx[lx0:lx1, ly0:ly1].copy()
    if near.empty:
        return
    # The place nearest the frame centre names the project's own area and is drawn first, in bold
    near["d"] = near.geometry.distance(shapely.Point(cx, cy))
    home = near["d"].idxmin()
    near["order"] = np.where(near.index == home, 0, near["rank"])
    taken = []
    for p in near.sort_values(["order", "d"]).itertuples():
        px, py = p.geometry.x, p.geometry.y
        if p.Index != home and x0 <= px <= x1 and y0 <= py <= y1:
            continue                                      # other places inside the frame box stay unlabelled
        is_home = p.Index == home
        t = ax.text(px, py, " " + p.name, fontsize=7.5 if is_home else 7 if p.rank < 3 else 6.5, color=INK,
                    va="center", ha="left", fontweight="bold" if is_home or p.rank == 1 else "normal",
                    path_effects=halo, zorder=9 if is_home else 8, clip_on=True)
        if is_home:
            ax.plot(px, py, "o", markersize=2.6, color=INK, zorder=9)
            bb = t.get_window_extent(renderer=renderer)
            taken += [[[bb.x0, bb.y0], [bb.x1, bb.y1]], ax.transData.transform([[x0, y0], [x1, y1]])]
            continue
        bb = t.get_window_extent(renderer=renderer)
        clash = any(not (bb.x1 < a[0][0] or bb.x0 > a[1][0] or bb.y1 < a[0][1] or bb.y0 > a[1][1]) for a in taken)
        if clash or bb.x1 > ax_box.x1 - 2 or bb.y0 < ax_box.y0 + 1 or bb.y1 > ax_box.y1 - 1:
            t.remove()
            continue
        ax.plot(px, py, "o", markersize=2.2, color=INK, zorder=8)
        taken.append([[bb.x0, bb.y0], [bb.x1, bb.y1]])


def map_page(pdf, r: pd.Series, page: str, pts: gpd.GeoDataFrame, kind: str, page_no: int, n_pages: int) -> None:
    """One map page: map on the left; summary, legend and locator on the right."""
    qr = qr_term(r["project_key"])
    titles = {"A": ("Eligibility changes in the project area", f"{cap(T_V4)} compared with the {T_V8}",
                    FABRIC_CLASSES, r["story_a"]),
              "B": ("Grantee submission changes", f"{cap(T_SUB4)} compared with the {qr}",
                    SUB_CLASSES, r["story_b"]),
              "C": ("Locations that need follow-up", "Listed locations not explained by the observed Fabric change",
                    FOLLOW_CLASSES, r["story_c"])}
    heading, sub, classes, story = titles[page]
    fig = plt.figure(figsize=(PAGE_W, PAGE_H))
    ax = fig.add_axes([fx(MAP_LEFT_IN), fy(MAP_BOTTOM_IN), fx(MAP_FRAME_W_IN), fy(MAP_FRAME_H_IN)])
    extent = extents[r["project_key"]]
    ax.set_xlim(extent[0], extent[2]), ax.set_ylim(extent[1], extent[3])
    has_tiles = add_basemap(ax, kind)
    page_frame(fig, f"{r['page_title']} · {heading}", sub, page_no, n_pages,
               credit=BASEMAPS[kind][1] if has_tiles else "")
    ax.set_xlim(extent[0], extent[2]), ax.set_ylim(extent[1], extent[3])
    area = areas.loc[areas["project_key"] == r["project_key"]].to_crs(MERC)
    area.boundary.plot(ax=ax, color="white" if kind == "imagery" else NAVY, linewidth=1.6, zorder=15)
    counts = draw_points(ax, pts, classes, kind == "imagery")
    if page == "A":
        counts = counts.drop(list(SHADED_CLASSES), errors="ignore")
        counts = pd.concat([counts, draw_density(ax, r["project_key"], kind == "imagery")])
    ax.set_xticks([]), ax.set_yticks([])
    ax.set_xlabel(""), ax.set_ylabel("")
    for s_ in ax.spines.values():
        s_.set_edgecolor("#999999"), s_.set_linewidth(0.6)
    lat = gpd.GeoSeries([shapely.Point((extent[0] + extent[2]) / 2, (extent[1] + extent[3]) / 2)], crs=MERC).to_crs(WGS84).y.iloc[0]
    scale_bar(ax, extent, lat)
    north_arrow(ax)
    if r["far_points"]:
        ax.text(0.99, 0.02, f"{r['far_points']:,} listed location(s) lie more than {FAR_POINT_KM} km away and are not shown",
                transform=ax.transAxes, ha="right", va="bottom", fontsize=7.5, color=INK, zorder=23,
                bbox=dict(boxstyle="round,pad=0.3", fc="white", ec="none", alpha=0.9))
    # Right-hand panel: summary (top), legend (middle), locator (bottom)
    top = MAP_BOTTOM_IN + MAP_FRAME_H_IN
    panel_heading(fig, "Description", PANEL_X_IN, top)
    desc = fit_text(fig, story, PANEL_X_IN, top - 0.3, PANEL_W_IN, 3.0, size=DESC_SIZE, min_size=8.5)
    desc_h = desc.get_window_extent(renderer=fig.canvas.get_renderer()).height / fig.dpi
    leg_top = min(top - 2.8, top - 0.3 - desc_h - 0.25)
    panel_heading(fig, "Legend", PANEL_X_IN, leg_top)
    fig.text(fx(PANEL_X_IN + PANEL_W_IN), fy(leg_top), "Locations", fontsize=8.5, color=QUIET, ha="right", va="top")
    legend_counts = counts.copy()
    y_end = draw_legend(fig, legend_counts, classes, PANEL_X_IN, leg_top - 0.32, PANEL_W_IN,
                        leg_top - 0.32 - (MAP_BOTTOM_IN + LOCATOR_MIN_H_IN + 0.35),
                        rename=lambda t: t.replace(T_SUB8_SHORT, qr))
    y_area = y_end - 0.12
    fig.add_artist(Line2D([fx(PANEL_X_IN + 0.02), fx(PANEL_X_IN + 0.18)], [fy(y_area)] * 2, color=NAVY, lw=1.6))
    fig.text(fx(PANEL_X_IN + 0.24), fy(y_area), "Project area outline" + (" (white on imagery)" if kind == "imagery" else ""),
             fontsize=8.5, color=INK, va="center")
    # The overview map takes the space left under the legend (between its minimum and full height)
    locator(fig, r["project_key"], extent, kind,
            h_in=float(np.clip(y_area - 0.25 - MAP_BOTTOM_IN, LOCATOR_MIN_H_IN, LOCATOR_H_IN)))
    pdf.savefig(fig)
    plt.close(fig)


def draw_table(fig, x_in: float, y_top_in: float, col_w: list[float], header: list[str], rows: list[list],
               size: float = 9.5, row_h: float = 0.27, swatches: list | None = None, header_h: float | None = None) -> float:
    """Simple ruled table (header text may hold a line break); returns the y of its bottom edge."""
    y = y_top_in
    hh = header_h or row_h
    fig.add_artist(Rectangle((fx(x_in), fy(y - hh)), fx(sum(col_w)), fy(hh), facecolor=NAVY, edgecolor="none"))
    cx = x_in
    for j, h in enumerate(header):
        fig.text(fx(cx + (0.08 if j == 0 else col_w[j] - 0.08)), fy(y - hh / 2), h, fontsize=size - 0.5, color="white",
                 fontweight="bold", va="center", ha="left" if j == 0 else "right", linespacing=1.15)
        cx += col_w[j]
    y -= hh - row_h
    for i, row in enumerate(rows):
        y -= row_h
        if i % 2:
            fig.add_artist(Rectangle((fx(x_in), fy(y - row_h)), fx(sum(col_w)), fy(row_h), facecolor="#f3f5f9",
                                     edgecolor="none"))
        cx = x_in
        for j, v in enumerate(row):
            pad = 0.08
            if j == 0 and swatches and swatches[i] is not None:
                fig.add_artist(Line2D([fx(cx + 0.14)], [fy(y - row_h / 2)], marker="o", markersize=7,
                                      markerfacecolor=rgb01(swatches[i]), markeredgecolor="white", linestyle="none"))
                pad = 0.3
            fig.text(fx(cx + (pad if j == 0 else col_w[j] - 0.08)), fy(y - row_h / 2), v, fontsize=size, color=INK,
                     va="center", ha="left" if j == 0 else "right")
            cx += col_w[j]
    y -= row_h
    fig.add_artist(Line2D([fx(x_in), fx(x_in + sum(col_w))], [fy(y)] * 2, color="#c0c0c0", lw=0.6))
    return y


def summary_page(pdf, r: pd.Series, n_pages: int) -> None:
    """Page 1: bottom line, key numbers, eligible-count bridge, list changes and follow-up counts."""
    key = r["project_key"]
    fig = plt.figure(figsize=(PAGE_W, PAGE_H))
    qr = qr_term(key)
    page_frame(fig, f"{r['page_title']} · Summary", f"FFA project location review: {T_V4} → {T_V8}", 1, n_pages)
    left_w = 6.3
    top = PAGE_H - 1.15
    # Bottom line
    box_h = 1.6
    fig.add_artist(FancyBboxPatch((fx(0.35), fy(top - box_h)), fx(left_w), fy(box_h), boxstyle="square,pad=0",
                                  facecolor="#fff8e6", edgecolor="#E6AB02", lw=0))
    fig.add_artist(Line2D([fx(0.35)] * 2, [fy(top - box_h), fy(top)], color="#E6AB02", lw=4))
    panel_heading(fig, "Bottom line", 0.55, top - 0.1)
    fit_text(fig, r["bottom_line"], 0.55, top - 0.36, left_w - 0.35, 0.8, size=DESC_SIZE, min_size=8)
    fit_text(fig, "Change pattern: " + r["patterns"], 0.55, top - 1.18, left_w - 0.35, 0.22, size=9, min_size=8,
             fontweight="bold")
    fit_text(fig, "Recommended action: " + r["actions"], 0.55, top - 1.38, left_w - 0.35, 0.22, size=9, min_size=8,
             fontweight="bold")
    # Key numbers
    rh, ts = 0.205, 9.0
    y = top - box_h - 0.14
    panel_heading(fig, "Key numbers", 0.35, y)
    widths = [2.6, 1.35, 1.35, 1.0]
    rows = [["Locations in the project area", f"{r['v4_total']:,}", f"{r['v8_total']:,}", f"{int(r['total_chg']):+,}"],
            ["Eligible locations in the area", f"{r['v4_elig']:,}", f"{r['v8_elig']:,}", f"{int(r['elig_chg']):+,}"]]
    y = draw_table(fig, 0.35, y - 0.27, widths, ["Project area", "Fabric version 4\neligibility data",
                                                 "Fabric version 8\neligibility data", "Change"], rows,
                   size=ts, row_h=rh, header_h=0.4)
    sub_rows = [("Locations listed", "listed"), ("Not found in the Fabric", "notfound"),
                ("Inside the area – eligible", "in_elig"), ("Inside the area – not eligible", "in_nelig"),
                ("Outside the area – eligible", "out_elig"), ("Outside the area – not eligible", "out_nelig")]
    rows = [[lab, f"{int(r[f'sub4_v4_{m}']):,}", f"{int(r[f'sub8_v8_{m}']):,}",
             f"{int(r[f'sub8_v8_{m}'] - r[f'sub4_v4_{m}']):+,}" if r[f'sub8_v8_{m}'] != r[f'sub4_v4_{m}'] else "0"]
            for lab, m in sub_rows]
    y = draw_table(fig, 0.35, y - 0.1, widths, ["Grantee submissions", "Initial grantee\nsubmission · v4",
                                                  qr.replace(" data", " data\n") + "· v8", "Change"], rows,
                   size=ts, row_h=rh, header_h=0.4)
    # Follow-up counts
    y -= 0.14
    panel_heading(fig, "Needs follow-up", 0.35, y)
    n = plain_b.loc[key]
    shown = [c for c in FOLLOW_UP_CLASSES if n[c]]          # categories with no location are left out
    fu_rows = [[SUB_CLASSES[c][0].replace(T_SUB8_SHORT, qr), f"{int(n[c]):,}"] for c in shown] or [["None", "0"]]
    fu_rows.append(["Total", f"{int(r['fu_total']):,}"])
    fu_rh = min(rh, (y - 0.27 - 0.72) / (len(fu_rows) + 1))   # rows shrink to stay above the footer
    y = draw_table(fig, 0.35, y - 0.27, [5.3, 1.0], ["Reason", "Locations"], fu_rows, size=min(ts, max(7.5, fu_rh * 45)),
                   row_h=fu_rh, swatches=[SUB_CLASSES[c][2] for c in shown] + [None] * (2 if not shown else 1))
    fig.text(fx(0.35), fy(y - 0.07), "v4 / v8: checked against the Fabric version 4 / version 8 eligibility data. "
             "Location IDs: page 5 and the project workbook. Full key numbers: page 2.", fontsize=7.5, color=QUIET, va="top")
    # Bridge chart: eligible count from Fabric version 4 to version 8 by cause (deltas reconcile exactly)
    ch = changes[changes["project_key"] == key]
    steps = ch.groupby("map_class")["delta_eligible"].sum()
    bars = [("Eligible – Fabric version 4 data", r["v4_elig"], "total")]
    bars += [(FABRIC_CLASSES[c][0], int(steps.get(c, 0)), c) for c in FABRIC_CLASSES if int(steps.get(c, 0))]
    bars += [("Eligible – Fabric version 8 data", r["v8_elig"], "total")]
    cx_in = 0.35 + left_w + 0.6
    cw_in = PAGE_W - cx_in - 0.45
    panel_heading(fig, "How the number of eligible locations changed", cx_in, top - 0.1)
    fig.text(fx(cx_in), fy(top - 0.38), f"{cap(T_V4)} → {T_V8}, step by step", fontsize=9, color=QUIET, va="top")
    ax = fig.add_axes([fx(cx_in + 2.75), fy(top - 3.4), fx(cw_in - 3.25), fy(2.75)])
    running = 0
    for i, (label, val, code) in enumerate(bars):
        if code == "total":
            ax.barh(i, val, color="#7f8fa6", height=0.6)
            running = val
            ax.text(val, i, f"  {val:,}", va="center", ha="left", fontsize=9, color=INK)
        else:
            start = running
            running += val
            ax.barh(i, val, left=start, color=rgb01(FABRIC_CLASSES[code][2]), height=0.6)
            ax.text(max(start, running), i, f"  {val:+,}", va="center", ha="left", fontsize=9, color=INK)
    ax.set_yticks(range(len(bars)))
    ax.set_yticklabels([b[0] for b in bars], fontsize=8.5)
    ax.invert_yaxis()
    ax.set_xlim(0, max(max(r["v4_elig"], r["v8_elig"]), 1) * 1.18)
    ax.tick_params(axis="x", labelsize=8, colors=QUIET)
    for s_ in ("top", "right"):
        ax.spines[s_].set_visible(False)
    ax.spines["left"].set_color("#c0c0c0"), ax.spines["bottom"].set_color("#c0c0c0")
    # Grantee list changes
    panel_heading(fig, "How the grantee's submissions changed", cx_in, 3.5)
    fig.text(fx(cx_in), fy(3.22), f"{cap(T_SUB4)} → {qr}", fontsize=9, color=QUIET, va="top")
    lst = [(f"In {T_SUB4}", r["sub4_n"], "#7f8fa6"), ("In both, same ID", r["s_kept"], rgb01(RGB["grey"])),
           (f"Not carried into {qr}", r["s_dropped"], rgb01(RGB["slate"])),
           (f"New in {qr}", r["s_added"], rgb01(RGB["blue"])), (f"In {qr}", r["sub8_n"], "#7f8fa6")]
    ax2 = fig.add_axes([fx(cx_in + 2.75), fy(0.85), fx(cw_in - 3.25), fy(2.0)])
    for i, (label, val, col) in enumerate(lst):
        ax2.barh(i, val, color=col, height=0.6)
        ax2.text(val, i, f"  {val:,}", va="center", ha="left", fontsize=9, color=INK)
    ax2.set_yticks(range(len(lst)))
    ax2.set_yticklabels([b[0] for b in lst], fontsize=8.5)
    ax2.invert_yaxis()
    ax2.set_xlim(0, max(max(v for _, v, _ in lst), 1) * 1.18)
    ax2.tick_params(axis="x", labelsize=8, colors=QUIET)
    for s_ in ("top", "right"):
        ax2.spines[s_].set_visible(False)
    ax2.spines["left"].set_color("#c0c0c0"), ax2.spines["bottom"].set_color("#c0c0c0")
    pdf.savefig(fig)
    plt.close(fig)


def project_workbook(path: Path, r: pd.Series) -> None:
    """Per-project workbook: summary, follow-up list (IDs + reason), counts, locations, data cleaning."""
    key = r["project_key"]
    ch, sb = changes[changes["project_key"] == key], subs[subs["project_key"] == key]
    t = location_tables(ch, sb)
    facts = pd.DataFrame([
        ("Project", r["label"]), ("PR number", r["pr_number"]), ("Bottom line", r["bottom_line"]),
        (f"Locations in area – {T_V4}", r["v4_total"]), (f"Locations in area – {T_V8}", r["v8_total"]),
        (f"Eligible in area – {T_V4}", r["v4_elig"]), (f"Eligible in area – {T_V8}", r["v8_elig"]),
        (f"Locations in {T_SUB4}", r["sub4_n"]), (f"Locations in {qr_term(key)}", r["sub8_n"]),
        (f"Of those in {qr_term(key)}, eligible in {T_V8}", r["sub8_elig"]), ("Needs follow-up (total)", r["fu_total"]),
        *[(f"Follow-up: {SUB_CLASSES[c][0]}", r[f"fu_{c.lower()}"]) for c in FOLLOW_UP_CLASSES],
        ("Change pattern", r["patterns"]), ("Pattern notes", r["pattern_notes"]),
        ("How the change came together", r["story_why"]),
        ("Eligibility summary", r["story_a"]), ("Grantee submission summary", r["story_b"]),
        ("2026 data file used", r["sub_v8_file"])], columns=["Item", "Value"])
    readme = [(f"{r['page_title']}: {T_V4} → {T_V8} review", ""), (f"Run on {RUN_DATE}", ""), ("", ""),
              ("Tab", "What it holds"), ("Summary", "Bottom line, key numbers and plain-English summaries."),
              ("Fabric in the area", "Fabric locations inside the project area, version 4 and version 8, with change."),
              ("Grantee lists", "Each grantee list checked against both Fabric versions: matched, inside or outside the "
                                "area, eligible or not, distance outside, boundary, coordinates."),
              ("Coverage", "Share of the eligible locations inside the area that the grantee lists."),
              ("From-to tables", "Status of each location in version 4 against its status in version 8."),
              ("Follow-up list", "Listed locations to raise with the grantee: location ID and reason."),
              ("Category counts", "Locations per category on each map page."),
              ("Eligibility changes", "FCC locations in the area whose eligibility, ID or position changed."),
              ("Submitted locations", f"Every location in the {T_SUB4} or the {qr_term(key)}."),
              ("Data cleaning", "Rows skipped and duplicates removed before counting."), ("", ""),
              ("Key term", "Meaning")] + GLOSSARY
    kt = key_tables(key)
    write_workbook(path, readme, {"Summary": facts, **kt, **t,
                                  "Data cleaning": cleaning_table(cleaning[cleaning["project_key"] == key])})


WRITE_GPKG = False          # True adds a GeoPackage copy of the project layers (for non-ArcGIS users)
WRITE_FULL_PROJECT_GDB = False   # True also writes every technical layer and table, filtered, per project


def project_gdb(path: Path, key: str) -> list[str]:
    """Per-project file geodatabase: every unified layer and table, filtered to this project."""
    if path.exists():
        if path.suffix.lower() != ".gdb":
            raise ValueError(f"Refusing to delete {path}: not a .gdb folder")
        shutil.rmtree(path)
    written = []
    for name, df in GDB_LAYERS.items():
        part = df[df["project_key"] == key] if "project_key" in df.columns else df
        if len(part):
            write(part, name, gdb=path, quiet=True)
            written.append(name)
    return written


def project_gpkg(path: Path, key: str) -> None:
    """Optional GeoPackage with the project's map layers."""
    if path.exists():
        path.unlink()
    for i, name in enumerate(("project_areas", "page_extents", "map_A_eligibility_changes", "map_B_list_changes",
                              "map_C_follow_up")):
        part = GDB_LAYERS[name][GDB_LAYERS[name]["project_key"] == key]
        pyogrio.write_dataframe(gdb_ready(part), path, layer=name, driver="GPKG", append=i > 0)


def key_numbers_page(pdf, r: pd.Series, n_pages: int) -> None:
    """Page 2: how the change came together, then the key-number tables and the from-to tables."""
    key = r["project_key"]
    t = key_tables(key)
    fig = plt.figure(figsize=(PAGE_W, PAGE_H))
    page_frame(fig, f"{r['page_title']} · Key numbers", f"{cap(T_V4)} → {T_V8}; {T_SUB4} → {qr_term(key)}", 2, n_pages)
    top = PAGE_H - 1.15
    panel_heading(fig, "How the change came together", 0.35, top)
    fit_text(fig, r["story_why"], 0.35, top - 0.28, PAGE_W - 0.7, 0.72, size=9.5, min_size=7.5)
    y0 = top - 1.12
    rh = 0.185
    size = 8.0

    def table(df: pd.DataFrame, x: float, y: float, label_w: float, col_w: float, heads: list[str], fmts: list,
              hh: float = 0.42) -> float:
        rows = [[str(row.iloc[0])] + [f(v) for f, v in zip(fmts, row.iloc[1:])] for _, row in df.iterrows()]
        return draw_table(fig, x, y, [label_w] + [col_w] * (len(df.columns) - 1), [df.columns[0]] + heads, rows,
                          size=size, row_h=rh, header_h=hh)

    SHORT = {S_EI: "Eligible,\ninside", S_NI: "Not\neligible,\ninside", S_EO: "Eligible,\noutside",
             S_NO: "Not\neligible,\noutside", S_NOT: "Not in\nthe Fabric", S_CONF: "Conflicting\nsource\nrecords",
             S_NB: "Not\nserviceable", "Total": "Total"}

    def short(c: str) -> str:
        return SHORT.get(c, c)

    qd = qr_term(key)[:7]                                  # "2026 Q3"
    # Left column: Fabric in the area, coverage, Fabric from-to
    xl, wl = 0.35, 6.25
    A = t["Fabric in the area"]
    A = A[~A.iloc[:, 0].str.startswith("Units")]
    y = table(A, xl, y0, 2.95, 1.1, ["Version 4", "Version 8", "Change"], [fmt_n, fmt_n, fmt_d])
    D = t["Coverage"]
    y = table(D, xl, y - 0.14, 2.95, 1.1, ["Version 4\n(initial)", f"Version 8\n({qd})", "Change"], [fmt_n, fmt_n, fmt_d])
    F = t["From-to Fabric"]
    y = table(F.set_axis(["Fabric in the area: v4 → v8"] + list(F.columns[1:]), axis=1), xl, y - 0.14, 1.75,
              (wl - 1.75) / (len(F.columns) - 1), [short(c) for c in F.columns[1:]], [fmt_n] * (len(F.columns) - 1),
              hh=0.56)
    # Right column: grantee lists against both versions, then where the initial list's locations went
    xr = 6.85
    wr = PAGE_W - 0.35 - xr
    B = t["Grantee lists"]
    hb = ["Initial\nsubmission, v4", f"{qd}\ndata, v8", "Change", "Initial\nsubmission, v8", f"{qd}\ndata, v4"]
    yb = table(B, xr, y0, 2.55, (wr - 2.55) / 5, hb, [fmt_n, fmt_n, fmt_d, fmt_n, fmt_n])
    fig.text(fx(xr), fy(yb - 0.05), "Last two columns (check): each list looked up in the other Fabric version.",
             fontsize=7.5, color=QUIET, va="top")
    I = t["From-to initial list"]
    table(I.set_axis(["Initial submission: v4 → v8"] + list(I.columns[1:]), axis=1), xr, yb - 0.3, 1.75,
          (wr - 1.75) / (len(I.columns) - 1), [short(c) for c in I.columns[1:]], [fmt_n] * (len(I.columns) - 1), hh=0.56)
    pdf.savefig(fig)
    plt.close(fig)


PROJECTS_DIR.mkdir(parents=True, exist_ok=True)
package_files = {}
for _, r in summary.iterrows():
    key = r["project_key"]
    folder = PROJECTS_DIR / f"{key}_{r['pr_number']}"
    folder.mkdir(parents=True, exist_ok=True)
    stem = f"{key}_{r['pr_number']}"
    pts = project_points(key)
    with PdfPages(folder / f"{stem}_brief.pdf") as pdf:
        summary_page(pdf, r, 5)
        key_numbers_page(pdf, r, 5)
        for i, page in enumerate(("A", "B", "C"), start=3):
            map_page(pdf, r, page, pts[page], "gray", i, 5)
    with PdfPages(folder / f"{stem}_maps_imagery.pdf") as pdf:
        for i, page in enumerate(("A", "B", "C"), start=1):
            map_page(pdf, r, page, pts[page], "imagery", i, 3)
    results_workbook(folder / f"{stem}_results.xlsx", [key])
    if WRITE_MARKDOWN:
        write_markdown(folder / f"{stem}_results.md", [key])
    if WRITE_CSV:
        write_csvs(folder / "csv", [key])
    write_analyst_gdb(folder / f"{stem}_analyst.gdb", key)
    if WRITE_FULL_PROJECT_GDB:
        project_gdb(folder / f"{stem}_technical.gdb", key)
    one = summary[summary["project_key"] == key].drop(columns="geometry", errors="ignore").iloc[0]
    if WRITE_CSV:
        pd.DataFrame({"Measure": [FIELD_GUIDE.get(c, (c, c))[1] for c in one.index], "Value": one.to_numpy()}).to_csv(
            folder / f"{stem}_summary.csv", index=False, encoding="utf-8-sig")
    if WRITE_GPKG:
        project_gpkg(folder / f"{stem}.gpkg", key)
    package_files[key] = sorted(p.name for p in folder.iterdir())
    log.info("package → %s", folder)
pd.Series(package_files)

# %%
GLANCE_PDF = OUTPUT_DIR / "FFA_V4_V8_all_projects_at_a_glance.pdf"
GLANCE_W = [1.4, 0.95, 1.35, 1.2, 1.05, 0.9, 0.85, 0.9, 1.15, 1.0, 1.05, 1.5]    # inches, sums to page width − margins


def glance_page(path: Path) -> None:
    """All six projects on one landscape page; cells wrap and each row is as tall as its longest cell."""
    fig = plt.figure(figsize=(PAGE_W, PAGE_H))
    page_frame(fig, "All projects at a glance", f"{cap(T_V4)} → {T_V8}; {T_SUB4} → {T_SUB8}", 1, 1)
    renderer = fig.canvas.get_renderer()
    pad = 0.07
    probe = fig.text(0, 0, "", alpha=0, linespacing=1.25)
    size = 8.5

    def wrap(text: str, w: float, bold: bool = False) -> list[str]:
        probe.set_fontweight("bold" if bold else "normal")
        probe.set_fontsize(size)
        out = []
        for para in str(text).split("\n"):
            cur = ""
            for word in para.split():
                trial = (cur + " " + word).strip()
                probe.set_text(trial)
                if cur and probe.get_window_extent(renderer=renderer).width / fig.dpi > w - 2 * pad:
                    out.append(cur)
                    cur = word
                else:
                    cur = trial
            out.append(cur)
        return out

    # Largest font (8.5 down to 7 pt) at which the whole table fits above the notes
    for size in (8.5, 8.0, 7.5, 7.0):
        probe.set_fontsize(size)
        probe.set_text("Ag\nAg\nAg\nAg")
        line = probe.get_window_extent(renderer=renderer).height / fig.dpi / 4
        heads = [wrap(h, w, bold=True) for h, w in zip(GLANCE_COLS, GLANCE_W)]
        rows_w = [[wrap(v, w, bold=(j == 0)) for j, (v, w) in enumerate(zip(row, GLANCE_W))]
                  for row in glance.itertuples(index=False)]
        total = (max(len(h) for h in heads) + sum(max(len(c) for c in cells) for cells in rows_w)) * line \
            + 2 * pad * (len(rows_w) + 1)
        if PAGE_H - 1.1 - total > 1.2:
            break
    x0, y = 0.35, PAGE_H - 1.1
    hh = max(len(h) for h in heads) * line + 2 * pad
    fig.add_artist(Rectangle((fx(x0), fy(y - hh)), fx(sum(GLANCE_W)), fy(hh), facecolor=NAVY, edgecolor="none"))
    x = x0
    for h, w in zip(heads, GLANCE_W):
        fig.text(fx(x + pad), fy(y - pad), "\n".join(h), fontsize=size, fontweight="bold", color="white", va="top",
                 linespacing=1.25)
        x += w
    y -= hh
    for i, cells in enumerate(rows_w):
        rh = max(len(c) for c in cells) * line + 2 * pad
        if i % 2:
            fig.add_artist(Rectangle((fx(x0), fy(y - rh)), fx(sum(GLANCE_W)), fy(rh), facecolor="#f3f5f9", edgecolor="none"))
        x = x0
        for j, (c, w) in enumerate(zip(cells, GLANCE_W)):
            fig.text(fx(x + pad), fy(y - pad), "\n".join(c), fontsize=size, color=INK, va="top", linespacing=1.25,
                     fontweight="bold" if j == 0 else "normal")
            x += w
        y -= rh
        fig.add_artist(Line2D([fx(x0), fx(x0 + sum(GLANCE_W))], [fy(y)] * 2, color="#d0d0d0", lw=0.5))
    probe.remove()
    fit_text(fig, "Main reasons: the largest changes behind the eligible count, from the from-to table on page 2 of "
             "each brief. Suggested next step: generated from the change patterns; review it before sending it to a "
             "grantee.", x0, y - 0.15, sum(GLANCE_W), 0.5, size=7.5, min_size=7, color=QUIET)
    if y < 1.1:
        log.warning("All-projects table is taller than the page; check the PDF")
    with PdfPages(path) as pdf:
        pdf.savefig(fig)
    plt.close(fig)


glance_page(GLANCE_PDF)
log.info("All projects at a glance → %s", GLANCE_PDF)

# %% [markdown]
# ## 10. Integrity checks

# %%
s = summary.set_index("project_key")
# Comparison 4 category deltas reconcile exactly to each project's net Fabric change
recon = cmp4.groupby("project_key")[["delta_total", "delta_eligible"]].sum().reindex(s.index)
assert (recon["delta_total"] == s["total_chg"]).all(), "Fabric total change does not reconcile"
assert (recon["delta_eligible"] == s["elig_chg"]).all(), "Fabric eligible change does not reconcile"
# Comparison 3 categories reconcile to the submission counts
assert (s["s_kept"] + s["s_dropped"] == s["sub4_n"]).all(), "Initial grantee submission does not reconcile"
assert (s["s_kept"] + s["s_added"] == s["sub8_n"]).all(), "2026 Q2 and Q3 data does not reconcile"
# One row per location
assert not changes.duplicated(["project_key", "location_id"]).any()
assert not subs.duplicated(["project_key", "location_id"]).any()
# Scope: only the six projects
assert set(summary["project_key"]) == set(PROJECTS)
# Cleaning counts reconcile: rows read = skipped + duplicates removed + locations counted
assert (cleaning["rows_read"] == cleaning["non_numeric_rows_skipped"] + cleaning["duplicate_rows_removed"]
        + cleaning["unique_ids_used"]).all(), "Cleaning counts do not reconcile"
# Locations counted in each list = the list sizes used in the comparisons
for ds, col in (("Initial grantee submission", "sub4_n"), ("2026 Q2 and Q3 data", "sub8_n")):
    used = cleaning[cleaning["dataset"] == ds].set_index("project_key")["unique_ids_used"].reindex(s.index)
    assert (used == s[col]).all(), f"{ds} counts do not match the comparison"
# Every location has a plain map class
assert changes["map_class"].notna().all() and subs["map_class"].notna().all()
print("All integrity checks passed.")
