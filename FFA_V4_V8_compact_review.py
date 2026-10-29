"""FFA project locations – compact Fabric version 4 → version 8 review (six CPUC FFA projects).

For each project the script answers two questions, both checked against the awarded project polygon:
  1. Fabric: how did the locations inside the project area change from Fabric version 4 to version 8
     (eligible / not eligible, lost or gained eligibility, moved in or out, ID in one version only)?
  2. Grantee lists: how does the initial grantee submission compare with the 2026 quarterly list
     (IDs found in the Fabric, inside or outside the area, eligible, duplicated, missing or invalid IDs)?

Outputs in OUTPUT_DIR, one folder per project (e.g. SLO_PR-5578):
  <project>_locations.gdb   project_area, submitted_locations (Yes/No flags per location ID), fabric_changes
  <project>_results.xlsx    Summary, For review, List quality, Fabric changes, Definitions
  <project>_results.md      the same summary and review list as text
  <project>_maps_gray.pdf / _maps_imagery.pdf   three map pages each
plus FFA_V4_V8_all_projects_overview.xlsx (one row per project).

Run:  python FFA_V4_V8_compact_review.py      (edit only the Configuration section)
Needs: duckdb, geopandas, pyogrio, shapely, pandas, numpy, openpyxl, matplotlib (with Pillow), requests.
Basemaps are Esri tiles fetched over the internet; without access the maps draw on a plain background.
"""
from __future__ import annotations


import logging
import re
import shutil
import warnings
from pathlib import Path

import duckdb
import geopandas as gpd
import numpy as np
import pandas as pd
import pyogrio
import requests
import shapely

logging.basicConfig(level=logging.INFO, format="%(asctime)s  %(levelname)s  %(message)s", datefmt="%H:%M:%S")
log = logging.getLogger("ffa_v4_v8")
warnings.filterwarnings("ignore", category=FutureWarning)   # pandas deprecation notices only
pd.set_option("display.max_columns", 60)
pd.set_option("display.width", 200)

# =====================================================================================================================
# Configuration
# =====================================================================================================================
# %%
# ---------------------------------------------------------------- folders
BASE_DIR = Path(r"C:\Change_Analysis")
INPUT_DIR = BASE_DIR / "inputs"
OUTPUT_DIR = BASE_DIR / "results"             # one folder per project, plus the all-projects overview

# ---------------------------------------------------------------- maps
WRITE_MAPS = True                               # False skips the map PDFs (fastest run)
BASEMAP_KINDS = ("gray", "imagery")             # one map PDF per basemap
FAR_POINT_MI = 3                                # listed locations farther than this from the area are left off the map

# ---------------------------------------------------------------- the six projects (only these are analysed)
PROJECTS = {
    "SLO":       {"label": "San Luis Obispo Affordable Broadband Network",
                  "name_like": "%San Luis Obispo Affordable Broadband Network%"},
    "OAK":       {"label": "Oakland Connect - Last Mile Connectivity for Oaklanders",
                  "name_like": "%Oakland Connect%Last Mile Connectivity for Oaklanders%"},
    "ED_COLOMA": {"label": "Coloma Broadband Project",
                  "name_like": "%Coloma Broadband Project%"},
    "ED_INDCRK": {"label": "Indian Creek Broadband Project",
                  "name_like": "%Indian Creek Broadband Project%"},
    "ED_GRDNVL": {"label": "Garden Valley Broadband Project",
                  "name_like": "%Garden Valley Broadband Project%"},
    "ED_GRGTWN": {"label": "Georgetown Broadband Project",
                  "name_like": "%Georgetown Broadband Project%"},
}
PR_NUMBERS = {"SLO": "PR-5578", "OAK": "PR-4651", "ED_COLOMA": "PR-5631",
              "ED_INDCRK": "PR-5676", "ED_GRDNVL": "PR-5674", "ED_GRGTWN": "PR-5644"}

# ---------------------------------------------------------------- Fabric eligibility inputs
# CSV or file geodatabase (.gdb → set the *_LAYER name)
V4_PATH = INPUT_DIR / "v4_eligibility_locations.csv"   # or .../Combined_EligibilityV4_Updated_02192025.gdb
V4_LAYER = None
V8_PATH = INPUT_DIR / "FFA_v8_D25_eligibility_full.csv"
V8_LAYER = None

# Field roles in each Fabric file ("address" and "units" may be None; a field missing from the file is skipped).
# V4 eligibility = FFA vintage December 2023. If the exact name differs, the notebook
# looks for a single field containing both "FFA" and "Dec23" and uses it.
V4_FIELDS = {"id": "location_id", "lat": "latitude", "lon": "longitude",
             "elig": "FFAv4Dec23", "bsl": "bsl_flag", "address": "address_primary", "units": "unit_count"}
V8_FIELDS = {"id": "location_id", "lat": "latitude", "lon": "longitude",
             "elig": "status_code", "bsl": "bsl_flag", "address": None, "units": "unit_count"}

# Values in the eligibility field that mean "eligible" (unserved).
# V4 FFAv4Dec23 and V8 status_code are both coded 1 = served, 0 = unserved.
# Section 3a prints the actual distinct values: confirm these sets against it.
V4_ELIGIBLE_VALUES = {"0"}
V8_ELIGIBLE_VALUES = {"0"}

# False = every Fabric location counts, whatever its bsl_flag.
# True  = only bsl_flag = True locations count in totals.
BSL_ONLY = False

# ---------------------------------------------------------------- applicant submissions
# V4 submission: GMS application locations (Fabric V4 IDs)
SUB_V4_PATH = INPUT_DIR / "All_448_PIDs_Combined_Filtered_Only_113_Reward_Projects.csv"
SUB_V4_FIELDS = {"id": "location_id", "pr": "GMS_PID", "title": "Project_Title",
                 "lat": "Latitude", "lon": "Longitude", "address": "Address_primary"}

# V8 submission: 2026 quarterly reports (Fabric V8 IDs), one file per project, found by
# PR number (or project name) in the file name. Prefixes are tried in order.
SUB_V8_DIR = INPUT_DIR
SUB_V8_PREFERENCE = ["Q3 2026", "Q3", "Q2 2026", "Q2"]
SUB_V8_FIELDS = {"id": "Fabric_ID_Number__c", "lat": "Latitude__c", "lon": "Longitude__c"}
HEADER_SEARCH_ROWS = 50        # rows searched for the header in each report sheet
# First row (spreadsheet row number) that holds location data, by the start of the report file name.
# Q3 2026 reports carry template rows (label, "Required", help text) in rows 5–7 under the header; their
# location IDs start in row 8. Rows above the first data row are not read. None = the row after the header.
SUB_V8_FIRST_DATA_ROW = {"Q3": 8, "Q2": None}

# ---------------------------------------------------------------- project boundaries
# Local file geodatabase / GeoPackage / shapefile, or an ArcGIS feature service layer URL
# ending in /FeatureServer/<n> or /MapServer/<n>.
BOUNDARY_SOURCE = INPUT_DIR / "FFA_Projects_Awarded.gdb"
BOUNDARY_LAYER = None                              # None = read every layer that matches a project (one-layer-per-project gdb)
BOUNDARY_NAME_FIELDS = ["Project_Name", "Name"]    # name_like is tested against each
ARCGIS_TOKEN = None                                # only for secured services

# ---------------------------------------------------------------- tuning
DUCKDB_MEMORY_LIMIT = None      # e.g. "8GB"; None = DuckDB default (80% of RAM)
BBOX_PAD_DEG = 0.05            # read-window padding around the project areas (degrees)
WGS84 = "EPSG:4326"

# =====================================================================================================================
# Helpers
# =====================================================================================================================
def like_to_regex(pattern: str) -> re.Pattern:
    """Converts a SQL LIKE pattern (% and _) into a case-insensitive regex."""
    out = []
    for ch in pattern:
        if ch == "%":
            out.append(".*")
        elif ch == "_":
            out.append(".")
        else:
            out.append(re.escape(ch))
    return re.compile("^" + "".join(out) + "$", re.IGNORECASE | re.DOTALL)


def to_location_id(series: pd.Series) -> pd.Series:
    """Converts IDs such as '1012345678', 1012345678.0 or ' 1012345678 ' to Int64."""
    cleaned = series.astype("string").str.strip().str.replace(r"\.0+$", "", regex=True)
    cleaned = cleaned.str.replace(",", "", regex=False)
    num = pd.to_numeric(cleaned, errors="coerce")
    whole = num.where(np.isfinite(num) & (num == np.floor(num)))
    return whole.astype("Int64")


BSL_TRUE = {"1", "true", "t", "y", "yes"}


CLEANING: list[dict] = []            # one record per cleaned list: rows read, IDs dropped, duplicates removed
DROPPED: list[dict] = []             # every grantee-list row without a usable ID (file row number and raw value)
DUPLICATES: list[pd.DataFrame] = []   # every duplicated ID with how often it appeared


def standardize_ids(df: pd.DataFrame, source: str, prefer_col: str | None = None,
                    dataset: str | None = None, project_key: str | None = None, note: str = "") -> pd.DataFrame:
    """Converts location_id to whole numbers and removes duplicates.

    When an ID appears more than once, the record where ``prefer_col`` is True
    (the eligible record) is kept. With ``dataset``, the counts and the
    duplicated IDs are recorded in CLEANING / DUPLICATES for the outputs.
    """
    n0 = len(df)
    df = df.copy()
    raw = df["location_id"].astype("string").str.strip()
    blank = int((raw.isna() | raw.eq("")).sum())                 # no ID at all
    df["location_id"] = to_location_id(df["location_id"])
    bad = int(df["location_id"].isna().sum())                    # blank + not a whole number
    if dataset is not None:
        for idx in df.index[df["location_id"].isna()]:
            DROPPED.append({"dataset": dataset, "project_key": project_key,
                            "row": int(df.at[idx, "_row"]) if "_row" in df.columns and pd.notna(df.at[idx, "_row"]) else None,
                            "raw": "" if pd.isna(raw.at[idx]) else str(raw.at[idx]),
                            "issue": "No ID" if pd.isna(raw.at[idx]) or raw.at[idx] == "" else "Invalid ID (not a whole number)"})
    df = df.dropna(subset=["location_id"])
    counts = df["location_id"].value_counts()
    dup_ids = counts[counts > 1]
    if prefer_col is not None:
        df = df.sort_values(prefer_col, ascending=False, kind="stable")
    df = df.drop_duplicates("location_id", keep="first")
    n_dup = n0 - bad - len(df)
    log.info("%s: %s rows → %s unique IDs (%s non-numeric IDs dropped, %s duplicates removed)",
             source, f"{n0:,}", f"{len(df):,}", bad, n_dup)
    if dataset is not None:
        CLEANING.append({"dataset": dataset, "project_key": project_key, "rows_read": n0,
                         "non_numeric_rows_skipped": bad, "blank_id_rows": blank, "invalid_id_rows": bad - blank,
                         "ids_listed_more_than_once": len(dup_ids),
                         "duplicate_rows_removed": n_dup, "unique_ids_used": len(df), "note": note})
        if len(dup_ids):
            DUPLICATES.append(pd.DataFrame({"dataset": dataset, "project_key": project_key,
                                            "location_id": dup_ids.index.astype("int64"),
                                            "times_listed": dup_ids.to_numpy()}))
    return df.reset_index(drop=True)


def points_from_lonlat(df: pd.DataFrame, lon: str = "lon", lat: str = "lat") -> gpd.GeoDataFrame:
    """Builds a WGS84 point GeoDataFrame; rows without coordinates get empty geometry."""
    x = pd.to_numeric(df[lon], errors="coerce")
    y = pd.to_numeric(df[lat], errors="coerce")
    geom = gpd.points_from_xy(x, y)
    geom[(x.isna() | y.isna()).to_numpy()] = None
    return gpd.GeoDataFrame(df, geometry=geom, crs=WGS84)


def check_gdb(path: Path) -> Path:
    """Validates a file geodatabase folder and returns the path GDAL can open.

    Handles the nested folder created when a .gdb.zip is extracted on Windows
    (X.gdb\\X.gdb) and explains the usual causes of "Permission denied".
    Non-.gdb paths are returned unchanged.
    """
    path = Path(path)
    if path.suffix.lower() != ".gdb":
        return path
    if path.is_file():
        raise ValueError(f"{path} is a file, not a geodatabase folder. If it is a zip archive, extract it first.")
    if not path.exists():
        zipped = path.with_name(path.name + ".zip")
        hint = f" A zipped copy exists ({zipped.name}): extract it." if zipped.exists() else ""
        raise FileNotFoundError(f"{path} does not exist.{hint}")
    try:
        if any(path.glob("*.gdbtable")):
            return path
        nested = [p for p in path.rglob("*.gdb") if p.is_dir() and any(p.glob("*.gdbtable"))]
        contents = sorted(p.name for p in path.iterdir())[:15]
    except PermissionError as e:
        raise PermissionError(f"Windows denies access to {path}. Check that the folder is fully downloaded "
                              "(OneDrive/SharePoint: right-click → Always keep on this device) and readable "
                              "by this user.") from e
    if len(nested) == 1:
        log.warning("Using nested geodatabase %s (typical after extracting a .gdb.zip)", nested[0])
        return nested[0]
    raise PermissionError(
        f"{path} contains no .gdbtable files, so GDAL cannot open it as a geodatabase "
        f"(Windows reports this as 'Permission denied'). Folder contents: {contents}. "
        "Typical causes: the zip was extracted into a nested or wrong folder, the folder is an "
        "online-only OneDrive/SharePoint copy, or the path points to the parent folder instead of the .gdb.")


def pct_change(old: float, new: float) -> float:
    """Percent change from old to new; NaN when old is zero."""
    return round((new - old) / old * 100, 2) if old else np.nan


def table_columns(path: Path, layer: str | None) -> list[str]:
    """Field names of a CSV or GDAL-readable table without reading its rows."""
    path = check_gdb(path)
    if path.suffix.lower() in {".gdb", ".gpkg", ".shp"}:
        return list(pyogrio.read_info(path, layer=layer)["fields"])
    return list(pd.read_csv(path, nrows=0, encoding="utf-8-sig").columns)


def resolve_elig_field(path: Path, layer: str | None, wanted: str, hint: tuple[str, ...] = ()) -> str:
    """Returns the eligibility field name, falling back to a unique case-insensitive hint match."""
    cols = table_columns(path, layer)
    if wanted in cols:
        return wanted
    lower = {c.lower(): c for c in cols}
    if wanted.lower() in lower:
        return lower[wanted.lower()]
    if hint:
        cands = [c for c in cols if all(h.lower() in c.lower().replace("_", "") for h in hint)]
        if len(cands) == 1:
            log.warning("Field %r not in %s; using %r instead", wanted, path.name, cands[0])
            return cands[0]
    ffa = [c for c in cols if "ffa" in c.lower()]
    raise KeyError(f"Eligibility field {wanted!r} not found in {path.name}. FFA-related fields: {ffa}")

# =====================================================================================================================
# 1. Project boundaries
# =====================================================================================================================
def read_feature_service(url: str, token: str | None = None) -> gpd.GeoDataFrame:
    """Reads every feature of an ArcGIS REST layer as WGS84 GeoJSON, page by page."""
    url = url.rstrip("/")
    feats, offset = [], 0
    while True:
        params = {"where": "1=1", "outFields": "*", "outSR": 4326, "f": "geojson",
                  "resultOffset": offset, "resultRecordCount": 1000}
        if token:
            params["token"] = token
        r = requests.get(f"{url}/query", params=params, timeout=120)
        r.raise_for_status()
        payload = r.json()
        if "error" in payload:
            raise RuntimeError(f"Feature service error: {payload['error']}")
        batch = payload.get("features", [])
        feats.extend(batch)
        exceeded = payload.get("exceededTransferLimit") or payload.get("properties", {}).get("exceededTransferLimit")
        if not batch or not exceeded:
            break
        offset += len(batch)
    return gpd.GeoDataFrame.from_features(feats, crs=WGS84)


def count_project_matches(df: pd.DataFrame) -> int:
    """Number of the six projects whose name_like matches a value in the name fields or "_layer"."""
    fields = [f for f in BOUNDARY_NAME_FIELDS + ["_layer"] if f in df.columns]
    values = {str(v) for f in fields for v in df[f].dropna()}
    return sum(any(like_to_regex(p["name_like"]).match(v) for v in values) for p in PROJECTS.values())


def read_boundary_layers(path: Path) -> gpd.GeoDataFrame:
    """Reads every polygon layer whose name fields or layer name match one of the six projects.

    Works both for one combined project layer and for geodatabases that store one
    layer per project (e.g. P3616_San_Luis_Obispo_Affordable_Broadband_Network).
    The layer name, with underscores as spaces, is kept in field "_layer" and matched too.
    """
    found, parts = [], []
    for name, geom_type in pyogrio.list_layers(path):
        if geom_type is None or "polygon" not in str(geom_type).lower():
            continue
        fields = [f for f in BOUNDARY_NAME_FIELDS if f in pyogrio.read_info(path, layer=name)["fields"]]
        lname = name.replace("_", " ")
        vals = pyogrio.read_dataframe(path, layer=name, columns=fields, read_geometry=False) if fields else pd.DataFrame()
        n = count_project_matches(pd.concat([vals, pd.DataFrame({"_layer": [lname]})], ignore_index=True))
        if n:
            gdf = gpd.read_file(path, layer=name)
            if gdf.crs is None:
                log.warning("Layer %r has no CRS; skipped", name)
                continue
            gdf["_layer"] = lname
            parts.append(gdf.to_crs(WGS84))
            found.append((name, n, len(gdf)))
    if not parts:
        raise ValueError(f"No polygon layer in {path.name} matches any project name in {BOUNDARY_NAME_FIELDS} "
                         "or in the layer names. Set BOUNDARY_LAYER / BOUNDARY_NAME_FIELDS in the Configuration cell.")
    print(pd.DataFrame(found, columns=["layer used", "projects_matched", "features"]).to_string(index=False))
    return gpd.GeoDataFrame(pd.concat(parts, ignore_index=True), crs=WGS84)


def polygonal_2d(geom) -> shapely.MultiPolygon:
    """Keeps only the polygon parts of a geometry, drops Z/M, and returns a MultiPolygon.

    make_valid / union_all can return a GeometryCollection holding stray lines or
    points; the file geodatabase driver rejects that geometry type.
    """
    polys, stack = [], [shapely.force_2d(geom)]
    while stack:
        g = stack.pop()
        if g is None or g.is_empty:
            continue
        if g.geom_type == "Polygon":
            polys.append(g)
        elif g.geom_type in ("MultiPolygon", "GeometryCollection"):
            stack.extend(shapely.get_parts(g))
    return shapely.MultiPolygon(polys)


def load_boundaries() -> gpd.GeoDataFrame:
    """Returns one dissolved WGS84 polygon per project key, with matched names."""
    src = str(BOUNDARY_SOURCE)
    if src.lower().startswith("http"):
        raw = read_feature_service(src, ARCGIS_TOKEN)
    else:
        src_path = check_gdb(Path(BOUNDARY_SOURCE))
        raw = gpd.read_file(src_path, layer=BOUNDARY_LAYER) if BOUNDARY_LAYER else read_boundary_layers(src_path)
    if raw.crs is None:
        raise ValueError("Boundary layer has no CRS; define it before running.")
    raw = raw.to_crs(WGS84)
    raw["geometry"] = shapely.make_valid(raw.geometry.values)
    name_fields = [f for f in BOUNDARY_NAME_FIELDS + ["_layer"] if f in raw.columns]
    if not name_fields:
        raise KeyError(f"None of {BOUNDARY_NAME_FIELDS} found in boundary fields: {list(raw.columns)}")

    rows = []
    for key, p in PROJECTS.items():
        rx = like_to_regex(p["name_like"])
        hit = pd.Series(False, index=raw.index)
        for f in name_fields:
            hit |= raw[f].map(lambda v: bool(rx.match(str(v))) if pd.notna(v) else False).astype(bool)
        sub = raw[hit]
        if sub.empty:
            log.error("%s: no boundary matched %r in %s", key, p["name_like"], name_fields)
            continue
        names = sorted({str(v) for f in name_fields for v in sub[f].dropna()})
        print(f"{key:<10} {len(sub)} part(s) matched → {names}")
        rows.append({"project_key": key, "label": p["label"], "pr_number": PR_NUMBERS[key],
                     "n_parts": len(sub), "matched_names": " | ".join(names)[:500],
                     "geometry": polygonal_2d(shapely.union_all(sub.geometry.values))})
    missing = set(PROJECTS) - {r["project_key"] for r in rows}
    if missing:
        raise ValueError(f"No boundary found for: {sorted(missing)}. Check name_like / BOUNDARY_NAME_FIELDS.")
    return gpd.GeoDataFrame(rows, crs=WGS84)


areas = load_boundaries()


# =====================================================================================================================
# 2. Grantee lists
# =====================================================================================================================
def pr_digits(value) -> str:
    """Extracts the numeric part of a PR number ('PR-5578' → '5578')."""
    m = re.search(r"(\d+)", str(value))
    return m.group(1) if m else ""


def read_table(path: Path, header_field: str | None = None) -> pd.DataFrame:
    """Reads a CSV or Excel table with every column as text.

    With ``header_field``, the header row is the first row (in any sheet) that
    contains that field name. This handles report templates with title, label,
    "Required" and help-text rows above or below the header; those extra rows
    are dropped later because their IDs are not numeric.
    """
    is_excel = path.suffix.lower() in {".xlsx", ".xlsm", ".xls"}
    if header_field is None:
        return (pd.read_excel(path, dtype=str) if is_excel
                else pd.read_csv(path, dtype=str, encoding="utf-8-sig", low_memory=False))
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message=".*extension is not supported.*")
        sheets = (pd.read_excel(path, sheet_name=None, header=None, dtype=str) if is_excel
                  else {"csv": pd.read_csv(path, header=None, dtype=str, encoding="utf-8-sig", low_memory=False)})
    for sheet, raw in sheets.items():
        top = raw.head(HEADER_SEARCH_ROWS)
        hits = top.index[top.apply(lambda r: r.astype("string").str.strip().eq(header_field).any(), axis=1)]
        if len(hits):
            h = hits[0]
            df = raw.iloc[h + 1:].copy()
            df.columns = [str(c).strip() for c in raw.iloc[h]]
            df["_row"] = df.index + 1                     # row number in the sheet (1 = first row)
            df = df.dropna(how="all")
            log.info("%s: header found on sheet %r, row %s", path.name, sheet, h + 1)
            return df.reset_index(drop=True)
    raise KeyError(f"Field {header_field!r} not found in the first {HEADER_SEARCH_ROWS} rows of any sheet in {path.name}")


def load_sub_v4() -> pd.DataFrame:
    """V4 applicant submission for the six projects: project_key, location_id, lat, lon, address."""
    raw = read_table(SUB_V4_PATH)
    raw["_row"] = raw.index + 2                           # row number in the file (row 1 = header)
    f = SUB_V4_FIELDS
    raw["_pr"] = raw[f["pr"]].map(pr_digits) if f.get("pr") in raw else ""
    frames = []
    for key, pr in PR_NUMBERS.items():
        sub = raw[raw["_pr"] == pr_digits(pr)]
        if sub.empty and f.get("title") in raw:
            rx = like_to_regex(PROJECTS[key]["name_like"])
            sub = raw[raw[f["title"]].fillna("").map(lambda t: bool(rx.match(t)))]
            log.info("%s: no rows for %s in %s; matched %s rows by project title", key, pr, f["pr"], len(sub))
        out = pd.DataFrame({"project_key": key, "location_id": sub[f["id"]],
                            "lat": pd.to_numeric(sub[f["lat"]], errors="coerce"),
                            "lon": pd.to_numeric(sub[f["lon"]], errors="coerce"),
                            "address": sub[f["address"]] if f.get("address") in sub else pd.NA, "_row": sub["_row"]})
        frames.append(standardize_ids(out, f"V4 submission {key} ({pr})", dataset="Initial grantee submission",
                                      project_key=key, note=SUB_V4_PATH.name))
    return pd.concat(frames, ignore_index=True)


def find_sub_v8_file(key: str) -> Path | None:
    """Picks the 2026 quarterly report for a project, honouring SUB_V8_PREFERENCE order."""
    pat = re.compile(rf"PR[-_ ]?0*{pr_digits(PR_NUMBERS[key])}(?!\d)", re.IGNORECASE)
    label_rx = like_to_regex(PROJECTS[key]["name_like"])
    files = [p for p in SUB_V8_DIR.iterdir()
             if p.suffix.lower() in {".csv", ".xlsx", ".xlsm", ".xls"} and p.resolve() != Path(SUB_V4_PATH).resolve()]
    cands = [p for p in files if pat.search(p.stem)] or [p for p in files if label_rx.match(p.stem)]
    for prefix in SUB_V8_PREFERENCE:
        hits = sorted(p for p in cands if p.stem.lower().startswith(prefix.lower()))
        if hits:
            return hits[-1]
    return sorted(cands)[-1] if cands else None


def load_sub_v8() -> tuple[pd.DataFrame, dict[str, str]]:
    """V8 applicant submission (2026 quarterly reports) plus the file used per project."""
    f = SUB_V8_FIELDS
    frames, used = [], {}
    for key in PROJECTS:
        path = find_sub_v8_file(key)
        if path is None:
            log.warning("%s: no 2026 quarterly report found in %s", key, SUB_V8_DIR)
            used[key] = ""
            continue
        used[key] = path.name
        raw = read_table(path, header_field=f["id"])
        first = next((r for pre, r in SUB_V8_FIRST_DATA_ROW.items() if path.stem.upper().startswith(pre.upper())), None)
        if first:
            skipped = int((raw["_row"] < first).sum())
            raw = raw[raw["_row"] >= first]
            log.info("%s: data read from row %s (%s template rows above it not read)", path.name, first, skipped)
        first_row[key] = int(raw["_row"].min()) if len(raw) else None
        out = pd.DataFrame({"project_key": key, "location_id": raw[f["id"]],
                            "lat": pd.to_numeric(raw.get(f["lat"]), errors="coerce"),
                            "lon": pd.to_numeric(raw.get(f["lon"]), errors="coerce"), "address": pd.NA,
                            "_row": raw["_row"]})
        frames.append(standardize_ids(out, f"V8 submission {key} ← {path.name}", dataset="2026 Q2 and Q3 data",
                                      project_key=key,
                                      note=f"{path.name}; data read from row {first_row[key]}"))
    return pd.concat(frames, ignore_index=True), used


first_row: dict[str, int | None] = {}    # first data row read from each 2026 report
sub_v4 = load_sub_v4()
sub_v8, sub_v8_files = load_sub_v8()
submitted_ids = set(sub_v4["location_id"].astype("int64")) | set(sub_v8["location_id"].astype("int64"))


# =====================================================================================================================
# 3. Fabric version 4 and version 8 eligibility data (DuckDB)
# =====================================================================================================================
read_windows = [(a - BBOX_PAD_DEG, b - BBOX_PAD_DEG, c + BBOX_PAD_DEG, d + BBOX_PAD_DEG)
                for a, b, c, d in areas.geometry.bounds.to_numpy()]

V4_FIELDS["elig"] = resolve_elig_field(V4_PATH, V4_LAYER, V4_FIELDS["elig"], hint=("ffa", "dec23"))
V8_FIELDS["elig"] = resolve_elig_field(V8_PATH, V8_LAYER, V8_FIELDS["elig"])
print("V4 eligibility field:", V4_FIELDS["elig"], "| V8 eligibility field:", V8_FIELDS["elig"])

con = duckdb.connect()
if DUCKDB_MEMORY_LIMIT:
    con.execute(f"SET memory_limit = '{DUCKDB_MEMORY_LIMIT}'")
con.register("submitted", pd.DataFrame({"location_id": sorted(submitted_ids)}, dtype="int64"))


def q(name: str) -> str:
    """Quotes a column name as a SQL identifier."""
    return '"' + name.replace('"', '""') + '"'


def lit(text: str) -> str:
    """Quotes a value as a SQL string literal."""
    return "'" + str(text).replace("'", "''") + "'"


def norm_sql(expr: str) -> str:
    """SQL expression for the lower-cased, trimmed text of a value, with a trailing '.0' removed."""
    return f"regexp_replace(lower(trim(CAST({expr} AS VARCHAR))), '\\.0+$', '')"


def fabric_source(path: Path, layer: str | None, cols: list[str]) -> str:
    """Returns a SQL FROM-source for a Fabric file (CSV, or gdb/gpkg via spatial / pyogrio)."""
    path = check_gdb(path)
    if path.suffix.lower() not in {".gdb", ".gpkg"}:
        return f"read_csv({lit(path.as_posix())}, header = true, all_varchar = true)"
    try:
        con.execute("INSTALL spatial; LOAD spatial;")
        layer_arg = f", layer = {lit(layer)}" if layer else ""
        return f"ST_Read({lit(path.as_posix())}{layer_arg})"
    except duckdb.Error as e:
        log.warning("DuckDB spatial extension unavailable (%s); reading %s with pyogrio", e, path.name)
        con.register("gdb_src", pyogrio.read_dataframe(path, layer=layer, columns=cols, read_geometry=False))
        return "gdb_src"


def load_fabric(path: Path, layer: str | None, fields: dict, eligible_values: set[str], version: str,
                ids_table: str | None = None) -> pd.DataFrame:
    """Reads one Fabric eligibility file → location_id, lon, lat, eligible, bsl, counted, address, units, elig_raw.

    Default: rows inside the read windows plus every submitted ID. With ids_table: only the IDs in that
    registered table (used to look up IDs that sit outside the read windows), without cleaning records.
    """
    present = set(table_columns(path, layer))
    for opt in ("bsl", "address", "units"):
        if fields.get(opt) and fields[opt] not in present:
            log.warning("%s: field '%s' (%s) not in %s; skipped", version, fields[opt], opt, path.name)
            fields[opt] = None
    cols = sorted({c for c in (fields["id"], fields["lat"], fields["lon"], fields["elig"],
                               fields.get("bsl"), fields.get("address"), fields.get("units")) if c})
    src = fabric_source(path, layer, cols)
    lon, lat = f"TRY_CAST({q(fields['lon'])} AS DOUBLE)", f"TRY_CAST({q(fields['lat'])} AS DOUBLE)"
    window_sql = " OR ".join(f"(lon BETWEEN {a} AND {c} AND lat BETWEEN {b} AND {d})" for a, b, c, d in read_windows)
    elig_list = ", ".join(lit(v.lower()) for v in eligible_values)
    bsl_list = ", ".join(lit(v) for v in BSL_TRUE)
    bsl_expr = f"{norm_sql(q(fields['bsl']))} IN ({bsl_list})" if fields.get("bsl") else "TRUE"
    addr_expr = f"CAST({q(fields['address'])} AS VARCHAR)" if fields.get("address") else "CAST(NULL AS VARCHAR)"
    units_expr = f"TRY_CAST({q(fields['units'])} AS DOUBLE)" if fields.get("units") else "CAST(NULL AS DOUBLE)"
    where_sql = (f"id_num = floor(id_num) AND CAST(id_num AS BIGINT) IN (SELECT location_id FROM {ids_table})"
                 if ids_table else
                 f"({window_sql}) OR (id_num = floor(id_num) AND CAST(id_num AS BIGINT) IN (SELECT location_id FROM submitted))")
    table = f"fab_{version}" + ("_extra" if ids_table else "")

    log.info("DuckDB scan of %s (%s) …", version, path.name)
    con.execute(f"""
        CREATE OR REPLACE TEMP TABLE {table} AS
        WITH src AS (
            SELECT TRY_CAST(replace(trim(CAST({q(fields['id'])} AS VARCHAR)), ',', '') AS DOUBLE) AS id_num,
                   {lon} AS lon, {lat} AS lat,
                   CAST({q(fields['elig'])} AS VARCHAR) AS elig_raw,
                   {norm_sql(q(fields['elig']))} IN ({elig_list}) AS eligible,
                   {bsl_expr} AS bsl,
                   {addr_expr} AS address,
                   {units_expr} AS units
            FROM {src}
        )
        SELECT * FROM src
        WHERE {where_sql}
    """)
    if ids_table:
        df = con.execute(f"""
            SELECT CAST(id_num AS BIGINT) AS location_id, lon, lat, elig_raw, eligible, bsl, address, units,
                   bool_or(eligible) OVER w AND NOT bool_and(eligible) OVER w AS conflict
            FROM {table}
            WINDOW w AS (PARTITION BY id_num)
            QUALIFY row_number() OVER (PARTITION BY id_num ORDER BY eligible) = 1""").df()
        df["location_id"] = df["location_id"].astype("Int64")
        df["eligible"] = df["eligible"].fillna(False).astype(bool)
        df["bsl"] = df["bsl"].fillna(False).astype(bool)
        df["conflict"] = df["conflict"].fillna(False).astype(bool)
        df["counted"] = (df["bsl"] | (not BSL_ONLY)) & ~df["conflict"]
        return df
    n_rows, n_bad, n_unique = con.execute(f"""
        SELECT count(*),
               count(*) FILTER (WHERE id_num IS NULL OR id_num <> floor(id_num)),
               count(DISTINCT id_num) FILTER (WHERE id_num = floor(id_num))
        FROM fab_{version}""").fetchone()
    df = con.execute(f"""
        SELECT CAST(id_num AS BIGINT) AS location_id, lon, lat, elig_raw, eligible, bsl, address, units,
               bool_or(eligible) OVER w AND NOT bool_and(eligible) OVER w AS conflict
        FROM fab_{version}
        WHERE id_num = floor(id_num)
        WINDOW w AS (PARTITION BY id_num)
        QUALIFY row_number() OVER (PARTITION BY id_num ORDER BY eligible) = 1
        ORDER BY location_id""").df()
    log.info("Fabric %s: %s rows in scope → %s unique IDs (%s non-numeric IDs dropped, %s duplicates removed)",
             version, f"{n_rows:,}", f"{n_unique:,}", n_bad, n_rows - n_bad - n_unique)
    fabric_dups[version] = con.execute(f"""
        SELECT CAST(id_num AS BIGINT) AS location_id, count(*) AS times_listed,
               bool_or(eligible) AND NOT bool_and(eligible) AS eligibility_conflict
        FROM fab_{version}
        WHERE id_num = floor(id_num)
        GROUP BY id_num HAVING count(*) > 1""").df()
    CLEANING.append({"dataset": f"Fabric version {version[1]} eligibility data (read window, all projects)", "project_key": None,
                     "rows_read": int(n_rows), "non_numeric_rows_skipped": int(n_bad),
                     "ids_listed_more_than_once": len(fabric_dups[version]),
                     "duplicate_rows_removed": int(n_rows - n_bad - n_unique), "unique_ids_used": int(n_unique),
                     "note": f"{path.name}; duplicated IDs whose copies disagree on eligibility are marked as "
                             "conflicting source records and left out of every count"})
    df["location_id"] = df["location_id"].astype("Int64")
    df["eligible"] = df["eligible"].fillna(False).astype(bool)
    df["bsl"] = df["bsl"].fillna(False).astype(bool)
    df["conflict"] = df["conflict"].fillna(False).astype(bool)
    df["counted"] = (df["bsl"] | (not BSL_ONLY)) & ~df["conflict"]       # conflicting records are never counted
    log.info("Fabric %s: %s IDs with conflicting source records (left out of the counts)", version,
             f"{int(df['conflict'].sum()):,}")
    return df


fabric_dups: dict[str, pd.DataFrame] = {}
fabric = {
    "v4": load_fabric(V4_PATH, V4_LAYER, V4_FIELDS, V4_ELIGIBLE_VALUES, "v4"),
    "v8": load_fabric(V8_PATH, V8_LAYER, V8_FIELDS, V8_ELIGIBLE_VALUES, "v8"),
}
ATTR_COLS = ["lon", "lat", "eligible", "bsl", "counted", "conflict", "address", "units"]
attrs = {v: df.set_index("location_id")[ATTR_COLS] for v, df in fabric.items()}

for v, df in fabric.items():
    fields, values = (V4_FIELDS, V4_ELIGIBLE_VALUES) if v == "v4" else (V8_FIELDS, V8_ELIGIBLE_VALUES)
    vc = (df.groupby(df["elig_raw"].astype("string").fillna("<null>"))["eligible"]
            .agg(records="size", counted_as_eligible="sum").sort_values("records", ascending=False))
    print(f"\n{v.upper()} field '{fields['elig']}' → eligible values {sorted(values)}")
    print(vc.head(25).to_string())
    if df["eligible"].sum() == 0:
        warnings.warn(f"{v.upper()}: no record matched the eligible-value set; check the configuration.")

# =====================================================================================================================
# 4. Inside / outside the project areas, list change and Fabric change
# =====================================================================================================================
fabric_pts = {v: points_from_lonlat(df) for v, df in fabric.items()}
in_area = {}
for v, pts in fabric_pts.items():
    j = gpd.sjoin(pts[["location_id", "geometry"]].dropna(subset=["geometry"]),
                  areas[["project_key", "geometry"]], predicate="intersects", how="inner")
    in_area[v] = j[["project_key", "location_id"]].drop_duplicates().reset_index(drop=True)
area_sets = {v: {k: set(g["location_id"]) for k, g in in_area[v].groupby("project_key")} for v in fabric}

# An ID inside a project area in one version but missing from the other version's read window is looked
# up in the whole file of the other version, so "not in the Fabric" always means absent statewide.
SOURCES = {"v4": (V4_PATH, V4_LAYER, V4_FIELDS, V4_ELIGIBLE_VALUES), "v8": (V8_PATH, V8_LAYER, V8_FIELDS, V8_ELIGIBLE_VALUES)}
for v, other in (("v4", "v8"), ("v8", "v4")):
    need = sorted(set().union(*area_sets[other].values()) - set(attrs[v].index.astype("int64")))
    if not need:
        continue
    con.register(f"need_{v}", pd.DataFrame({"location_id": need}, dtype="int64"))
    extra = load_fabric(*SOURCES[v], v, ids_table=f"need_{v}")
    log.info("Fabric %s: %s of %s IDs from %s project areas found outside the read window",
             v, len(extra), len(need), other)
    if len(extra):
        fabric[v] = pd.concat([fabric[v], extra], ignore_index=True)
        attrs[v] = fabric[v].set_index("location_id")[ATTR_COLS]



ST_NOT_IN = "Not in Fabric"
ST_NOT_BSL = "Not a BSL"
ST_ELIG = "Eligible"
ST_NOT_ELIG = "Not eligible"
ST_ELIG_OUT = "Eligible, outside project area"
ST_NOT_ELIG_OUT = "Not eligible, outside project area"
ST_CONFLICT = "Conflicting source records"


def fabric_status(key: str, lid: int, v: str) -> str:
    """Status of one location ID in one Fabric version, relative to one project area."""
    a = attrs[v]
    if lid not in a.index:
        return ST_NOT_IN
    if a.at[lid, "conflict"]:
        return ST_CONFLICT
    if not a.at[lid, "counted"]:
        return ST_NOT_BSL
    elig = bool(a.at[lid, "eligible"])
    if lid in area_sets[v].get(key, set()):
        return ST_ELIG if elig else ST_NOT_ELIG
    return ST_ELIG_OUT if elig else ST_NOT_ELIG_OUT


for df, v in ((sub_v4, "v4"), (sub_v8, "v8")):          # each list against its own Fabric version
    df["fabric_status"] = [fabric_status(k, int(i), v) for k, i in zip(df["project_key"], df["location_id"])]


SC_KEPT = "Kept (same ID)"
SC_DROPPED = "Dropped from submission"
SC_ADDED = "Added to submission"
SUB_CATEGORY_ORDER = [SC_KEPT, SC_DROPPED, SC_ADDED]


def with_fallback_coords(df: pd.DataFrame, v: str) -> pd.DataFrame:
    """Fills missing submission coordinates from the Fabric version the IDs belong to."""
    df = df.copy()
    miss = (df["lat"].isna() | df["lon"].isna()) & df["location_id"].isin(attrs[v].index)
    idx = df.loc[miss, "location_id"]
    df.loc[miss, "lat"] = attrs[v].loc[idx, "lat"].to_numpy()
    df.loc[miss, "lon"] = attrs[v].loc[idx, "lon"].to_numpy()
    return df


s4 = with_fallback_coords(sub_v4, "v4")
s8 = with_fallback_coords(sub_v8, "v8")

frames = []
for key in PROJECTS:
    a = s4[s4["project_key"] == key].set_index("location_id")
    b = s8[s8["project_key"] == key].set_index("location_id")
    ids = sorted(set(a.index) | set(b.index))
    d = pd.DataFrame({"project_key": key, "location_id": pd.array(ids, dtype="Int64")})
    d["in_sub_v4"] = d["location_id"].isin(a.index).to_numpy()
    d["in_sub_v8"] = d["location_id"].isin(b.index).to_numpy()
    for side, src in (("v4", a), ("v8", b)):
        r = src.reindex(d["location_id"])
        d[f"sub_lon_{side}"], d[f"sub_lat_{side}"] = r["lon"].to_numpy(), r["lat"].to_numpy()
    d["sub_address_v4"] = a["address"].reindex(d["location_id"]).to_numpy()

    d["sub_change"] = np.select([d["in_sub_v4"] & d["in_sub_v8"], d["in_sub_v4"]], [SC_KEPT, SC_DROPPED],
                                default=SC_ADDED)
    frames.append(d)

subs = pd.concat(frames, ignore_index=True)

CAT_ELIG_LOST = "Eligible -> not eligible"
CAT_ELIG_GAINED = "Not eligible -> eligible"
CAT_RETIRED = "ID only in older version (retired)"
CAT_NEW = "ID only in newer version (new)"
CAT_MOVED = "Moved into or out of the project area"
CAT_BSL = "BSL status changed"
CAT_SAME = "Unchanged"
CAT_CONFLICT = "Conflicting source records"
CATEGORY_ORDER = [CAT_CONFLICT, CAT_ELIG_LOST, CAT_ELIG_GAINED, CAT_RETIRED, CAT_NEW, CAT_MOVED, CAT_BSL, CAT_SAME]


def compare_fabric(key: str) -> pd.DataFrame:
    """One row per Fabric location counted in the project area in V4 and/or V8, with its category."""
    counted, conflicted = {}, {}
    for v in ("v4", "v8"):
        m = pd.Index(sorted(area_sets[v].get(key, set())), dtype="Int64")
        counted[v] = set(m[attrs[v].loc[m, "counted"].to_numpy()])
        conflicted[v] = set(m[attrs[v].loc[m, "conflict"].to_numpy()])
    ids = counted["v4"] | counted["v8"] | conflicted["v4"] | conflicted["v8"]
    df = pd.DataFrame({"project_key": key, "location_id": pd.array(sorted(ids), dtype="Int64")})
    for v in ("v4", "v8"):
        a = attrs[v].reindex(df["location_id"])
        df[f"in_{v}"] = df["location_id"].isin(attrs[v].index).to_numpy()
        df[f"in_area_{v}"] = df["location_id"].isin(area_sets[v].get(key, set())).to_numpy()
        df[f"counted_{v}"] = df["location_id"].isin(counted[v]).to_numpy()
        df[f"conflict_{v}"] = df["location_id"].isin(conflicted[v]).to_numpy()
        df[f"bsl_{v}"] = a["bsl"].fillna(False).astype(bool).to_numpy()
        df[f"elig_{v}"] = a["eligible"].fillna(False).astype(bool).to_numpy()
        df[f"lon_{v}"], df[f"lat_{v}"] = a["lon"].to_numpy(), a["lat"].to_numpy()
        df[f"address_{v}"] = a["address"].to_numpy()
    conds = [df["conflict_v4"] | df["conflict_v8"], df["in_v4"] & ~df["in_v8"], ~df["in_v4"] & df["in_v8"],
             df["in_area_v4"] != df["in_area_v8"], df["counted_v4"] != df["counted_v8"],
             df["elig_v4"] & ~df["elig_v8"], ~df["elig_v4"] & df["elig_v8"]]
    df["category"] = np.select(conds, [CAT_CONFLICT, CAT_RETIRED, CAT_NEW, CAT_MOVED, CAT_BSL, CAT_ELIG_LOST,
                                       CAT_ELIG_GAINED],
                               default=CAT_SAME)
    df["delta_total"] = df["counted_v8"].astype(int) - df["counted_v4"].astype(int)
    df["delta_eligible"] = ((df["counted_v8"] & df["elig_v8"]).astype(int)
                            - (df["counted_v4"] & df["elig_v4"]).astype(int))
    df["in_sub_v4"] = df["location_id"].isin(set(sub_v4.loc[sub_v4["project_key"] == key, "location_id"])).to_numpy()
    df["in_sub_v8"] = df["location_id"].isin(set(sub_v8.loc[sub_v8["project_key"] == key, "location_id"])).to_numpy()
    return df


changes = pd.concat([compare_fabric(k) for k in PROJECTS], ignore_index=True)


# =====================================================================================================================
# 5. Location facts – one row per submitted location ID, flags only
# =====================================================================================================================
YN = {True: "Yes", False: "No"}
LIST_KEPT, LIST_REMOVED, LIST_ADDED = "In both lists", "Removed in the 2026 list", "Added in the 2026 list"
LIST_WORDS = {SC_KEPT: LIST_KEPT, SC_DROPPED: LIST_REMOVED, SC_ADDED: LIST_ADDED}

# Fabric change for one location ID between version 4 and version 8
E_CONFLICT, E_NEITHER, E_ONLY4, E_ONLY8 = ("Conflicting Fabric records", "In neither Fabric version",
                                          "ID only in Fabric 4", "ID only in Fabric 8")
E_LOST, E_GAINED, E_BOTH, E_NONE = "Lost eligibility", "Became eligible", "Eligible in both", "Not eligible in both"
A_IN_BOTH, A_OUT_BOTH, A_MOVED_IN, A_MOVED_OUT, A_NA = ("Inside in both", "Outside in both", "Moved inside the area",
                                                        "Moved outside the area", "Not comparable")

# Review reasons, highest priority first: (code, short reason, question for the grantee)
REVIEW = [
    ("not_in_fabric8", "Not in Fabric 8", "What is the Fabric version 8 location ID for this location?"),
    ("conflict", "Conflicting Fabric records", "Fabric copies of this ID disagree on eligibility; which applies?"),
    ("outside", "Outside the project area", "Why is a location outside the awarded area in the 2026 list?"),
    ("not_eligible", "Not eligible in Fabric 8", "How does a location that is not eligible in Fabric 8 qualify?"),
    ("not_bsl", "Not a serviceable location in Fabric 8", "Is this a broadband serviceable location?"),
    ("removed_eligible", "Removed but still eligible", "Why was this eligible location taken off the 2026 list?"),
    ("added_eligible", "Added but already eligible in Fabric 4",
     "Why was this already-eligible location not in the initial submission?"),
]
FU_LABEL = {c: lab for c, lab, _ in REVIEW}
FU_QUESTION = {c: q_ for c, _, q_ in REVIEW}


def place_words(eligible: bool, inside: bool) -> str:
    return f"{'eligible' if eligible else 'not eligible'}, {'inside' if inside else 'outside'} the area"


def version_facts(key: str, ids: pd.Series, v: str) -> pd.DataFrame:
    """Yes/No facts of each ID in one Fabric version, relative to one project area."""
    a = attrs[v].reindex(ids)
    present = ids.isin(attrs[v].index).to_numpy()
    conflict = a["conflict"].fillna(False).astype(bool).to_numpy()
    return pd.DataFrame({
        f"in_fabric_{v[1]}": present,
        f"inside_area_{v[1]}": ids.isin(area_sets[v].get(key, set())).to_numpy() & present,
        f"eligible_{v[1]}": a["eligible"].fillna(False).astype(bool).to_numpy() & present & ~conflict,
        f"serviceable_{v[1]}": a["bsl"].fillna(False).astype(bool).to_numpy() & present,
        f"conflict_{v[1]}": conflict,
        f"lon_f{v[1]}": a["lon"].to_numpy(), f"lat_f{v[1]}": a["lat"].to_numpy()})


def dup_counts(dataset: str, key: str) -> pd.Series:
    d = [x for x in DUPLICATES if len(x) and x["dataset"].iat[0] == dataset and x["project_key"].iat[0] == key]
    return pd.concat(d).set_index("location_id")["times_listed"] if d else pd.Series(dtype="int64")


frames = []
for key in PROJECTS:
    s = subs[subs["project_key"] == key].reset_index(drop=True)
    ids = s["location_id"].astype("int64")
    f = pd.concat([version_facts(key, ids, "v4"), version_facts(key, ids, "v8")], axis=1)
    d = pd.DataFrame({"project_key": key, "location_id": ids,
                      "in_initial_list": s["in_sub_v4"].to_numpy(), "in_2026_list": s["in_sub_v8"].to_numpy(),
                      "list_change": s["sub_change"].map(LIST_WORDS).to_numpy()})
    d = pd.concat([d, f], axis=1)
    d["times_in_initial_list"] = ids.map(dup_counts("Initial grantee submission", key)).fillna(
        d["in_initial_list"].astype(int)).astype(int).to_numpy()
    d["times_in_2026_list"] = ids.map(dup_counts("2026 Q2 and Q3 data", key)).fillna(
        d["in_2026_list"].astype(int)).astype(int).to_numpy()
    d["sub_lon"] = s["sub_lon_v8"].fillna(s["sub_lon_v4"]).to_numpy()
    d["sub_lat"] = s["sub_lat_v8"].fillna(s["sub_lat_v4"]).to_numpy()
    frames.append(d)
loc = pd.concat(frames, ignore_index=True)

in4, in8 = loc["in_fabric_4"], loc["in_fabric_8"]
e4, e8 = loc["eligible_4"], loc["eligible_8"]
i4, i8 = loc["inside_area_4"], loc["inside_area_8"]
conflict_any = loc["conflict_4"] | loc["conflict_8"]
loc["eligibility_change"] = np.select(
    [conflict_any, ~in4 & ~in8, ~in8, ~in4, e4 & e8, ~e4 & ~e8, e4 & ~e8],
    [E_CONFLICT, E_NEITHER, E_ONLY4, E_ONLY8, E_BOTH, E_NONE, E_LOST], default=E_GAINED)
loc["area_change"] = np.select(
    [~(in4 & in8), i4 & i8, ~i4 & ~i8, ~i4 & i8], [A_NA, A_IN_BOTH, A_OUT_BOTH, A_MOVED_IN], default=A_MOVED_OUT)
loc["duplicated"] = (loc["times_in_initial_list"] > 1) | (loc["times_in_2026_list"] > 1)


def fabric_words(r: pd.Series, n: str) -> str:
    if not r[f"in_fabric_{n}"]:
        return "not in the Fabric"
    if r[f"conflict_{n}"]:
        return "conflicting records"
    return place_words(r[f"eligible_{n}"], r[f"inside_area_{n}"])


loc["fabric_4"] = loc.apply(lambda r: fabric_words(r, "4"), axis=1)
loc["fabric_8"] = loc.apply(lambda r: fabric_words(r, "8"), axis=1)

# Review flags (2026 list: current exceptions; removed / added: unexplained list changes)
cur = loc["in_2026_list"]
removed, added = loc["list_change"].eq(LIST_REMOVED), loc["list_change"].eq(LIST_ADDED)
usable8 = in8 & ~loc["conflict_8"]
flags = {
    "not_in_fabric8": cur & ~in8,
    "conflict": (cur & loc["conflict_8"]) | ((removed | added) & conflict_any),
    "outside": cur & usable8 & ~i8,
    "not_eligible": cur & usable8 & ~e8,
    "not_bsl": cur & usable8 & ~loc["serviceable_8"],
    "removed_eligible": removed & usable8 & i8 & e8,
    "added_eligible": added & in4 & ~loc["conflict_4"] & i4 & e4,
}
for c, m in flags.items():
    loc[f"fu_{c}"] = m.to_numpy()
fu_cols = [f"fu_{c}" for c, _, _ in REVIEW]
loc["for_review"] = loc[fu_cols].any(axis=1)
loc["review_reason"] = loc.apply(lambda r: "; ".join(FU_LABEL[c] for c, _, _ in REVIEW if r[f"fu_{c}"]), axis=1)
loc["primary_reason"] = loc.apply(lambda r: next((c for c, _, _ in REVIEW if r[f"fu_{c}"]), ""), axis=1)
loc["question"] = loc["primary_reason"].map(FU_QUESTION).fillna("")


def what_changed(r: pd.Series) -> str:
    """One plain sentence: list change, then the location's place in each Fabric version."""
    parts = [r["list_change"]]
    if r["eligibility_change"] in (E_LOST, E_GAINED, E_ONLY4, E_ONLY8, E_NEITHER, E_CONFLICT):
        e = r["eligibility_change"]
        parts.append(e if e.startswith("ID") else e[0].lower() + e[1:])
    if r["area_change"] in (A_MOVED_IN, A_MOVED_OUT):
        parts.append(r["area_change"].lower())
    return f"{'; '.join(parts)}. Fabric 4: {r['fabric_4']}. Fabric 8: {r['fabric_8']}."


loc["what_changed"] = loc.apply(what_changed, axis=1)
loc["status_2026"] = np.where(~cur, "Not in the 2026 list",
                              np.where(loc["for_review"] & cur, "For review", "Eligible, inside the area"))

# Map point: Fabric 8 point, else Fabric 4 point, else the grantee's own coordinates
loc["lon"] = loc["lon_f8"].fillna(loc["lon_f4"]).fillna(loc["sub_lon"])
loc["lat"] = loc["lat_f8"].fillna(loc["lat_f4"]).fillna(loc["sub_lat"])
loc["point_from"] = np.select([loc["lon_f8"].notna(), loc["lon_f4"].notna(), loc["sub_lon"].notna()],
                              ["Fabric 8", "Fabric 4", "Grantee coordinates"], default="No coordinates")

# =====================================================================================================================
# 6. Fabric changes inside each project area (every Fabric location, submitted or not)
# =====================================================================================================================
F_CONFLICT, F_ONLY4, F_ONLY8, F_IN, F_OUT, F_BSL, F_LOST, F_GAINED, F_SAME_E, F_SAME_N = (
    "Conflicting Fabric records", "ID only in Fabric 4", "ID only in Fabric 8", "Moved into the area",
    "Moved out of the area", "Serviceable-location status changed", "Lost eligibility", "Became eligible",
    "Unchanged – eligible", "Unchanged – not eligible")
fc = changes.copy()
fc["change"] = fc["category"].map({CAT_CONFLICT: F_CONFLICT, CAT_RETIRED: F_ONLY4, CAT_NEW: F_ONLY8, CAT_BSL: F_BSL,
                                   CAT_ELIG_LOST: F_LOST, CAT_ELIG_GAINED: F_GAINED})
fc.loc[fc["category"].eq(CAT_MOVED), "change"] = np.where(fc.loc[fc["category"].eq(CAT_MOVED), "in_area_v8"], F_IN, F_OUT)
fc.loc[fc["category"].eq(CAT_SAME), "change"] = np.where(fc.loc[fc["category"].eq(CAT_SAME), "elig_v8"], F_SAME_E, F_SAME_N)
fc["eligible_4"] = fc["counted_v4"] & fc["elig_v4"]
fc["eligible_8"] = fc["counted_v8"] & fc["elig_v8"]
fc["lon"], fc["lat"] = fc["lon_v8"].fillna(fc["lon_v4"]), fc["lat_v8"].fillna(fc["lat_v4"])
FABRIC_CHANGE_ORDER = [F_LOST, F_GAINED, F_IN, F_OUT, F_ONLY4, F_ONLY8, F_CONFLICT, F_BSL, F_SAME_E, F_SAME_N]

# =====================================================================================================================
# 7. Key facts per project (every table, map description and report reads from here)
# =====================================================================================================================
def cleaning_row(dataset: str, key: str) -> dict:
    rows = [c for c in CLEANING if c["dataset"] == dataset and c["project_key"] == key]
    return rows[0] if rows else {}


def list_check(key: str, df: pd.DataFrame) -> dict:
    """One grantee list against its own Fabric version (Fabric point decides inside / outside)."""
    st = df.loc[df["project_key"] == key, "fabric_status"]
    return {"listed": len(st), "found": int(st.ne(ST_NOT_IN).sum()), "not_found": int(st.eq(ST_NOT_IN).sum()),
            "in_elig": int(st.eq(ST_ELIG).sum()), "in_nelig": int(st.eq(ST_NOT_ELIG).sum()),
            "out_elig": int(st.eq(ST_ELIG_OUT).sum()), "out_nelig": int(st.eq(ST_NOT_ELIG_OUT).sum()),
            "conflict": int(st.eq(ST_CONFLICT).sum()), "not_bsl": int(st.eq(ST_NOT_BSL).sum())}


FACTS: dict[str, dict] = {}
for key in PROJECTS:
    k: dict = {}
    c = fc[fc["project_key"] == key]
    for n, v in (("4", "v4"), ("8", "v8")):
        k[f"f{n}_total"] = int(c[f"counted_{v}"].sum())
        k[f"f{n}_elig"] = int((c[f"counted_{v}"] & c[f"elig_{v}"]).sum())
        k[f"f{n}_nelig"] = k[f"f{n}_total"] - k[f"f{n}_elig"]
    k["f_conflict"] = int(c["change"].eq(F_CONFLICT).sum())
    k["f_changes"] = c["change"].value_counts().to_dict()
    k["f_elig_effect"] = c.groupby("change")["delta_eligible"].sum().to_dict()
    assert k["f4_elig"] + sum(k["f_elig_effect"].values()) == k["f8_elig"], f"{key}: eligible counts do not reconcile"
    for lst, df, ds in (("s4", sub_v4, "Initial grantee submission"), ("s8", sub_v8, "2026 Q2 and Q3 data")):
        cl = cleaning_row(ds, key)
        k[f"{lst}_rows"] = int(cl.get("rows_read", 0))
        k[f"{lst}_no_id"] = int(cl.get("blank_id_rows", 0))
        k[f"{lst}_bad_id"] = int(cl.get("invalid_id_rows", 0))
        k[f"{lst}_dup_rows"] = int(cl.get("duplicate_rows_removed", 0))
        k[f"{lst}_dup_ids"] = int(cl.get("ids_listed_more_than_once", 0))
        for m, val in list_check(key, df).items():
            k[f"{lst}_{m}"] = val
        k[f"{lst}_valid_rows"] = k[f"{lst}_rows"] - k[f"{lst}_no_id"] - k[f"{lst}_bad_id"]
        # entries reconcile: entries − no ID − invalid ID − repeat entries = unique location IDs
        assert k[f"{lst}_valid_rows"] - k[f"{lst}_dup_rows"] == k[f"{lst}_listed"], f"{key} {lst}: entries do not reconcile"
    l_ = loc[loc["project_key"] == key]
    k["kept"], k["removed"], k["added"] = (int(l_["list_change"].eq(x).sum()) for x in (LIST_KEPT, LIST_REMOVED, LIST_ADDED))
    kept = l_[l_["list_change"].eq(LIST_KEPT)]
    k["kept_lost"] = int(kept["eligibility_change"].eq(E_LOST).sum())
    k["kept_moved_out"] = int(kept["area_change"].eq(A_MOVED_OUT).sum())
    k["fu"] = {c_: int(l_[f"fu_{c_}"].sum()) for c_, _, _ in REVIEW}
    k["fu_total"] = int(l_["for_review"].sum())
    k["cov_listed"], k["cov_eligible"] = k["s8_in_elig"], k["f8_elig"]
    FACTS[key] = k


def n_(x: int, word: str) -> str:
    """Count with its noun, plural when needed ('1 row', '3 rows', '2 entries')."""
    plural = word[:-1] + "ies" if word.endswith("y") and word[-2:-1] not in "aeiou" else word + "s"
    return f"{x:,} {word if x == 1 else plural}"


def signed(x: int) -> str:
    return f"{x:+,}" if x else "0"


def quarter(key: str) -> str:
    """'2026 Q3' or '2026 Q2' from the quarterly report used."""
    name = sub_v8_files.get(key, "").upper()
    return "2026 Q3" if "Q3" in name else "2026 Q2" if "Q2" in name else "2026"


def cap_first(text: str) -> str:
    return text[:1].upper() + text[1:]


def story(key: str) -> dict[str, str]:
    """Short plain sentences: Fabric, lists, current list, data quality, review."""
    k, qr = FACTS[key], quarter(key)
    ch, ef = k["f_changes"], k["f_elig_effect"]
    drivers = [f"{ch.get(F_LOST, 0):,} lost eligibility" if ch.get(F_LOST) else "",
               f"{ch.get(F_GAINED, 0):,} became eligible" if ch.get(F_GAINED) else "",
               f"{-ef.get(F_ONLY4, 0):,} eligible IDs are only in Fabric 4" if ef.get(F_ONLY4) else "",
               f"{ef.get(F_ONLY8, 0):,} eligible IDs are only in Fabric 8" if ef.get(F_ONLY8) else "",
               f"{n_(ch.get(F_IN, 0), 'point')} moved into the area" if ch.get(F_IN) else "",
               f"{n_(ch.get(F_OUT, 0), 'point')} moved out of the area" if ch.get(F_OUT) else "",
               f"{ch.get(F_CONFLICT, 0):,} {'has' if ch.get(F_CONFLICT) == 1 else 'have'} conflicting records"
               if ch.get(F_CONFLICT) else ""]
    drivers = [d for d in drivers if d]
    fabric = (f"Inside the project area, eligible locations went from {k['f4_elig']:,} in Fabric 4 to {k['f8_elig']:,} in "
              f"Fabric 8 ({signed(k['f8_elig'] - k['f4_elig'])})." + (" Of the changes: " + "; ".join(drivers) + "."
                                                                     if drivers else " No location changed."))
    def entries_words(l: str, name: str) -> str:
        extra = [n_(k[l + '_dup_rows'], 'repeat entry') if k[l + "_dup_rows"] else "",
                 f"{k[l + '_no_id']:,} without an ID" if k[l + "_no_id"] else "",
                 f"{k[l + '_bad_id']:,} with an invalid ID" if k[l + "_bad_id"] else ""]
        extra = [e for e in extra if e]
        return (f"the {name} has {n_(k[l + '_rows'], 'entry')} for "
                f"{n_(k[l + '_listed'], 'unique location ID')}" + (f" ({'; '.join(extra)})" if extra else ""))

    lists = (cap_first(entries_words("s4", "initial list") + "; " + entries_words("s8", f"{qr} list")) + ". "
             f"Comparing unique IDs: {k['kept']:,} in both, {k['removed']:,} removed, {k['added']:,} added."
             + (f" Of the locations in both, {k['kept_lost']:,} lost eligibility and {k['kept_moved_out']:,} moved outside "
                "the area." if k["kept_lost"] or k["kept_moved_out"] else ""))
    out8 = k["s8_out_elig"] + k["s8_out_nelig"]
    current = (f"Checked against Fabric 8, {k['s8_in_elig']:,} of the {k['s8_listed']:,} locations in the {qr} list are "
               f"eligible inside the area; {out8:,} {'is' if out8 == 1 else 'are'} outside it, {k['s8_in_nelig']:,} are inside but not eligible"
               f" and {k['s8_not_found']:,} are not in Fabric 8."
               f" The list covers {k['cov_listed']:,} of the {k['cov_eligible']:,} eligible locations in the area.")
    dq_items = [f"{n_(k[l + '_no_id'], 'entry')} without an ID in the {w}" for l, w in (("s4", "initial list"), ("s8", f"{qr} list"))
                if k[l + "_no_id"]]
    dq_items += [f"{n_(k[l + '_bad_id'], 'entry')} with an invalid ID in the {w}" for l, w in (("s4", "initial list"), ("s8", f"{qr} list"))
                 if k[l + "_bad_id"]]
    dq_items += [f"{n_(k[l + '_dup_ids'], 'ID')} listed on more than one row in the {w}" for l, w in (("s4", "initial list"), ("s8", f"{qr} list"))
                 if k[l + "_dup_ids"]]
    quality = ("List quality: " + "; ".join(dq_items) + "." if dq_items else "List quality: no missing, invalid or "
               "duplicated IDs.")
    top = [FU_LABEL[c][0].lower() + FU_LABEL[c][1:] for c, n in sorted(k["fu"].items(), key=lambda x: -x[1]) if n][:2]
    follow = (f"{n_(k['fu_total'], 'location')} {'is' if k['fu_total'] == 1 else 'are'} flagged for review"
              + (f", mostly: {' and '.join(top)}." if top else "."))
    return {"fabric": fabric, "lists": lists, "current": current, "quality": quality, "follow": follow}


STORIES = {key: story(key) for key in PROJECTS}
for key in PROJECTS:
    log.info("%s: %s", key, STORIES[key]["follow"])


# =====================================================================================================================
# 8. Output tables (shared by the geodatabase, the workbook and the Markdown report)
# =====================================================================================================================
RUN_DATE = pd.Timestamp.today().strftime("%B %d, %Y")
V4_FIELD_NAME, V8_FIELD_NAME = V4_FIELDS["elig"], V8_FIELDS["elig"]
LABEL = {k: v["label"] for k, v in PROJECTS.items()}
FOLDER = {k: f"{k}_{PR_NUMBERS[k]}" for k in PROJECTS}

LOCATION_FIELDS = {   # geodatabase field: column in `loc` (Yes/No fields are written as text)
    "location_id": "location_id", "in_initial_list": "in_initial_list", "in_2026_list": "in_2026_list",
    "list_change": "list_change", "times_in_initial_list": "times_in_initial_list",
    "times_in_2026_list": "times_in_2026_list", "listed_more_than_once": "duplicated",
    "in_fabric_4": "in_fabric_4", "inside_area_fabric_4": "inside_area_4", "eligible_fabric_4": "eligible_4",
    "in_fabric_8": "in_fabric_8", "inside_area_fabric_8": "inside_area_8", "eligible_fabric_8": "eligible_8",
    "serviceable_location_fabric_8": "serviceable_8", "conflicting_fabric_records": "conflict_any",
    "eligibility_change": "eligibility_change", "area_change": "area_change", "fabric_4": "fabric_4",
    "fabric_8": "fabric_8", "what_changed": "what_changed", "status_2026_list": "status_2026",
    "for_review": "for_review", "review_reason": "review_reason", "question_for_grantee": "question",
    "point_from": "point_from", "lon": "lon", "lat": "lat"}
loc["conflict_any"] = loc["conflict_4"] | loc["conflict_8"]
YES_NO_FIELDS = {"in_initial_list", "in_2026_list", "listed_more_than_once", "in_fabric_4", "inside_area_fabric_4",
                 "eligible_fabric_4", "in_fabric_8", "inside_area_fabric_8", "eligible_fabric_8",
                 "serviceable_location_fabric_8", "conflicting_fabric_records", "for_review"}


def yes_no(df: pd.DataFrame, cols) -> pd.DataFrame:
    df = df.copy()
    for c in cols:
        if c in df.columns:
            df[c] = df[c].fillna(False).astype(bool).map(YN)
    return df


def location_layer(key: str) -> gpd.GeoDataFrame:
    d = loc[loc["project_key"] == key].copy()
    d["_prio"] = d["primary_reason"].map({c: i for i, (c, _, _) in enumerate(REVIEW)}).fillna(99)
    d = d.sort_values(["_prio", "location_id"])
    out = d[list(LOCATION_FIELDS.values())].set_axis(list(LOCATION_FIELDS), axis=1)
    return points_from_lonlat(yes_no(out, YES_NO_FIELDS))


def fabric_change_table(key: str, changed_only: bool = True) -> pd.DataFrame:
    c = fc[fc["project_key"] == key]
    if changed_only:
        c = c[~c["change"].isin([F_SAME_E, F_SAME_N])]
    c = c.assign(_o=c["change"].map({x: i for i, x in enumerate(FABRIC_CHANGE_ORDER)})).sort_values(["_o", "location_id"])
    out = pd.DataFrame({"location_id": c["location_id"].astype("int64"), "change": c["change"],
                        "eligible_fabric_4": c["eligible_4"], "eligible_fabric_8": c["eligible_8"],
                        "inside_area_fabric_4": c["in_area_v4"] & c["in_v4"],
                        "inside_area_fabric_8": c["in_area_v8"] & c["in_v8"],
                        "in_initial_list": c["in_sub_v4"], "in_2026_list": c["in_sub_v8"],
                        "lon": c["lon"], "lat": c["lat"]})
    return yes_no(out, [x for x in out.columns if x.startswith(("eligible", "inside", "in_"))])


def list_quality(key: str) -> pd.DataFrame:
    """Rows that cannot be placed on a map: no ID, invalid ID; plus duplicated IDs and IDs in neither Fabric."""
    qr = quarter(key)
    names = {"Initial grantee submission": "Initial list", "2026 Q2 and Q3 data": f"{qr} list"}
    rows = [{"List": names[d["dataset"]], "Issue": d["issue"], "Row in file": d["row"], "Value in the ID field": d["raw"],
             "Location ID": None, "Times listed": None}
            for d in DROPPED if d["project_key"] == key]
    for ds, nm in names.items():
        for lid, n in dup_counts(ds, key).items():
            rows.append({"List": nm, "Issue": "Listed more than once (counted once)", "Row in file": None,
                         "Value in the ID field": None, "Location ID": int(lid), "Times listed": int(n)})
    l_ = loc[(loc["project_key"] == key) & ~loc["in_fabric_4"] & ~loc["in_fabric_8"]]
    for r in l_.itertuples():
        rows.append({"List": "Both lists" if r.in_initial_list and r.in_2026_list else
                     ("Initial list" if r.in_initial_list else f"{qr} list"), "Issue": "ID in neither Fabric version",
                     "Row in file": None, "Value in the ID field": None, "Location ID": int(r.location_id),
                     "Times listed": None})
    return pd.DataFrame(rows, columns=["List", "Issue", "Row in file", "Value in the ID field", "Location ID",
                                       "Times listed"])


def review_table(key: str) -> pd.DataFrame:
    d = loc[(loc["project_key"] == key) & loc["for_review"]].copy()
    d["_prio"] = d["primary_reason"].map({c: i for i, (c, _, _) in enumerate(REVIEW)})
    d = d.sort_values(["_prio", "location_id"])
    return pd.DataFrame({"Location ID": d["location_id"].astype("int64"), "Reason": d["review_reason"],
                         "What changed": d["what_changed"], "Question for the grantee": d["question"]})


def summary_tables(key: str) -> dict[str, pd.DataFrame]:
    """The four small tables of the Summary tab and the Markdown report."""
    k, qr = FACTS[key], quarter(key)
    fabric = pd.DataFrame([("Locations", k["f4_total"], k["f8_total"]), ("Eligible", k["f4_elig"], k["f8_elig"]),
                           ("Not eligible", k["f4_nelig"], k["f8_nelig"])], columns=["Inside the project area", "Fabric 4", "Fabric 8"])
    fabric["Change"] = fabric["Fabric 8"] - fabric["Fabric 4"]
    bridge = pd.DataFrame([(x, k["f_changes"].get(x, 0), int(k["f_elig_effect"].get(x, 0)))
                           for x in FABRIC_CHANGE_ORDER if k["f_changes"].get(x, 0)],
                          columns=["What happened to the location", "Locations", "Effect on the eligible count"])
    bridge.loc[len(bridge)] = ["Total", int(bridge["Locations"].sum()), int(bridge["Effect on the eligible count"].sum())]
    entries = pd.DataFrame([
        ("Entries (rows) in the file", k["s4_rows"], k["s8_rows"]),
        ("– Entries without an ID", k["s4_no_id"], k["s8_no_id"]),
        ("– Entries with an invalid ID (not a whole number)", k["s4_bad_id"], k["s8_bad_id"]),
        ("= Entries with a valid location ID", k["s4_valid_rows"], k["s8_valid_rows"]),
        ("– Repeat entries (an ID already listed in an earlier row)", k["s4_dup_rows"], k["s8_dup_rows"]),
        ("= Unique location IDs (used in this analysis)", k["s4_listed"], k["s8_listed"]),
        ("Location IDs that appear on more than one row", k["s4_dup_ids"], k["s8_dup_ids"])],
        columns=["Entries in each grantee list", "Initial list", f"{qr} list"])
    lists = pd.DataFrame([
        ("Unique location IDs", k["s4_listed"], k["s8_listed"]),
        ("Not in the Fabric", k["s4_not_found"], k["s8_not_found"]),
        ("Conflicting Fabric records", k["s4_conflict"], k["s8_conflict"]),
        ("Inside the area – eligible", k["s4_in_elig"], k["s8_in_elig"]),
        ("Inside the area – not eligible", k["s4_in_nelig"], k["s8_in_nelig"]),
        ("Outside the area – eligible", k["s4_out_elig"], k["s8_out_elig"]),
        ("Outside the area – not eligible", k["s4_out_nelig"], k["s8_out_nelig"])],
        columns=["Unique location IDs checked in the Fabric", "Initial list (checked in Fabric 4)",
                 f"{qr} list (checked in Fabric 8)"])
    change = pd.DataFrame([(LIST_KEPT, k["kept"]), (LIST_REMOVED, k["removed"]), (LIST_ADDED, k["added"])],
                          columns=["Initial list → " + qr + " list", "Locations"])
    fu = pd.DataFrame([(FU_LABEL[c], k["fu"][c]) for c, _, _ in REVIEW if k["fu"][c]],
                      columns=["Flagged for review (a location can have several reasons)", "Locations"])
    fu.loc[len(fu)] = ["Locations for review", k["fu_total"]]
    return {"Fabric locations inside the project area": fabric, "Why the eligible count changed": bridge,
            "Grantee list entries": entries, "Grantee lists checked in the Fabric": lists, "How the list changed": change, "Reasons for review": fu}


DEFINITIONS = [
    ("Fabric 4 / Fabric 8", f"FCC location data used for FFA eligibility: version 4 (field {V4_FIELD_NAME}) and version 8 "
                            f"(field {V8_FIELD_NAME}); 0 = unserved (eligible), 1 = served."),
    ("Eligible", "Classified as unserved (no cable or fiber at 25/3 Mbps or better). Not eligible = classified as served."),
    ("Lost eligibility / Became eligible", "Same location ID, classified eligible in one Fabric version and not "
                                           "eligible in the other. This is a change in classification."),
    ("ID only in Fabric 4 / Fabric 8", "The location ID exists in one version only. It is not matched to a nearby ID: "
                                       "being close does not show two IDs are the same building."),
    ("Inside / outside the area", "Whether the Fabric point of the location ID falls inside the awarded project polygon."),
    ("Moved into / out of the area", "Same ID, its Fabric point is inside the area in one version and outside in the "
                                     "other (a map coordinate change)."),
    ("Conflicting Fabric records", "One Fabric file lists the ID more than once with different eligibility. The ID is left "
                                   "out of the eligible counts and never treated as eligible."),
    ("Initial list", "The locations in the initial grantee submission (Fabric 4 IDs)."),
    ("2026 list", "The grantee's planned locations in the 2026 quarterly report (Q3, or Q2 when Q3 is missing; Fabric 8 IDs)."),
    ("In both lists / Removed / Added", "Location ID in both lists, only in the initial list, or only in the 2026 list."),
    ("Entries", "Rows of location data in a grantee file (for Q3 2026 reports, from row 8 down; the template rows "
                "above are not read)."),
    ("Unique location IDs", "Entries with a valid location ID, each ID counted once. Entries − entries without an ID − "
                            "entries with an invalid ID − repeat entries = unique location IDs."),
    ("Repeat entry", "A row whose location ID was already listed in an earlier row of the same file; counted once."),
    ("Rows without an ID / invalid ID", "Rows with an empty ID field, or text that is not a whole number; they cannot be "
                                        "placed on a map and are listed in the List quality tab."),
    ("For review", "2026 list: not in Fabric 8, conflicting records, outside the area, not eligible, or not a "
                        "broadband serviceable location. Removed although still eligible inside the area. Added although "
                        "already eligible inside the area in Fabric 4."),
    ("Coverage", "Eligible locations inside the area in Fabric 8 that the 2026 list includes."),
]

# =====================================================================================================================
# 9. Geodatabase, workbook and Markdown report per project
# =====================================================================================================================
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

GDB_OPTS = {"TARGET_ARCGIS_VERSION": "ARCGIS_PRO_3_2_OR_LATER"}
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
HEAD_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14, color="1F3864")
SECTION_FONT = Font(bold=True, size=11, color="1F3864")
PLAIN_NUMBER_COLS = {"Location ID", "Row in file"}             # identifiers: no thousands separator
SIGNED_COLS = {"Change", "Effect on the eligible count"}        # differences: always with a sign
THIN = Border(bottom=Side(style="thin", color="D0D0D0"))
WRAP = Alignment(wrap_text=True, vertical="top")


def write_layer(df: gpd.GeoDataFrame | pd.DataFrame, gdb: Path, layer: str) -> None:
    out = df.copy()
    for c in out.columns:
        if c != "geometry" and (out[c].dtype == object or str(out[c].dtype).startswith("string")):
            out[c] = out[c].astype("string").str.slice(0, 2000)
    pyogrio.write_dataframe(out, gdb, layer=layer, driver="OpenFileGDB", **GDB_OPTS)


def project_gdb(key: str, path: Path) -> None:
    if path.exists():
        shutil.rmtree(path)
    area = areas.loc[areas["project_key"] == key, ["geometry"]].assign(project=LABEL[key], pr_number=PR_NUMBERS[key])
    write_layer(gpd.GeoDataFrame(area, crs=WGS84)[["project", "pr_number", "geometry"]], path, "project_area")
    write_layer(location_layer(key), path, "submitted_locations")
    write_layer(points_from_lonlat(fabric_change_table(key)), path, "fabric_changes")


def put_table(ws, df: pd.DataFrame, row: int, title: str | None = None, widths: dict | None = None) -> int:
    """Writes a titled table at `row`; returns the next free row."""
    if title:
        ws.cell(row, 1, title).font = SECTION_FONT
        row += 1
    for j, c in enumerate(df.columns, 1):
        cell = ws.cell(row, j, c)
        cell.fill, cell.font, cell.alignment = HEAD_FILL, HEAD_FONT, Alignment(wrap_text=True, vertical="center")
    for i, r in enumerate(df.itertuples(index=False), row + 1):
        for j, v in enumerate(r, 1):
            v = None if (isinstance(v, float) and np.isnan(v)) else (v.item() if hasattr(v, "item") else v)
            cell = ws.cell(i, j, v)
            cell.border, cell.alignment = THIN, WRAP
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                col = df.columns[j - 1]
                cell.number_format = ("0" if col in PLAIN_NUMBER_COLS else
                                      "+#,##0;-#,##0;0" if col in SIGNED_COLS else "#,##0")
        if str(r[0]).startswith(("Total", "Locations for review")):
            for j in range(1, len(df.columns) + 1):
                ws.cell(i, j).font = Font(bold=True)
    return row + len(df) + 2


def print_setup(wb) -> None:
    """Landscape, one page wide, so printed sheets keep every column."""
    for ws in wb.worksheets:
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True


def sheet_from(wb, name: str, df: pd.DataFrame, widths: list[float]) -> None:
    ws = wb.create_sheet(name)
    put_table(ws, df, 1)
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"
    if len(df):
        ws.auto_filter.ref = f"A1:{get_column_letter(len(df.columns))}{len(df) + 1}"


def project_workbook(key: str, path: Path) -> None:
    st = STORIES[key]
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = f"{LABEL[key]} ({PR_NUMBERS[key]})"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = (f"Fabric version 4 → version 8 and initial grantee submission → {quarter(key)} list. "
                f"2026 file: {sub_v8_files.get(key) or 'none found'}"
                + (f" (data read from row {first_row[key]})" if first_row.get(key) else "") + f". Run on {RUN_DATE}.")
    ws["A2"].font = Font(italic=True, color="5F6368")
    ws["A4"] = "What changed"
    ws["A4"].font = SECTION_FONT
    row = 5
    for txt in st.values():
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        ws.cell(row, 1, txt).alignment = WRAP
        ws.row_dimensions[row].height = 15 * (1 + len(txt) // 105)
        row += 1
    row += 1
    for title, df in summary_tables(key).items():
        row = put_table(ws, df, row, title)
    for col, w in zip("ABCD", (46, 30, 30, 22)):
        ws.column_dimensions[col].width = w
    sheet_from(wb, "For review", review_table(key), [14, 40, 70, 60])
    sheet_from(wb, "List quality", list_quality(key), [16, 36, 12, 24, 14, 12])
    fct = fabric_change_table(key).drop(columns=["lon", "lat"]).rename(columns={
        "location_id": "Location ID", "change": "What happened", "eligible_fabric_4": "Eligible in Fabric 4",
        "eligible_fabric_8": "Eligible in Fabric 8", "inside_area_fabric_4": "Inside the area in Fabric 4",
        "inside_area_fabric_8": "Inside the area in Fabric 8", "in_initial_list": "In the initial list",
        "in_2026_list": "In the 2026 list"})
    sheet_from(wb, "Fabric changes", fct, [14, 30, 12, 12, 14, 14, 12, 12])
    sheet_from(wb, "Definitions", pd.DataFrame(DEFINITIONS, columns=["Term", "Meaning"]), [34, 110])
    print_setup(wb)
    wb.save(path)


def fmt_cell(v, col: str) -> str:
    if v is None or (isinstance(v, float) and np.isnan(v)) or v is pd.NA:
        return ""
    if isinstance(v, (int, np.integer, float)) and not isinstance(v, bool):
        v = int(v)
        return str(v) if col in PLAIN_NUMBER_COLS else (signed(v) if col in SIGNED_COLS else f"{v:,}")
    return str(v).replace("|", "/")


def md_table(df: pd.DataFrame) -> str:
    return ("| " + " | ".join(map(str, df.columns)) + " |\n|" + "---|" * len(df.columns) + "\n"
            + "".join("| " + " | ".join(fmt_cell(v, c) for v, c in zip(r, df.columns)) + " |\n"
                      for r in df.itertuples(index=False)))


def project_markdown(key: str, path: Path) -> None:
    st, fu, lq = STORIES[key], review_table(key), list_quality(key)
    out = [f"# {LABEL[key]} ({PR_NUMBERS[key]})\n",
           f"Fabric version 4 → version 8 and initial grantee submission → {quarter(key)} list. Run on {RUN_DATE}.\n",
           "## What changed\n", *[f"- {t}" for t in st.values()], ""]
    for title, df in summary_tables(key).items():
        out += [f"## {title}\n", md_table(df)]
    out += [f"## Locations for review ({len(fu):,})\n",
            md_table(fu.drop(columns="What changed")) if len(fu) else "None.\n",
            f"## List quality ({len(lq):,} rows)\n", md_table(lq.dropna(axis=1, how="all")) if len(lq) else "None.\n",
            "Every changed Fabric location ID is in the Fabric changes tab of the workbook and in the `fabric_changes` "
            "layer of the geodatabase; every submitted location ID with its flags is in `submitted_locations`.\n",
            "## Definitions\n", md_table(pd.DataFrame(DEFINITIONS, columns=["Term", "Meaning"]))]
    path.write_text("\n".join(out), encoding="utf-8")


def overview_workbook(path: Path) -> None:
    rows = []
    for key in PROJECTS:
        k = FACTS[key]
        rows.append({"Project": LABEL[key], "PR number": PR_NUMBERS[key], "2026 list used": quarter(key),
                     "Eligible in the area, Fabric 4": k["f4_elig"], "Eligible in the area, Fabric 8": k["f8_elig"],
                     "Change in eligible": k["f8_elig"] - k["f4_elig"],
                     "Lost eligibility": k["f_changes"].get(F_LOST, 0), "Became eligible": k["f_changes"].get(F_GAINED, 0),
                     "Initial list: entries": k["s4_rows"], "Initial list: unique location IDs": k["s4_listed"],
                     "2026 list: entries": k["s8_rows"], "2026 list: unique location IDs": k["s8_listed"],
                     "In both lists": k["kept"],
                     "Removed": k["removed"], "Added": k["added"],
                     "2026 list: eligible inside": k["s8_in_elig"],
                     "2026 list: outside the area": k["s8_out_elig"] + k["s8_out_nelig"],
                     "2026 list: inside, not eligible": k["s8_in_nelig"],
                     "2026 list: not in Fabric 8": k["s8_not_found"],
                     "Rows without or with an invalid ID": k["s4_no_id"] + k["s4_bad_id"] + k["s8_no_id"] + k["s8_bad_id"],
                     "Repeat entries (ID already listed)": k["s4_dup_rows"] + k["s8_dup_rows"],
                     "Locations for review": k["fu_total"],
                     "Coverage (2026 list of eligible in the area)": f"{k['cov_listed']:,} of {k['cov_eligible']:,}"})
    df = pd.DataFrame(rows)
    wb = Workbook()
    ws = wb.active
    ws.title = "Projects"
    put_table(ws, df, 1)
    ws.row_dimensions[1].height = 45
    for j in range(1, len(df.columns) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 34 if j == 1 else 14
    ws.freeze_panes = "B2"
    sheet_from(wb, "Definitions", pd.DataFrame(DEFINITIONS, columns=["Term", "Meaning"]), [34, 110])
    print_setup(wb)
    wb.save(path)
    log.info("overview → %s", path)


# =====================================================================================================================
# 10. Map pages (PDF): Fabric eligibility change, grantee list change, locations for review – on gray canvas and on imagery
# =====================================================================================================================
import textwrap

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
from matplotlib.lines import Line2D
from matplotlib.patches import Rectangle

import io
import math

from PIL import Image

MERC = "EPSG:3857"
PAGE_W, PAGE_H = 14.0, 8.5                            # inches, legal landscape
MAP_LEFT_IN, MAP_BOTTOM_IN = 0.35, 0.75
MAP_FRAME_W_IN, MAP_FRAME_H_IN = 9.0, 6.75
PANEL_X_IN = MAP_LEFT_IN + MAP_FRAME_W_IN + 0.35
PANEL_W_IN = PAGE_W - PANEL_X_IN - 0.35
PAD_FRACTION = 0.06
NAVY, INK, QUIET = "#1F3864", "#1d1d1f", "#5f6368"
# Esri tile services: (URL with {z}/{y}/{x}, highest zoom level, credit)
BASEMAPS = {"gray": ("https://server.arcgisonline.com/ArcGIS/rest/services/Canvas/World_Light_Gray_Base/MapServer/tile/"
                     "{z}/{y}/{x}", 16, "Basemap: Esri, HERE, Garmin, © OpenStreetMap contributors"),
            "imagery": ("https://server.arcgisonline.com/ArcGIS/rest/services/World_Imagery/MapServer/tile/{z}/{y}/{x}",
                        19, "Basemap: Esri, Maxar, Earthstar Geographics")}
basemap_ok = {"gray": True, "imagery": True}
basemap_error: dict[str, str] = {}
TILE_PIXELS_ACROSS = 1800          # basemap detail: tile pixels across the map frame
MAX_TILES = 120                    # tiles per map at most (zoom drops until the map fits)
_tile_cache: dict[tuple, Image.Image] = {}
WORLD = 2 * math.pi * 6378137      # Web Mercator world width (m)
DESC_SIZE, LEGEND_SIZE, LEGEND_MARKER = 10.0, 9.0, 7.5
POINT_SIZE = {"main": 26, "context": 9}               # marker area (pt²)
LOCATOR_H_IN, LOCATOR_MIN_H_IN = 1.3, 0.65
HEX_ACROSS = 160                                      # hexagons across the map width for the shaded group
SHADE_RGB = {1: (204, 204, 204), 2: (170, 170, 170), 3: (130, 130, 130)}

RGB = {"gold": (230, 171, 2), "teal": (1, 133, 113), "lilac": (175, 141, 195), "plum": (120, 98, 160),
       "brown": (140, 81, 10), "blue": (33, 102, 172), "olive": (128, 128, 0), "slate": (99, 99, 99),
       "skyblue": (86, 180, 233), "grey": (150, 150, 150), "lightgrey": (200, 200, 200), "purple": (118, 42, 131),
       "pink": (197, 27, 125), "darkred": (128, 0, 38), "charcoal": (45, 45, 45), "darkorange": (217, 95, 2)}

# Page classes: label → (colour, group, size); "shaded" draws as hexagon density instead of dots
PAGE_CLASSES = {
    "fabric": {F_LOST: ("gold", "Changed", "main"), F_GAINED: ("teal", "Changed", "main"),
               F_IN: ("lilac", "Changed", "main"), F_OUT: ("plum", "Changed", "main"),
               F_ONLY4: ("brown", "Changed", "main"), F_ONLY8: ("blue", "Changed", "main"),
               F_CONFLICT: ("olive", "Changed", "main"), F_BSL: ("slate", "Changed", "main"),
               F_SAME_E: ("skyblue", "No change", "context"), F_SAME_N: ("lightgrey", "No change", "shaded")},
    "lists": {LIST_REMOVED: ("purple", "Grantee lists", "main"), LIST_ADDED: ("pink", "Grantee lists", "main"),
              LIST_KEPT: ("grey", "Grantee lists", "main")},
    "follow": {**{FU_LABEL[c]: (col, "For review (main reason)", "main") for (c, _, _), col in zip(REVIEW, (
        "darkred", "olive", "charcoal", "darkorange", "slate", "purple", "pink"))},
               "Other listed location": ("lightgrey", "Not flagged", "context")},
}
FOOTER = (f"Sources: Fabric version 4 (field {V4_FIELD_NAME}) and version 8 (field {V8_FIELD_NAME}) eligibility data, "
          "0 = unserved (eligible), 1 = served; CPUC FFA awarded project areas; initial grantee submission and 2026 "
          "quarterly report. Points sit at the Fabric location of each ID.")


def fx(x_in: float) -> float:
    return x_in / PAGE_W


def fy(y_in: float) -> float:
    return y_in / PAGE_H


def rgb01(name: str) -> tuple:
    return tuple(v / 255 for v in RGB[name])


def fit_text(fig, text: str, x_in: float, y_top_in: float, w_in: float, h_in: float,
             size: float = 10.5, min_size: float = 7.0, **kw):
    """Draws wrapped text inside a box (inches), shrinking the font until it fits both ways."""
    renderer = fig.canvas.get_renderer()
    color = kw.pop("color", INK)
    s = size
    while s >= min_size:
        chars = int(w_in * 72 / (s * 0.50))
        while chars > 15:
            wrapped = "\n".join(textwrap.fill(p, chars) if p else "" for p in text.split("\n"))
            t = fig.text(fx(x_in), fy(y_top_in), wrapped, fontsize=s, va="top", ha="left", linespacing=1.3,
                         color=color, **kw)
            bb = t.get_window_extent(renderer=renderer)
            if bb.width / fig.dpi <= w_in:
                if bb.height / fig.dpi <= h_in:
                    return t
                t.remove()
                break
            t.remove()
            chars -= 3
        s -= 0.5
    return fig.text(fx(x_in), fy(y_top_in), textwrap.shorten(text, int(w_in * h_in * 20), placeholder=" …"),
                    fontsize=min_size, va="top", ha="left", wrap=True, color=color)


def page_frame(fig, title: str, subtitle: str, page_no: int, n_pages: int, credit: str = "") -> None:
    """Title band and footer shared by every page."""
    renderer = fig.canvas.get_renderer()
    for text, y_in, size, min_size, weight, colour in ((title, PAGE_H - 0.42, 17, 11, "bold", NAVY),
                                                       (subtitle, PAGE_H - 0.72, 10.5, 8, "normal", QUIET)):
        t = fig.text(fx(0.35), fy(y_in), text, fontsize=size, fontweight=weight, color=colour, va="center")
        while t.get_window_extent(renderer=renderer).width / fig.dpi > PAGE_W - 0.7 and size > min_size:
            size -= 0.5
            t.set_fontsize(size)
    fig.add_artist(Line2D([fx(0.35), fx(PAGE_W - 0.35)], [fy(PAGE_H - 0.92)] * 2, color=NAVY, lw=1.2))
    fig.add_artist(Line2D([fx(0.35), fx(PAGE_W - 0.35)], [fy(0.52)] * 2, color="#d0d0d0", lw=0.6))
    fit_text(fig, FOOTER + (" " + credit + "." if credit else "") + f" Created {RUN_DATE}.",
             0.35, 0.45, PAGE_W - 1.6, 0.36, size=7, min_size=6, color=QUIET)
    fig.text(fx(PAGE_W - 0.35), fy(0.3), f"Page {page_no} of {n_pages}", fontsize=8, color=QUIET, ha="right", va="center")


def panel_heading(fig, text: str, x_in: float, y_in: float) -> None:
    fig.text(fx(x_in), fy(y_in), text, fontsize=11.5, fontweight="bold", color=NAVY, va="top")


def get_tile(kind: str, z: int, x: int, y: int) -> Image.Image:
    """One 256-pixel Esri tile (cached for the run)."""
    k = (kind, z, x, y)
    if k not in _tile_cache:
        r = requests.get(BASEMAPS[kind][0].format(z=z, y=y, x=x), timeout=30,
                         headers={"User-Agent": "FFA-V4-V8-review/1.0"})
        r.raise_for_status()
        _tile_cache[k] = Image.open(io.BytesIO(r.content)).convert("RGB")
    return _tile_cache[k]


def tile_mosaic(kind: str, extent, pixels_across: int) -> tuple[Image.Image, tuple]:
    """Stitched tiles covering a Web Mercator extent, and the extent of the stitched image."""
    x0, y0, x1, y1 = extent
    z = int(np.clip(round(math.log2(WORLD * pixels_across / (256 * (x1 - x0)))), 0, BASEMAPS[kind][1]))
    while True:
        size = WORLD / 2 ** z                                            # tile width (m)
        tx0, tx1 = int((x0 + WORLD / 2) // size), int((x1 + WORLD / 2) // size)
        ty0, ty1 = int((WORLD / 2 - y1) // size), int((WORLD / 2 - y0) // size)
        if (tx1 - tx0 + 1) * (ty1 - ty0 + 1) <= MAX_TILES or z == 0:
            break
        z -= 1
    img = Image.new("RGB", ((tx1 - tx0 + 1) * 256, (ty1 - ty0 + 1) * 256))
    for tx in range(tx0, tx1 + 1):
        for ty in range(ty0, ty1 + 1):
            img.paste(get_tile(kind, z, tx % 2 ** z, ty), ((tx - tx0) * 256, (ty - ty0) * 256))
    ext = (tx0 * size - WORLD / 2, (tx1 + 1) * size - WORLD / 2, WORLD / 2 - (ty1 + 1) * size, WORLD / 2 - ty0 * size)
    return img, ext


def add_basemap(ax, kind: str, pixels_across: int = TILE_PIXELS_ACROSS) -> bool:
    """Esri tiles under the map; False (plain background) when tiles cannot be fetched."""
    ax.set_facecolor("#f4f4f2")
    if not basemap_ok.get(kind):
        return False
    xlim, ylim = ax.get_xlim(), ax.get_ylim()
    try:
        img, ext = tile_mosaic(kind, (xlim[0], ylim[0], xlim[1], ylim[1]), pixels_across)
    except Exception as e:                              # offline, proxy or tile server refused
        log.warning("Basemap '%s' unavailable (%s); maps draw on a plain background", kind, e)
        basemap_ok[kind] = False
        basemap_error[kind] = str(e)
        return False
    ax.imshow(img, extent=ext, interpolation="bilinear", zorder=0)
    ax.set_xlim(xlim), ax.set_ylim(ylim)
    return True


def scale_bar(ax, extent, lat: float) -> None:
    """Scale bar in miles, lower left of the map, on a white backing."""
    x0, y0, x1, y1 = extent
    ground = np.cos(np.radians(lat))                    # Web Mercator metres → ground metres
    target = (x1 - x0) * 0.2 * ground / 1609.344
    exp = 10 ** np.floor(np.log10(max(target, 0.1)))
    length_mi = max(v * exp for v in (1, 2, 5) if v * exp <= max(target, 0.1))
    length = length_mi * 1609.344 / ground
    bx, by = x0 + (x1 - x0) * 0.03, y0 + (y1 - y0) * 0.04
    ax.add_patch(Rectangle((bx - (x1 - x0) * 0.012, by - (y1 - y0) * 0.022), length + (x1 - x0) * 0.024,
                           (y1 - y0) * 0.075, facecolor="white", alpha=0.85, edgecolor="none", zorder=20))
    ax.plot([bx, bx + length], [by, by], color=INK, lw=2.2, solid_capstyle="butt", zorder=21)
    ax.text(bx + length / 2, by + (y1 - y0) * 0.012, f"{length_mi:g} mi", ha="center", va="bottom", fontsize=7.5,
            color=INK, zorder=21)


def north_arrow(ax) -> None:
    ax.annotate("N", xy=(0.965, 0.955), xytext=(0.965, 0.885), xycoords="axes fraction", ha="center", va="center",
                fontsize=9, fontweight="bold", color=INK, arrowprops=dict(arrowstyle="-|>", color=INK, lw=1.2),
                bbox=dict(boxstyle="round,pad=0.25", fc="white", ec="none", alpha=0.85), zorder=22)


def hex_cells(x: np.ndarray, y: np.ndarray, size: float) -> pd.DataFrame:
    """Counts points per pointy-top hexagon of circumradius `size`; returns centre x/y and count per cell."""
    q = (np.sqrt(3) / 3 * x - y / 3) / size
    r = (2 / 3 * y) / size
    cx, cz = q, r
    cy = -cx - cz
    rx, ry, rz = np.round(cx), np.round(cy), np.round(cz)
    dx, dy, dz = np.abs(rx - cx), np.abs(ry - cy), np.abs(rz - cz)
    fix_x = (dx > dy) & (dx > dz)
    fix_y = ~fix_x & (dy > dz)
    rx = np.where(fix_x, -ry - rz, rx)
    rz = np.where(~fix_x & ~fix_y, -rx - ry, rz)
    cells = pd.DataFrame({"q": rx.astype(int), "r": rz.astype(int)}).value_counts().rename("locations").reset_index()
    cells["x"] = size * np.sqrt(3) * (cells["q"] + cells["r"] / 2)
    cells["y"] = size * 1.5 * cells["r"]
    return cells


def draw_shaded(ax, pts: gpd.GeoDataFrame, width: float, imagery: bool) -> None:
    """Hexagon density (three shades of grey) for a large background group."""
    size = width / HEX_ACROSS / np.sqrt(3)
    cells = hex_cells(pts.geometry.x.to_numpy(), pts.geometry.y.to_numpy(), size)
    q1, q2 = np.quantile(cells["locations"], [1 / 3, 2 / 3])
    cells["shade"] = 1 + (cells["locations"] > q1).astype(int) + (cells["locations"] > q2).astype(int)
    ang = np.radians(np.arange(6) * 60 + 30)
    for shade, grp in cells.groupby("shade"):
        polys = [shapely.Polygon(np.column_stack([x + size * np.cos(ang), y + size * np.sin(ang)]))
                 for x, y in zip(grp["x"], grp["y"])]
        gpd.GeoSeries(polys, crs=MERC).plot(ax=ax, color=tuple(v / 255 for v in SHADE_RGB[shade]), edgecolor="none",
                                           alpha=0.6 if imagery else 0.75, zorder=8)


def draw_legend(fig, counts: dict, classes: dict, x_in: float, y_top_in: float, w_in: float, h_in: float) -> float:
    """Legend by group with counts (only classes present); labels wrap by measured width; font shrinks to fit."""
    rows, group = [], None
    for lab, (col, grp, kind) in classes.items():
        if counts.get(lab, 0):
            if grp != group:
                rows.append(("group", grp))
                group = grp
            rows.append(("item", lab))
    if not rows:
        return y_top_in
    label_w = w_in - 0.24 - 0.55
    probe = fig.text(0, 0, "", alpha=0)
    renderer = fig.canvas.get_renderer()

    def wrap(text: str, size: float) -> list[str]:
        words, lines, cur_ = text.split(), [], ""
        for w in words:
            trial = (cur_ + " " + w).strip()
            probe.set_text(trial), probe.set_fontsize(size)
            if cur_ and probe.get_window_extent(renderer=renderer).width / fig.dpi > label_w:
                lines.append(cur_)
                cur_ = w
            else:
                cur_ = trial
        return lines + [cur_]

    for size in np.arange(LEGEND_SIZE, 6.9, -0.25):
        wrapped = {lab: wrap(lab + (" (shaded, darker = more)" if classes[lab][2] == "shaded" else ""), size)
                   for k_, lab in rows if k_ == "item"}
        line = size * 1.45 / 72
        if sum(line * (len(wrapped[l]) + 0.15) if k_ == "item" else line * 1.2 for k_, l in rows) <= h_in:
            break
    probe.remove()
    y = y_top_in
    for kind, lab in rows:
        if kind == "group":
            fig.text(fx(x_in), fy(y), lab, fontsize=size, fontweight="bold", color=QUIET, va="top")
            y -= line * 1.2
            continue
        col, _, style = classes[lab]
        fig.add_artist(Line2D([fx(x_in + 0.09)], [fy(y - line * 0.42)], marker="h" if style == "shaded" else "o",
                              markersize=LEGEND_MARKER * (1.3 if style == "shaded" else 1.0),
                              markerfacecolor=tuple(v / 255 for v in SHADE_RGB[2]) if style == "shaded" else rgb01(col),
                              markeredgecolor="white", linestyle="none"))
        fig.text(fx(x_in + 0.24), fy(y), "\n".join(wrapped[lab]), fontsize=size, color=INK, va="top", linespacing=1.2)
        fig.text(fx(x_in + w_in), fy(y), f"{counts[lab]:,}", fontsize=size, color=INK, va="top", ha="right")
        y -= line * len(wrapped[lab]) + line * 0.15
    return y


# Places labelled on the small overview map: (name, latitude, longitude, rank); rank 1 is drawn first
PLACES = [
    ("San Francisco", 37.7749, -122.4194, 1), ("Oakland", 37.8044, -122.2712, 1), ("San Jose", 37.3382, -121.8863, 1),
    ("Sacramento", 38.5816, -121.4944, 1), ("Stockton", 37.9577, -121.2908, 1), ("Santa Rosa", 38.4405, -122.7144, 1),
    ("Fresno", 36.7378, -119.7871, 1), ("Bakersfield", 35.3733, -119.0187, 1), ("Santa Barbara", 34.4208, -119.6982, 1),
    ("Reno", 39.5296, -119.8138, 1),
    ("Berkeley", 37.8716, -122.2727, 2), ("Hayward", 37.6688, -122.0808, 2), ("Fremont", 37.5485, -121.9886, 2),
    ("Walnut Creek", 37.9101, -122.0652, 2), ("Richmond", 37.9358, -122.3477, 2), ("Vallejo", 38.1041, -122.2566, 2),
    ("Concord", 37.9780, -122.0311, 2), ("San Mateo", 37.5630, -122.3255, 2), ("Daly City", 37.6879, -122.4702, 2),
    ("San Leandro", 37.7249, -122.1561, 2), ("Alameda", 37.7652, -122.2416, 2), ("Folsom", 38.6780, -121.1761, 2),
    ("Roseville", 38.7521, -121.2880, 2), ("Auburn", 38.8966, -121.0769, 2), ("Placerville", 38.7296, -120.7985, 2),
    ("South Lake Tahoe", 38.9399, -119.9772, 2), ("Truckee", 39.3280, -120.1833, 2), ("Jackson", 38.3488, -120.7741, 2),
    ("Grass Valley", 39.2191, -121.0611, 2), ("Davis", 38.5449, -121.7405, 2), ("Elk Grove", 38.4088, -121.3716, 2),
    ("San Luis Obispo", 35.2828, -120.6596, 2), ("Santa Maria", 34.9530, -120.4357, 2),
    ("Paso Robles", 35.6266, -120.6910, 2), ("Atascadero", 35.4894, -120.6707, 2), ("Lompoc", 34.6391, -120.4579, 2),
    ("Emeryville", 37.8313, -122.2852, 3), ("Piedmont", 37.8244, -122.2317, 3), ("Castro Valley", 37.6941, -122.0864, 3),
    ("Orinda", 37.8771, -122.1797, 3), ("Lafayette", 37.8858, -122.1180, 3), ("San Lorenzo", 37.6810, -122.1244, 3),
    ("Morro Bay", 35.3658, -120.8499, 3), ("Los Osos", 35.3111, -120.8324, 3), ("Pismo Beach", 35.1428, -120.6413, 3),
    ("Arroyo Grande", 35.1186, -120.5907, 3), ("Grover Beach", 35.1217, -120.6210, 3), ("Avila Beach", 35.1797, -120.7319, 3),
    ("Cayucos", 35.4428, -120.8921, 3), ("Templeton", 35.5497, -120.7060, 3), ("Santa Margarita", 35.3897, -120.6088, 3),
    ("Nipomo", 35.0428, -120.4760, 3), ("Coloma", 38.7999, -120.8891, 3), ("Georgetown", 38.9068, -120.8388, 3),
    ("Garden Valley", 38.8524, -120.8580, 3), ("Cool", 38.8874, -121.0147, 3), ("Pollock Pines", 38.7616, -120.5861, 3),
    ("Diamond Springs", 38.6946, -120.8149, 3), ("Shingle Springs", 38.6657, -120.9241, 3),
    ("Cameron Park", 38.6688, -120.9872, 3), ("El Dorado Hills", 38.6857, -121.0822, 3), ("Foresthill", 39.0202, -120.8180, 3),
    ("Camino", 38.7385, -120.6750, 3), ("Somerset", 38.6477, -120.6844, 3),
]
_pl = gpd.GeoDataFrame({"name": [p[0] for p in PLACES], "rank": [p[3] for p in PLACES]},
                       geometry=gpd.points_from_xy([p[2] for p in PLACES], [p[1] for p in PLACES]), crs=WGS84).to_crs(MERC)
LOCATOR_H_IN, LOCATOR_MIN_H_IN = 1.3, 0.65


def locator(fig, key: str, extent, kind: str, h_in: float = LOCATOR_H_IN) -> None:
    """Small overview map with labelled towns and the main map's frame (plain background when tiles are missing)."""
    import matplotlib.patheffects as pe
    w_in = PANEL_W_IN
    ax = fig.add_axes([fx(PANEL_X_IN), fy(MAP_BOTTOM_IN), fx(w_in), fy(h_in)])
    x0, y0, x1, y1 = extent
    cx, cy = (x0 + x1) / 2, (y0 + y1) / 2
    span_h = max((y1 - y0) * 6, (x1 - x0) * 6 * h_in / w_in)
    lx0, lx1 = cx - span_h * w_in / h_in / 2, cx + span_h * w_in / h_in / 2
    ly0, ly1 = cy - span_h / 2, cy + span_h / 2
    ax.set_xlim(lx0, lx1), ax.set_ylim(ly0, ly1)
    add_basemap(ax, kind, pixels_across=500)
    ax.set_xlim(lx0, lx1), ax.set_ylim(ly0, ly1)
    ax.add_patch(Rectangle((x0, y0), x1 - x0, y1 - y0, fill=False, edgecolor=NAVY, lw=1.6, zorder=6))
    ax.set_xticks([]), ax.set_yticks([])
    for s_ in ax.spines.values():
        s_.set_edgecolor("#999999"), s_.set_linewidth(0.6)
    # Labels: biggest places first; a label is skipped when it would overlap another label, the frame
    # box or the edge of the overview map
    renderer = fig.canvas.get_renderer()
    ax_box = ax.get_window_extent(renderer=renderer)
    halo = [pe.withStroke(linewidth=2.2, foreground="white")]
    near = _pl.cx[lx0:lx1, ly0:ly1].copy()
    if near.empty:
        return
    # The place nearest the frame centre names the project's own area and is drawn first, in bold
    near["d"] = near.geometry.distance(shapely.Point(cx, cy))
    home = near["d"].idxmin()
    near["order"] = np.where(near.index == home, 0, near["rank"])
    taken = []
    for p in near.sort_values(["order", "d"]).itertuples():
        px, py = p.geometry.x, p.geometry.y
        if p.Index != home and x0 <= px <= x1 and y0 <= py <= y1:
            continue                                      # other places inside the frame box stay unlabelled
        is_home = p.Index == home
        t = ax.text(px, py, " " + p.name, fontsize=7.5 if is_home else 7 if p.rank < 3 else 6.5, color=INK,
                    va="center", ha="left", fontweight="bold" if is_home or p.rank == 1 else "normal",
                    path_effects=halo, zorder=9 if is_home else 8, clip_on=True)
        if is_home:
            ax.plot(px, py, "o", markersize=2.6, color=INK, zorder=9)
            bb = t.get_window_extent(renderer=renderer)
            taken += [[[bb.x0, bb.y0], [bb.x1, bb.y1]], ax.transData.transform([[x0, y0], [x1, y1]])]
            continue
        bb = t.get_window_extent(renderer=renderer)
        clash = any(not (bb.x1 < a[0][0] or bb.x0 > a[1][0] or bb.y1 < a[0][1] or bb.y0 > a[1][1]) for a in taken)
        if clash or bb.x1 > ax_box.x1 - 2 or bb.y0 < ax_box.y0 + 1 or bb.y1 > ax_box.y1 - 1:
            t.remove()
            continue
        ax.plot(px, py, "o", markersize=2.2, color=INK, zorder=8)
        taken.append([[bb.x0, bb.y0], [bb.x1, bb.y1]])



def page_points(key: str) -> dict[str, gpd.GeoDataFrame]:
    """Points per map page in Web Mercator, with the class label in `cls`."""
    f = fc[fc["project_key"] == key].assign(cls=lambda d: d["change"])
    l_ = loc[loc["project_key"] == key]
    lists = l_.assign(cls=l_["list_change"])
    follow = l_.assign(cls=np.where(l_["for_review"], l_["primary_reason"].map(FU_LABEL), "Other listed location"))
    out = {}
    for name, df in (("fabric", f), ("lists", lists), ("follow", follow)):
        g = points_from_lonlat(df[["cls", "lon", "lat"]].reset_index(drop=True)).dropna(subset=["geometry"])
        out[name] = g[~g.geometry.is_empty].to_crs(MERC)
    return out


def page_extent(key: str, pts: dict) -> tuple[tuple, int]:
    """Map extent covering the project area and every point within FAR_POINT_MI of it; also the count left out."""
    area = areas.loc[areas["project_key"] == key].to_crs(MERC).geometry.iloc[0]
    allp = pd.concat([pts["lists"].geometry, pts["fabric"].geometry])
    ground = np.cos(np.radians(areas.loc[areas["project_key"] == key].geometry.iloc[0].centroid.y))
    miles = allp.distance(area) * ground / 1609.344
    near = allp[miles <= FAR_POINT_MI]
    x0, y0, x1, y1 = gpd.GeoSeries(list(near) + [area], crs=MERC).total_bounds
    w, h = max(x1 - x0, 500) * (1 + 2 * PAD_FRACTION), max(y1 - y0, 500) * (1 + 2 * PAD_FRACTION)
    cx, cy = (x0 + x1) / 2, (y0 + y1) / 2
    aspect = MAP_FRAME_W_IN / MAP_FRAME_H_IN
    w, h = (w, w / aspect) if w / h > aspect else (h * aspect, h)
    far = int((pts["lists"].geometry.distance(area) * ground / 1609.344 > FAR_POINT_MI).sum())
    return (cx - w / 2, cy - h / 2, cx + w / 2, cy + h / 2), far


TOTALS_H_IN = 0.95                                  # height of the totals block under the legend


def draw_totals(fig, key: str, x_in: float, y_top_in: float, w_in: float) -> float:
    """Totals of the four datasets, the same on every page; returns the y under the block."""
    k, qr = FACTS[key], quarter(key)
    rows = [(f"Fabric 4, inside the area ({k['f4_elig']:,} eligible)", k["f4_total"]),
            (f"Fabric 8, inside the area ({k['f8_elig']:,} eligible)", k["f8_total"]),
            ("Initial list: entries / unique location IDs", f"{k['s4_rows']:,} / {k['s4_listed']:,}"),
            (f"{qr} list: entries / unique location IDs", f"{k['s8_rows']:,} / {k['s8_listed']:,}")]
    fig.text(fx(x_in), fy(y_top_in), "Totals", fontsize=8.5, fontweight="bold", color=QUIET, va="top")
    y = y_top_in - 0.19
    for lab, n in rows:
        fig.text(fx(x_in + 0.24), fy(y), lab, fontsize=8.5, color=INK, va="top")
        fig.text(fx(x_in + w_in), fy(y), n if isinstance(n, str) else f"{n:,}", fontsize=8.5, color=INK, va="top",
                 ha="right")
        y -= 0.17
    return y


def map_page(pdf, key: str, page: str, pts: gpd.GeoDataFrame, extent, far: int, no_xy: int, kind: str,
             page_no: int) -> None:
    """One map page: map on the left; description, legend and overview map on the right."""
    qr, st = quarter(key), STORIES[key]
    heading, sub, desc = {
        "fabric": ("Fabric eligibility change in the project area",
                   "Fabric version 4 compared with Fabric version 8, every location inside the area in either version",
                   st["fabric"]),
        "lists": ("Grantee list change", f"Initial grantee submission compared with the {qr} list",
                  st["lists"] + " " + st["current"]),
        "follow": ("Locations for review", "Listed locations flagged for review, by main reason",
                   st["follow"] + " " + st["quality"] + " Location IDs and questions: For review tab of the workbook."),
    }[page]
    classes = PAGE_CLASSES[page]
    imagery = kind == "imagery"
    fig = plt.figure(figsize=(PAGE_W, PAGE_H))
    ax = fig.add_axes([fx(MAP_LEFT_IN), fy(MAP_BOTTOM_IN), fx(MAP_FRAME_W_IN), fy(MAP_FRAME_H_IN)])
    ax.set_xlim(extent[0], extent[2]), ax.set_ylim(extent[1], extent[3])
    tiles = add_basemap(ax, kind)
    page_frame(fig, f"{LABEL[key]} ({PR_NUMBERS[key]}) · {heading}", sub, page_no, 3,
               credit=BASEMAPS[kind][2] if tiles else "")
    ax.set_xlim(extent[0], extent[2]), ax.set_ylim(extent[1], extent[3])
    areas.loc[areas["project_key"] == key].to_crs(MERC).boundary.plot(
        ax=ax, color="white" if imagery else NAVY, linewidth=1.6, zorder=15)
    counts = pts["cls"].value_counts().to_dict()
    # Background groups first, then changes on top in legend order reversed (first-listed drawn last)
    for lab in reversed(list(classes)):
        col, _, style = classes[lab]
        g = pts[pts["cls"] == lab]
        if g.empty:
            continue
        if style == "shaded":
            draw_shaded(ax, g, extent[2] - extent[0], imagery)
        else:
            ax.scatter(g.geometry.x, g.geometry.y, s=POINT_SIZE[style], c=[rgb01(col)], edgecolors="white",
                       linewidths=0.35 if imagery else 0.25, zorder=12 if style == "main" else 10)
    ax.set_xticks([]), ax.set_yticks([])
    ax.set_xlabel(""), ax.set_ylabel("")
    for s_ in ax.spines.values():
        s_.set_edgecolor("#999999"), s_.set_linewidth(0.6)
    lat = gpd.GeoSeries([shapely.Point((extent[0] + extent[2]) / 2, (extent[1] + extent[3]) / 2)],
                        crs=MERC).to_crs(WGS84).y.iloc[0]
    scale_bar(ax, extent, lat)
    north_arrow(ax)
    notes = []
    if far and page != "fabric":
        notes.append(f"{far:,} listed location(s) more than {FAR_POINT_MI} miles from the area are not shown")
    if no_xy and page != "fabric":
        notes.append(f"{no_xy:,} listed location(s) without coordinates are not shown")
    if notes:
        ax.text(0.99, 0.02, "\n".join(notes),
                transform=ax.transAxes, ha="right", va="bottom", fontsize=7.5, color=INK, zorder=23,
                bbox=dict(boxstyle="round,pad=0.3", fc="white", ec="none", alpha=0.9))
    top = MAP_BOTTOM_IN + MAP_FRAME_H_IN
    panel_heading(fig, "Description", PANEL_X_IN, top)
    t = fit_text(fig, desc, PANEL_X_IN, top - 0.3, PANEL_W_IN, 2.5, size=DESC_SIZE, min_size=8)
    desc_h = t.get_window_extent(renderer=fig.canvas.get_renderer()).height / fig.dpi
    leg_top = top - 0.3 - desc_h - 0.3
    panel_heading(fig, "Legend", PANEL_X_IN, leg_top)
    fig.text(fx(PANEL_X_IN + PANEL_W_IN), fy(leg_top), "Locations", fontsize=8.5, color=QUIET, ha="right", va="top")
    y_end = draw_legend(fig, counts, classes, PANEL_X_IN, leg_top - 0.32, PANEL_W_IN,
                        leg_top - 0.32 - (MAP_BOTTOM_IN + LOCATOR_MIN_H_IN + 0.35 + TOTALS_H_IN))
    y_area = y_end - 0.12
    if imagery:                                         # white line on a dark edge, as drawn on the imagery
        fig.add_artist(Line2D([fx(PANEL_X_IN + 0.02), fx(PANEL_X_IN + 0.18)], [fy(y_area)] * 2, color=INK, lw=3.2))
    fig.add_artist(Line2D([fx(PANEL_X_IN + 0.02), fx(PANEL_X_IN + 0.18)], [fy(y_area)] * 2,
                          color="white" if imagery else NAVY, lw=1.6))
    fig.text(fx(PANEL_X_IN + 0.24), fy(y_area), "Project area" + (" (white on imagery)" if imagery else ""),
             fontsize=8.5, color=INK, va="center")
    y_tot = draw_totals(fig, key, PANEL_X_IN, y_area - 0.22, PANEL_W_IN)
    locator(fig, key, extent, kind, h_in=float(np.clip(y_tot - 0.2 - MAP_BOTTOM_IN, LOCATOR_MIN_H_IN, LOCATOR_H_IN)))
    pdf.savefig(fig)
    plt.close(fig)


def project_maps(key: str, folder: Path) -> list[Path]:
    pts = page_points(key)
    extent, far = page_extent(key, pts)
    no_xy = int((loc["project_key"].eq(key) & (loc["lon"].isna() | loc["lat"].isna())).sum())
    paths = []
    for kind in BASEMAP_KINDS:
        path = folder / f"{FOLDER[key]}_maps_{kind}.pdf"
        with PdfPages(path) as pdf:
            for i, page in enumerate(("fabric", "lists", "follow"), 1):
                map_page(pdf, key, page, pts[page], extent, far, no_xy, kind, i)
        paths.append(path)
    return paths


# =====================================================================================================================
# 11. Write everything
# =====================================================================================================================
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
for key in PROJECTS:
    folder = OUTPUT_DIR / FOLDER[key]
    folder.mkdir(parents=True, exist_ok=True)
    project_gdb(key, folder / f"{FOLDER[key]}_locations.gdb")
    project_workbook(key, folder / f"{FOLDER[key]}_results.xlsx")
    project_markdown(key, folder / f"{FOLDER[key]}_results.md")
    if WRITE_MAPS:
        project_maps(key, folder)
    log.info("%s → %s", key, folder)
overview_workbook(OUTPUT_DIR / "FFA_V4_V8_all_projects_overview.xlsx")

# Checks: every submitted ID appears once per project; map classes match the flags
assert not loc.duplicated(["project_key", "location_id"]).any()
for key in PROJECTS:
    k = FACTS[key]
    assert k["kept"] + k["removed"] == k["s4_listed"] and k["kept"] + k["added"] == k["s8_listed"], key
for kind, err in basemap_error.items():
    log.warning("No %s basemap on the maps: %s. Check internet or proxy access to server.arcgisonline.com.", kind, err)
print("Done. Results in", OUTPUT_DIR)
